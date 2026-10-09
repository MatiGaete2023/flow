"""Sondas de auditoría con datos sintéticos; no corrige ni accede a instituciones.

Ejecutar desde el checkout revisado con dependencias dev/Descargador instaladas.
La salida describe observaciones, no representa aceptación institucional ni tests
de regresión aprobados. Todos los productos temporales se eliminan al terminar.
"""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from pathlib import Path
import random
import string
import sys
from tempfile import TemporaryDirectory

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'src'), str(ROOT / 'descargador'), str(ROOT / 'tests')]

from openpyxl import load_workbook
from nurus.capture import read_capture, write_capture
from nurus.capture.acceptance import template, validate_acta
from nurus.capture.fixtures import fixture
from nurus.capture.requests import read_request, write_request
from nurus.personal.capture_input import import_capture
from nurus.personal.capture_products import evidence_text, request_document
from nurus.personal.config import defaults
from nurus.personal.record_edits import apply

CUT = date(2026, 10, 8)


def run(folder):
    findings = {}
    metadata, tables = fixture()
    first = write_capture(folder / 'base.zip', metadata, tables)
    work = import_capture(first, defaults(), as_of=CUT)

    # Un texto ASCII supera el límite sin exceder los límites del contrato ZIP.
    long_meta, long_tables = fixture(1)
    raw = json.loads(long_tables['registros'][0]['normalized_json'])
    text = ''.join(random.Random(20261009).choices(string.ascii_letters, k=33000))
    raw['NOMBRE'] = text
    long_tables['registros'][0].update(raw_json=json.dumps(raw), normalized_json=json.dumps(raw))
    long_work = import_capture(write_capture(folder / 'long.zip', long_meta, long_tables), defaults(), as_of=CUT)
    book = load_workbook(BytesIO(long_work.content))
    headers = {c.value: c.column for c in book.active[1]}
    saved = book.active.cell(2, headers[long_work.mapping['nombre']]).value
    findings['H01'] = {'input_characters': len(text), 'workbook_characters': len(saved),
                      'silent_truncation': saved != text, 'import_raised_error': False}
    book.close()

    changed = deepcopy(work)
    old_rules = list(changed.rows[0].rules)
    apply(changed, [changed.rows[0].id], overrides={'espera': '0'})
    current_rules = changed.rows[0].rules
    stored_rules = [f['rule_id'] for f in changed.capture_context[changed.rows[0].id]['findings']
                    if not f['rule_id'].startswith('CT-')]
    findings['H02'] = {'before_rules': old_rules, 'after_rules': current_rules,
                      'stored_findings': stored_rules, 'outdated_findings': stored_rules != current_rules}

    reviewed = deepcopy(work)
    reviewed.rows[0].review['OBSERVACION'] = 'Revisión humana ficticia conservable'
    reviewed.rows[0].excluded = True
    replacement = deepcopy(metadata)
    replacement['downloader']['commit'] = 'OTRO_BUILD_FICTICIO'
    new_work = import_capture(write_capture(folder / 'new.zip', replacement, tables),
                              defaults(), as_of=CUT, previous=reviewed)
    findings['H03'] = {'same_capture_id': reviewed.capture_id == new_work.capture_id,
                      'same_record_ids': [r.id for r in reviewed.rows] == [r.id for r in new_work.rows],
                      'same_sources': reviewed.capture_sources == new_work.capture_sources,
                      'different_zip_digest': reviewed.capture_digest != new_work.capture_digest,
                      'human_review_carried_forward': new_work.rows[0].review.get('OBSERVACION') == reviewed.rows[0].review['OBSERVACION'],
                      'exclusion_carried_forward': new_work.rows[0].excluded,
                      'previous_object_preserved': reviewed.rows[0].excluded}

    # Cobertura parcial de la consulta/fila contradice el agregado de tabla.
    partial_meta, partial_tables = fixture()
    partial_tables['consultas'][0].update(coverage='PARCIAL', expected_pages='2', saved_pages='1')
    for name in ('registros', 'bitacoras'):
        for row in partial_tables[name]:
            if row['query_id'] == 'q1': row['coverage'] = 'PARCIAL'
    partial_path = write_capture(folder / 'contradictory.zip', partial_meta, partial_tables)
    partial_capture = read_capture(partial_path)
    partial_work = import_capture(partial_path, defaults(), as_of=CUT)
    requests = request_document(partial_work)['requests']
    findings['H04'] = {'package_accepted': True,
                      'query_coverage': partial_capture.rows('consultas')[0]['coverage'],
                      'diary_row_coverage': partial_capture.rows('bitacoras')[0]['coverage'],
                      'declared_diary_table': partial_capture.manifest['table_states']['bitacoras'],
                      'diary_requests_for_partial_ingreso': sum(r['source'] == 'bitacoras' and r['ingreso_id'] == 'ingreso-ficticio-0' for r in requests)}

    from test_resultados import fixture as lot_fixture
    from resultados import load_lot
    from export_csmp_capture import export_capture
    lot_folder = lot_fixture(folder / 'calendar', mode='Informes', complete=False)
    manifest_path = lot_folder / 'verificacion 001.json'
    manifest = json.loads(manifest_path.read_text())
    manifest['seleccion']['screen'] = 'calendario_informes'
    manifest.update(estado='INCOMPLETA', paginas_esperadas=2)
    manifest_path.write_text(json.dumps(manifest), encoding='utf-8')
    summary_path = lot_folder / 'resumen.json'
    summary = json.loads(summary_path.read_text())
    summary.update(consultas=[], consultas_completadas=0, plan=[manifest['seleccion']])
    summary_path.write_text(json.dumps(summary), encoding='utf-8')
    exported = read_capture(export_capture([load_lot(lot_folder, partial=True)], folder / 'calendar.zip'))
    findings['H05'] = {'query_coverage': exported.rows('consultas')[0]['coverage'],
                      'calendar_row_coverages': sorted({r['coverage'] for r in exported.rows('calendarios')}),
                      'calendar_table_state': exported.manifest['table_states']['calendarios']}

    acta = template({'source_commit': 'a' * 40, 'assistant_version': '0.5.0.dev4',
                     'downloader_version': '2.7.2', 'extension_version': '2.7.2'}, cases=['A-T01'])
    record = acta['records'][0]
    record.update(gate='G5', evidence_kind='SINTETICA_LOCAL', environment='LINUX',
                  result='PASÓ', executed_at='2026-10-08T15:00:00Z',
                  counts={'expected': 1, 'observed': 1, 'matched': 200000, 'failed': 0, 'unknown': 0},
                  evidence_retained_locally=True, evidence_sha256=['a' * 64], operator_reference='OP-000000000001')
    validate_acta(acta)
    findings['H06'] = {'contradictory_counts_accepted': True, 'counts': record['counts'], 'result': record['result']}

    request = {'schema_id': 'CSMP_REQUEST', 'schema_version': '1.0.0', 'capture_sha256': 'a' * 64,
               'as_of': '2026-10-08', 'requests': [
                   {'record_id': f'r{i}', 'source_system': 'RUS', 'court_code': 'FICTICIO',
                    'cause_id': 'causa-ficticia', 'ingreso_id': f'ingreso-ficticio-{i}',
                    'source': 'bitacoras', 'reason': 'x' * 2000} for i in range(2200)]}
    request_path = write_request(folder / 'request.json', request)
    try:
        read_request(request_path)
    except ValueError as exc:
        findings['H07'] = {'write_succeeded': True, 'bytes': request_path.stat().st_size,
                          'read_succeeded': False, 'error': str(exc)}
    else:
        findings['H07'] = {'write_succeeded': True, 'read_succeeded': True}

    from test_bitacora_lote import manifest as diary_fixture
    from nurus.capture.diaries import import_lote
    diary_folder = folder / 'diary'; diary_folder.mkdir()
    diary_path = diary_fixture(diary_folder)
    diary_meta = json.loads(diary_path.read_text())
    diary_meta['registros'] = diary_meta['registros'][:1]
    diary_meta['registros'][0]['nombre'] = 'Otro nombre visible ficticio'
    diary_path.write_text(json.dumps(diary_meta), encoding='utf-8')
    parsed = import_lote(diary_path)
    normal_folder = lot_fixture(folder / 'normal')
    diary_capture = read_capture(export_capture([load_lot(normal_folder)], folder / 'diary.zip', diaries=diary_path))
    findings['H08'] = {'source_warnings': parsed['queries'][0]['warnings'],
                      'exported_incident_codes': [r['code'] for r in diary_capture.rows('incidencias')],
                      'exported_diary_query_error': next(q['error'] for q in diary_capture.rows('consultas') if q['screen'] == 'bitacoras'),
                      'exported_identity_states': sorted({r['identity_status'] for r in diary_capture.rows('bitacoras')})}
    incident_meta, incident_tables = fixture()
    incident_tables['incidencias'].append({
        'capture_id': incident_meta['capture_id'], 'query_id': 'q1', 'record_id': 'r0',
        'incident_id': 'incident-ficticio', 'phase': 'LECTURA', 'code': 'COTEJO_FICTICIO',
        'severity': 'ADVERTENCIA', 'message': 'MARCADOR_FICTICIO_DE_INCIDENCIA', 'recoverable': 'true'})
    incident_meta['table_states']['incidencias'] = 'COMPLETA'
    incident_path = write_capture(folder / 'incident.zip', incident_meta, incident_tables)
    incident_work = import_capture(incident_path, defaults(), as_of=CUT)
    incident_work.save(folder / 'session')
    findings['H08'].update(
        input_incident_message='MARCADOR_FICTICIO_DE_INCIDENCIA',
        incident_present_in_saved_work='MARCADOR_FICTICIO_DE_INCIDENCIA' in (folder / 'session/trabajo.json').read_text(encoding='utf-8'),
        incident_present_in_evidence_view='MARCADOR_FICTICIO_DE_INCIDENCIA' in evidence_text(incident_work, incident_work.rows[0]))
    return findings


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    with TemporaryDirectory(prefix='flow-review-') as directory:
        observations = run(Path(directory))
    print(json.dumps({'data_kind': 'SINTETICA', 'institutional_access': False,
                      'observations': observations}, ensure_ascii=False, indent=2))
