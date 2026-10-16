"""Capturas auténticas y práctica del manual 2.7.4, exclusivamente ficticias.

Ejecutar con el entorno de desarrollo, Tk y DISPLAY disponible:
  DISPLAY=:91 python tools/manual_descargador_capturas.py --demo-dir /tmp/manual-demo
No consulta SITFA ni abre datos de usuarios. Rechaza destinos de práctica existentes.
"""

from __future__ import annotations

import argparse
import asyncio
import json
from pathlib import Path
import sys
import tempfile
import threading
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'src'), str(ROOT / 'descargador')]
IMAGES = ROOT / 'docs/manual_descargador_2_7_4/imagenes'


def desktop(demo: Path):
    import tkinter as tk
    from PIL import ImageGrab
    from descargador import Application, TribunalPicker
    from test_lotes import CATALOG, dispose_interface
    from test_reutilizacion_local import bound_lot, diaries
    from test_lectura_bitacoras import Bridge as DiaryFixture
    from lectura_bitacoras import DiaryRunner
    from lotes import Batch
    from ventana_resultados import ResultsWindow

    first = bound_lot(demo / '01_Cumplimiento')
    second = bound_lot(demo / '02_Informes_Por_vencer', 'Informes')
    previous = diaries(demo / '03_Bitacoras_anteriores', 800)
    exports = demo / '04_Productos'
    exports.mkdir()
    root = tk.Tk()
    app = Application(root, start_worker=False)
    app.code.set('CÓDIGO DE EJEMPLO · NO VÁLIDO')
    app.dest.set(str(demo))
    app.sound.set(False)
    root.geometry('1060x1300+15+15')

    def settle(window=None, timeout=30):
        deadline = time.monotonic() + timeout
        while True:
            root.update()
            if window is None or not window.busy:
                root.update_idletasks()
                root.update()
                break
            if time.monotonic() > deadline:
                raise TimeoutError('Operación local no terminada')
            time.sleep(0.02)
        time.sleep(0.1)
        root.update()
        # El puntero fuera de los botones evita un resaltado accidental por hover.
        root.event_generate('<Motion>', warp=True, x=3, y=3)
        root.update()

    def shot(name, widget, through=None):
        widget.lift() if isinstance(widget, tk.Toplevel) else None
        settle()
        x, y = widget.winfo_rootx(), widget.winfo_rooty()
        bottom = through.winfo_rooty() + through.winfo_height() if through is not None else y + widget.winfo_height()
        ImageGrab.grab(bbox=(x, y, x + widget.winfo_width(), bottom)).save(IMAGES / (name + '.png'))

    def panel(text):
        def search(widget):
            for child in widget.winfo_children():
                try:
                    if child.cget('text') == text:
                        return child
                except tk.TclError:
                    pass
                found = search(child)
                if found is not None:
                    return found
            return None

        widget = search(root)
        if widget is None:
            raise ValueError(text)
        return widget

    try:
        app.process_event('connected', 'Conectado · demostración ficticia')
        app.process_event('catalog', CATALOG)
        app.select_tribunals(('113',))
        app.modality.set('Ambulatorio')
        app.tab.set('Cumplimiento')
        app.refresh()
        settle()
        shot('01_ventana_principal', root)
        shot('04_conexion_descargador', panel('1. Tu Chrome habitual'))
        picker = TribunalPicker(root, CATALOG['tribunales'], ('113',), app.select_tribunals)
        picker.geometry('600x520+100+100')
        shot('05_tribunales', picker)
        picker.destroy()
        app.tab.set('Espera')
        app.use_dates.set(False)
        app.refresh()
        settle()
        shot('06_consulta_espera', panel('2. Elige qué descargar'))
        app.tab.set('Informes')
        app.report.set('Por vencer')
        app.refresh()
        settle()
        shot('07_consulta_informes', panel('2. Elige qué descargar'))
        app.tab.set('Cumplimiento')
        app.refresh()
        app.process_event('folder', str(first.folder))
        app.process_event(
            'query', {'index': 1, 'total': 1, 'label': 'Cumplimiento · Ambulatorio · Laja · ejemplo ficticio'}
        )
        app.process_event('progress', {'pagina': 1, 'total': 1, 'registros': 2, 'archivo': 'pagina.xls'})
        app.process_event('done', {'folder': str(first.folder), 'queries': 1, 'files': 1, 'records': 2})
        body = app.progress.master.master
        status_widget = next(w for w in body.winfo_children() if str(w.grid_info().get('row')) == '6')
        shot('08_descarga_finalizada', app.progress.master, status_widget)
        result = DiaryRunner(DiaryFixture(), CATALOG, app.process_event, threading.Event(), pause=0).run(
            Batch(('113',), ('2',), ('Cumplimiento',)), demo / '05_Lectura_nueva'
        )
        app.process_event('diary_done', result)
        shot('09_lectura_finalizada', app.progress.master, status_widget)

        # Todas las acciones son las de la interfaz real; solo se eligen rutas ficticias.
        app.last_folder = None
        win = ResultsWindow(app)
        win.geometry('1450x860+15+15')
        settle(win)
        # Ajustes de presentación, equivalentes a ensanchar la ventana/columnas.
        # Capturar cada región directamente evita reducir letras en el manual.
        for col, width in [('fecha', 355), ('tipo', 170), ('estado', 170), ('consultas', 370)]:
            win.history.column(col, width=width, stretch=False)
        for col, width in [('consulta', 360), ('estado', 155), ('filas', 120), ('aptitud', 430)]:
            win.detail.column(col, width=width, stretch=False)

        def results_top(name):
            settle(win)
            x, y = win.winfo_rootx() + 12, win.winfo_rooty() + 12
            bottom = win.detail.winfo_rooty() + 115
            ImageGrab.grab(bbox=(x, y, x + 1100, bottom)).save(IMAGES / (name + '.png'))

        def results_bottom(name):
            settle(win)
            button = next(
                w
                for row in win.winfo_children()
                for w in row.winfo_children()
                if isinstance(w, tk.ttk.Button) and w.cget('text') == 'Consolidar XLSX'
            )
            x, y = win.winfo_rootx() + 12, button.master.winfo_rooty()
            ImageGrab.grab(bbox=(x, y, x + 1100, win.winfo_rooty() + win.winfo_height() - 5)).save(
                IMAGES / (name + '.png')
            )

        with patch('ventana_resultados.filedialog.askopenfilenames', return_value=(str(previous),)):
            win.add_diaries()
        settle(win)
        results_top('10_bitacoras_recuperadas')
        excel = exports / 'Bitacoras_RUS.xlsx'
        with patch('ventana_resultados.filedialog.asksaveasfilename', return_value=str(excel)):
            win.diary_excel()
        settle(win)
        results_bottom('11_excel_creado')
        # No incluir la lectura de prueba como fuente del paquete de los 800 ingresos.
        app.prefs.data['historial_bitacoras'] = [str(previous.resolve())]
        for lot in (first, second):
            with patch('ventana_resultados.filedialog.askdirectory', return_value=str(lot.folder)):
                win.choose()
            settle(win)
        results_bottom('12_hojas_reunidas')
        target = exports / 'Captura_CSMP.zip'
        with patch('ventana_resultados.filedialog.asksaveasfilename', return_value=str(target)):
            win.capture_package()
        settle(win)
        results_bottom('13_paquete_creado')
        # Comprobar los productos mostrados, sin contactar ningún sistema remoto.
        from nurus.capture import read_capture
        from openpyxl import load_workbook

        capture = read_capture(target)
        assert len(capture.rows('registros')) == 4
        assert len(capture.rows('bitacoras')) == 2
        book = load_workbook(excel, read_only=True)
        try:
            assert book['Ingresos'].max_row == 801
            sheets = {name: book[name].max_row for name in book.sheetnames}
        finally:
            book.close()
        (demo / 'VERIFICACION_FICTICIA.json').write_text(
            json.dumps(
                {
                    'datos': 'EXCLUSIVAMENTE FICTICIOS',
                    'peticiones_reales_sitfa': 0,
                    'bitacoras_anteriores': 800,
                    'registros_hojas': 4,
                    'bitacoras_correspondientes': 2,
                    'no_correspondientes': 798,
                    'hojas_excel': sheets,
                },
                indent=2,
                ensure_ascii=False,
            ),
            encoding='utf-8',
        )
        win.close()
    finally:
        dispose_interface(app, root)


async def browser(browser_path=None):
    from playwright.async_api import async_playwright, expect
    from puente import Bridge
    from formulario_ficticio import master

    bridge = Bridge()
    try:
        with tempfile.TemporaryDirectory(prefix='manual-chrome-ficticio-') as profile:
            async with async_playwright() as pw:
                extension = ROOT / 'descargador/extension'
                context = await pw.chromium.launch_persistent_context(
                    profile,
                    headless=True,
                    executable_path=str(browser_path or pw.chromium.executable_path),
                    viewport={'width': 1150, 'height': 850},
                    locale='es-CL',
                    args=[
                        '--lang=es',
                        '--disable-extensions-except=' + str(extension),
                        '--load-extension=' + str(extension),
                    ],
                )
                try:

                    async def route(r):
                        if r.request.url.endswith('/fixture'):
                            await r.fulfill(
                                body='<html><iframe src="/fixture-form"></iframe></html>', content_type='text/html'
                            )
                        elif r.request.url.endswith('/fixture-form'):
                            await r.fulfill(body=master(), content_type='text/html')
                        else:
                            await r.abort()
                            raise AssertionError('No se permiten peticiones institucionales: ' + r.request.url)

                    await context.route('https://familia.pjud.cl/**', route)
                    worker = (
                        context.service_workers[0]
                        if context.service_workers
                        else await context.wait_for_event('serviceworker')
                    )
                    extension_id = worker.url.split('/')[2]
                    manager = await context.new_page()
                    await manager.goto('chrome://extensions/')
                    # Chrome usa shadow DOM; Playwright atraviesa las raíces abiertas.
                    toggle = manager.locator('#devMode')
                    if not await toggle.get_attribute('aria-pressed') == 'true':
                        await toggle.click()
                    await expect(manager.locator('#loadUnpacked')).to_be_visible()
                    await expect(manager.locator('extensions-item')).to_have_count(1)
                    await manager.wait_for_timeout(450)  # Terminar la animación de la barra de desarrollador.
                    await manager.screenshot(
                        path=str(IMAGES / '02_extension_chrome.png'),
                        clip={'x': 0, 'y': 0, 'width': 1150, 'height': 450},
                    )
                    page = await context.new_page()
                    await page.goto('https://familia.pjud.cl/fixture')
                    ids = await worker.evaluate(
                        "async()=> (await chrome.tabs.query({url:'https://familia.pjud.cl/*'})).map(t=>t.id)"
                    )
                    connector = await context.new_page()
                    await connector.goto(f'chrome-extension://{extension_id}/conexion.html?tab={ids[0]}')
                    await connector.locator('#code').fill(bridge.connection_code)
                    await connector.locator('#connect').click()
                    await expect(connector.locator('#status')).to_contain_text('Conectado.')
                    # Sustituir el valor solo en la vista de demostración antes de capturar:
                    # el código efímero del puente jamás debe quedar en una imagen publicada.
                    await connector.locator('#code').evaluate("el => el.value = 'CÓDIGO DE EJEMPLO · NO VÁLIDO'")
                    await connector.locator('section').first.screenshot(path=str(IMAGES / '03_conexion_extension.png'))
                finally:
                    await context.close()
    finally:
        bridge.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--demo-dir', type=Path, required=True)
    parser.add_argument('--browser', type=Path, help='Chrome/Chromium compatible con extensiones')
    args = parser.parse_args()
    from contrato_integral import VERSION

    if VERSION != '2.7.4':
        parser.error(
            'Las imágenes archivadas corresponden a 2.7.4. Para reproducirlas usa el checkout f8a1fcc; no reemplaces el manual antiguo con capturas de otra versión.'
        )
    demo = args.demo_dir.resolve()
    if demo.exists():
        parser.error('El directorio de práctica debe ser nuevo: nunca se sobrescriben lotes existentes.')
    demo.mkdir(parents=True)
    IMAGES.mkdir(parents=True, exist_ok=True)
    desktop(demo)
    asyncio.run(browser(args.browser))
    print('Capturas y productos ficticios verificados. No se consultó SITFA.')


if __name__ == '__main__':
    main()
