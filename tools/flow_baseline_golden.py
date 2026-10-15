"""P0.05: genera productos ficticios del motor vigente, sin GUI, red ni Outlook.

python tools/flow_baseline_golden.py --output /carpeta/nueva
python tools/flow_baseline_golden.py --output /otra/carpeta --expected docs/evidencia_p0/golden/semantic.json
Compara contenidos, tipos de celda, estilos, filtros, textos y decisiones.
Los ZIP Office se estabilizan solo en estas fixtures; nunca se editan originales reales.
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from hashlib import sha256
import json
from pathlib import Path
import re
import sys
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED, ZipInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


class FixtureDate(date):
    @classmethod
    def today(cls):
        return cls(2026, 10, 8)


def canonical_office(path):
    with ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    if "docProps/core.xml" in members:
        members["docProps/core.xml"] = re.sub(
            rb"(<dcterms:(?:created|modified)[^>]*>).*?(</dcterms:(?:created|modified)>)",
            rb"\g<1>2026-10-08T00:00:00Z\g<2>",
            members["docProps/core.xml"],
        )
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
        for name, data in sorted(members.items()):
            item = ZipInfo(name, date_time=(2026, 10, 8, 0, 0, 0))
            item.compress_type = ZIP_DEFLATED
            target.writestr(item, data)


def workbook_snapshot(path, folder):
    from openpyxl import load_workbook

    book = load_workbook(path, data_only=False)
    try:
        result = []
        for sheet in book:
            cells = []
            for row in sheet:
                values = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    elif isinstance(value, str):
                        value = value.replace(str(folder), "<OUTPUT>").replace(str(ROOT), "<REPO>")
                    values.append(
                        {
                            "value": value,
                            "type": cell.data_type,
                            "format": cell.number_format,
                            "style": list(cell._style) if cell.has_style else None,
                        }
                    )
                cells.append(values)
            result.append(
                {
                    "sheet": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "cells": cells,
                    "filter": sheet.auto_filter.ref,
                    "freeze": str(sheet.freeze_panes or ""),
                    "merged": sorted(str(v) for v in sheet.merged_cells.ranges),
                    "widths": {k: v.width for k, v in sheet.column_dimensions.items()},
                }
            )
        return result
    finally:
        book.close()


def source_book(folder, mode):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = mode
    sheet.append(
        [
            "RIT",
            "TRIBUNAL",
            "NOMBRE",
            "RUT",
            "DERIVACION",
            "T ESPERA",
            "DIAS DE CUMPLIMIENTO",
            "DIAS PARA EGRESAR",
            "FEC.EGRESO PROYECTADO",
            "FECHA VENCIMIENTO",
            "FEC. RESOLUCION",
            "DURACION",
        ]
    )
    for number, court, program, days in [
        (1, "LAJA", "AFT FICTICIO", 30),
        (2, "MULCHEN", "DCE FICTICIO", 30),
        (3, "TOME", "AFT FICTICIO", 60),
    ]:
        sheet.append(
            [
                f"X-{number}-2026",
                court,
                f"Persona ficticia {number}",
                f"00{number}11111-1",
                program,
                days,
                50,
                90,
                date(2027, 1, 6),
                date(2026, 10, 20),
                date(2026, 8, 1),
                "6 meses",
            ]
        )
    sheet.column_dimensions["A"].width = 25
    sheet.auto_filter.ref = "A1:L4"
    extra = book.create_sheet("Original intacto")
    extra["A1"] = "=1+2"
    extra["B1"] = "Solo datos ficticios"
    path = folder / ("Origen_" + mode + ".xlsx")
    book.save(path)
    book.close()
    canonical_office(path)
    return path


CURADORES_FICTICIOS = (
    "(00411111-1)ABOGADO FICTICIO A - Institución: Programa Ficticio Uno#"
    "(00511111-1)ABOGADO FICTICIO B - Institución: Corporación Ficticia Dos"
)


def rus_rules_snapshot():
    """Reglas RUS de la decisión u (E10, T5, J2, R6): corren sin historia y por el camino Excel.
    Los libros van a una carpeta temporal: no agregan productos al golden."""
    from tempfile import TemporaryDirectory
    from openpyxl import Workbook
    from nurus.personal.config import defaults
    from nurus.personal.work import Work

    books = {
        "ESPERA": (
            ["RIT", "TRIBUNAL", "NOMBRE", "RUT", "DERIVACION", "T ESPERA", "CURADOR"],
            [
                ["X-11-2026", "LAJA", "Persona ficticia 11", "00611111-1", "RVA FICTICIO", 120, CURADORES_FICTICIOS],
                ["X-12-2026", "LAJA", "Persona ficticia 12", "00711111-1", "AFT FICTICIO", 10, ""],
                ["X-12-2026", "LAJA", "Persona ficticia 12", "00711111-1", "AFT FICTICIO", 10, ""],
                ["X-12-2026", "LAJA", "Persona ficticia 13", "00811111-1", "AFT FICTICIO", 10, ""],
            ],
        ),
        "INFORMES": (
            ["RIT", "TRIBUNAL", "NOMBRE", "RUT", "DERIVACION", "FECHA VENCIMIENTO", "CURADOR"],
            [
                ["X-21-2026", "LAJA", "Persona ficticia 21", "00911111-1", "AFT FICTICIO", date(2026, 8, 20), ""],
                ["X-21-2026", "LAJA", "Persona ficticia 21", "00911111-1", "AFT FICTICIO", date(2026, 9, 20), ""],
                [
                    "X-22-2026",
                    "LAJA",
                    "Persona ficticia 22",
                    "00311111-1",
                    "AFT FICTICIO",
                    date(2026, 9, 20),
                    CURADORES_FICTICIOS,
                ],
            ],
        ),
    }
    result = {}
    with TemporaryDirectory() as temp:
        for mode, (headers, rows) in books.items():
            book = Workbook()
            sheet = book.active
            sheet.title = mode
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            path = Path(temp) / f"reglas_rus_{mode}.xlsx"
            book.save(path)
            work = Work(defaults()).analyze(path, mode, as_of=date(2026, 10, 8))
            result[mode] = [
                {
                    "observation": row.observation,
                    "rules": row.rules,
                    "actions": row.actions,
                    "avisos": [list(a) for a in row.avisos],
                }
                for row in work.rows
            ]
    return result


def generate(folder):
    from nurus.personal.config import defaults, BASE
    from nurus.personal.work import Work
    from nurus.personal.outputs import prepare_drafts
    from nurus.personal.resolutions import prepare_projects, generate_projects
    from docx import Document

    semantic = {
        "fixture": "P0-ficticio-v1",
        "as_of": "2026-10-08",
        "modes": {},
        "config": defaults(),
        "network_calls": 0,
        "outlook_executed": False,
    }
    generated = []
    for mode, kind in [("ESPERA", "programa_espera"), ("CUMPLIMIENTO", "medidas"), ("INFORMES", "programa_por_vencer")]:
        source = source_book(folder, mode)
        before = sha256(source.read_bytes()).hexdigest()
        work = Work(defaults()).analyze(source, mode, as_of=date(2026, 10, 8))
        exported = Path(
            work.export(folder / ("Observaciones_" + mode + ".xlsx"), backend="portable", reduced_fidelity=True)
        )
        with patch("nurus.personal.outputs.date", FixtureDate):
            drafts = prepare_drafts(
                work, kind, directory=folder / ("Nominas_" + mode), modalities="Ambulatoria", period="octubre 2026"
            )
            selections = [(row.id, "PC_IE") for row in work.rows if "PC_IE" in row.actions]
            projects, errors = prepare_projects(work, selections, BASE / "plantillas_word")
            words = []
            if projects:
                target = folder / ("Proyectos_" + mode + ".docx")
                generate_projects(work, projects, target)
                doc = Document(target)
                words = [p.text for p in doc.paragraphs]
                generated.append(target)
        semantic["modes"][mode] = {
            "source_sha256": before,
            "rows": [
                {
                    "values": row.values,
                    "observation": row.observation,
                    "rules": row.rules,
                    "actions": row.actions,
                    "warnings": row.warnings,
                    "excluded": row.excluded,
                }
                for row in work.rows
            ],
            "exported_workbook": workbook_snapshot(exported, folder),
            "drafts": [
                {
                    "subject": d.subject,
                    "body": d.body,
                    "to": d.to,
                    "cc": d.cc,
                    "program": d.program,
                    "record_ids": d.record_ids,
                    "attachments": [Path(a).name for a in d.attachments],
                }
                for d in drafts
            ],
            "rosters": [workbook_snapshot(Path(a), folder) for d in drafts for a in d.attachments],
            "project_texts": words,
            "project_errors": errors,
        }
        assert len(work.rows) == 3, "La cardinalidad del baseline cambió"
        assert sha256(source.read_bytes()).hexdigest() == before, "Se modificó el original ficticio"
        assert all(not d.to or d.to.endswith(".test") for d in drafts), "Destinatario no ficticio"
        assert "PC_IE" not in work.rows[1].actions if mode == "ESPERA" else True, "DCE Espera no puede producir PC_IE"
        generated.extend([source, exported, *[Path(a) for d in drafts for a in d.attachments]])
    semantic["reglas_rus"] = rus_rules_snapshot()
    # Reutiliza únicamente la fixture sintética del baseline; no captura HTML remoto.
    sys.path.insert(0, str(ROOT / "tests"))
    from test_bitacora_lote import manifest
    from nurus.personal.bitacora_lote import export_lote

    diary_folder = folder / "Bitacora_ficticia"
    diary_folder.mkdir()
    diary_source = manifest(diary_folder)
    diary_excel = folder / "Bitacoras_RUS_baseline.xlsx"
    export_lote(diary_source, diary_excel, date(2026, 6, 8), date(2026, 10, 8), today=date(2026, 10, 8))
    semantic["diary_workbook"] = workbook_snapshot(diary_excel, folder)
    generated.append(diary_excel)
    for path in generated:
        canonical_office(path)
    return json.loads(json.dumps(semantic, ensure_ascii=False, default=str))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--expected", type=Path)
    args = parser.parse_args()
    folder = args.output.resolve()
    folder.mkdir(parents=True, exist_ok=False)
    semantic = generate(folder)
    (folder / "semantic.json").write_text(json.dumps(semantic, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    hashes = {
        p.relative_to(folder).as_posix(): sha256(p.read_bytes()).hexdigest()
        for p in sorted(folder.rglob("*"))
        if p.is_file()
    }
    (folder / "hashes.json").write_text(json.dumps(hashes, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.expected:
        expected = json.loads(args.expected.read_text(encoding="utf-8"))
        if semantic != expected:
            changed = [key for key in semantic if semantic.get(key) != expected.get(key)]
            raise SystemExit("PARIDAD FALLIDA: " + ", ".join(changed))
        print("PARIDAD PASÓ: contenidos, observaciones, reglas, productos y estilos coinciden")
    print(f"Productos ficticios: {len(hashes)}; sin consultas institucionales ni guardado Outlook")


if __name__ == "__main__":
    main()
