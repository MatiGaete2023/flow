"""Regresiones de la clasificación y escritura (base v2.4), por el camino que usa el producto.

Ejecutar desde la raíz del paquete:  python -m pytest tests -q
Todos los datos son ficticios. No usar archivos reales del Poder Judicial.

Cada bloque referencia el hallazgo de auditoría que protege (F-01, F-02, H-02, H-03, M-01,
M-05, A-01…) para que el motivo de cada prueba sea rastreable. Desde 2.3 no existe la línea
de comandos propia de generate_meta4: las pruebas pasan por meta4_bitacoras.fuentes/generar.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE_DIR / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import generate_meta4 as g  # noqa: E402
import meta4_bitacoras as app  # noqa: E402
import biff_escritor  # noqa: E402

CORTE = dt.date(2026, 10, 16)

ESP_HDR = [
    "RIT", "TRIBUNAL", "RUT", "NOMBRE", "NACIONALIDAD / PAIS", "TIPO", "SEXO",
    "DERIVACIÓN", "CURADOR", "T ESPERA", "FEC. RESOLUCIÓN", "REVISIÓN", "DURACIÓN",
    "80 BIS", "FEC. NACIMIENTO", "EDAD", "FEC. ACT. F. INDIVIDUAL",
    "FEC. ACT. F. AMBULATORIA", "FEC. ACT. F. FAE/FAS", "FEC. OÍDO", "PROXS. AUDS.",
]

CUM_HDR = [
    "RIT", "TRIBUNAL", "RUT", "NOMBRE", "NACIONALIDAD / PAÍS", "TIPO", "SEXO",
    "DERIVACIÓN", "CURADOR", "FEC. RESOLUCIÓN", "FEC. INGRESO EFECTIVO",
    "FEC. EGRESO PROYECTADO", "DÍAS DE CUMPLIMIENTO", "DÍAS PARA EGRESAR", "80 BIS",
    "FEC. NACIMIENTO", "EDAD", "FEC. ACT. F. INDIVIDUAL", "FEC. ACT. F. AMBULATORIA",
    "FEC. ACT. F. FAE/FAS", "FEC. OÍDO",
]


def make_excel(
    path: Path, tribunal: str, estado: str, filas: list[dict], emision: str = "01/07/2026",
    when: dt.datetime | None = None,
) -> Path:
    headers = ESP_HDR if estado == "Espera" else CUM_HDR
    titulo = (
        "Informe de Ingresos en Lista de Espera"
        if estado == "Espera"
        else "Informe de Ingresos en Lista de Cumplimiento"
    )
    wb = Workbook()
    ws = wb.active
    ws["A1"] = f"Juzgado de Familia de {tribunal}"
    ws["A2"] = titulo
    ws["A3"] = f"Fecha Emisión   : {emision}"
    for col, header in enumerate(headers, 1):
        ws.cell(5, col, header)
    for offset, fila in enumerate(filas):
        base = {
            "RIT": fila["rit"], "TRIBUNAL": fila.get("tribunal", tribunal), "RUT": fila.get("rut", "11111111-1"),
            "NOMBRE": "FICTICIO NNA", "TIPO": "MP", "SEXO": "F",
            "DERIVACIÓN": fila["der"], "FEC. RESOLUCIÓN": dt.date(2026, 6, 15),
        }
        if estado == "Espera":
            base["T ESPERA"] = fila.get("tesp", 13)
        else:
            base["FEC. INGRESO EFECTIVO"] = dt.date(2026, 6, 20)
        for col, header in enumerate(headers, 1):
            ws.cell(6 + offset, col, base.get(header))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    if when:
        os.utime(path, (when.timestamp(), when.timestamp()))
    return path


def make_blob(path: Path, when: dt.datetime | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"%PDF-ficticio")
    if when:
        os.utime(path, (when.timestamp(), when.timestamp()))
    return path


def generar(folder: Path, salida: Path) -> dict:
    """Producto completo sobre una carpeta de reportes independientes, sin bitácoras leídas."""
    (folder / "bitacoras_meta4.json").write_text(json.dumps(dict(
        tipo="META4_BITACORAS_MES", desde="2026-09-16", hasta="2026-10-16", estado="COMPLETA", registros=[])))
    return app.generar(folder, salida, CORTE)


def matrix_cell(path: Path, tribunal: str, column: int):
    ws = load_workbook(path)["MATRIZ_COBERTURA"]
    return next(ws.cell(r, column).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == tribunal)


# --------------------------------------------------------------------------
# Clasificación de tribunal por ruta (F-03, F-04)
# --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "ruta,esperado",
    [
        ("informes de la semana/espera fae x.pdf", None),
        ("INFORMES DE LA SEMANA/ESPERA FAE.pdf", None),
        ("ACTA DE LA REUNION.pdf", None),
        ("LA LISTA DE ESPERA.pdf", None),
        ("ESPERA FAE LA.pdf", "LOS ANGELES"),
        ("LA FAE CUMPLIMIENTO JULIO.xlsx", "LOS ANGELES"),
        ("CUMPLIMIENTO RES LA 2026.pdf", "LOS ANGELES"),
        ("espera fae la.pdf", "LOS ANGELES"),
        ("L.A. ESPERA FAE.pdf", "LOS ANGELES"),
        ("PROVINCIA CONCEPCION/cumplimiento residencia tome.pdf", "TOMÉ"),
        ("PROVINCIA DE ARAUCO - LEBU/espera fae.pdf", None),
        ("TOMEMOS NOTA.pdf", None),
        ("espera_fae_tome2026.xlsx", "TOMÉ"),
        ("LOS ANGELES/espera fae.pdf", "LOS ANGELES"),
        ("THNO/cumplimiento fae.pdf", "TALCAHUANO"),
        ("CCP/espera residencia.pdf", "CONCEPCIÓN"),
    ],
)
def test_tribunal_from_path(ruta, esperado):
    assert g.tribunal_from_path(ruta)[0] == esperado


def test_tribunal_name_celda():
    assert g.tribunal_name("Juzgado de Familia de Los Angeles") == "LOS ANGELES"
    assert g.tribunal_name("1° Juzgado de Familia de Concepción") == "CONCEPCIÓN"
    assert g.tribunal_name("Tribunal desconocido") is None


# --------------------------------------------------------------------------
# Estado y modalidad (F-06, F-07)
# --------------------------------------------------------------------------

def test_estado_ambiguo_no_clasifica():
    assert g.status_from_path("ESPERA Y CUMPLIMIENTO/x.pdf") is None
    assert g.status_from_path("cumplimiento/x.pdf") == "Cumplimiento"
    assert g.status_from_path("espera/x.pdf") == "Espera"


def test_modalidad_plural_y_ambigua():
    assert g.category_from_path("RESIDENCIAS/espera yumbel.pdf") == "Residencia"
    assert g.category_from_path("FAE/espera.pdf") == "FAE"
    assert g.category_from_path("FAE Y RESIDENCIA/espera.pdf") is None


# --------------------------------------------------------------------------
# F-01: sin duplicación de filas, huérfanas fuera de alcance
# --------------------------------------------------------------------------

def test_f01_sin_duplicacion_de_filas(tmp_path):
    make_excel(
        tmp_path / "FAE" / "lista espera coronel.xlsx", "CORONEL", "Espera",
        [{"rit": f"A-{i}-2026", "der": "DAM DIAGNOSTICO"} for i in range(1, 5)],
    )
    files, _, selected, _ = app.fuentes(tmp_path)
    assert g.consolidated_rows(selected)["Espera"] == []
    assert sum(len(f.orphan_rows) for f in files) == 4


def test_f01_mixto_conserva_solo_alcance(tmp_path):
    make_excel(
        tmp_path / "cumplimiento tome.xlsx", "TOMÉ", "Cumplimiento",
        [
            {"rit": "C-1-2026", "der": "FAE PRO"},
            {"rit": "C-2-2026", "der": "RESIDENCIA RLP"},
            {"rit": "C-3-2026", "der": "OPD SEGUIMIENTO"},
        ],
    )
    files, _, selected, _ = app.fuentes(tmp_path)
    rows = g.consolidated_rows(selected)["Cumplimiento"]
    assert {r.values["RIT"] for r in rows} == {"C-1-2026", "C-2-2026"}
    assert sum(len(f.orphan_rows) for f in files) == 1


# --------------------------------------------------------------------------
# F-02: un PDF nunca desplaza a un Excel con registros
# --------------------------------------------------------------------------

def test_f02_excel_vence_a_pdf_mas_reciente(tmp_path):
    make_excel(
        tmp_path / "espera fae lebu.xlsx", "LEBU", "Espera",
        [{"rit": f"B-{i}-2026", "der": "FAE PROGRAMA"} for i in range(1, 4)],
        emision="01/07/2026",
        when=dt.datetime(2026, 7, 1, 9, 0),  # explícitamente ANTES que el PDF
    )
    make_blob(tmp_path / "espera fae lebu sin registros.pdf", dt.datetime(2026, 7, 30, 12, 0))
    _, _, selected, _ = app.fuentes(tmp_path)
    chosen = selected[("LEBU", "FAE", "Espera")]
    assert len(selected) == 1 and chosen.kind == "EXCEL" and len(chosen.rows) == 3


def test_f02_error_si_seleccionado_pierde_registros(tmp_path):
    make_excel(
        tmp_path / "viejo" / "espera fae arauco.xlsx", "ARAUCO", "Espera",
        [{"rit": f"V-{i}-2026", "der": "FAE X"} for i in range(1, 4)], emision="01/06/2026",
    )
    make_excel(tmp_path / "nuevo" / "espera fae arauco.xlsx", "ARAUCO", "Espera", [], emision="01/07/2026")
    _, issues, _, _ = app.fuentes(tmp_path)
    assert any(i.level == "ERROR" and "no aporta registros" in i.message for i in issues)


# --------------------------------------------------------------------------
# Continuidad, cobertura e idempotencia
# --------------------------------------------------------------------------

def test_faltantes_no_detienen_generacion(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "cumplimiento residencia laja.xlsx", "LAJA", "Cumplimiento",
        [{"rit": "G-1-2026", "der": "RESIDENCIA RLP"}],
    )
    salida = tmp_path / "out"
    report = generar(folder, salida)
    assert report["registros"] == 1
    assert (salida / "META 4 - actualizado 2026-10.xlsx").is_file()
    assert load_workbook(salida / "Registros.xlsx")["MATRIZ_COBERTURA"]["A2"].value.startswith(
        "Cobertura: 1/68 · Faltantes: 67")


def test_reejecucion_no_lee_sus_salidas(tmp_path):
    folder = tmp_path / "in"
    make_excel(folder / "espera fae yumbel.xlsx", "YUMBEL", "Espera", [{"rit": "Y-1-2026", "der": "FAE UNO"}])
    primera = generar(folder, folder / "Resultado META4 1")
    # Las salidas copiadas junto a las fuentes tampoco se leen como reportes.
    for name in ("Registros.xlsx", "META 4 - actualizado 2026-10.xlsx", "Registros consolidados.xlsx",
                 "Bitacoras_del_mes.xlsx"):
        (folder / name).write_bytes((folder / "Resultado META4 1" / name).read_bytes())
    segunda = generar(folder, folder / "Resultado META4 2")
    assert primera["registros"] == segunda["registros"] == 1


def test_pdf_incompleto_no_cubre_combinacion(tmp_path):
    make_blob(tmp_path / "espera lebu.pdf")  # falta modalidad
    files, _, selected, _ = app.fuentes(tmp_path)
    assert selected == {} and files[0].result == "NO CLASIFICADO"


def test_excel_ilegible_detiene_sin_planilla_parcial(tmp_path):
    # El generador v2.4 continuaba; el producto se detiene para no emitir un META 4 incompleto.
    (tmp_path / "roto.xlsx").write_bytes(b"esto no es un excel")
    make_excel(tmp_path / "espera fae tome.xlsx", "TOMÉ", "Espera", [{"rit": "T-1-2026", "der": "FAE UNO"}])
    with pytest.raises(ValueError, match="planilla parcial.*roto.xlsx"):
        app.fuentes(tmp_path)


def test_formato_t_espera_es_numero(tmp_path):
    folder = tmp_path / "in"
    make_excel(folder / "espera fae laja.xlsx", "LAJA", "Espera", [{"rit": "L-1-2026", "der": "FAE UNO", "tesp": 13}])
    generar(folder, tmp_path / "out")
    ws = load_workbook(tmp_path / "out" / "META 4 - actualizado 2026-10.xlsx")["FAE Espera"]
    celdas = [c for row in ws.iter_rows(min_col=6, max_col=6) for c in row if isinstance(c.value, (int, float))]
    assert celdas, "no se encontró la celda T ESPERA"
    assert all("YY" not in (c.number_format or "").upper() for c in celdas)


# --------------------------------------------------------------------------
# H-03: un Excel con más de un tribunal no debe contradecir la matriz
# --------------------------------------------------------------------------

def test_h03_excel_multitribunal_coherencia_matriz_meta(tmp_path):
    folder = tmp_path / "in"
    make_excel(
        folder / "espera fae tome.xlsx", "TOMÉ", "Espera",
        [{"rit": "X-1-2026", "der": "FAE UNO"},
         {"rit": "X-2-2026", "der": "FAE UNO", "tribunal": "Juzgado de Familia de Lebu"}],
    )
    report = generar(folder, tmp_path / "out")
    assert report["totales_meta4"]["FAE Espera"] == 2 and report["excluidos_meta4"] == 0
    for tribunal in ("TOMÉ", "LEBU"):
        assert matrix_cell(tmp_path / "out" / "Registros.xlsx", tribunal, 2) == "EXCEL · 1 registro(s)"


def test_h03_archivo_uniforme_sin_cambio_de_comportamiento(tmp_path):
    make_excel(
        tmp_path / "espera fae laja.xlsx", "LAJA", "Espera",
        [{"rit": f"U-{i}-2026", "der": "FAE UNO"} for i in range(1, 4)],
    )
    files, _, selected, _ = app.fuentes(tmp_path)
    assert len(files[0].candidates) == 1 and len(selected[("LAJA", "FAE", "Espera")].rows) == 3


# --------------------------------------------------------------------------
# A-01 (auditoría 2.3): filas de otro tribunal no desplazan al archivo propio
# --------------------------------------------------------------------------

def test_filas_de_otro_tribunal_se_suman_sin_desplazar_ni_duplicar(tmp_path):
    make_excel(tmp_path / "espera fae talcahuano.xlsx", "Talcahuano", "Espera",
               [{"rit": f"T-{i}-2026", "der": "FAE A"} for i in range(5)], emision="01/07/2026")
    make_excel(tmp_path / "espera fae concepcion.xlsx", "Concepción", "Espera",
               [{"rit": f"C-{i}-2026", "der": "FAE A"} for i in range(3)]
               + [{"rit": "T-9-2026", "der": "FAE A", "tribunal": "Juzgado de Familia de Talcahuano"},
                  {"rit": "T-0-2026", "der": "FAE A", "tribunal": "Juzgado de Familia de Talcahuano"}],
               emision="02/07/2026")
    _, issues, selected, groups = app.fuentes(tmp_path)
    talcahuano = selected[("TALCAHUANO", "FAE", "Espera")]
    assert g.matrix_cell_text(("TALCAHUANO", "FAE", "Espera"), selected, groups) == "EXCEL · 6 registro(s)"
    assert talcahuano.file_result.relative_path == "espera fae talcahuano.xlsx"
    assert sorted(r.values["RIT"] for r in talcahuano.rows) == [f"T-{i}-2026" for i in range(5)] + ["T-9-2026"]
    assert len(selected[("CONCEPCIÓN", "FAE", "Espera")].rows) == 3
    assert any("se incorporaron 1; 1 ya estaban" in i.message for i in issues)


# --------------------------------------------------------------------------
# A-02 / A-03 (auditoría 2.3): exclusión visible y caracteres de control
# --------------------------------------------------------------------------

def test_fila_con_tribunal_no_reconocido_queda_informada(tmp_path):
    folder = tmp_path / "in"
    make_excel(folder / "espera fae lebu.xlsx", "Lebu", "Espera",
               [{"rit": "L-1-2026", "der": "FAE A", "tribunal": "1° JF"}])
    report = generar(folder, tmp_path / "out")
    assert report["excluidos_meta4"] == 1
    ws = load_workbook(tmp_path / "out" / "Registros.xlsx")["ERRORES_Y_ADVERTENCIAS"]
    mensajes = [r[3] for r in ws.iter_rows(min_row=4, values_only=True)]
    assert any("Registro excluido de META 4" in str(m) and "1° JF" in str(m) for m in mensajes)


def test_caracteres_de_control_no_abortan_la_generacion(tmp_path):
    folder = tmp_path / "in"
    folder.mkdir()
    headers = ["RIT", "TRIBUNAL", "RUT", "NOMBRE", "DERIVACIÓN", "TIPO", "SEXO", "T ESPERA", "FEC. RESOLUCIÓN"]
    rows = [["Juzgado de Familia de Laja"], ["Informe de Ingresos en Lista de Espera"],
            ["Fecha Emisión : 01/07/2026"], [], headers,
            ["X-1-2026", "Laja", "1-9", "NOMBRE\x02FICTICIO", "FAE A", "MP", "F", 3.0, "01/06/2026"]]
    (folder / "espera fae laja.xls").write_bytes(biff_escritor.xls([("Hoja1", rows)]))
    generar(folder, tmp_path / "out")
    ws = load_workbook(tmp_path / "out" / "Registros.xlsx")["ESPERA"]
    assert ws.cell(2, 4).value == "NOMBRE FICTICIO"
    assert g.clean_text("A\x00B\x1fC") == "A B C"


# --------------------------------------------------------------------------
# H-02 opción A: Excel con datos vence a un PDF más reciente, marcado para revisión
# --------------------------------------------------------------------------

def test_h02a_excel_vence_y_queda_marcado_requiere_revision(tmp_path):
    make_excel(
        tmp_path / "espera fae lebu.xlsx", "LEBU", "Espera",
        [{"rit": f"B-{i}-2026", "der": "FAE PROGRAMA"} for i in range(1, 4)],
        emision="01/07/2026",
        when=dt.datetime(2026, 7, 1, 9, 0),
    )
    make_blob(tmp_path / "espera fae lebu sin registros.pdf", dt.datetime(2026, 7, 30, 12, 0))
    _, issues, selected, groups = app.fuentes(tmp_path)
    revision = [i for i in issues if i.level == "REQUIERE_REVISION"]
    assert len(revision) == 1
    assert "espera fae lebu.xlsx" in revision[0].message and "espera fae lebu sin registros.pdf" in revision[0].message
    assert "REQUIERE REVISIÓN" in g.matrix_cell_text(("LEBU", "FAE", "Espera"), selected, groups)


# --------------------------------------------------------------------------
# M-05 / M-01: formatos y rutas
# --------------------------------------------------------------------------

def test_xlsm_se_reconoce_y_procesa(tmp_path):
    make_excel(tmp_path / "espera fae laja.xlsm", "LAJA", "Espera", [{"rit": "M-1-2026", "der": "FAE UNO"}])
    _, _, selected, _ = app.fuentes(tmp_path)
    assert len(selected) == 1 and len(selected[("LAJA", "FAE", "Espera")].rows) == 1


def test_xls_binario_biff_se_reconoce_y_procesa(tmp_path):
    rows = [["Juzgado de Familia de Laja"], ["Informe de Ingresos en Lista de Espera"],
            ["Fecha Emisión   : 01/07/2026"], [], ESP_HDR,
            ["W-1-2026", "LAJA", "11111111-1", "PERSONA FICTICIA", None, "MP", "F", "FAE UNO", None, 13,
             ("date", 46188)]]
    (tmp_path / "espera fae laja.xls").write_bytes(biff_escritor.xls([("Hoja1", rows)]))
    _, _, selected, _ = app.fuentes(tmp_path)
    row = selected[("LAJA", "FAE", "Espera")].rows[0]
    assert row.values["T ESPERA"] == 13 and row.values["FEC. RESOLUCIÓN"] == dt.date(2026, 6, 15)


def test_nombre_con_codificacion_dañada_usa_contenido_del_excel(tmp_path):
    # Nombre de archivo corrompido (carácter de reemplazo Unicode): el tribunal sale del contenido.
    make_excel(tmp_path / "espera fae �ome.xlsx", "TOMÉ", "Espera", [{"rit": "D-1-2026", "der": "FAE UNO"}])
    _, _, selected, _ = app.fuentes(tmp_path)
    assert set(selected) == {("TOMÉ", "FAE", "Espera")}


def test_conflicto_carpeta_archivo_end_to_end(tmp_path):
    # M-01: el nombre del archivo (más específico) gana a la carpeta.
    make_blob(tmp_path / "ESPERA" / "cumplimiento residencia laja.pdf")
    _, _, selected, _ = app.fuentes(tmp_path)
    assert set(selected) == {("LAJA", "Residencia", "Cumplimiento")}
    assert selected[("LAJA", "Residencia", "Cumplimiento")].kind == "PDF"
