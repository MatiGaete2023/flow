"""Casos ficticios: fechas, identidades, páginas y alerta de etapa."""
from datetime import date
import hashlib
import json
import re
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import meta4_bitacoras as app
from test_generate_meta4 import make_excel

CORTE = date(2026, 10, 16)


def entrada(day, stage='Espera'):
    return dict(fecha=day, etapa=stage, texto='Observación ficticia', fecha_respuesta='2026-10-16T15:00:00')


def diary(rit='X-1-2026', ingreso='201', entries=None):
    return dict(tribunal_nombre='Juzgado de Familia de Laja', tribunal_codigo='101',
                causa_id='301', ingreso_id=ingreso, rit=rit, rut='11111111-1',
                nombre='FICTICIO NNA', programa='FAE FICTICIO', estado='LEIDA',
                capturada='2026-10-16T12:00:00', advertencias=[],
                snapshot=dict(tribunal_codigo='101', causa_id='301', ingreso_id=ingreso,
                              rit=rit, persona='FICTICIO NNA', centro='FAE FICTICIO',
                              entries=entries if entries is not None else [entrada('2026-10-10T10:00:00')]))


def save_diaries(root, records):
    (root / 'bitacoras_meta4.json').write_text(json.dumps(dict(tipo='META4_BITACORAS_MES', version=1,
        desde='2026-09-16', hasta='2026-10-16', registros=records)), encoding='utf-8')


def test_month_uses_calendar_and_clamps_day():
    assert app.periodo(date(2026, 3, 31))[0] == date(2026, 2, 28)
    assert app.periodo(date(2026, 1, 10))[0] == date(2025, 12, 10)
    assert app.periodo(date(2024, 3, 31))[0] == date(2024, 2, 29)


def test_last_ignores_reports_old_and_future_uses_center_date():
    events = [entrada('2026-09-15T23:59:59'), entrada('2026-09-16T00:00:00', 'Cumplimiento'),
              entrada('2026-10-09T10:00:00'), entrada('2026-10-15T10:00:00', 'Informes'),
              entrada('2026-10-17T10:00:00', 'Cumplimiento')]
    assert app.ultima(events, CORTE)[:2] == (date(2026, 10, 9), 'Espera')


def test_ties_require_review_instead_of_inventing_order():
    day, stage, control = app.ultima([entrada('2026-10-10'), entrada('2026-10-10', 'Cumplimiento')], CORTE)
    assert day == date(2026, 10, 10) and not stage and control.startswith('REVISAR')


def test_no_monthly_observation_is_blank():
    assert app.ultima([entrada('2026-08-01'), entrada('2026-10-10', 'Informes')], CORTE)[0] is None


def prepare(root, records=None):
    make_excel(root / 'CUMP FAE LAJA.xlsx', 'Laja', 'Cumplimiento', [{'rit': 'X-1-2026', 'der': 'FAE FICTICIO'}])
    save_diaries(root, records or [diary()])


def test_generation_alert_in_meta_registros_and_consolidado(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source, [diary(entries=[entrada('2026-10-09T10:00:00'), entrada('2026-10-15', 'Informes')])])
    target = tmp_path / 'resultado'
    result = app.generar(source, target, CORTE)
    assert result['registros'] == 1 and result['alertas'] == 1
    for filename, sheet in [('Registros.xlsx', 'CUMPLIMIENTO'), ('Registros consolidados.xlsx', 'Cumplimiento')]:
        wb = load_workbook(target / filename)
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        assert ws.cell(2, headers.index('FECHA OBS.') + 1).value.date() == date(2026, 10, 9)
        assert ws.cell(2, headers.index('ALERTA') + 1).value == app.ALERTA
        assert ws.cell(2, 1).fill.fgColor.rgb.endswith('FFC7CE')
    wb = load_workbook(target / 'META 4 - actualizado 2026-10.xlsx')
    ws = wb['FAE Cumplimiento']
    cells = [c for row in ws for c in row if c.value == 'X-1-2026']
    assert len(cells) == 1
    cell = cells[0]
    assert ws.cell(cell.row, 8).value.date() == date(2026, 10, 9)
    assert ws.cell(cell.row, 8).comment.text == app.ALERTA
    assert cell.fill.fgColor.rgb.endswith('FFC7CE')


def test_multiple_ingresos_never_match_by_rit_only(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source, [diary(), diary(ingreso='202')])
    target = tmp_path / 'resultado'
    report = app.generar(source, target, CORTE)
    assert report['alertas'] == 0 and report['revision_bitacora'] == 1
    ws = load_workbook(target / 'Registros consolidados.xlsx')['Cumplimiento']
    headers = [c.value for c in ws[1]]
    assert ws.cell(2, headers.index('FECHA OBS.') + 1).value is None
    assert ws.cell(2, headers.index('CONTROL BITÁCORA') + 1).value == 'IDENTIDAD AMBIGUA'


def test_empty_and_failed_diaries_visible(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    failed = diary()
    failed['estado'] = 'FALLIDA'
    prepare(source, [failed])
    result = app.generar(source, tmp_path / 'resultado', CORTE)
    assert result['revision_bitacora'] == 1


def test_persisted_period_cannot_be_extended(tmp_path):
    save_diaries(tmp_path, [diary()])
    with pytest.raises(ValueError, match='no cubren'):
        app.cargar_bitacoras(tmp_path, date(2026, 10, 17))


def test_native_multipage_and_html_xls(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    pages = []
    for index in (1, 2):
        headers = ['RIT', 'TRIBUNAL', 'RUT', 'NOMBRE', 'DERIVACIÓN', 'TIPO', 'SEXO', 'FEC. RESOLUCIÓN', 'FEC. INGRESO EFECTIVO']
        values = [f'X-{index}-2026', 'Laja', '11111111-1', 'FICTICIO NNA', 'FAE FICTICIO', 'MP', 'F', '01/09/2026', '02/09/2026']
        data = ('<html><table><tr><td>Juzgado de Familia de Laja</td></tr><tr><td>Informe de Ingresos en Lista de Cumplimiento</td></tr><tr><td>Fecha Emisión : 16/10/2026</td></tr><tr>' + ''.join(f'<td>{v}</td>' for v in headers) + '</tr><tr>' + ''.join(f'<td>{v}</td>' for v in values) + '</tr></table></html>').encode()
        filename = f'CUMP FAE LAJA {index}.xls'
        (source / filename).write_bytes(data)
        pages.append(dict(archivo=filename, sha256=hashlib.sha256(data).hexdigest()))
    (source / 'verificacion.json').write_text(json.dumps(dict(estado='VALIDADA', paginas_esperadas=2, paginas_verificadas=2, archivos=pages)))
    (source / 'resumen.json').write_text(json.dumps(dict(estado='VALIDADA', consultas=[dict(verificacion='verificacion.json', seleccion=dict(tab='Cumplimiento', modality='3'))])))
    save_diaries(source, [diary(rit='X-1-2026'), diary(rit='X-2-2026', ingreso='202')])
    result = app.generar(source, tmp_path / 'resultado', CORTE)
    assert result['registros'] == 2 and result['alertas'] == 2


def test_changed_source_rejected(tmp_path):
    (tmp_path / 'p.xls').write_bytes(b'changed')
    (tmp_path / 'v.json').write_text(json.dumps(dict(estado='VALIDADA', paginas_esperadas=1, paginas_verificadas=1, archivos=[dict(archivo='p.xls', sha256='bad')])) )
    (tmp_path / 'resumen.json').write_text(json.dumps(dict(estado='VALIDADA', consultas=[dict(verificacion='v.json', seleccion=dict(tab='Espera', modality='3'))])))
    with pytest.raises(ValueError, match='cambió'):
        app.fuentes(tmp_path)


def test_remote_identity_narrows_matching_ingreso():
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {
        'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': '11111111-1', 'DERIVACIÓN': 'FAE FICTICIO',
        'NOMBRE': 'FICTICIO NNA', '_REMOTE': ('101', '301', '202')}, '', 1)
    records = [dict(diary(ingreso=str(i)), control='OK', entries=[entrada('2026-10-10', s)])
               for i, s in [(201, 'Cumplimiento'), (202, 'Espera')]]
    app.enriquecer({'Cumplimiento': [row]}, records, CORTE)
    assert row.values['ALERTA'] == app.ALERTA


def test_same_rit_different_person_or_program_not_associated():
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {
        'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': '22222222-2', 'DERIVACIÓN': 'FAE FICTICIO',
        'NOMBRE': 'OTRA PERSONA'}, '', 1)
    d = dict(diary(), control='OK', entries=[entrada('2026-10-10')])
    assert app.buscar(row, [d]) == (None, 'SIN BITÁCORA ASOCIADA')


def test_last_cumplimiento_does_not_raise_alert():
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {
        'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': '11111111-1', 'DERIVACIÓN': 'FAE FICTICIO',
        'NOMBRE': 'FICTICIO NNA'}, '', 1)
    d = dict(diary(), control='OK', entries=[entrada('2026-10-10'), entrada('2026-10-11', 'Cumplimiento')])
    app.enriquecer({'Cumplimiento': [row]}, [d], CORTE)
    assert row.values['FECHA OBS.'] == date(2026, 10, 11)
    assert not row.values['ALERTA']


def test_month_export_contains_only_center_dates_in_window(tmp_path):
    d = dict(diary(), control='OK', entries=[entrada('2026-10-10T12:34:56')])
    app.exportar_bitacoras(tmp_path / 'bitacoras.xlsx', [d])
    ws = load_workbook(tmp_path / 'bitacoras.xlsx')['Observaciones del mes']
    assert ws.max_row == 2
    assert ws.cell(2, 8).value.strftime('%H:%M:%S') == '12:34:56'


# --------------------------------------------------------------------------
# Regresiones 1.1
# --------------------------------------------------------------------------

def lote_html(source, filas, tab='Cumplimiento'):
    """Lote del descargador con una consulta y una página XLS HTML por elemento de `filas`."""
    headers = ['RIT', 'TRIBUNAL', 'RUT', 'NOMBRE', 'DERIVACIÓN', 'TIPO', 'SEXO', 'FEC. RESOLUCIÓN', 'FEC. INGRESO EFECTIVO']
    pages = []
    for index, fila in enumerate(filas, 1):
        values = {'RIT': f'X-{index}-2026', 'TRIBUNAL': 'Laja', 'RUT': '11111111-1', 'NOMBRE': 'FICTICIO NNA',
                  'DERIVACIÓN': 'FAE FICTICIO', 'TIPO': 'MP', 'SEXO': 'F', 'FEC. RESOLUCIÓN': '01/09/2026',
                  'FEC. INGRESO EFECTIVO': '02/09/2026', **fila}
        data = ('<html><table><tr><td>Juzgado de Familia de Laja</td></tr><tr><td>Informe de Ingresos en Lista de ' + tab +
                '</td></tr><tr><td>Fecha Emisión : 16/10/2026</td></tr><tr>' + ''.join(f'<td>{h}</td>' for h in headers) +
                '</tr><tr>' + ''.join(f'<td>{values[h]}</td>' for h in headers) + '</tr></table></html>').encode()
        filename = f'CUMP FAE LAJA {index}.xls'
        (source / filename).write_bytes(data)
        pages.append(dict(archivo=filename, sha256=hashlib.sha256(data).hexdigest()))
    (source / 'verificacion.json').write_text(json.dumps(dict(estado='VALIDADA', paginas_esperadas=len(pages), paginas_verificadas=len(pages), archivos=pages)))
    (source / 'resumen.json').write_text(json.dumps(dict(estado='VALIDADA', consultas=[dict(verificacion='verificacion.json', seleccion=dict(tab=tab, modality='3'))])))


def test_text_starting_with_equals_is_kept_as_text(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{'TIPO': '=MP'}])
    save_diaries(source, [diary(rit='X-1-2026')])
    target = tmp_path / 'resultado'
    app.generar(source, target, CORTE)
    ws = load_workbook(target / 'Registros.xlsx')['CUMPLIMIENTO']
    headers = [c.value for c in ws[1]]
    cell = ws.cell(2, headers.index('TIPO') + 1)
    assert cell.value == '=MP' and cell.data_type == 's'
    ws = load_workbook(target / 'META 4 - actualizado 2026-10.xlsx')['FAE Cumplimiento']
    row = next(c.row for r in ws for c in r if c.value == 'X-1-2026')
    assert ws.cell(row, 3).value == '=MP' and ws.cell(row, 3).data_type == 's'


def test_pages_of_one_query_are_not_duplicates(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{}, {}])
    save_diaries(source, [diary(rit='X-1-2026'), diary(rit='X-2-2026', ingreso='202')])
    target = tmp_path / 'resultado'
    app.generar(source, target, CORTE)
    ws = load_workbook(target / 'Registros.xlsx')['MATRIZ_COBERTURA']
    text = next(ws.cell(r, 3).value for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == 'LAJA')
    assert text == 'EXCEL · 2 registro(s)'
    assert 'duplicadas: 0' in ws['A2'].value


def test_existing_output_is_rejected_before_processing(tmp_path):
    target = tmp_path / 'resultado'
    target.mkdir()
    with pytest.raises(ValueError, match='ya existe'):
        app.generar(tmp_path, target, CORTE)


def test_failed_verification_leaves_no_partial_output(tmp_path, monkeypatch):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source)
    monkeypatch.setattr(app.meta, 'verify_meta_output', lambda path: ({}, {}, ['simulado']))
    with pytest.raises(ValueError, match='simulado'):
        app.generar(source, tmp_path / 'resultado', CORTE)
    assert sorted(p.name for p in tmp_path.iterdir()) == ['fuentes']


def test_missing_template_is_an_error(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source)
    with pytest.raises(ValueError, match='plantilla'):
        app.generar(source, tmp_path / 'resultado', CORTE, template=tmp_path / 'no-existe.xlsx')


def test_report_declares_reading_state_and_controls(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source)
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert report['lectura_bitacoras'] == ['SIN ESTADO']
    assert report['controles_bitacora'] == {'VERIFICADA EN EL PERÍODO': 1}


PARAMS = {'tipo_popUp': '11', 'CRR_IdCausa': '301', 'COD_Tribunal': '101', 'TIP_Consulta': '3',
          'ID_Ingreso': '201', 'COD_Etapa': '5', 'FLG_MejorNinez': '0'}


def bitacora_html(fecha='10/10/2026 10:00', etapa='Espera'):
    headers = ['Tipo Observación', 'Etapa', 'Fecha Centro', 'Usuario Centro', 'Observación Centro',
               'Fecha Tribunal', 'Usuario Tribunal', 'Observación Tribunal', 'Acción']
    hidden = ''.join(f'<input type="hidden" name="{n}77" value="{v}">' for n, v in
                     [('HOBS_Centro_', 'Texto ficticio'), ('HOBS_Tribunal_', '---'), ('HCOD_Estado_', '1'),
                      ('HFUN_Tribunal_', ''), ('HFLG_Envio_', 'N')])
    inputs = ''.join(f'<input name="{k}" value="{v}">' for k, v in [
        ('COD_Tribunal', '101'), ('CRR_IdCausa', '301'), ('ID_Ingreso', '201'), ('COD_Etapa', '5'),
        ('TIP_Consulta', '3'), ('tipo_popUp', '11'), ('RIT_Causa', 'X-1-2026'), ('GLS_Nombre', 'FICTICIO NNA'),
        ('GLS_Centro', 'FAE FICTICIO')])
    cells = ['Observación', etapa, fecha, 'Usuario', 'Texto ficticio', '---', '', '---']
    return ('<html><body><form name="InformesPpalForm" action="/SITFAWEB/InformesDAction.do">' + inputs +
            '<table id="TablaInforme"><tr>' + ''.join(f'<td>{h}</td>' for h in headers) + '</tr><tr>' +
            ''.join(f'<td>{c}</td>' for c in cells) + '<td><img id="77" onclick="ver()">' + hidden +
            '</td></tr></table></form></body></html>').encode('utf-8')


def legacy_lote(source, body):
    folder = source / 'Bitacoras 2026-10-16'
    folder.mkdir()
    (folder / 'b1.html').write_bytes(body)
    record = {k: v for k, v in diary().items() if k != 'snapshot'}
    record.update(archivo='b1.html', sha256=hashlib.sha256(body).hexdigest(), params=PARAMS)
    (folder / 'bitacoras.json').write_text(json.dumps(dict(tipo='BITACORAS_LECTURA', version=1, estado='COMPLETA',
                                                            registros=[record])), encoding='utf-8')
    return folder / 'b1.html'


def test_legacy_downloader_html_diaries_are_verified(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    make_excel(source / 'CUMP FAE LAJA.xlsx', 'Laja', 'Cumplimiento', [{'rit': 'X-1-2026', 'der': 'FAE FICTICIO'}])
    legacy_lote(source, bitacora_html())
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert report['alertas'] == 1 and report['lectura_bitacoras'] == ['COMPLETA']


def test_legacy_html_changed_after_download_is_not_used(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    page = legacy_lote(source, bitacora_html())
    page.write_bytes(bitacora_html(etapa='Cumplimiento'))
    records = app.cargar_bitacoras(source, CORTE)
    assert records[0]['control'] == 'BITÁCORA NO VERIFICABLE' and records[0]['entries'] == []


def test_cli_rejects_future_cutoff_and_bad_month(tmp_path):
    import subprocess
    script = Path(app.__file__)
    future = subprocess.run([sys.executable, str(script), '--folder', str(tmp_path), '--output-dir', str(tmp_path / 'r'),
                             '--cutoff', '2999-01-01'], capture_output=True, text=True)
    assert future.returncode == 1 and 'futuro' in future.stderr
    month = subprocess.run([sys.executable, str(script), '--folder', str(tmp_path), '--output-dir', str(tmp_path / 'r'),
                            '--month', '2026-13'], capture_output=True, text=True)
    assert month.returncode == 2 and 'AAAA-MM' in month.stderr


# --------------------------------------------------------------------------
# 2.0: solo lotes del Descargador
# --------------------------------------------------------------------------

def test_diaries_in_separate_folder(tmp_path):
    source, diaries = tmp_path / 'registros', tmp_path / 'bitacoras'
    source.mkdir()
    diaries.mkdir()
    make_excel(source / 'CUMP FAE LAJA.xlsx', 'Laja', 'Cumplimiento', [{'rit': 'X-1-2026', 'der': 'FAE FICTICIO'}])
    legacy_lote(diaries, bitacora_html())
    with pytest.raises(ValueError, match='Faltan bitácoras'):
        app.generar(source, tmp_path / 'r1', CORTE)
    report = app.generar(source, tmp_path / 'r2', CORTE, bitacoras=diaries)
    assert report['alertas'] == 1


def test_filtered_lot_is_accepted_and_reported(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{}])
    save_diaries(source, [diary()])
    data = json.loads((source / 'resumen.json').read_text())
    data.update(filtro_fecha={'desde': '01/09/2026', 'hasta': '30/09/2026'}, otros_filtros='formulario')
    data['consultas'][0].update(consulta='Laja / Cumplimiento / FAE')
    data['consultas'][0]['seleccion']['report'] = '4'
    (source / 'resumen.json').write_text(json.dumps(data))
    target = tmp_path / 'resultado'
    report = app.generar(source, target, CORTE)
    assert report['registros'] == 1 and len(report['filtros_registros']) == 3
    assert report['filtros_registros'][0].startswith('Filtro de fechas 01/09/2026 a 30/09/2026')
    assert 'Laja / Cumplimiento / FAE' in report['filtros_registros'][2]
    control = load_workbook(target / 'META 4 - actualizado 2026-10.xlsx')['Control bitácoras']
    assert control['A5'].value == 'Filtros del lote de registros' and 'Filtro de fechas' in control['B5'].value
    issues = [r[3] for r in load_workbook(target / 'Registros.xlsx')['ERRORES_Y_ADVERTENCIAS'].iter_rows(min_row=4, values_only=True)]
    assert any(str(m).startswith('Filtro de fechas') for m in issues)


def test_unfiltered_lot_declares_no_filters(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{}])
    save_diaries(source, [diary()])
    target = tmp_path / 'resultado'
    assert app.generar(source, target, CORTE)['filtros_registros'] == []
    control = load_workbook(target / 'META 4 - actualizado 2026-10.xlsx')['Control bitácoras']
    assert control['B5'].value == 'Sin filtros'


def test_read_before_cutoff_requires_review(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source, [dict(diary(), capturada='2026-10-10T12:00:00')])
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert report['controles_bitacora'] == {app.LEIDA_ANTES: 1}
    assert app.fecha_lectura(source) == date(2026, 10, 10)


def test_failed_read_in_one_lot_does_not_hide_verified_read_in_another(tmp_path):
    failed = diary(entries=[entrada('2026-10-01', 'Cumplimiento')])
    failed['estado'] = 'FALLIDA'
    ok = diary(entries=[entrada('2026-10-10', 'Espera')])
    for name, record in (('a', failed), ('b', ok)):
        (tmp_path / name).mkdir()
        save_diaries(tmp_path / name, [record])
    records = app.cargar_bitacoras(tmp_path, CORTE)
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {
        'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': '11111111-1', 'DERIVACIÓN': 'FAE FICTICIO',
        'NOMBRE': 'FICTICIO NNA'}, '', 1)
    app.enriquecer({'Cumplimiento': [row]}, records, CORTE)
    assert row.values['ALERTA'] == app.ALERTA


def test_tampered_read_still_blocks_even_with_other_lot(tmp_path):
    bad = dict(diary(), snapshot=dict(diary()['snapshot'], ingreso_id='999'))
    for name, record in (('a', bad), ('b', diary())):
        (tmp_path / name).mkdir()
        save_diaries(tmp_path / name, [record])
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {
        'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': '11111111-1', 'DERIVACIÓN': 'FAE FICTICIO',
        'NOMBRE': 'FICTICIO NNA'}, '', 1)
    assert app.buscar(row, app.cargar_bitacoras(tmp_path, CORTE)) == (None, app.NO_VERIFICABLE)


def test_resumed_lot_previous_reads_are_used(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    make_excel(source / 'CUMP FAE LAJA.xlsx', 'Laja', 'Cumplimiento', [{'rit': 'X-1-2026', 'der': 'FAE FICTICIO'}])
    legacy_lote(source, bitacora_html())
    path = next(source.rglob('bitacoras.json'))
    data = json.loads(path.read_text())
    data['anteriores_no_reenumeradas'], data['registros'] = data['registros'], []
    path.write_text(json.dumps(data))
    assert app.generar(source, tmp_path / 'resultado', CORTE)['alertas'] == 1


def test_queries_without_diaries_are_reported(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{}])
    data = json.loads((source / 'resumen.json').read_text())
    data['consultas'][0].update(estado='VALIDADA')
    data['consultas'][0]['seleccion'].update(tribunal='101', tribunal_nombre='Juzgado de Familia de Laja', modalidad_nombre='FAE')
    (source / 'resumen.json').write_text(json.dumps(data))
    (source / 'b').mkdir()
    (source / 'b' / 'bitacoras_meta4.json').write_text(json.dumps(dict(
        tipo='META4_BITACORAS_MES', desde='2026-09-16', hasta='2026-10-16', registros=[diary()],
        consultas=[dict(tribunal='101', tab='Espera', modality='3', estado='ENUMERADA')])))
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert report['consultas_sin_bitacoras'] == ['Juzgado de Familia de Laja / Cumplimiento / FAE']


# --------------------------------------------------------------------------
# 2.1: resumen por tribunal y filtro de tribunales
# --------------------------------------------------------------------------

def dos_tribunales(source):
    make_excel(source / 'CUMP FAE LAJA.xlsx', 'Laja', 'Cumplimiento', [{'rit': 'X-1-2026', 'der': 'FAE FICTICIO'}])
    make_excel(source / 'ESP RES TOME.xlsx', 'Tomé', 'Espera',
               [{'rit': 'X-2-2026', 'der': 'RESIDENCIA FICTICIA'}, {'rit': 'X-3-2026', 'der': 'RESIDENCIA FICTICIA'}])
    tome = dict(diary(rit='X-9-2026', ingreso='209'), tribunal_nombre='Juzgado de Familia de Tomé')
    save_diaries(source, [diary(entries=[entrada('2026-10-09T10:00:00')]), tome])


def test_summary_sheet_by_tribunal(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    dos_tribunales(source)
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    wb = load_workbook(tmp_path / 'resultado' / 'Registros consolidados.xlsx')
    assert wb.sheetnames[0] == 'Resumen por tribunal'
    rows = {r[0]: dict(zip(app.RESUMEN, r)) for r in wb['Resumen por tribunal'].iter_rows(min_row=2, values_only=True)}
    assert list(rows) == ['LAJA', 'TOMÉ', 'TOTAL']
    assert rows['LAJA']['FAE CUMPLIMIENTO'] == 1 and rows['LAJA']['ALERTAS'] == 1 and rows['LAJA']['SIN FECHA OBS.'] == 0
    assert rows['TOMÉ']['RESIDENCIA ESPERA'] == 2 and rows['TOMÉ']['REVISAR BITÁCORA'] == 2
    assert rows['TOTAL']['TOTAL REGISTROS'] == 3 and report['tribunales'] == 'TODOS'
    assert report['resumen_tribunales'][-1]['TOTAL REGISTROS'] == 3


def test_filter_generates_partial_meta4_for_selected_tribunals(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    dos_tribunales(source)
    target = tmp_path / 'resultado'
    report = app.generar(source, target, CORTE, tribunales=['Juzgado de Familia de Laja'])
    assert report['tribunales'] == ['LAJA'] and report['registros'] == 1
    wb = load_workbook(target / 'META 4 - actualizado 2026-10 - parcial.xlsx')
    assert [c.value for c in wb['FAE Espera']['A'] if c.value in app.meta.TRIBUNALS] == ['LAJA']
    assert wb['Control bitácoras']['B4'].value == 'LAJA'
    matrix = load_workbook(target / 'Registros.xlsx')['MATRIZ_COBERTURA']
    tome = next(r for r in range(5, matrix.max_row + 1) if matrix.cell(r, 1).value == 'TOMÉ')
    assert matrix.cell(tome, 4).value == 'NO INCLUIDO'
    assert 'Cobertura: 1/4' in matrix['A2'].value
    lecturas = load_workbook(target / 'Bitacoras_del_mes.xlsx')['Control de lecturas']
    assert [r[0] for r in lecturas.iter_rows(min_row=2, values_only=True)] == ['Juzgado de Familia de Laja']


def test_filter_rejects_unknown_or_empty_tribunals(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    dos_tribunales(source)
    with pytest.raises(ValueError, match='no reconocido'):
        app.generar(source, tmp_path / 'r1', CORTE, tribunales=['Santiago'])
    with pytest.raises(ValueError, match='tribunales elegidos'):
        app.generar(source, tmp_path / 'r2', CORTE, tribunales=['Lebu'])
    assert not (tmp_path / 'r2').exists()


def test_available_tribunals_from_lot(tmp_path):
    lote_html(tmp_path, [{}])
    data = json.loads((tmp_path / 'resumen.json').read_text())
    data['consultas'][0]['seleccion']['tribunal_nombre'] = 'Juzgado de Familia de Laja'
    (tmp_path / 'resumen.json').write_text(json.dumps(data))
    assert app.tribunales_disponibles(tmp_path) == ['LAJA']


# --------------------------------------------------------------------------
# Auditoría 2.3: rutas vivas que no tenían prueba
# --------------------------------------------------------------------------

def vincular(source, ingreso):
    """Agrega a la página el vínculo remoto que el Descargador registra para cada fila."""
    path = source / 'verificacion.json'
    data = json.loads(path.read_text())
    data['archivos'][0]['vinculos_ingreso'] = [dict(rit='X-1-2026', rut='11111111-1', nombre='FICTICIO NNA',
                                                    programa='FAE FICTICIO', tribunal_codigo='101',
                                                    causa_id='301', ingreso_id=ingreso)]
    path.write_text(json.dumps(data))


def test_downloader_binding_resolves_ambiguous_ingreso_end_to_end(tmp_path):
    records = [diary(ingreso='201', entries=[entrada('2026-10-10', 'Cumplimiento')]),
               diary(ingreso='202', entries=[entrada('2026-10-10', 'Espera')])]
    sin = tmp_path / 'sin'
    sin.mkdir()
    lote_html(sin, [{}])
    save_diaries(sin, records)
    assert app.generar(sin, tmp_path / 'r1', CORTE)['controles_bitacora'] == {'IDENTIDAD AMBIGUA': 1}
    con = tmp_path / 'con'
    con.mkdir()
    lote_html(con, [{}])
    vincular(con, '202')
    save_diaries(con, records)
    report = app.generar(con, tmp_path / 'r2', CORTE)
    assert report['alertas'] == 1 and report['controles_bitacora'] == {app.VERIFICADA: 1}


def test_empty_query_pdf_evidence_covers_combination(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    lote_html(source, [{}])
    (source / 'ESP FAE Jdo Familia Laja.pdf').write_bytes(b'%PDF-ficticio')
    (source / 'verificacion 2.json').write_text(json.dumps(dict(
        estado='SIN_RESULTADOS', paginas_esperadas=0, paginas_verificadas=0, archivos=[],
        evidencia=dict(archivo='ESP FAE Jdo Familia Laja.pdf', registros=0))))
    data = json.loads((source / 'resumen.json').read_text())
    data['consultas'].append(dict(verificacion='verificacion 2.json', estado='SIN_RESULTADOS',
                                  seleccion=dict(tab='Espera', modality='3')))
    (source / 'resumen.json').write_text(json.dumps(data))
    save_diaries(source, [diary()])
    target = tmp_path / 'resultado'
    app.generar(source, target, CORTE)
    ws = load_workbook(target / 'Registros.xlsx')['MATRIZ_COBERTURA']
    laja = next(r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value == 'LAJA')
    assert ws.cell(laja, 2).value == 'PDF · 0 registros' and ws.cell(laja, 3).value == 'EXCEL · 1 registro(s)'


@pytest.mark.parametrize('cambio,mensaje', [
    (lambda r, v: r.update(estado='INCOMPLETA'), 'incompleta'),
    (lambda r, v: v.update(paginas_verificadas=0), 'Faltan páginas'),
    (lambda r, v: v.update(estado='EN_CURSO'), 'no está validada'),
])
def test_downloader_lot_integrity_checks(tmp_path, cambio, mensaje):
    lote_html(tmp_path, [{}])
    resumen = json.loads((tmp_path / 'resumen.json').read_text())
    check = json.loads((tmp_path / 'verificacion.json').read_text())
    cambio(resumen, check)
    (tmp_path / 'resumen.json').write_text(json.dumps(resumen))
    (tmp_path / 'verificacion.json').write_text(json.dumps(check))
    with pytest.raises(ValueError, match=mensaje):
        app.fuentes(tmp_path)


def test_same_query_in_two_lots_is_rejected(tmp_path):
    for name in ('a', 'b'):
        (tmp_path / name).mkdir()
        lote_html(tmp_path / name, [{}])
    with pytest.raises(ValueError, match=r'LAJA \| FAE \| Cumplimiento aparece en varios lotes \(a, b\)'):
        app.fuentes(tmp_path)


def visible(rit='X-1-2026', persona='FICTICIO NNA', advertencia=None):
    record = diary()
    record['snapshot'].update(rit=rit, persona=persona)
    if advertencia:
        record['advertencias'] = [f'Advertencia: {advertencia} visible de bitácora distinto al listado; revisar el ingreso.']
    return record


@pytest.mark.parametrize('record,control', [
    (visible(advertencia='RIT'), app.IDENTIDAD_VISIBLE),
    (visible(rit='X-9-2026'), app.IDENTIDAD_VISIBLE),
    (visible(persona='NNA FICTICIO', advertencia='nombre'), 'OK'),
    (visible(persona='FICTICIÓ  NNA', advertencia='nombre'), 'OK'),
    (visible(persona='FICTICIO NNA SEGUNDO', advertencia='nombre'), 'OK'),
    (visible(persona='OTRA PERSONA', advertencia='nombre'), app.NOMBRE_VISIBLE),
])
def test_visible_identity_rit_blocks_name_only_requires_review(tmp_path, record, control):
    save_diaries(tmp_path, [record])
    assert app.cargar_bitacoras(tmp_path, CORTE)[0]['control'] == control


def test_name_visible_difference_keeps_date_and_marks_yellow(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source, [visible(persona='OTRA PERSONA', advertencia='nombre')])
    target = tmp_path / 'resultado'
    report = app.generar(source, target, CORTE)
    ws = load_workbook(target / 'Registros.xlsx')['CUMPLIMIENTO']
    headers = [c.value for c in ws[1]]
    row = {h: ws.cell(2, i + 1).value for i, h in enumerate(headers)}
    assert row['FECHA OBS.'].date() == date(2026, 10, 10)
    assert row['CONTROL BITÁCORA'] == f'{app.VERIFICADA} · {app.NOMBRE_VISIBLE}'
    assert report['revision_bitacora'] == 1


def test_visible_rit_difference_still_leaves_date_blank(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source, [visible(advertencia='RIT')])
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert report['controles_bitacora'] == {app.IDENTIDAD_VISIBLE: 1}


def test_filtered_diary_lot_is_reported(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    prepare(source)
    path = source / 'bitacoras_meta4.json'
    data = json.loads(path.read_text())
    data['seleccion'] = {'dates': {'start': '01/10/2026', 'end': '16/10/2026'}, 'keep_filters': False}
    path.write_text(json.dumps(data))
    report = app.generar(source, tmp_path / 'resultado', CORTE)
    assert len(report['avisos_bitacoras']) == 1 and 'filtros' in report['avisos_bitacoras'][0]


def test_match_by_name_without_rut_and_program_must_match():
    base = {'TRIBUNAL': 'Laja', 'RIT': 'X-1-2026', 'RUT': None, 'NOMBRE': 'FICTICIO NNA', 'DERIVACIÓN': 'FAE FICTICIO'}
    d = dict(diary(), control='OK', entries=[])
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', dict(base), '', 1)
    assert app.buscar(row, [d])[1] == 'OK'
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', dict(base, NOMBRE='OTRA PERSONA'), '', 1)
    assert app.buscar(row, [d]) == (None, 'SIN BITÁCORA ASOCIADA')
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', dict(base, DERIVACIÓN='OTRO PROGRAMA'), '', 1)
    assert app.buscar(row, [d]) == (None, 'SIN BITÁCORA ASOCIADA')


def test_alert_marking_refuses_misaligned_rows():
    row = app.meta.SourceRow('Cumplimiento', 'FAE', 'LAJA', {'RIT': 'X-1-2026', 'ALERTA': '', app.EXTRA[1]: app.VERIFICADA}, '', 1)
    records = {name: {t: [] for t in app.meta.TRIBUNALS} for name in app.meta.SHEET_DEFS}
    records['FAE Cumplimiento']['LAJA'] = [['X-9-2026']]
    wb = app.meta.build_default_template()
    with pytest.raises(ValueError, match='alinear'):
        app.marcar_planilla(wb, records, {'Espera': [], 'Cumplimiento': [row]})


def test_grid_frame_matches_previous_xlsx_roundtrip_semantics():
    grid = [['', None, 'A1'], [], ['RIT', 'RUT (->RCEI)', 3.0, 2.5, date(2026, 1, 1), '=1+1'], ['', '']]
    df = app.hoja_desde_grilla(grid)
    assert df.shape == (3, 6)
    assert df.iloc[0].tolist()[:3] == [None, None, 'A1'] and all(v is None for v in df.iloc[1])
    header = df.iloc[2].tolist()
    assert header[:2] == ['RIT', 'RUT'] and header[2] == 3 and isinstance(header[2], int) and header[3] == 2.5
    assert header[4].hour == 0 and header[4].year == 2026 and header[5] == '=1+1'


def test_diary_dates_are_iso_and_invalid_dates_are_rejected():
    assert app.fecha('2026-10-10T12:34:56Z').hour == 12
    with pytest.raises(ValueError, match='no reconocida'):
        app.fecha('10/10/2026')


def test_legacy_record_with_null_params_is_not_verifiable(tmp_path):
    legacy_lote(tmp_path, bitacora_html())
    path = next(tmp_path.rglob('bitacoras.json'))
    data = json.loads(path.read_text())
    data['registros'][0]['params'] = None
    path.write_text(json.dumps(data))
    assert app.cargar_bitacoras(tmp_path, CORTE)[0]['control'] == app.NO_VERIFICABLE


# --------------------------------------------------------------------------
# 2.4: varios lotes de registros en un solo Meta 4
# --------------------------------------------------------------------------

MODALIDADES = {'FAE': ('3', 'FAE FICTICIO'), 'Residencia': ('1', 'RESIDENCIA FICTICIA')}


def lote_lista(source, tab, modalidad, rits, tribunal='101'):
    """Lote del Descargador con una sola consulta (tribunal × lista × modalidad)."""
    source.mkdir(parents=True)
    code, programa = MODALIDADES[modalidad]
    headers = ['RIT', 'TRIBUNAL', 'RUT', 'NOMBRE', 'DERIVACIÓN', 'TIPO', 'SEXO', 'FEC. RESOLUCIÓN', 'FEC. INGRESO EFECTIVO']
    filas = ''.join('<tr>' + ''.join(f'<td>{v}</td>' for v in (rit, 'Laja', '11111111-1', 'FICTICIO NNA', programa,
                                                                 'MP', 'F', '01/09/2026', '02/09/2026')) + '</tr>' for rit in rits)
    data = ('<html><table><tr><td>Juzgado de Familia de Laja</td></tr><tr><td>Informe de Ingresos en Lista de ' + tab +
            '</td></tr><tr><td>Fecha Emisión : 16/10/2026</td></tr><tr>' + ''.join(f'<td>{h}</td>' for h in headers) +
            '</tr>' + filas + '</table></html>').encode()
    (source / 'pagina 1.xls').write_bytes(data)
    (source / 'verificacion.json').write_text(json.dumps(dict(estado='VALIDADA', paginas_esperadas=1, paginas_verificadas=1,
        archivos=[dict(archivo='pagina 1.xls', sha256=hashlib.sha256(data).hexdigest())])))
    seleccion = dict(tab=tab, modality=code, tribunal=tribunal, tribunal_nombre='Juzgado de Familia de Laja',
                     modalidad_nombre=modalidad)
    (source / 'resumen.json').write_text(json.dumps(dict(estado='VALIDADA', consultas=[dict(
        verificacion='verificacion.json', estado='VALIDADA', consulta=f'Laja / {tab} / {modalidad}', seleccion=seleccion)])))
    return source


def diario(rit, ingreso, programa):
    return dict(diary(rit=rit, ingreso=ingreso), programa=programa)


def cuatro_lotes(root):
    lotes = [lote_lista(root / 'FAE Espera', 'Espera', 'FAE', ['E-1-2026']),
             lote_lista(root / 'FAE Cumplimiento', 'Cumplimiento', 'FAE', ['C-1-2026', 'C-2-2026']),
             lote_lista(root / 'RES Espera', 'Espera', 'Residencia', ['E-2-2026']),
             lote_lista(root / 'RES Cumplimiento', 'Cumplimiento', 'Residencia', ['C-3-2026'])]
    diaries = root / 'bitacoras'
    diaries.mkdir()
    save_diaries(diaries, [diario('E-1-2026', '1', 'FAE FICTICIO'), diario('C-1-2026', '2', 'FAE FICTICIO'),
                           diario('C-2-2026', '3', 'FAE FICTICIO'), diario('E-2-2026', '4', 'RESIDENCIA FICTICIA'),
                           diario('C-3-2026', '5', 'RESIDENCIA FICTICIA')])
    return lotes, diaries


def filas_meta4(path):
    wb = load_workbook(path)
    return {name: sorted(c.value for r in wb[name] for c in r if re.fullmatch(r'[EC]-\d-2026', str(c.value or '')))
            for name in ('FAE Espera', 'FAE Cumplimiento', 'Residencia Espera', 'Residencia Cumplimiento')}


ESPERADO = {'FAE Espera': ['E-1-2026'], 'FAE Cumplimiento': ['C-1-2026', 'C-2-2026'],
            'Residencia Espera': ['E-2-2026'], 'Residencia Cumplimiento': ['C-3-2026']}


def test_four_separate_lots_make_one_meta4(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    target = tmp_path / 'resultado'
    report = app.generar(lotes, target, CORTE, bitacoras=diaries)
    assert report['registros'] == 5 and report['totales_meta4'] == {k: len(v) for k, v in ESPERADO.items()}
    assert report['lotes_registros'] == ['FAE Espera', 'FAE Cumplimiento', 'RES Espera', 'RES Cumplimiento']
    assert report['controles_bitacora'] == {app.VERIFICADA: 5}
    assert filas_meta4(target / 'META 4 - actualizado 2026-10.xlsx') == ESPERADO
    control = load_workbook(target / 'META 4 - actualizado 2026-10.xlsx')['Control bitácoras']
    assert control['A6'].value == 'Lotes de registros' and control['B6'].value == 4
    resumen = report['resumen_tribunales'][0]
    assert resumen['TRIBUNAL'] == 'LAJA' and resumen['FAE CUMPLIMIENTO'] == 2 and resumen['RESIDENCIA ESPERA'] == 1


def test_parent_folder_with_several_lots_is_combined(tmp_path):
    cuatro_lotes(tmp_path / 'descargas')
    report = app.generar(tmp_path / 'descargas', tmp_path / 'resultado', CORTE)
    assert report['registros'] == 5 and len(report['lotes_registros']) == 4
    assert app.tribunales_disponibles(tmp_path / 'descargas') == ['LAJA']


def test_lot_chosen_twice_or_nested_is_read_once(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    report = app.generar(lotes + [lotes[0], tmp_path], tmp_path / 'resultado', CORTE, bitacoras=[diaries, diaries])
    assert report['registros'] == 5 and len(report['lotes_registros']) == 4


def test_lot_with_espera_and_cumplimiento_combines_with_other_modality(tmp_path):
    fae = lote_lista(tmp_path / 'FAE', 'Espera', 'FAE', ['E-1-2026'])
    otra = lote_lista(tmp_path / 'tmp', 'Cumplimiento', 'FAE', ['C-1-2026'])
    data = json.loads((fae / 'resumen.json').read_text())
    (otra / 'pagina 1.xls').rename(fae / 'pagina 2.xls')
    check = json.loads((otra / 'verificacion.json').read_text())
    check['archivos'][0]['archivo'] = 'pagina 2.xls'
    (fae / 'verificacion 2.json').write_text(json.dumps(check))
    data['consultas'].append(dict(json.loads((otra / 'resumen.json').read_text())['consultas'][0], verificacion='verificacion 2.json'))
    (fae / 'resumen.json').write_text(json.dumps(data))
    for name in ('resumen.json', 'verificacion.json'):
        (otra / name).unlink()
    otra.rmdir()
    lote_lista(tmp_path / 'RES', 'Cumplimiento', 'Residencia', ['C-3-2026'])
    save_diaries(tmp_path, [diario('E-1-2026', '1', 'FAE FICTICIO'), diario('C-1-2026', '2', 'FAE FICTICIO'),
                            diario('C-3-2026', '5', 'RESIDENCIA FICTICIA')])
    report = app.generar([tmp_path / 'FAE', tmp_path / 'RES'], tmp_path / 'resultado', CORTE, bitacoras=tmp_path)
    assert report['totales_meta4'] == {'FAE Espera': 1, 'FAE Cumplimiento': 1, 'Residencia Espera': 0, 'Residencia Cumplimiento': 1}


def test_same_query_downloaded_in_two_lots_is_rejected_before_output(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    lote_lista(tmp_path / 'FAE Espera otra vez', 'Espera', 'FAE', ['E-1-2026'])
    target = tmp_path / 'resultado'
    with pytest.raises(ValueError, match=r'Laja / Espera / FAE está en dos lotes \(FAE Espera y FAE Espera otra vez\)'):
        app.generar(tmp_path, target, CORTE, bitacoras=diaries)
    assert not target.exists() and not list(tmp_path.glob('resultado*'))


def test_folder_without_lot_cannot_be_mixed_with_lots(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    with pytest.raises(ValueError, match='no contiene un lote del Descargador'):
        app.fuentes(lotes + [diaries])
    sueltos = [tmp_path / 'x', tmp_path / 'y']
    for folder in sueltos:
        folder.mkdir()
    with pytest.raises(ValueError, match='Excel sueltos elige una sola carpeta'):
        app.fuentes(sueltos)


def test_incomplete_lot_names_the_lot(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    data = json.loads((lotes[2] / 'resumen.json').read_text())
    (lotes[2] / 'resumen.json').write_text(json.dumps(dict(data, estado='INCOMPLETA')))
    with pytest.raises(ValueError, match=r'incompleta \(RES Espera\)'):
        app.fuentes(lotes)


def test_filters_name_their_lot(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    data = json.loads((lotes[1] / 'resumen.json').read_text())
    (lotes[1] / 'resumen.json').write_text(json.dumps(dict(data, otros_filtros='formulario')))
    report = app.generar(lotes, tmp_path / 'resultado', CORTE, bitacoras=diaries)
    assert report['filtros_registros'] == ['Filtros del formulario de SITFA conservados: el lote puede no incluir '
                                           'todos los registros. [FAE Cumplimiento]']


def test_cli_accepts_repeated_folder(tmp_path, monkeypatch, capsys):
    lotes, diaries = cuatro_lotes(tmp_path)
    argv = ['meta4']
    for lote in lotes:
        argv += ['--folder', str(lote)]
    argv += ['--diaries', str(diaries), '--output-dir', str(tmp_path / 'resultado'), '--cutoff', '2026-10-16']
    monkeypatch.setattr(sys, 'argv', argv)
    app.main()
    assert json.loads(capsys.readouterr().out)['registros'] == 5


# --------------------------------------------------------------------------
# 2.5: regenerar desde Registros.xlsx revisado
# --------------------------------------------------------------------------

def editar_registros(path, cambios, borrar=()):
    """`cambios`: {(hoja, RIT): {columna: valor}}. `borrar`: [(hoja, RIT)]."""
    wb = load_workbook(path)
    for sheet in ('ESPERA', 'CUMPLIMIENTO'):
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        for row in range(ws.max_row, 1, -1):
            rit = ws.cell(row, headers.index('RIT') + 1).value
            if (sheet, rit) in borrar:
                ws.delete_rows(row)
                continue
            for column, value in cambios.get((sheet, rit), {}).items():
                ws.cell(row, headers.index(column) + 1, value)
    wb.save(path)


def fecha_meta4(path, sheet, rit):
    ws = load_workbook(path)[sheet]
    row = next(c.row for r in ws for c in r if c.value == rit)
    return ws.cell(row, 8).value, ws.cell(row, 1).fill.fgColor.rgb


def test_regenerate_uses_user_corrections(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    first = tmp_path / 'Resultado META4 1'
    app.generar(lotes, first, CORTE, bitacoras=diaries)
    editar_registros(first / 'Registros.xlsx', {
        ('ESPERA', 'E-1-2026'): {'FECHA OBS.': '05-10-2026', 'CONTROL BITÁCORA': 'REVISADO manualmente'},
        ('CUMPLIMIENTO', 'C-1-2026'): {'ETAPA ÚLTIMA OBS.': 'Cumplimiento', 'CONTROL BITÁCORA': 'pendiente'},
    }, borrar=[('CUMPLIMIENTO', 'C-3-2026')])
    original = (first / 'META 4 - actualizado 2026-10.xlsx').read_bytes()
    second = tmp_path / 'Resultado META4 2 regenerado'
    report = app.regenerar(first / 'Registros.xlsx', second)
    assert report['registros'] == 4 and report['alertas'] == 1 and report['revision_bitacora'] == 1
    assert report['totales_meta4'] == {'FAE Espera': 1, 'FAE Cumplimiento': 2, 'Residencia Espera': 1,
                                       'Residencia Cumplimiento': 0}
    assert report['lotes_registros'] == ['FAE Espera', 'FAE Cumplimiento', 'RES Espera', 'RES Cumplimiento']
    assert report['regenerado_desde'] == str((first / 'Registros.xlsx').resolve())
    out = second / 'META 4 - actualizado 2026-10.xlsx'
    day, _ = fecha_meta4(out, 'FAE Espera', 'E-1-2026')
    assert day.date() == date(2026, 10, 5)
    assert fecha_meta4(out, 'FAE Cumplimiento', 'C-1-2026')[1].endswith(app.AMARILLO)
    assert fecha_meta4(out, 'FAE Cumplimiento', 'C-2-2026')[1].endswith(app.ROJO)
    assert not fecha_meta4(out, 'FAE Espera', 'E-1-2026')[1].endswith((app.ROJO, app.AMARILLO))
    control = {r[0]: r[1:] for r in load_workbook(out)['Control bitácoras'].iter_rows(values_only=True)}
    assert 'Regenerado desde Registros.xlsx revisado' in control
    ws = load_workbook(second / 'Registros.xlsx')['CUMPLIMIENTO']
    headers = [c.value for c in ws[1]]
    assert [ws.cell(r, headers.index('ALERTA') + 1).value for r in (2, 3)] == [None, app.ALERTA]
    assert load_workbook(second / 'Registros.xlsx').sheetnames[2:] == ['INVENTARIO_ARCHIVOS', 'MATRIZ_COBERTURA',
                                                                        'ERRORES_Y_ADVERTENCIAS']
    assert (second / 'Bitacoras_del_mes.xlsx').is_file() and (second / 'Registros consolidados.xlsx').is_file()
    assert (first / 'META 4 - actualizado 2026-10.xlsx').read_bytes() == original


def test_regenerate_rejects_bad_date_without_output(tmp_path):
    lotes, diaries = cuatro_lotes(tmp_path)
    first = tmp_path / 'r1'
    app.generar(lotes, first, CORTE, bitacoras=diaries)
    editar_registros(first / 'Registros.xlsx', {('ESPERA', 'E-1-2026'): {'FECHA OBS.': 'ayer'}})
    with pytest.raises(ValueError, match='FECHA OBS. no reconocida en la hoja ESPERA, fila 2'):
        app.regenerar(first / 'Registros.xlsx', tmp_path / 'r2')
    assert not list(tmp_path.glob('r2*'))


def test_regenerate_without_report_needs_cutoff_and_keeps_partial_filter(tmp_path):
    source = tmp_path / 'fuentes'
    source.mkdir()
    dos_tribunales(source)
    first = tmp_path / 'r1'
    app.generar(source, first, CORTE, tribunales=['Laja'])
    report = app.regenerar(first / 'Registros.xlsx', tmp_path / 'r2')
    assert report['tribunales'] == ['LAJA'] and (tmp_path / 'r2' / 'META 4 - actualizado 2026-10 - parcial.xlsx').is_file()
    (first / 'Reporte META4.json').unlink()
    with pytest.raises(ValueError, match='indica la fecha de corte'):
        app.regenerar(first / 'Registros.xlsx', tmp_path / 'r3')
    assert app.regenerar(first / 'Registros.xlsx', tmp_path / 'r3', corte=CORTE)['registros'] == 1


def test_cli_regenerate(tmp_path, monkeypatch, capsys):
    lotes, diaries = cuatro_lotes(tmp_path)
    app.generar(lotes, tmp_path / 'r1', CORTE, bitacoras=diaries)
    monkeypatch.setattr(sys, 'argv', ['meta4', '--regenerar', str(tmp_path / 'r1' / 'Registros.xlsx'),
                                      '--output-dir', str(tmp_path / 'r2')])
    app.main()
    assert json.loads(capsys.readouterr().out)['registros'] == 5
