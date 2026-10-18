"""Lectura de los Excel del Descargador SITFA: XLS BIFF, OOXML, XML Spreadsheet 2003 y HTML tabular.

Extraído de descargador/motor.py 2.7.6 (solo read_excel y lo que necesita). Cambios respecto del
original: el HTML sin charset declarado se decodifica como UTF-8 o Windows-1252 (antes ISO-8859-1,
que deformaba las tildes de un archivo UTF-8) y el OOXML ignora una dimensión declarada desactualizada
(antes se perdían filas sin aviso).
"""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
import io
import re
import unicodedata
import zipfile

from lxml import etree, html


class PocError(Exception):
    """Solo mensajes controlados: nunca valores del servidor ni excepciones crudas."""


def normalized(value):
    text = ("" if value is None else str(value)).replace("\xa0", " ")
    text = "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))
    return " ".join(text.upper().split())


def record_key(rit, name):
    r = normalized(rit).replace("–", "-").replace("—", "-")
    match = re.search(r"\b[A-Z]+\s*-\s*\d+\s*-\s*\d{4}\b", r)
    if not match or not normalized(name):
        raise PocError("No se reconoce la estructura RIT/Nombre del listado; descarga detenida.")
    return re.sub(r"\s+", "", match.group(0)), normalized(name)


WINDOWS_1252 = {'iso-8859-1', 'iso8859-1', 'latin1', 'latin-1', 'windows-1252', 'cp1252', 'us-ascii', 'ascii'}


def html_text(data):
    """Texto del HTML. Sin charset declarado: UTF-8 si es válido, si no Windows-1252.

    Un charset Latin-1 declarado se lee como Windows-1252, igual que un navegador."""
    if data.startswith(b'\xef\xbb\xbf'):
        return data.decode('utf-8-sig')
    if data.lstrip().startswith(b'<?xml'):
        return data  # lxml respeta la declaración XML
    declared = re.search(rb'<meta\b[^>]*\bcharset\s*=\s*["\']?\s*([a-zA-Z0-9_-]+)', data[:8192], re.I)
    charset = declared.group(1).decode('ascii').lower() if declared else ''
    if charset and charset not in WINDOWS_1252 and charset not in ('utf-8', 'utf8'):
        return data  # otra codificación declarada: la resuelve lxml
    if charset not in WINDOWS_1252:
        try:
            return data.decode('utf-8')
        except UnicodeDecodeError:
            pass
    try:
        return data.decode('cp1252')
    except UnicodeDecodeError:
        return data.decode('latin-1')  # bytes sin carácter en Windows-1252 (0x81, 0x8D, …)


def root_html(data):
    try:
        return html.fromstring(html_text(data) if isinstance(data, bytes) else data,
                               parser=html.HTMLParser(no_network=True))
    except (etree.ParserError, ValueError):
        raise PocError("Respuesta HTML vacia o ilegible.") from None


def has_login(root):
    return bool(root.xpath('//input[translate(@type,"PASSWORD","password")="password"]'))


@dataclass(frozen=True)
class RecordColumns:
    case: tuple[int, ...]
    name: int


def record_columns(row):
    """Encabezados observados en XLS reales; no confunde Nombre Centro con el NNA."""
    values = [normalized(v) for v in row]
    if "RIT" in values:
        case_labels = ("RIT",)
    elif all(label in values for label in ("TIPO CAUSA", "ROL CAUSA", "ERA CAUSA")):
        case_labels = ("TIPO CAUSA", "ROL CAUSA", "ERA CAUSA")
    else:
        return None
    names = [i for i, value in enumerate(values) if value in ("NOMBRE", "NOMBRE MENOR")]
    if not names:
        return None
    if len(names) != 1 or any(values.count(label) != 1 for label in case_labels):
        raise PocError("El listado contiene columnas de causa o nombre ambiguas.")
    return RecordColumns(tuple(values.index(label) for label in case_labels), names[0])


def record_from_row(row, columns):
    if len(row) <= max(*columns.case, columns.name):
        return None
    parts = [
        normalized(int(row[i]) if isinstance(row[i], float) and row[i].is_integer() else row[i]) for i in columns.case
    ]
    if len(parts) == 3:
        if not (
            re.fullmatch(r"[A-Z]", parts[0]) and re.fullmatch(r"\d+", parts[1]) and re.fullmatch(r"\d{4}", parts[2])
        ):
            return None
        rit = "-".join(parts)
    else:
        rit = parts[0]
    if not re.search(r"\b[A-Z]+\s*-\s*\d+\s*-\s*\d{4}\b", rit):
        return None
    return record_key(rit, row[columns.name])


def records_in_grid(rows):
    header = None
    result = Counter()
    for row in rows:
        columns = record_columns(row)
        if columns is not None:
            header = columns
            continue
        if header is None:
            continue
        key = record_from_row(row, header)
        if key is not None:
            result[key] += 1
    return result


def table_grid(table):
    result = []
    cells = 0
    for row in table.xpath('./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr'):
        values = [cell.text_content() for cell in row.xpath('./th | ./td')]
        cells += len(values)
        if len(result) >= 200_000 or len(values) > 200 or cells > 2_000_000:
            raise PocError('La tabla supera los límites de filas/celdas.')
        result.append(values)
    return result


@dataclass
class ExcelInfo:
    format: str
    records: Counter
    grid: list | None = None


def read_excel(data):
    if not data:
        raise PocError("El XLS esta vacio.")
    if len(data) > 80 * 1024 * 1024:
        raise PocError('El Excel supera 80 MB.')
    grids = []
    kind = ""
    try:
        if data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
            import xlrd

            wb = xlrd.open_workbook(file_contents=data, logfile=io.StringIO())
            try:
                if sum(s.nrows * s.ncols for s in wb.sheets()) > 2_000_000 or any(s.ncols > 200 for s in wb.sheets()):
                    raise PocError('El Excel BIFF excede los límites de celdas.')
                grids = []
                for sheet in wb.sheets():
                    grid = []
                    for i in range(sheet.nrows):
                        values = sheet.row_values(i)
                        for j, cell_type in enumerate(sheet.row_types(i)):
                            if cell_type == xlrd.XL_CELL_DATE:
                                values[j] = xlrd.xldate_as_datetime(values[j], wb.datemode)
                        grid.append(values)
                    grids.append(grid)
            finally:
                wb.release_resources()
            kind = "XLS-BIFF"
        elif data.startswith(b"PK\x03\x04"):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(item.file_size for item in entries) > 160 * 1024 * 1024:
                    raise PocError('El Excel comprimido excede los límites de lectura.')
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                grids = []
                cells = 0
                for sheet in wb.worksheets:
                    if sheet.max_row and sheet.max_row > 200_000 or sheet.max_column and sheet.max_column > 200:
                        raise PocError('Las dimensiones del Excel exceden los límites de lectura.')
                    # La dimensión declarada puede estar desactualizada: se leen todas las filas reales.
                    sheet.reset_dimensions()
                    grid = []
                    for row in sheet.iter_rows(values_only=True):
                        cells += len(row)
                        if cells > 2_000_000 or len(row) > 200 or len(grid) >= 200_000:
                            raise PocError('El Excel excede los límites de lectura.')
                        grid.append(row)
                    # Mismo ancho por fila que con la dimensión declarada: celdas finales vacías = None.
                    width = max((len(row) for row in grid), default=0)
                    grids.append([tuple(row) + (None,) * (width - len(row)) for row in grid])
            finally:
                wb.close()
            kind = "OOXML-con-extension-xls"
        elif b"urn:schemas-microsoft-com:office:spreadsheet" in data[:4096]:
            root = etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True))
            ns = {"ss": "urn:schemas-microsoft-com:office:spreadsheet"}
            for table in root.xpath('//ss:Table', namespaces=ns):
                grid = []
                for row in table.xpath('./ss:Row', namespaces=ns):
                    row_index = row.get('{' + ns['ss'] + '}Index')
                    if row_index:
                        target = int(row_index)
                        if not len(grid) < target <= 200_000:
                            raise PocError('Índice de fila XML inválido.')
                        grid.extend([[] for _ in range(target - 1 - len(grid))])
                    vals = []
                    for cell in row.xpath('./ss:Cell', namespaces=ns):
                        index = cell.get('{' + ns['ss'] + '}Index')
                        if index:
                            if not 1 <= int(index) <= 200:
                                raise PocError('El índice XML excede los límites de columnas.')
                            vals.extend([""] * (int(index) - 1 - len(vals)))
                        value = "".join(cell.xpath('./ss:Data//text()', namespaces=ns))
                        data_cell = cell.find('ss:Data', namespaces=ns)
                        if data_cell is not None and data_cell.get('{' + ns['ss'] + '}Type') == 'DateTime':
                            try:
                                value = datetime.fromisoformat(value.rstrip('Z')).date()
                            except ValueError:
                                pass
                        vals.append(value)
                        if len(vals) > 200:
                            raise PocError('La fila XML excede los límites de columnas.')
                    grid.append(vals)
                    if len(grid) > 200_000:
                        raise PocError('La tabla XML excede los límites de filas.')
                grids.append(grid)
            kind = "XML-Spreadsheet-2003"
        else:
            root = root_html(data)
            if has_login(root):
                raise PocError("Sesion expirada: la descarga contiene un formulario de acceso.")
            if root.xpath('//form | //script'):
                raise PocError("La descarga es una pagina web activa, no un Excel verificable.")
            grids = [table_grid(t) for t in root.xpath('//table')]
            kind = "HTML-tabular-con-extension-xls"
        candidates = [(records_in_grid(grid), grid) for grid in grids]
        candidates = [(r, g) for r, g in candidates if r]
        if len(candidates) != 1:
            raise PocError(
                "Se recibio el Excel, pero no se reconoce un unico listado de causas y nombres. No es una consulta sin resultados."
            )
        return ExcelInfo(kind, candidates[0][0], candidates[0][1])
    except PocError:
        raise
    except Exception:
        raise PocError("No se pudo abrir/parsear el archivo Excel; requiere revision local.") from None
