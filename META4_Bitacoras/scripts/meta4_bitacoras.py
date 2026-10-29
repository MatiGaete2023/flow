"""Meta 4 a partir de los lotes del Descargador SITFA: registros y bitácoras de un mes móvil.

No se conecta a SITFA ni a RUS: solo lee carpetas ya descargadas."""
from __future__ import annotations
import argparse
import calendar
from collections import Counter, defaultdict
from collections.abc import Mapping
from datetime import date, datetime
import hashlib
import json
import os
from pathlib import Path
import re
import secrets
import shutil
import sys
from dataclasses import dataclass, field, replace

BASE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BASE / 'src'))
import generate_meta4 as meta
import lector_excel
from nurus.bitacora_html import read_snapshot
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = '2.5'
ALERTA = 'ALERTA: última observación en Espera; registro en Cumplimiento'
EXTRA = ['ETAPA ÚLTIMA OBS.', 'CONTROL BITÁCORA', 'ALERTA']
VERIFICADA = 'VERIFICADA EN EL PERÍODO'
ROJO = 'FFC7CE'
AMARILLO = 'FFF2CC'
SALIDAS = ('Resultado META4', 'Salida META4')
EN_CURSO = '.en-curso-'
ALIAS_COLUMNAS = {'RUT RCEI': 'RUT', 'NOMBRE MENOR': 'NOMBRE'}
IDENTIDAD = ('tribunal_codigo', 'causa_id', 'ingreso_id')
NO_COMPLETADA = 'LECTURA NO COMPLETADA'
NO_VERIFICABLE = 'BITÁCORA NO VERIFICABLE'
IDENTIDAD_VISIBLE = 'REVISAR IDENTIDAD VISIBLE'
NOMBRE_VISIBLE = 'REVISAR NOMBRE VISIBLE (fecha tomada de la bitácora)'
LEIDA_ANTES = 'REVISAR: bitácora leída antes de la fecha de corte'
BLOQUEANTES = (NO_VERIFICABLE, IDENTIDAD_VISIBLE)
CONSULTA_CUBIERTA = ('ENUMERADA', 'SIN_REGISTROS')


def periodo(corte):
    year, month = (corte.year - 1, 12) if corte.month == 1 else (corte.year, corte.month - 1)
    return date(year, month, min(corte.day, calendar.monthrange(year, month)[1])), corte


def fecha(value):
    # Toda entrada nace de bitacora_html._date(...).isoformat(): otro formato es un JSON alterado.
    try:
        return datetime.fromisoformat(str(value or '').strip().replace('Z', '+00:00')).replace(tzinfo=None)
    except ValueError:
        raise ValueError('Fecha de bitácora no reconocida.') from None


def etapa(entry):
    # La etapa se toma de la columna de la bitácora, nunca del texto libre.
    value = meta.norm(entry.get('etapa'))
    if value in ('ESPERA', 'EN ESPERA', 'LISTA DE ESPERA', 'LISTA ESPERA'):
        return 'Espera'
    if value in ('CUMPLIMIENTO', 'EN CUMPLIMIENTO'):
        return 'Cumplimiento'
    return None


def entradas_periodo(entries, corte):
    inicio, fin = periodo(corte)
    return [e for e in entries if inicio <= fecha(e['fecha']).date() <= fin]


def ultima(entries, corte):
    valid = [(fecha(e['fecha']), etapa(e)) for e in entradas_periodo(entries, corte) if etapa(e)]
    if not valid:
        return None, '', 'SIN OBSERVACIÓN DE ESPERA/CUMPLIMIENTO EN EL MES'
    last = max(item[0] for item in valid)
    stages = {s for d, s in valid if d == last}
    if len(stages) > 1:
        return last.date(), '', 'REVISAR: etapas distintas en la misma fecha y hora'
    return last.date(), stages.pop(), VERIFICADA


def dentro(folder, name):
    target = (folder / name).resolve()
    if not target.is_relative_to(folder.resolve()):
        raise ValueError('Referencia de bitácora fuera de la carpeta de origen.')
    return target


def es_salida(path, origen):
    """Carpetas de resultados (terminados o en curso) nunca son fuentes."""
    return any(part.startswith(SALIDAS) or EN_CURSO in part for part in path.relative_to(origen).parts)


def raices(carpetas):
    """Una carpeta o varias, resueltas y sin repetir, en el orden elegido."""
    items = [carpetas] if isinstance(carpetas, (str, os.PathLike)) else list(carpetas or ())
    result = []
    for item in items:
        path = Path(item).expanduser().resolve()
        if path not in result:
            result.append(path)
    return result


def buscar_archivos(roots, name):
    """Archivos `name` bajo cada carpeta, sin resultados de Meta 4 ni repeticiones por carpetas anidadas."""
    found = []
    for root in roots:
        for path in sorted(root.rglob(name)):
            if not es_salida(path, root) and path.resolve() not in found:
                found.append(path.resolve())
    return found


def base_comun(roots):
    """Carpeta desde la que se muestran las rutas de las fuentes: con varios lotes, su carpeta común."""
    if len(roots) == 1:
        return roots[0]
    try:
        return Path(os.path.commonpath(roots))
    except ValueError:  # Unidades distintas en Windows.
        return Path(roots[0].anchor)


def etiqueta_lote(path, base):
    folder = path.parent
    return str(folder.relative_to(base)) if folder != base and folder.is_relative_to(base) else folder.name


@dataclass
class Lectura:
    """Bitácoras leídas desde uno o más lotes del Descargador."""
    registros: list = field(default_factory=list)
    estados: list = field(default_factory=list)
    cubiertas: set = field(default_factory=set)
    declara_alcance: bool = False
    avisos: list = field(default_factory=list)


def _tokens(value):
    return set(meta.norm(value).replace(',', ' ').split())


def nombres_compatibles(a, b):
    """Mismo nombre con otro orden, tildes o un nombre/apellido de menos (al menos dos palabras en común)."""
    x, y = _tokens(a), _tokens(b)
    return bool(x) and bool(y) and (x == y or ((x <= y or y <= x) and len(x & y) >= 2))


def control_identidad(record, snapshot):
    """Compara lo visible en la bitácora con la fila del listado. La identidad remota
    (tribunal, causa, ingreso) ya coincide: un RIT visible distinto bloquea; un nombre
    escrito distinto solo pide revisión y la fecha se usa igual."""
    avisos = ' '.join(map(str, record.get('advertencias') or ()))
    if 'RIT visible' in avisos or meta.norm(snapshot.get('rit')) != meta.norm(record.get('rit')):
        return IDENTIDAD_VISIBLE
    if nombres_compatibles(snapshot.get('persona'), record.get('nombre')):
        return None
    return NOMBRE_VISIBLE


def dia_captura(value):
    """Día local de la lectura. La extensión registra la hora en UTC."""
    moment = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
    return (moment.astimezone() if moment.tzinfo else moment).date()


def archivos_bitacoras(root):
    roots = raices(root)
    # Formato propio de Meta 4 1.0/1.1: se sigue aceptando para lotes ya generados.
    return buscar_archivos(roots, 'bitacoras.json') + buscar_archivos(roots, 'bitacoras_meta4.json')


def _consulta(item):
    return str(item.get('tribunal', '')), str(item.get('tab', '')), str(item.get('modality', ''))


def _leer_bitacoras(root, corte):
    """Registros de bitácora del período, con el estado y el alcance de cada lote de lectura."""
    lectura = Lectura()
    paths = archivos_bitacoras(root)
    if not paths:
        raise ValueError('Faltan bitácoras. Selecciona la carpeta de bitácoras del Descargador (bitacoras.json y sus HTML).')
    inicio = periodo(corte)[0]
    for path in paths:
        if path.stat().st_size > 64 * 1024 * 1024:
            raise ValueError('Un lote de bitácoras excede 64 MB; revisa la carpeta elegida.')
        data = json.loads(path.read_text(encoding='utf-8'))
        own = data.get('tipo') == 'META4_BITACORAS_MES'
        if own and (inicio < date.fromisoformat(data['desde']) or corte > date.fromisoformat(data['hasta'])):
            raise ValueError('Las bitácoras guardadas no cubren este período. Usa su fecha de corte o vuelve a descargarlas.')
        if not own and (data.get('tipo') != 'BITACORAS_LECTURA' or data.get('version') != 1):
            raise ValueError('Formato de bitácoras no reconocido.')
        lectura.estados.append(str(data.get('estado') or 'SIN ESTADO'))
        if 'consultas' in data:
            lectura.declara_alcance = True
            lectura.cubiertas.update(_consulta(q) for q in data['consultas'] if q.get('estado') in CONSULTA_CUBIERTA)
        seleccion = data.get('seleccion') or {}
        if seleccion.get('dates') or seleccion.get('keep_filters'):
            lectura.avisos.append(f'El lote de bitácoras {path.parent.name} usó filtros de fecha o del formulario; '
                                  'puede no incluir todos los ingresos.')
        # Un lote retomado conserva lecturas previas aún no reenumeradas: también valen.
        for original in data.get('registros', []) + data.get('anteriores_no_reenumeradas', []):
            record = dict(original)
            record['entries'] = []
            record['control'] = NO_COMPLETADA
            if record.get('estado') == 'LEIDA':
                try:
                    if own:
                        snapshot = record['snapshot']
                    else:
                        body = dentro(path.parent, record['archivo']).read_bytes()
                        if hashlib.sha256(body).hexdigest() != record['sha256']:
                            raise ValueError('La bitácora cambió después de su descarga.')
                        snapshot = read_snapshot(body, record['params'])
                    if tuple(str(snapshot[k]) for k in IDENTIDAD) != tuple(str(record[k]) for k in IDENTIDAD):
                        raise ValueError('La bitácora pertenece a otro ingreso.')
                    leida = dia_captura(record['capturada'])
                    record['entries'] = entradas_periodo(snapshot['entries'], corte)
                    identidad = control_identidad(record, snapshot)
                    if identidad:
                        record['control'] = identidad
                    elif leida < corte:
                        # Lo escrito entre la lectura y el corte no está en la bitácora guardada.
                        record['control'] = LEIDA_ANTES
                    else:
                        record['control'] = 'OK'
                except (KeyError, ValueError, OSError, TypeError):
                    record['control'] = NO_VERIFICABLE
            lectura.registros.append(record)
    return lectura


def cargar_bitacoras(root, corte):
    return _leer_bitacoras(root, corte).registros


def fecha_lectura(root):
    """Último día con lecturas en los lotes de bitácoras; sirve para proponer la fecha de corte."""
    days = []
    for path in archivos_bitacoras(root):
        try:
            data = json.loads(path.read_text(encoding='utf-8'))
            for record in data.get('registros', []) + data.get('anteriores_no_reenumeradas', []):
                if record.get('estado') == 'LEIDA' and record.get('capturada'):
                    days.append(dia_captura(record['capturada']))
        except (OSError, ValueError, TypeError, AttributeError):
            continue
    return max(days) if days else None


def normal(value):
    return meta.norm(value)


def _tribunal(value):
    return meta.tribunal_name(value) or normal(value)


def _rut(value):
    return normal(value).replace(' ', '')


def indexar_bitacoras(diaries):
    """Índice por (tribunal, RIT): cada fila compara solo con sus bitácoras posibles."""
    index = defaultdict(list)
    for diary in diaries:
        index[(_tribunal(diary.get('tribunal_nombre')), normal(diary.get('rit')))].append(diary)
    return index


def buscar(row, diaries):
    index = diaries if isinstance(diaries, Mapping) else indexar_bitacoras(diaries)
    values = row.values
    rut = _rut(values.get('RUT'))
    nombre = normal(values.get('NOMBRE'))
    programa = normal(values.get('DERIVACIÓN'))
    remote = values.get('_REMOTE')
    candidates = []
    for diary in index.get((_tribunal(values.get('TRIBUNAL')), normal(values.get('RIT'))), ()):
        if rut:
            if _rut(diary.get('rut')) != rut:
                continue
        elif not nombre or normal(diary.get('nombre')) != nombre:
            continue
        if not programa or normal(diary.get('programa')) != programa:
            continue
        if remote and tuple(str(diary.get(k, '')) for k in IDENTIDAD) != remote:
            continue
        candidates.append(diary)
    identities = {tuple(r.get(k) for k in IDENTIDAD) for r in candidates}
    if len(identities) != 1 or not all(next(iter(identities), ())):
        return None, 'IDENTIDAD AMBIGUA' if candidates else 'SIN BITÁCORA ASOCIADA'
    # Un archivo alterado o una identidad visible distinta siempre detienen la asociación.
    blocking = next((r['control'] for r in candidates if r.get('control') in BLOQUEANTES), None)
    if blocking:
        return None, blocking
    # La misma identidad puede venir de ambas pestañas o de varios lotes (p. ej. por el límite
    # de lecturas): una lectura fallida en un lote no anula la lectura verificada de otro.
    read = [r for r in candidates if r.get('control') in ('OK', LEIDA_ANTES, NOMBRE_VISIBLE)]
    if not read:
        return None, candidates[0].get('control', NO_COMPLETADA)
    newest = max(read, key=lambda r: str(r.get('capturada', '')))
    if newest['control'] == LEIDA_ANTES:
        return None, newest['control']
    return newest, newest['control']


def enriquecer(rows, diaries, corte):
    index = indexar_bitacoras(diaries)
    for group in rows.values():
        for row in group:
            diary, control = buscar(row, index)
            day, stage = None, ''
            if diary:
                day, stage, resultado = ultima(diary['entries'], corte)
                # El aviso de nombre visible se conserva junto al resultado: la fila queda en amarillo.
                control = resultado if control == 'OK' else f'{resultado} · {control}'
            row.values['FECHA OBS.'] = day
            row.values[EXTRA[0]] = stage
            row.values[EXTRA[1]] = control
            row.values['ALERTA'] = ALERTA if row.status == 'Cumplimiento' and stage == 'Espera' else ''


def revisada(control):
    """Control sin pendiente: verificada en el período o marcada REVISADO por el usuario."""
    return control == VERIFICADA or normal(control).startswith('REVISADO')


def requiere_revision(values):
    return bool(values['ALERTA']) or not revisada(values[EXTRA[1]])


def tabla(ws, headers, records, formatos=None):
    """Hoja simple. `formatos` fija el formato numérico de columnas por encabezado."""
    formatos = formatos or {}
    ws.append(headers)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(color='FFFFFF', bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    ws.freeze_panes = 'A2'
    for index in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(index)].width = 24
    alert_col = headers.index('ALERTA') if 'ALERTA' in headers else None
    column_formats = [formatos.get(h) for h in headers]
    memo = meta.StyleMemo()
    alert_fill = PatternFill('solid', fgColor=ROJO)
    for row_idx, record in enumerate(records, 2):
        row = [ws.cell(row_idx, col, meta.excel_value(record.get(h))) for col, h in enumerate(headers, 1)]
        alert = alert_col is not None and bool(row[alert_col].value)
        for cell, forced in zip(row, column_formats):
            meta.guard_text(cell)
            number_format = forced or ('dd-mm-yyyy' if isinstance(cell.value, (date, datetime)) else None)
            if not number_format and not alert:
                continue

            def setter(target, number_format=number_format, alert=alert):
                if number_format:
                    target.number_format = number_format
                if alert:
                    target.fill = alert_fill

            memo.apply(cell, (number_format, alert), setter)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(records) + 1}'


RESUMEN = ['TRIBUNAL', 'FAE ESPERA', 'FAE CUMPLIMIENTO', 'RESIDENCIA ESPERA', 'RESIDENCIA CUMPLIMIENTO',
           'TOTAL REGISTROS', 'ALERTAS', 'REVISAR BITÁCORA', 'SIN FECHA OBS.', 'CONSULTAS SIN BITÁCORAS']
SIN_TRIBUNAL = 'SIN TRIBUNAL RECONOCIDO'


def tribunal_de(row):
    return row.tribunal or meta.tribunal_name(row.values.get('TRIBUNAL')) or SIN_TRIBUNAL


def resumen_tribunales(rows, faltan=(), tribunales=None):
    """Conteos por tribunal: registros por lista, alertas, filas por revisar y consultas sin bitácoras.

    `faltan` son pares (etiqueta, tribunal). Se listan los tribunales del filtro o los que tienen datos."""
    counts = defaultdict(Counter)
    for status, group in rows.items():
        for r in group:
            c = counts[tribunal_de(r)]
            c[f'{(r.category or "").upper()} {status.upper()}'] += 1
            c['TOTAL REGISTROS'] += 1
            c['ALERTAS'] += bool(r.values['ALERTA'])
            c['REVISAR BITÁCORA'] += not revisada(r.values[EXTRA[1]])
            c['SIN FECHA OBS.'] += r.values.get('FECHA OBS.') is None
    for _, tribunal in faltan:
        counts[tribunal or SIN_TRIBUNAL]['CONSULTAS SIN BITÁCORAS'] += 1
    order = [t for t in meta.TRIBUNALS if (tribunales is None and t in counts) or (tribunales and t in tribunales)]
    order += [t for t in counts if t not in meta.TRIBUNALS]
    result = [dict({h: counts[t][h] for h in RESUMEN[1:]}, TRIBUNAL=t) for t in order]
    total = Counter()
    for item in result:
        total.update({h: item[h] for h in RESUMEN[1:]})
    result.append(dict({h: total[h] for h in RESUMEN[1:]}, TRIBUNAL='TOTAL'))
    return result


def consolidado(path, rows, resumen=None):
    wb = Workbook()
    wb.remove(wb.active)
    if resumen is not None:
        ws = wb.create_sheet('Resumen por tribunal')
        tabla(ws, RESUMEN, resumen)
        ws.column_dimensions['A'].width = 30
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for status in ('Espera', 'Cumplimiento'):
        tabla(wb.create_sheet(status), meta.SOURCE_HEADERS[status] + EXTRA, [r.values for r in rows[status]])
    groups = defaultdict(list)
    for status, group in rows.items():
        for r in group:
            key = (_rut(r.values.get('RUT')), normal(r.values.get('NOMBRE')))
            if all(key):
                groups[key].append(dict(r.values, UBICACIÓN=status))
    duplicates = [r for group in groups.values() if len(group) > 1 for r in group]
    titles = ['UBICACIÓN'] + meta.ESPERA_HEADERS + meta.CUMPLIMIENTO_HEADERS + EXTRA
    canonical = {}
    for h in titles:
        canonical.setdefault(normal(h), h)
    headers = list(canonical.values())

    def unified(records):
        return [{canonical.get(normal(k), k): v for k, v in record.items()} for record in records]

    tabla(wb.create_sheet('Duplicados'), headers, unified(duplicates))
    alert_rows = [dict(r.values, UBICACIÓN=status) for status, group in rows.items() for r in group if requiere_revision(r.values)]
    tabla(wb.create_sheet('Alertas y revisión'), headers, unified(alert_rows))
    meta.save_workbook_atomic(wb, path)


def marcar_planilla(wb, records, rows, tribunales=None):
    # Mismo orden por modalidad y tribunal que read_registros/rebuild; no se coteja solo por RIT.
    flagged = defaultdict(list)
    for status, group in rows.items():
        for r in group:
            sheet = f'{r.category} {status}'
            if sheet in records and r.tribunal in records[sheet]:
                flagged[(sheet, r.tribunal)].append(r)
    fills = {color: PatternFill('solid', fgColor=color) for color in (ROJO, AMARILLO)}
    for name in meta.SHEET_DEFS:
        ws = wb[name]
        pos = 1
        for court in [t for t in meta.TRIBUNALS if tribunales is None or t in tribunales]:
            pos += 4
            source = flagged[(name, court)]
            written = records[name][court]
            # Control de alineación: la marca solo se aplica si cada fila coincide con su RIT escrito.
            if len(source) != len(written) or any(
                normal(r.values.get('RIT')) != normal(rec[0]) for r, rec in zip(source, written)
            ):
                raise ValueError(f'No se pudo alinear la marca de bitácoras con la planilla ({name}, {court}).')
            if written:
                pos += 1
                for r in source:
                    if requiere_revision(r.values):
                        alert = r.values['ALERTA']
                        ws.cell(pos, 8).comment = Comment(alert or r.values[EXTRA[1]], 'META 4')
                        fill = fills[ROJO if alert else AMARILLO]
                        for cell in ws[pos][:8]:
                            cell.fill = fill
                    pos += 1
            else:
                pos += 1
            pos += 1


def exportar_bitacoras(path, diaries):
    wb = Workbook()
    wb.remove(wb.active)
    observations, controls = [], []
    for d in diaries:
        common = {'TRIBUNAL': d.get('tribunal_nombre'), 'RIT': d.get('rit'), 'RUT': d.get('rut'),
                  'NOMBRE': d.get('nombre'), 'PROGRAMA': d.get('programa'),
                  'CAUSA ID': d.get('causa_id'), 'INGRESO ID': d.get('ingreso_id')}
        controls.append(dict(common, LECTURA=d.get('estado'), CONTROL=d.get('control')))
        for e in d['entries']:
            observations.append(dict(common, **{'FECHA CENTRO': fecha(e['fecha']), 'ETAPA': e.get('etapa'),
                                               'TIPO OBSERVACIÓN': e.get('tipo'), 'OBSERVACIÓN CENTRO': e.get('texto')}))
    headers = ['TRIBUNAL', 'RIT', 'RUT', 'NOMBRE', 'PROGRAMA', 'CAUSA ID', 'INGRESO ID']
    tabla(wb.create_sheet('Observaciones del mes'),
          headers + ['FECHA CENTRO', 'ETAPA', 'TIPO OBSERVACIÓN', 'OBSERVACIÓN CENTRO'], observations,
          {'FECHA CENTRO': 'dd-mm-yyyy hh:mm:ss'})
    tabla(wb.create_sheet('Control de lecturas'), headers + ['LECTURA', 'CONTROL'], controls)
    meta.save_workbook_atomic(wb, path)


def _celda(value):
    """Mismo valor que entregaba pandas al releer el XLSX intermedio, salvo textos con '='."""
    if value is None or (isinstance(value, str) and value == ''):
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, datetime.min.time())
    return value


def hoja_desde_grilla(grid):
    """DataFrame de la grilla validada por el descargador, sin escribir un XLSX temporal."""
    rows = []
    for row in grid:
        values = [_celda(v) for v in row]
        if 'RIT' in [normal(v) for v in values]:
            values = [ALIAS_COLUMNAS.get(normal(v), v) for v in values]
        while values and values[-1] is None:
            values.pop()
        rows.append(values)
    while rows and not rows[-1]:
        rows.pop()
    return pd.DataFrame(rows, dtype=object) if rows else pd.DataFrame()


def _agrupar_paginas(files):
    """Las páginas de una misma consulta forman un solo candidato; no son duplicados."""
    groups = defaultdict(list)
    for f in files:
        for c in f.candidates:
            groups[c.key].append(c)
    selected, merged_groups = {}, {}
    for key, group in groups.items():
        for c in group:
            c.selected = True
        base = next((c for c in group if c.kind == 'EXCEL'), group[0])
        pages = {id(c.file_result): c.file_result for c in group}.values()
        view = replace(base.file_result, orphan_rows=[r for page in pages for r in page.orphan_rows])
        merged = replace(base, file_result=view, rows=[r for c in group for r in c.rows])
        selected[key] = merged
        merged_groups[key] = [merged]
    return selected, merged_groups


def fuentes(origen):
    return _fuentes(origen)[:4]


def _lote(path, label, sources, bindings, consultas, filtros, vistas):
    """Valida un lote del Descargador y agrega sus páginas verificadas."""
    summary = json.loads(path.read_text(encoding='utf-8'))
    if summary.get('estado') != 'VALIDADA':
        raise ValueError(f'La descarga de registros está incompleta ({label}). Completa el lote antes de generar Meta 4.')
    # Un lote filtrado se acepta, pero puede dejar registros fuera de Meta 4: se informa siempre.
    rango = summary.get('filtro_fecha')
    if rango:
        desde = rango.get('desde', '?') if isinstance(rango, dict) else '?'
        hasta = rango.get('hasta', '?') if isinstance(rango, dict) else '?'
        filtros.append(f'Filtro de fechas {desde} a {hasta}: el lote incluye solo los ingresos de ese rango. [{label}]')
    if summary.get('otros_filtros', 'todos los registros') != 'todos los registros':
        filtros.append(f'Filtros del formulario de SITFA conservados: el lote puede no incluir todos los registros. [{label}]')
    for query in summary['consultas']:
        seleccion = query['seleccion']
        if seleccion.get('tab') not in ('Espera', 'Cumplimiento') or seleccion.get('modality') not in ('1', '3'):
            continue
        if seleccion.get('tab') == 'Cumplimiento' and seleccion.get('report') not in (None, '', '0'):
            filtros.append(f"Filtro de vencimiento en {query.get('consulta') or 'Cumplimiento'}: "
                           f'incluye solo esas medidas. [{label}]')
        # La misma consulta en dos lotes duplicaría sus registros: se exige elegir uno.
        vista = (str(seleccion.get('tribunal') or meta.tribunal_name(seleccion.get('tribunal_nombre')) or ''),
                 seleccion['tab'], seleccion['modality'])
        if vista[0] and vista in vistas and vistas[vista] != label:
            raise ValueError(f"La consulta {query.get('consulta') or ' / '.join(vista)} está en dos lotes "
                             f'({vistas[vista]} y {label}). Quita uno de los dos para no duplicar registros.')
        vistas[vista] = label
        consultas.append(query)
        manifest = dentro(path.parent, query['verificacion'])
        check = json.loads(manifest.read_text(encoding='utf-8'))
        if check.get('estado') not in ('VALIDADA', 'SIN_RESULTADOS', 'COMPLETA', 'COMPLETO'):
            raise ValueError(f'Una consulta del descargador no está validada ({label}).')
        if check.get('paginas_verificadas') != check.get('paginas_esperadas'):
            raise ValueError(f'Faltan páginas del descargador ({label}).')
        if not check['archivos'] and check.get('evidencia'):
            sources.append((dentro(path.parent, check['evidencia']['archivo']), None, label))
        for page in check['archivos']:
            source = dentro(path.parent, page['archivo'])
            body = source.read_bytes()
            if hashlib.sha256(body).hexdigest() != page['sha256']:
                raise ValueError(f'Un Excel del descargador cambió después de su verificación ({label}).')
            sources.append((source, body, label))
            bindings[source] = page.get('vinculos_ingreso', [])


def _fuentes(origen):
    """Normaliza XLS BIFF, XML, HTML tabular y OOXML sin modificar originales.

    `origen` es una carpeta o una lista de carpetas. Cada `resumen.json` encontrado es un
    lote del Descargador (por ejemplo, FAE Espera, FAE Cumplimiento, Residencia Espera y
    Residencia Cumplimiento por separado) y todos se combinan en un solo Meta 4.
    El manifiesto de cada lote identifica TODAS las páginas de una consulta.
    Devuelve además las consultas de Meta 4 incluidas, los filtros y los lotes leídos.
    """
    roots = raices(origen)
    base = base_comun(roots)
    summaries = buscar_archivos(roots, 'resumen.json')
    sources = []
    bindings = {}
    consultas = []
    filtros = []
    lotes = []
    if summaries:
        sin_lote = [r for r in roots if not any(s.is_relative_to(r) for s in summaries)]
        if sin_lote:
            raise ValueError('Esta carpeta no contiene un lote del Descargador (resumen.json): ' + ', '.join(map(str, sin_lote)))
        vistas = {}
        for path in summaries:
            label = etiqueta_lote(path, base)
            lotes.append(label)
            _lote(path, label, sources, bindings, consultas, filtros, vistas)
    elif len(roots) > 1:
        raise ValueError('Varias carpetas solo se combinan si son lotes del Descargador (con resumen.json). '
                         'Para reportes Excel sueltos elige una sola carpeta.')
    else:
        origen = roots[0]
        sources = [(p, None, '') for p in sorted(origen.rglob('*'))
                   if p.suffix.lower() in meta.EXCEL_SUFFIXES | meta.PDF_SUFFIXES and p.is_file()
                   and not p.is_symlink() and not meta.output_file_reason(p)
                   and not p.name.startswith(('Bitacoras_', 'Registros consolidados')) and not es_salida(p, origen)]
    origen = base
    lote_de = {}
    files, issues = [], []
    for path, body, label in sources:
        if path.suffix.lower() in meta.PDF_SUFFIXES:
            files.append(meta.inspect_zero_record_file(path, origen, issues))
        else:
            try:
                grid = lector_excel.read_excel(body if body is not None else path.read_bytes()).grid
            except lector_excel.PocError:
                files.append(meta.inspect_excel(path, origen, path, issues))
            else:
                files.append(meta.inspect_excel(path, origen, None, issues, frames=[('Sheet', hoja_desde_grilla(grid))]))
        lote_de[id(files[-1])] = label
    broken = [f for f in files if f.result == 'ERROR']
    if broken:
        raise ValueError('No se generó una planilla parcial: hay reportes que no se pudieron leer. Revisa: ' + ', '.join(f.relative_path for f in broken))
    for f in files:
        index = defaultdict(set)
        for b in bindings.get(f.path, []):
            identity = (normal(b.get('rit')), _rut(b.get('rut')), normal(b.get('nombre')), normal(b.get('programa')))
            index[identity].add(tuple(str(b.get(k, '')) for k in IDENTIDAD))
        if not index:
            continue
        for candidate in f.candidates:
            for row in candidate.rows:
                v = row.values
                remotes = index.get((normal(v.get('RIT')), _rut(v.get('RUT')), normal(v.get('NOMBRE')), normal(v.get('DERIVACIÓN'))), ())
                if len(remotes) == 1 and all(next(iter(remotes))):
                    v['_REMOTE'] = next(iter(remotes))
    if summaries:
        origenes = defaultdict(set)
        for f in files:
            for c in f.candidates:
                origenes[c.key].add(lote_de[id(f)])
        for key, labels in origenes.items():
            if len(labels) > 1:
                raise ValueError(f'{meta.combination_text(key)} aparece en varios lotes ({", ".join(sorted(labels))}). '
                                 'Quita uno de los lotes para no duplicar registros.')
        selected, groups = _agrupar_paginas(files)
    else:
        selected, groups = meta.select_candidates(files, issues)
    return files, issues, selected, groups, consultas, filtros, lotes


def sin_bitacoras(consultas, lectura):
    """Consultas con registros que ningún lote de bitácoras enumeró: pares (etiqueta, tribunal)."""
    if not lectura.declara_alcance:
        return []
    faltan = []
    for query in consultas:
        seleccion = query['seleccion']
        if query.get('estado') == 'VALIDADA' and _consulta(seleccion) not in lectura.cubiertas:
            label = ' / '.join(str(seleccion.get(k) or '') for k in ('tribunal_nombre', 'tab', 'modalidad_nombre'))
            faltan.append((label, meta.tribunal_name(seleccion.get('tribunal_nombre'))))
    return faltan


def tribunales_disponibles(origen):
    """Tribunales de Meta 4 consultados en los lotes de registros (vacío si no hay resumen.json)."""
    found = set()
    for path in buscar_archivos(raices(origen), 'resumen.json'):
        try:
            summary = json.loads(path.read_text(encoding='utf-8'))
            for query in summary.get('consultas', []):
                name = meta.tribunal_name((query.get('seleccion') or {}).get('tribunal_nombre'))
                if name:
                    found.add(name)
        except (OSError, ValueError, AttributeError):
            continue
    return [t for t in meta.TRIBUNALS if t in found]


def normalizar_tribunales(values):
    """Nombres canónicos de Meta 4; None significa todos."""
    if not values:
        return None
    result = []
    for value in values:
        name = meta.tribunal_name(value)
        if not name:
            raise ValueError(f'Tribunal no reconocido en el alcance de Meta 4: {value}.')
        if name not in result:
            result.append(name)
    return [t for t in meta.TRIBUNALS if t in result]


def aplicar_filtro(tribunales, files, issues, selected, groups, faltan):
    """Deja solo las combinaciones de los tribunales elegidos; el resto queda NO INCLUIDO."""
    if tribunales is None:
        return selected, groups, faltan
    for f in files:
        for c in f.candidates:
            if c.tribunal not in tribunales:
                c.selected = False
                c.selection_note = 'Fuera del filtro de tribunales'
    issues[:] = [i for i in issues if not (i.level == 'FALTANTE' and i.combination.split(' | ')[0] not in tribunales)]
    selected = {k: v for k, v in selected.items() if k[0] in tribunales}
    groups = {k: v for k, v in groups.items() if k[0] in tribunales}
    return selected, groups, [item for item in faltan if item[1] in tribunales]


def lineas_control(inicio, fin, estados, faltan, tribunales, filtros, lotes, extra=()):
    lines = [['Período de bitácoras', inicio, fin],
             ['Lectura de bitácoras', ', '.join(sorted(set(estados)))],
             ['Consultas sin bitácoras', len(faltan)],
             ['Tribunales incluidos', 'Todos' if tribunales is None else ', '.join(tribunales)],
             ['Filtros del lote de registros', ' | '.join(filtros) if filtros else 'Sin filtros'],
             ['Lotes de registros', len(lotes), ' | '.join(lotes)]]
    lines += [list(e) for e in extra]
    lines += [['Rojo: última observación en Espera y registro en Cumplimiento. Amarillo: revisar control de bitácora.'],
              ['Sin observación en el mes: FECHA OBS. vacía. Informes no participa.']]
    return lines


def _publicar(destino, registros, rows, corte, emission, template, tribunales, faltan, control_lines):
    """META 4, Registros consolidados y verificación final a partir de un Registros.xlsx ya escrito."""
    records, counts, unmatched = meta.read_registros(registros)
    wb = meta.load_template(template)
    summary = meta.rebuild(wb, records, emission or corte, tribunales)
    marcar_planilla(wb, records, rows, tribunales)
    control = wb.create_sheet('Control bitácoras')
    for line in control_lines:
        control.append(line)
    for cell in control[1][1:]:
        cell.number_format = 'dd-mm-yyyy'
    parcial = '' if tribunales is None else ' - parcial'
    out = destino / f'META 4 - actualizado {(emission or corte):%Y-%m}{parcial}.xlsx'
    meta.save_workbook_atomic(wb, out)
    resumen = resumen_tribunales(rows, faltan, tribunales)
    consolidado(destino / 'Registros consolidados.xlsx', rows, resumen)
    totals, counted, problems = meta.verify_meta_output(out)
    problems = list(problems)
    if totals != summary or counted != summary:
        problems.append('Los totales escritos no coinciden con los registros generados.')
    if sum(summary.values()) != sum(len(v) for v in rows.values()) - len(unmatched):
        problems.append('Faltan registros en la planilla final.')
    if problems:
        raise ValueError('La planilla requiere revisión: ' + '; '.join(problems))
    return summary, unmatched, resumen


def _conteos(rows):
    every = [r for group in rows.values() for r in group]
    return {
        'registros': len(every),
        'alertas': sum(bool(r.values['ALERTA']) for r in every),
        'revision_bitacora': sum(not revisada(r.values[EXTRA[1]]) for r in every),
        'controles_bitacora': dict(Counter(r.values[EXTRA[1]] for r in every)),
    }


def _escribir(destino, corte, emission, template, lectura, faltan, files, issues, selected, groups, rows, tribunales,
              filtros=(), lotes=()):
    estados = lectura.estados
    diaries = lectura.registros
    if tribunales is not None:
        diaries = [d for d in diaries if meta.tribunal_name(d.get('tribunal_nombre')) in tribunales]
    registros = destino / 'Registros.xlsx'
    meta.write_registros_workbook(registros, rows, files, selected, groups, issues, extra_headers=EXTRA,
                                  row_fill=lambda r: ROJO if r.values['ALERTA'] else None, scope=tribunales)
    inicio, fin = periodo(corte)
    summary, unmatched, resumen = _publicar(destino, registros, rows, corte, emission, template, tribunales, faltan,
                                            lineas_control(inicio, fin, estados, faltan, tribunales, filtros, lotes))
    exportar_bitacoras(destino / 'Bitacoras_del_mes.xlsx', diaries)
    report = {
        'version': VERSION,
        'desde': str(inicio),
        'hasta': str(fin),
        'mes_emision': f'{(emission or corte):%Y-%m}',
        **_conteos(rows),
        'lectura_bitacoras': sorted(set(estados)),
        'consultas_sin_bitacoras': [label for label, _ in faltan],
        'tribunales': 'TODOS' if tribunales is None else tribunales,
        'resumen_tribunales': resumen,
        'avisos_bitacoras': lectura.avisos,
        'lotes_registros': list(lotes),
        'filtros_registros': list(filtros),
        'totales_meta4': summary,
        'excluidos_meta4': len(unmatched),
        'incidencias_fuentes': [i.message for i in issues],
    }
    (destino / 'Reporte META4.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


def _en_carpeta_nueva(destino, trabajo_fn):
    """Escribe en una carpeta temporal hermana: si algo falla, no queda una planilla parcial."""
    if destino.exists():
        raise ValueError('La carpeta de salida ya existe; indica una carpeta nueva.')
    trabajo = destino.with_name(destino.name + EN_CURSO + secrets.token_hex(4))
    trabajo.mkdir(parents=True)
    try:
        report = trabajo_fn(trabajo)
        os.replace(trabajo, destino)
    except BaseException:
        shutil.rmtree(trabajo, ignore_errors=True)
        raise
    return report


HOJAS_REGISTROS = {'ESPERA': 'Espera', 'CUMPLIMIENTO': 'Cumplimiento'}


def leer_registros_editados(path):
    """Filas de ESPERA y CUMPLIMIENTO de un Registros.xlsx revisado por el usuario.

    Se conservan los valores escritos (FECHA OBS., ETAPA ÚLTIMA OBS., CONTROL BITÁCORA);
    la ALERTA se recalcula con la etapa, porque depende de ella."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    rows = {'Espera': [], 'Cumplimiento': []}
    for sheet, status in HOJAS_REGISTROS.items():
        name = next((n for n in wb.sheetnames if normal(n) == sheet), None)
        if name is None:
            raise ValueError(f'Registros.xlsx no tiene la hoja {sheet}.')
        expected = meta.SOURCE_HEADERS[status] + EXTRA
        canonical = {normal(h): h for h in expected}
        grid = list(wb[name].iter_rows(values_only=True))
        if not grid:
            raise ValueError(f'La hoja {sheet} de Registros.xlsx está vacía.')
        headers = [canonical.get(normal(h)) for h in grid[0]]
        missing = [h for h in expected if h not in headers]
        if missing:
            raise ValueError(f'En la hoja {sheet} de Registros.xlsx faltan columnas: ' + ', '.join(missing))
        for index, values in enumerate(grid[1:], 2):
            record = {h: v for h, v in zip(headers, values) if h}
            if not str(record.get('RIT') or '').strip():
                continue
            for h, v in record.items():
                if normal(h) in meta.DATE_HEADERS:
                    record[h] = meta.as_date(v)
            if record['FECHA OBS.'] is not None and not isinstance(record['FECHA OBS.'], date):
                raise ValueError(f'FECHA OBS. no reconocida en la hoja {sheet}, fila {index}: '
                                 f"{record['FECHA OBS.']!r}. Usa el formato DD-MM-AAAA.")
            record[EXTRA[1]] = str(record.get(EXTRA[1]) or '').strip()
            record[EXTRA[0]] = str(record.get(EXTRA[0]) or '').strip()
            stage = etapa({'etapa': record[EXTRA[0]]})
            record['ALERTA'] = ALERTA if status == 'Cumplimiento' and stage == 'Espera' else ''
            rows[status].append(meta.SourceRow(status, meta.classify_derivation(record.get('DERIVACIÓN')),
                                               meta.tribunal_name(record.get('TRIBUNAL')), record, name, index))
    return wb, rows


def regenerar(registros, destino, corte=None, emission=None, template=None):
    """Vuelve a generar META 4 y sus consolidados desde un Registros.xlsx revisado.

    Toma el período, los tribunales, los lotes y los avisos del `Reporte META4.json` de la
    misma carpeta; copia Bitacoras_del_mes.xlsx. Nunca modifica la carpeta original."""
    registros = Path(registros).expanduser().resolve()
    destino = Path(destino).expanduser().resolve()
    if not registros.is_file():
        raise ValueError('Elige el archivo Registros.xlsx revisado.')
    template = Path(template).expanduser().resolve() if template else None
    if template and not template.is_file():
        raise ValueError('No existe la plantilla indicada.')
    previo_path = registros.parent / 'Reporte META4.json'
    previo = json.loads(previo_path.read_text(encoding='utf-8')) if previo_path.is_file() else {}
    if corte is None:
        if 'hasta' not in previo:
            raise ValueError('No está el Reporte META4.json junto a Registros.xlsx: indica la fecha de corte.')
        corte = date.fromisoformat(previo['hasta'])
    if emission is None and previo.get('mes_emision'):
        emission = date.fromisoformat(previo['mes_emision'] + '-01')
    tribunales = None if previo.get('tribunales', 'TODOS') == 'TODOS' else normalizar_tribunales(previo['tribunales'])
    faltan = [(label, meta.tribunal_name(label.split(' / ')[0])) for label in previo.get('consultas_sin_bitacoras', [])]
    wb, rows = leer_registros_editados(registros)
    for status, group in rows.items():
        if tribunales is not None:
            fuera = [r for r in group if r.tribunal not in tribunales]
            if fuera:
                raise ValueError(f'Registros.xlsx tiene filas de tribunales fuera del filtro original ({fuera[0].values.get("TRIBUNAL")}).')
    inicio, fin = periodo(corte)
    editado = datetime.fromtimestamp(registros.stat().st_mtime)

    def trabajo(carpeta):
        # Registros.xlsx: las hojas revisadas se reescriben con su ALERTA recalculada; el resto se conserva.
        for sheet, status in HOJAS_REGISTROS.items():
            name = next(n for n in wb.sheetnames if normal(n) == sheet)
            position = wb.sheetnames.index(name)
            wb.remove(wb[name])
            ws = wb.create_sheet(sheet, position)
            meta.style_data_sheet(ws, meta.SOURCE_HEADERS[status] + EXTRA, rows[status],
                                  lambda r: ROJO if r.values['ALERTA'] else None, len(EXTRA))
        wb.active = 0
        nuevo = carpeta / 'Registros.xlsx'
        meta.save_workbook_atomic(wb, nuevo)
        extra = [('Regenerado desde Registros.xlsx revisado', editado.strftime('%d-%m-%Y %H:%M'), str(registros))]
        summary, unmatched, resumen = _publicar(
            carpeta, nuevo, rows, corte, emission, template, tribunales, faltan,
            lineas_control(inicio, fin, previo.get('lectura_bitacoras', ['SIN DATO']), faltan, tribunales,
                           previo.get('filtros_registros', []), previo.get('lotes_registros', []), extra))
        bitacoras = registros.parent / 'Bitacoras_del_mes.xlsx'
        if bitacoras.is_file():
            shutil.copy2(bitacoras, carpeta / 'Bitacoras_del_mes.xlsx')
        report = dict(previo, **{
            'version': VERSION,
            'desde': str(inicio),
            'hasta': str(fin),
            'mes_emision': f'{(emission or corte):%Y-%m}',
            **_conteos(rows),
            'tribunales': 'TODOS' if tribunales is None else tribunales,
            'resumen_tribunales': resumen,
            'totales_meta4': summary,
            'excluidos_meta4': len(unmatched),
            'regenerado_desde': str(registros),
            'regenerado_el': datetime.now().isoformat(timespec='seconds'),
        })
        report.setdefault('lectura_bitacoras', ['SIN DATO'])
        for key in ('consultas_sin_bitacoras', 'avisos_bitacoras', 'lotes_registros', 'filtros_registros'):
            report.setdefault(key, [])
        (carpeta / 'Reporte META4.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        return report

    return _en_carpeta_nueva(destino, trabajo)


def generar(origen, destino, corte, emission=None, template=None, bitacoras=None, tribunales=None):
    """`origen`: una o varias carpetas con lotes de registros del Descargador; todos los lotes
    se combinan en un solo Meta 4. `bitacoras`: una o varias carpetas de lotes de bitácoras;
    si se omite, se buscan dentro de `origen`. `tribunales`: nombres a incluir (por defecto,
    todos); con filtro, la planilla se marca como parcial."""
    tribunales = normalizar_tribunales(tribunales)
    origen, destino = raices(origen), Path(destino).resolve()
    if not origen:
        raise ValueError('Elige al menos una carpeta de registros del Descargador.')
    if not all(p.is_dir() for p in origen):
        raise ValueError('El origen debe ser una carpeta del descargador.')
    raiz_bitacoras = raices(bitacoras) if bitacoras else origen
    if not all(p.is_dir() for p in raiz_bitacoras):
        raise ValueError('La carpeta de bitácoras no existe.')
    if destino.exists():
        raise ValueError('La carpeta de salida ya existe; indica una carpeta nueva.')
    template = Path(template).expanduser().resolve() if template else None
    if template and not template.is_file():
        raise ValueError('No existe la plantilla indicada.')
    lectura = _leer_bitacoras(raiz_bitacoras, corte)
    files, issues, selected, groups, consultas, filtros, lotes = _fuentes(origen)
    for filtro in filtros:
        issues.append(meta.Issue('ADVERTENCIA', '', '', filtro))
    if not any(f.candidates for f in files):
        raise ValueError('No se encontraron reportes de Espera/Cumplimiento reconocidos para FAE o Residencia.')
    selected, groups, faltan = aplicar_filtro(tribunales, files, issues, selected, groups,
                                              sin_bitacoras(consultas, lectura))
    if tribunales is not None and not selected:
        raise ValueError('El lote no contiene registros de los tribunales elegidos.')
    for label, _ in faltan:
        issues.append(meta.Issue('ADVERTENCIA', '', label, 'Ningún lote de bitácoras enumeró esta consulta; '
                                 'sus registros quedan sin FECHA OBS.'))
    for aviso in lectura.avisos:
        issues.append(meta.Issue('ADVERTENCIA', '', '', aviso))
    rows = meta.consolidated_rows(selected)
    for group in rows.values():
        for r in group:
            if meta.tribunal_name(r.values.get('TRIBUNAL')) is None:
                issues.append(meta.Issue('ADVERTENCIA', f'{r.source_path} fila {r.excel_row}', '',
                                         f"Registro excluido de META 4: la columna TRIBUNAL ({r.values.get('TRIBUNAL')!r}) "
                                         'no identifica un tribunal. Queda en Registros.'))
    enriquecer(rows, lectura.registros, corte)
    return _en_carpeta_nueva(destino, lambda trabajo: _escribir(
        trabajo, corte, emission, template, lectura, faltan, files, issues, selected, groups, rows, tribunales,
        filtros, lotes))


def _fecha_arg(text):
    try:
        return date.fromisoformat(text)
    except ValueError:
        raise argparse.ArgumentTypeError('usa el formato AAAA-MM-DD') from None


def _mes_arg(text):
    if not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', text):
        raise argparse.ArgumentTypeError('usa el formato AAAA-MM')
    return date.fromisoformat(text + '-01')


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--folder', action='append',
                   help='Lote de registros del Descargador (resumen.json y Excel); repetir para combinar varios lotes')
    p.add_argument('--diaries', action='append',
                   help='Carpeta con lotes de bitácoras; repetible; por defecto, dentro de --folder')
    p.add_argument('--output-dir', required=True)
    p.add_argument('--cutoff', type=_fecha_arg, help='Fecha de corte AAAA-MM-DD; por defecto, hoy')
    p.add_argument('--month', type=_mes_arg, help='Mes de emisión AAAA-MM; no cambia el período de bitácoras')
    p.add_argument('--template')
    p.add_argument('--tribunal', action='append', help='Incluir solo este tribunal; repetir para varios')
    p.add_argument('--regenerar', metavar='REGISTROS_XLSX',
                   help='Regenera META 4 y consolidados desde un Registros.xlsx revisado (no usa --folder)')
    args = p.parse_args()
    if args.cutoff and args.cutoff > date.today():
        p.exit(1, 'La fecha de corte no puede estar en el futuro.\n')
    if not args.regenerar and not args.folder:
        p.exit(1, 'Indica --folder (uno o varios lotes) o --regenerar Registros.xlsx.\n')
    try:
        if args.regenerar:
            report = regenerar(args.regenerar, args.output_dir, args.cutoff, args.month, args.template)
        else:
            report = generar(args.folder, args.output_dir, args.cutoff or date.today(), args.month, args.template,
                             args.diaries, args.tribunal)
        print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
    except (ValueError, OSError) as exc:
        p.exit(1, str(exc) + '\n')


if __name__ == '__main__':
    main()
