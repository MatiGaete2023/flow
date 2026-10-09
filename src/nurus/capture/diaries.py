"""Lotes de bitácoras y Excel íntegro compartidos; sin Asistente ni red."""
from pathlib import Path
from hashlib import sha256
import json
from datetime import datetime, timezone
from nurus.bitacora_html import read_snapshot, visible_identity_warnings

def _moment(query):
    try:
        value=datetime.fromisoformat(query['captured_at'].replace('Z','+00:00'))
        return value.astimezone(timezone.utc) if value.tzinfo else None
    except (ValueError,KeyError,AttributeError):return None


def _merge_snapshot(old,new):
    if old is None or old['coverage']=='FALLIDA':return new
    if new['coverage']=='FALLIDA':return old
    def content(query):
        return {**{k:query.get(k) for k in ('tribunal_codigo','causa_id','ingreso_id','rit','persona','centro','stage_code','row_metadata')},
                'entries':[{k:v for k,v in entry.items() if k!='fuente'} for entry in query['entries']]}
    identical=old['html_sha256']==new['html_sha256'] or content(old)==content(new)
    a,b=_moment(old),_moment(new)
    if identical:
        chosen=new if a is not None and b is not None and b>a else old
        chosen['captures']=old['captures']+new['captures']
        chosen['warnings']=list(dict.fromkeys(old['warnings']+new['warnings']))
        return chosen
    if a is None or b is None or a==b:
        raise ValueError('Hay copias distintas del mismo ingreso sin fecha inequívoca. Selecciona un solo lote de bitácoras para comparar.')
    return new if b>a else old

def import_lote(manifest):
    manifest=Path(manifest).resolve()
    if manifest.stat().st_size>32*1024*1024:raise ValueError('Lote de bitácoras demasiado grande.')
    raw=manifest.read_bytes();data=json.loads(raw.decode('utf-8'))
    if data.get('tipo')!='BITACORAS_LECTURA' or data.get('version')!=1:
        raise ValueError('El archivo no es un lote de lectura de bitácoras.')
    records=data.get('registros');scopes=data.get('consultas')
    if not isinstance(records,list) or not isinstance(scopes,list):raise ValueError('Falta el control de lecturas del lote.')
    queries={};ledger=[]
    for record in records:
        identity=tuple(record.get(k,'') for k in ('tribunal_codigo','causa_id','ingreso_id'))
        state=record.get('estado','FALLIDA');error=record.get('error','');query=None
        if state=='LEIDA':
            try:
                name=record['archivo'];path=(manifest.parent/name).resolve()
                if path.parent!=manifest.parent or path.suffix!='.html' or path.stat().st_size>8*1024*1024:
                    raise ValueError('Archivo de bitácora fuera de alcance.')
                body=path.read_bytes()
                if sha256(body).hexdigest()!=record['sha256']:raise ValueError('La copia HTML cambió después de su verificación.')
                query=read_snapshot(body,record['params'],source=name,captured_at=record['capturada'])
                if tuple(query[k] for k in ('tribunal_codigo','causa_id','ingreso_id'))!=identity:
                    raise ValueError('La bitácora pertenece a otro ingreso.')
                notes=visible_identity_warnings(query,record['rit'],record['nombre'])
                if notes:
                    query['warnings'].extend(notes)
                    error=' '.join(notes)
            except (ValueError,KeyError,OSError,TypeError) as exc:
                state='FALLIDA';error=str(exc) if isinstance(exc,ValueError) else 'No se pudo abrir una copia verificable.';query=None
        if all(identity):
            if query is None:
                query={'tribunal_codigo':identity[0],'causa_id':identity[1],'ingreso_id':identity[2],
                    'rit':record.get('rit',''),'persona':record.get('nombre',''),'centro':record.get('programa',''),
                    'coverage':'FALLIDA','error':error or 'No se completó la lectura.', 'entries':[],
                    'stage_code':'','captured_at':'','captures':0,'html_sha256':'','warnings':[],'row_metadata':[]}
            queries[identity]=_merge_snapshot(queries.get(identity),query)
        ledger.append([record.get(k,'') for k in ('tribunal_nombre','modalidad_nombre','pestana','pagina','fila','rit','nombre','programa','causa_id','ingreso_id')]+[state,error,record.get('archivo',''),record.get('sha256',''),record.get('capturada',''),record.get('diagnostico',{}).get('codigo_extension','')])
    scope_rows=[[s.get(k,'') for k in ('tribunal_nombre','modalidad_nombre','tab','estado','paginas','registros','error')] for s in scopes]
    expected=data.get('seleccion',{})
    planned={(t,mod,tab) for t in expected.get('tribunals',[]) for mod in expected.get('modalities',[]) for tab in expected.get('tabs',[])}
    attempted={(s.get('tribunal'),s.get('modality'),s.get('tab')) for s in scopes}
    for tribunal,modality,tab in sorted(planned-attempted):scope_rows.append([tribunal,modality,tab,'NO_CONSULTADA',0,0,'La ejecución terminó antes de esta consulta.'])
    extra=[('Lecturas',('Tribunal','Modalidad','Pestaña','Página','Fila','RIT','Persona','Centro','Causa RUS','Ingreso RUS','Estado lectura','Detalle','Archivo','SHA-256','Fecha captura','Código extensión'),ledger),
           ('Consultas',('Tribunal','Modalidad','Pestaña','Estado','Páginas recorridas','Registros enumerados','Detalle'),scope_rows),
           ('Lote',('Dato','Valor'),[['Estado del recorrido',data.get('estado','INCOMPLETA')],['Fecha',data.get('fecha','')],
                ['Error del recorrido',data.get('error','')],
                ['Descargador',json.dumps(data.get('origen',{}),ensure_ascii=False)],
                ['Extensión',json.dumps(data.get('extension',{}),ensure_ascii=False)],
                ['Selección',json.dumps(expected,ensure_ascii=False)],['Historial','Copia de la tabla capturada; no acredita historial remoto completo.']])]
    if data.get('anteriores_no_reenumeradas'):
        rows=[[r.get(k,'') for k in ('tribunal_nombre','modalidad_nombre','pestana','rit','nombre','causa_id','ingreso_id','estado','archivo','sha256')]
              +['No se reenumeró en esta ejecución; no se acredita desaparición del registro.'] for r in data['anteriores_no_reenumeradas']]
        extra.append(('Lecturas anteriores',('Tribunal','Modalidad','Pestaña','RIT','Persona','Causa RUS','Ingreso RUS','Estado previo','Archivo previo','SHA-256 previo','Alcance'),rows))
    return {'queries':list(queries.values()),'sources':[{'file':manifest.name,'sha256':sha256(raw).hexdigest(),
        'requests':len(records),'openings':sum(r[10]=='LEIDA' for r in ledger),'detail':data.get('estado','')}],'extra_sheets':extra}


def diary_paths(paths):
    if isinstance(paths,(str,Path)): paths=[paths]
    return list(dict.fromkeys(Path(p).resolve() for p in paths or []))


def import_lotes(paths, *, identities=None):
    """Reutiliza originales comprobados; nunca empareja por RIT o nombre."""
    paths=diary_paths(paths)
    if not paths: raise ValueError('Añade un control bitacoras.json con sus HTML originales.')
    queries={};sources=[];extra={};excluded=0;duplicates=0
    for path in paths:
        capture=import_lote(path);sources.extend(capture['sources'])
        for title,headers,rows in capture['extra_sheets']:
            if title not in extra:extra[title]=(headers,[])
            extra[title][1].extend(rows)
        for query in capture['queries']:
            query['manifest_path']=str(path)
            key=tuple(query[k] for k in ('tribunal_codigo','causa_id','ingreso_id'))
            if identities is not None and key not in identities:excluded+=1;continue
            old=queries.get(key)
            if old is None:queries[key]=query;continue
            duplicates+=1
            queries[key]=_merge_snapshot(old,query)
    return {'queries':list(queries.values()),'sources':sources,
        'extra_sheets':[(title,headers,rows) for title,(headers,rows) in extra.items()],
        'reuse':{'manifests':len(paths),'matched':len(queries),'excluded':excluded,'duplicates':duplicates}}


def text_parts(value,limit=30000):
    parts=[];piece=[];units=0
    for char in str(value):
        size=len(char.encode('utf-16-le'))//2
        if units+size>limit:parts.append(''.join(piece));piece=[];units=0
        piece.append(char);units+=size
    if piece:parts.append(''.join(piece))
    return parts or ['']


def export_diaries(manifest,destination):
    """Copia de todas las entradas, sin ventana de cuatro meses ni truncamiento."""
    from openpyxl import Workbook
    import os
    import tempfile
    capture=import_lotes(manifest);destination=Path(destination)
    if destination.exists():raise ValueError('El destino ya existe; elige otra copia.')
    destination.parent.mkdir(parents=True,exist_ok=True)
    book=Workbook();book.remove(book.active)
    copy=book.create_sheet('Copia íntegra')
    columns=('tribunal_codigo','causa_id','ingreso_id','entry_id','fecha','autor','tipo','etapa','texto','fecha_respuesta','respuesta','fuente','fila_fuente')
    copy.append(columns)
    parts_sheet=book.create_sheet('Textos íntegros');parts_sheet.append(('Clave','Campo','Parte','Total','SHA256 texto','Texto'))
    next_rows={}
    def append(sheet,values):
        safe=[]
        for value in values:
            if isinstance(value,(dict,list,tuple)):value=json.dumps(value,ensure_ascii=False)
            safe.append(value)
        sheet.append(safe)
        number=next_rows.get(sheet.title,2);next_rows[sheet.title]=number+1
        for column,value in enumerate(safe,1):
            if isinstance(value,str):sheet.cell(number,column).data_type='s'
    sources=book.create_sheet('Fuentes');sources.append(('Archivo','SHA256','Peticiones','Aperturas','Estado'))
    for source in capture['sources']:append(sources,[source[k] for k in ('file','sha256','requests','openings','detail')])
    summary=book.create_sheet('Ingresos');summary.append(('Tribunal','Causa','Ingreso','Alcance','Entradas','Advertencias'))
    for query in capture['queries']:
        append(summary,[query['tribunal_codigo'],query['causa_id'],query['ingreso_id'],query['coverage'],len(query['entries']),'\n'.join(query['warnings'])])
        for entry in query['entries']:
            values=[]
            key='|'.join(str(entry[k]) for k in ('tribunal_codigo','causa_id','ingreso_id','entry_id'))
            for field in columns:
                value=entry.get(field,'');parts=text_parts(value)
                if len(parts)>1:
                    digest=sha256(str(value).encode('utf-8')).hexdigest()
                    for number,part in enumerate(parts,1):append(parts_sheet,[key,field,number,len(parts),digest,part])
                    value=f'Texto íntegro en hoja Textos íntegros; clave {key}; campo {field}; SHA256 {digest}'
                values.append(value)
            append(copy,values)
    for name,headers,rows in capture['extra_sheets']:
        sheet=book.create_sheet(name);sheet.append(headers)
        for row in rows:append(sheet,row)
    for sheet in book:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    fd,temp=tempfile.mkstemp(prefix='.diaries-',suffix='.xlsx',dir=destination.parent);os.close(fd)
    try:book.save(temp);os.link(temp,destination)
    finally:Path(temp).unlink(missing_ok=True);book.close()
    return destination
