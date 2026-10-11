"""Corte vertical para el EXE autónomo; no red, Chrome, Office ni otros binarios."""

from datetime import date
from hashlib import sha256
import json
from pathlib import Path
import zipfile


def require(condition, message):
    if not condition:
        raise ValueError('Verificación captura: ' + message)


def check_pipeline(folder):
    from openpyxl import load_workbook
    from docx import Document
    from nurus.capture import CaptureError, write_capture
    from nurus.capture.fixtures import fixture
    from nurus.capture.requests import read_request
    from .capture_input import import_capture
    from .capture_products import export_requests, export_evidence, institutional_statement_issue
    from .config import defaults, BASE
    from .outputs import prepare_required_drafts
    from .product_state import stamp, stale
    from .resolutions import prepare_projects, generate_projects
    from .roster import generate
    from .work import Work

    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=False)
    metadata, tables = fixture(partial=True)
    # Una sola captura con tres modos y cuatro modalidades, sin inferir historial.
    for index, mode in enumerate(('ESPERA', 'CUMPLIMIENTO', 'INFORMES', 'ESPERA')):
        source = tables['registros'][index]
        raw = json.loads(source['normalized_json'])
        if mode == 'CUMPLIMIENTO':
            raw.update({'DIAS DE CUMPLIMIENTO': '50', 'DIAS PARA EGRESAR': '10', 'FEC.EGRESO PROYECTADO': '2026-10-18'})
        if mode == 'INFORMES':
            raw['FECHA VENCIMIENTO'] = '2026-10-15'
        source.update(record_type=mode, raw_json=json.dumps(raw), normalized_json=json.dumps(raw))
    package = write_capture(folder / 'Captura_CSMP.zip', metadata, tables)
    digest = sha256(package.read_bytes()).hexdigest()
    config = defaults()
    config['contactos'] = {}
    for court in config['correos']['tribunales'].values():
        court['para'] = []
    # El CC fijo heredado tampoco es un destinatario autorizado de esta prueba.
    config['correos']['cc_adicional'] = ''
    work = import_capture(package, config, as_of=date(2026, 10, 8))
    require(len(work.rows) == 4, 'faltan registros.')
    require(all(c['class'] == 'NO_EVALUABLE' for c in work.capture_context.values()), 'se activó contexto no aceptado.')
    # CT-01/CT-02 solo se muestran (D5): ninguna vista genera hallazgos contextuales aplicables.
    require(
        all(
            f['status'] != 'APLICA'
            for c in work.capture_context.values()
            for f in c['findings']
            if f['rule_id'].startswith('CT-')
        ),
        'la vista de cronología generó hallazgos.',
    )
    row = work.rows[0]
    row.review = {'OBSERVACION': 'Edición humana ficticia', 'TT': 'T', 'CC': '1', 'RES': 'PC_IE'}
    work.rows[1].excluded = True
    work.export(folder / 'Observaciones.xlsx', backend='portable', reduced_fidelity=True)
    book = load_workbook(work.output)
    try:
        sheet = book['Trabajo']
        headers = {cell.value: cell.column for cell in sheet[1]}
        require(
            sheet.cell(row.source_row, headers['OBSERVACION']).value == 'Edición humana ficticia',
            'se perdió texto humano en Excel.',
        )
        sheet.cell(row.source_row, headers['OBSERVACION'], 'Edición Excel ficticia')
        book.save(work.output)
    finally:
        book.close()
    require(work.refresh(), 'no se leyó la edición externa.')
    row = work.rows[0]
    require(row.review['OBSERVACION'] == 'Edición Excel ficticia', 'no se recuperó texto Excel.')
    export_evidence(work, folder / 'Cobertura.xlsx')
    book = load_workbook(folder / 'Cobertura.xlsx')
    try:
        states = {r[0]: r[1] for r in book['Cobertura'].iter_rows(min_row=2, values_only=True)}
        require(states['bitacoras'] == 'NO_CONSULTADA', 'se inventó ausencia de bitácora.')
    finally:
        book.close()
    export_requests(work, folder / 'solicitud_antecedentes.json')
    request = read_request(folder / 'solicitud_antecedentes.json')
    require(len(request['requests']) == 8, 'solicitud incompleta o duplicada.')
    require(request['capture_sha256'] == digest, 'solicitud sin vínculo al paquete.')
    drafts = prepare_required_drafts(
        work, modalities='Modalidades ficticias', period='Prueba ficticia', defer_attachments=True
    )
    require(bool(drafts), 'no se prepararon borradores.')
    # Solo descriptores: no guardar en Outlook ni enviar. Contactos siempre vacíos.
    for draft in drafts:
        draft.to = ''
        draft.cc = ''
    descriptors = [
        {
            'kind': d.kind,
            'record_ids': d.record_ids,
            'to': d.to,
            'cc': d.cc,
            'requires_human_review': True,
            'statement_requires_review': bool(institutional_statement_issue(work, d)),
        }
        for d in drafts
    ]
    (folder / 'Borradores.json').write_text(json.dumps(descriptors, ensure_ascii=False, indent=2), encoding='utf-8')
    rosters = []
    for draft in drafts:
        if draft.roster:
            rosters.extend(generate(draft, folder))
    require(bool(rosters), 'no se generó nómina revisable.')
    for path in rosters:
        book = load_workbook(path)
        try:
            require(book.active.max_row > 1, 'nómina vacía.')
        finally:
            book.close()
    projects, errors = prepare_projects(work, [(row.id, 'PC_IE')], BASE / 'plantillas_word')
    require(bool(projects) and not errors, 'no se preparó matriz Word.')
    generate_projects(work, projects, folder / 'Proyectos_revisables.docx')
    doc = Document(folder / 'Proyectos_revisables.docx')
    require('FICT-0-2026' in '\n'.join(p.text for p in doc.paragraphs), 'Word sin causa ficticia.')
    work.save(folder / 'sesion')
    restored = Work.load(folder / 'sesion', config)
    require(
        restored.rows[0].review == row.review and restored.rows[1].excluded, 'recuperación perdió decisiones humanas.'
    )
    require(not restored.activity_reviewed and not restored.signed_activity, 'marcas locales usadas como fuente.')
    require(restored.capture_digest == digest, 'recuperación cambió captura.')
    stamp(restored, drafts[0])
    updated = import_capture(package, config, as_of=date(2026, 10, 9), previous=restored)
    require(
        updated.rows[0].review == row.review and updated.rows[1].excluded, 'reimportación perdió decisiones humanas.'
    )
    require(stale(updated, drafts[0]), 'producto anterior vigente tras cambiar corte.')
    damaged = folder / 'Captura_corrupta.zip'
    damaged.write_bytes(package.read_bytes())
    with zipfile.ZipFile(damaged, 'a') as archive:
        archive.writestr('archivo_no_declarado.txt', 'ficticio')
    try:
        import_capture(damaged, config, previous=updated)
    except CaptureError:
        pass
    else:
        raise ValueError('Verificación captura: se aceptó paquete corrupto.')
    require(updated.capture_digest == digest and updated.rows[0].review == row.review, 'fallo alteró trabajo previo.')
    require(sha256(package.read_bytes()).hexdigest() == digest, 'se modificó original.')
    products = {
        name: sha256((folder / name).read_bytes()).hexdigest()
        for name in (
            'Observaciones.xlsx',
            'Cobertura.xlsx',
            'solicitud_antecedentes.json',
            'Borradores.json',
            'Proyectos_revisables.docx',
        )
    }
    return {
        'ok': True,
        'data_kind': 'SINTETICA',
        'rows': 4,
        'modalities': 4,
        'modes': 3,
        'human_edits_preserved': True,
        'external_excel_edit_recovered': True,
        'session_recovered': True,
        'reimport_preserved': True,
        'stale_products_rejected': True,
        'corrupt_capture_rejected': True,
        'original_preserved': True,
        'production_context_disabled': True,
        'chronology_display_only': True,
        'request_items': len(request['requests']),
        'draft_descriptors': len(drafts),
        'rosters': len(rosters),
        'word_projects': len(projects),
        'product_sha256': products,
        'institutional_reads': 0,
        'institutional_writes': 0,
        'outlook_saved': 0,
    }
