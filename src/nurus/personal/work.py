"""Un trabajo local compartido: analizar, exportar y recuperar cambios humanos."""

from copy import deepcopy
from dataclasses import dataclass, field, asdict
from datetime import date
from hashlib import sha256
from pathlib import Path
import json

from nurus.rus.columns import normalize
from nurus.rus.reader import read_workbook
from nurus.rus.util import as_date, as_int, tribunal, historical_match
from nurus.services.exports import _export_preserved_payload
from .config import atomic_json
from .motor.contexto import current
from .motor.composicion import Incidencias
from .motor.reglas_espera import generar_observacion_espera
from .motor.reglas_cumplimiento import generar_observacion_cumplimiento
from .motor.reglas_informes import generar_observacion_informes


@dataclass
class Row:
    id: str
    source_row: int
    values: dict
    observation: str
    rules: list
    actions: list
    warnings: list
    excluded: bool = False
    review: dict = field(default_factory=dict)
    sync_base: dict | None = None
    sync_local: dict | None = None
    sync_state: str = ''
    decisions: dict = field(default_factory=dict)
    overrides: dict = field(default_factory=dict)
    word_overrides: dict = field(default_factory=dict)
    avisos: list = field(default_factory=list)


def actions_for(events):
    """Selección por identificador de rama, nunca por palabras de la observación."""
    actions = []
    if any(e.endswith(('E05_SOLO_CORREO', 'E05_PROYECTO_Y_CORREO')) for e in events):
        actions.append('programa_espera')
    if 'ESPERA.E05_PROYECTO_Y_CORREO' in events:
        actions.append('PC_IE')
    if any(e.startswith('INFORMES.I01_') for e in events):
        actions.extend(['programa_vencido', 'PC_INFO'])
    # I02 por vencer solo genera correo preventivo; nunca PC_INFO.
    if any(e.startswith('INFORMES.I02_') for e in events):
        actions.append('programa_por_vencer')
    if any(e.startswith(('CUMPLIMIENTO.C04_', 'CUMPLIMIENTO.C05_')) for e in events):
        actions.append('medidas')
    return actions


def _explicit(value):
    if value is None:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    return True


SERVICE_SHEETS = {'Resumen', 'SITFA_ARCHIVOS', 'Vencimientos'}
SOURCE_DATE_FIELDS = frozenset(
    {'nacimiento', 'oido', 'resolucion', 'prox_aud', 'ingreso', 'egreso_proy', 'ficha_ind', 'ficha_fae', 'vencimiento'}
)
SOURCE_NUMBER_FIELDS = frozenset({'espera', 'dias_cumpl', 'dias_egresar'})


def _source_cell_key(field, value):
    """Compara representaciones de la misma celda sin cambiar cómo el motor las lee."""
    from .motor.utilidades import get_date, get_int

    if field in SOURCE_DATE_FIELDS:
        parsed = get_date(value)
        if parsed is not None:
            return ('date', parsed.isoformat())
    elif field in SOURCE_NUMBER_FIELDS:
        parsed = get_int(value)
        if parsed is not None:
            return ('number', parsed)
    return ('text', '' if value is None else str(value).strip())


REDUCED_FIDELITY_NOTICE = (
    'Copia de fidelidad reducida: el XLSX consolidado del paquete no sirve como base (por ejemplo, trae una hoja '
    'por modalidad o un registro repetido). El consolidado se conserva intacto dentro del paquete.'
)


def _consolidated_base(work, notices=None):
    """Libro consolidado del paquete (D6) como base de la copia, o None si no sirve.

    Si el paquete trae consolidado pero no sirve, agrega a notices el aviso de fidelidad reducida.

    Sirve solo si el paquete sigue siendo el importado y cada registro aparece una única vez, por
    (SITFA_SHA256, SITFA_ARCHIVO, SITFA_FILA), en una misma hoja de datos. Ante cualquier duda
    devuelve None y la copia queda portable de fidelidad reducida, como antes.
    """
    from io import BytesIO
    from openpyxl import load_workbook
    from nurus.capture import read_capture

    try:
        capture = read_capture(work.path)
        if capture.digest != work.capture_digest or not capture.consolidated:
            return None
        book = load_workbook(BytesIO(capture.consolidated), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        base = _map_consolidated(work, capture, book)
    except Exception:
        base = None  # libro inesperado: la copia queda portable, como sin consolidado
    finally:
        book.close()
    if base is None and notices is not None:
        notices.append(REDUCED_FIDELITY_NOTICE)
    return base


def _map_consolidated(work, capture, book):
    """(bytes, hoja, filas) si cada registro está una sola vez en una misma hoja de datos; si no, None."""
    from nurus.rus.columns import map_columns, normalize
    from .sync import IDENTITY_FIELDS

    index = {}
    duplicated = set()
    identity_fields = [key for key in IDENTITY_FIELDS if work.mapping.get(key)]
    modes = set(getattr(work, 'capture_modes', {}).values()) or {work.mode}
    for sheet in book.worksheets:
        if sheet.title in SERVICE_SHEETS:
            continue
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not {'SITFA_ARCHIVO', 'SITFA_FILA', 'SITFA_SHA256'} <= set(header):
            continue
        headers = [
            str(value).strip() if value is not None and str(value).strip() else f'__col_{i}'
            for i, value in enumerate(header)
        ]
        if len({normalize(value) for value in headers}) != len(headers):
            continue
        mapping = {}
        for mode in modes:
            mapping.update(map_columns(headers, mode))
        if not all(mapping.get(key) for key in {*identity_fields, 'rit', 'tribunal', 'nombre'}):
            continue
        source_columns = {field: headers.index(column) for field, column in mapping.items() if field in work.mapping}
        file_col, row_col, hash_col = (header.index(name) for name in ('SITFA_ARCHIVO', 'SITFA_FILA', 'SITFA_SHA256'))
        for number, values in enumerate(rows, 2):
            key = (str(values[hash_col]), str(values[file_col]), str(values[row_col]))
            if key in index:
                duplicated.add(key)
            index[key] = (
                sheet.title,
                number,
                {field: _source_cell_key(field, values[col]) for field, col in source_columns.items()},
            )
    positions = {}
    sheets = set()
    for row in work.rows:
        record = work.capture_sources.get(row.id)
        key = (record['source_sha256'], record['original_file'], record['original_row']) if record else None
        if key is None or key not in index or key in duplicated:
            return None
        sheet, number, source_values = index[key]
        # Las modalidades mixtas comparten encabezados, pero cada fila conserva solo sus campos de origen.
        # Las ediciones humanas (review/overrides) se aplican después; no forman parte de esta comparación.
        if any(
            source_values.get(field) != _source_cell_key(field, row.values[column])
            for field, column in work.mapping.items()
            if column in row.values
        ):
            return None
        sheets.add(sheet)
        positions[row.id] = number
    if len(sheets) != 1:
        return None
    return capture.consolidated, sheets.pop(), positions


def _human_reviewed(row):
    """Distingue una edición humana real de columnas técnicas reimportadas vacías."""
    review = row.review or {}
    if 'OBSERVACION' in review and str(review.get('OBSERVACION', '')) != str(row.observation or ''):
        return True
    return any(_explicit(review.get(key)) for key in ('FECHA_OBS', 'TT', 'CC', 'RES'))


def _source_value(row, key, default=''):
    wanted = normalize(key)
    for column, value in row.values.items():
        if normalize(column) == wanted:
            return value
    return default


def _review_value(row, key, fallback=None):
    if key in (row.review or {}):
        return row.review.get(key, '')
    if fallback is not None:
        return fallback
    return _source_value(row, key, '')


def _sync_resolution_warning(row):
    from .resolutions import resolution_review_issue

    prefix = 'RES no reconocido:'
    row.warnings = [item for item in row.warnings if not str(item).startswith(prefix)]
    issue = resolution_review_issue((row.review or {}).get('RES'))
    if issue:
        row.warnings.append(issue)


RECEIPT_JOURNAL = 'recibos.diario.json'


def _discard_journal(directory, work=None):
    """Retira el diario tras un guardado completo. Si no se puede (bloqueo en Windows), queda obsoleto: su base ya no
    coincide con la trabajo.json nueva y load lo ignora, así que no debe hacer fallar un guardado ya escrito.

    Un diario ilegible es evidencia de borradores posiblemente guardados: se decide por el archivo, no por la
    instancia (E3 WP-94 C1), y la instancia que guarda queda bloqueada igual que la que lo cargó.
    """
    try:
        _receipt_journal(directory, None)
    except JournalUnreadable as exc:
        if work is not None:
            work._save_journal_blocked = str(exc)
        return
    try:
        (Path(directory) / RECEIPT_JOURNAL).unlink(missing_ok=True)
    except OSError:
        pass


def journal_blocked_message():
    return (
        'No se pudo leer el diario de recibos de la sesión (recibos.diario.json). Puede haber borradores ya '
        'guardados en Outlook que esta sesión no ve. Revisa la carpeta Borradores; para continuar, corrige el '
        'acceso al archivo o muévelo fuera de la carpeta de la sesión y vuelve a abrirla.'
    )


class JournalUnreadable(Exception):
    """Existe un diario de recibos que no se puede leer o validar: podría contener borradores ya guardados."""


def _receipt_journal(directory, base):
    """Recibos más recientes que trabajo.json, solo si el diario se escribió sobre esa misma versión.

    Ausente u obsoleto (otra base) → None. Existente pero ilegible o con estructura inválida → JournalUnreadable:
    nunca se trata como ausente, porque perder sus recibos permitiría repetir un borrador (revisión externa H2).
    """
    path = Path(directory) / RECEIPT_JOURNAL
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise JournalUnreadable(str(path)) from exc
    receipts = data.get('receipts') if isinstance(data, dict) else None
    journal_base = data.get('base') if isinstance(data, dict) else None
    if (
        not isinstance(data, dict)
        or data.get('version') != 1
        or not isinstance(receipts, dict)
        or not all(isinstance(key, str) and isinstance(value, dict) for key, value in receipts.items())
        or not isinstance(journal_base, str)
        or len(journal_base) != 64
    ):
        raise JournalUnreadable(str(path))  # estructura inválida (E3 WP-94 C5)
    if data.get('base') != base:
        return None
    return data['receipts']


class Work:
    def __init__(self, config):
        self.config = deepcopy(config)
        self.rows = []
        self.path = ''
        self.output = ''
        self.warnings = []
        self.exception = ''
        self.output_hash = ''
        self.receipts = {}
        self.revision = 0
        self.session_version = 2
        self.session = {}
        self.original_hash = ''
        self.original_content = b''
        # Campos conservados para recuperar sesiones antiguas. En el flujo vigente
        # la ausencia del cruce de Cumplimiento es solo una advertencia.
        self.needs_cross = False
        self.cross_missing = False
        self.cross_dates = {}
        self.signed_activity = {}
        self.activity_reviewed = {}
        self.source_scope = 'SOLO_PLANILLA'
        self.capture_changes = None

    def analyze(
        self,
        path,
        mode,
        *,
        sheet=None,
        cross_sheet=None,
        as_of=None,
        batch=None,
        cross_by_record=None,
        historia_by_record=None,
        bitacora_by_record=None,
    ):
        from .operations import progress, checkpoint

        checkpoint()
        progress('Leyendo Excel y verificando su estructura')
        batch = (
            batch if batch is not None else read_workbook(path, mode, sheet_name=sheet, cross_sheet_name=cross_sheet)
        )
        self.path = batch.source_path
        self.mode = str(batch.mode)
        self.as_of = (as_of or date.today()).isoformat()
        self.sheet = batch.primary_sheet
        self.header = batch.header_row
        self.mapping = dict(batch.column_mapping)
        self.source_hash = batch.workbook_sha256
        self.content = batch.source_bytes
        self.warnings = list(batch.warnings)
        from .sitfa import inspect_batch, source_index

        self.compatibility = inspect_batch(batch, date.fromisoformat(self.as_of))
        self.warnings.extend(self.compatibility['mensajes'])
        self.sitfa_sources = source_index(batch.source_bytes)
        self.needs_cross = False
        self.cross_missing = self.mode == 'CUMPLIMIENTO' and not batch.cross_records
        self.rows = []
        index = {}
        key_fields = ('rit', 'rut', 'nombre', 'tribunal', 'programa')
        # Catálogo Asistente G-08: última fila futura para la identidad completa.
        for row in batch.cross_records:
            cm = batch.cross_mapping
            due = as_date(row.values.get(cm.get('vencimiento', '')), batch.excel_epoch)
            if due and due > date.fromisoformat(self.as_of) and all(cm.get(k) for k in key_fields):
                key = tuple(historical_match(row.values.get(cm[k], '')) for k in key_fields)
                if key in index and index[key] != due:
                    self.warnings.append('Cruce duplicado: se usa la última fila según el perfil Asistente.')
                index[key] = due
        # Paquete de captura: el cruce llega por registro (calendarios verificados), no como hoja 2.
        if (
            self.mode == 'CUMPLIMIENTO'
            and cross_by_record is None
            and (not batch.cross_records or not all(batch.cross_mapping.get(k) for k in (*key_fields, 'vencimiento')))
        ):
            self.cross_missing = True
            self.needs_cross = False
            self.warnings.append(
                'Sin cruce utilizable: no se evalúa C-10. El resto del análisis y la exportación continúan sin bloqueo.'
            )
        cols = dict(self.mapping)
        fn = {
            'ESPERA': generar_observacion_espera,
            'CUMPLIMIENTO': generar_observacion_cumplimiento,
            'INFORMES': generar_observacion_informes,
        }[self.mode]
        duplicity, accumulated = self._prepass(batch, cols)
        for number, record in enumerate(batch.records):
            if number % 100 == 0:
                checkpoint()
                progress('Analizando registros', number, len(batch.records))
            values = dict(record.values)
            for name in (
                'nacimiento',
                'resolucion',
                'ingreso',
                'egreso_proy',
                'oido',
                'prox_aud',
                'ficha_ind',
                'ficha_fae',
                'vencimiento',
            ):
                column = cols.get(name)
                if column and isinstance(values.get(column), (int, float)):
                    values[column] = as_date(values[column], batch.excel_epoch)
            historia = (historia_by_record or {}).get(record.record_id)
            corte = date.fromisoformat(self.as_of)
            ctx = {
                'config': self.config,
                'events': [],
                'date': corte,
                'invalid': [],
                'historia': historia.al_corte(corte) if historia is not None else None,
                'bitacora': tuple(o for o in (bitacora_by_record or {}).get(record.record_id, ()) if o.fecha <= corte),
                'marcas': [],
                'duplicidad': duplicity.get(record.record_id),
                'acumulado': accumulated.get(record.record_id),
            }
            token = current.set(ctx)
            inc = Incidencias()
            warnings = []
            try:
                if self.mode == 'CUMPLIMIENTO':
                    end = as_date(values.get(cols.get('egreso_proy', '')), batch.excel_epoch)
                    counts = [as_int(values.get(cols.get(k, ''))) for k in ('dias_cumpl', 'dias_egresar')]
                    if end and end > ctx['date'] and any(n is not None and n < 0 for n in counts):
                        warnings.append(
                            'Días negativos y egreso futuro: verifica la contradicción; no se afirma que la medida esté vencida.'
                        )
                        ctx['invalid'].append('CUMPLIMIENTO.C04_VENCIDA')
                court = tribunal(values.get(cols.get('tribunal', ''), '')) or ''
                if not court:
                    warnings.append(
                        'Tribunal fuera de Laja, Mulchén y Tomé: sin reglas configuradas; requiere revisión manual.'
                    )
                key = tuple(historical_match(values.get(cols.get(k, ''), '')) for k in key_fields)
                args = {'incidencias': inc, 'fila_excel': record.source.row_number}
                if self.mode == 'CUMPLIMIENTO':
                    args['fecha_hoja2'] = (
                        cross_by_record.get(record.record_id) if cross_by_record is not None else index.get(key)
                    )
                observation = fn(values, court, cols, **args) if court else ''
                warnings.extend(i['MOTIVO'] for i in inc._items)
            except (ValueError, TypeError, KeyError) as exc:
                observation = ''
                warnings.append('No se pudo evaluar esta fila: ' + str(exc))
            finally:
                current.reset(token)
            events = [e for e in ctx['events'] if e not in self.config['desactivadas']]
            excluded = 'COMUN.NO_SEGUIMIENTO' in events
            identity_key = tuple(historical_match(values.get(cols.get(k, ''), '')) for k in key_fields)
            self.cross_dates[record.record_id] = (
                cross_by_record.get(record.record_id) if cross_by_record is not None else index.get(identity_key)
            )
            self.rows.append(
                Row(
                    record.record_id,
                    record.source.row_number,
                    dict(record.values),
                    observation,
                    list(dict.fromkeys(events)),
                    [] if excluded else actions_for(events),
                    warnings,
                    excluded,
                )
            )
            self.rows[-1].decisions = {'structured': True, 'mail': {}, 'resolution': 'auto'}
            self.rows[-1].avisos = list(dict.fromkeys(ctx['marcas']))
        progress('Análisis completo', len(batch.records), len(batch.records))
        return self

    def _prepass(self, batch, cols, *, mode=None):
        """Duplicidad de ingreso (Espera/Cumplimiento) e informes vencidos acumulados (Informes), por causa.
        Solo usa datos de RUS: corre igual con Excel o con paquete, haya o no historia (decisión u, 14-oct)."""
        mode = mode or self.mode
        if mode not in ('ESPERA', 'CUMPLIMIENTO', 'INFORMES'):
            return {}, {}
        from .motor.analisis_historia import duplicidades

        def cell(record, field):
            return str(record.values.get(cols.get(field, ''), '') or '').strip()

        if mode != 'INFORMES':
            rows = [
                (r.record_id, cell(r, 'tribunal'), cell(r, 'rit'), cell(r, 'rut'), cell(r, 'programa'))
                for r in batch.records
            ]
            return duplicidades(rows), {}
        today = date.fromisoformat(self.as_of)
        groups = {}
        for r in batch.records:
            due = as_date(r.values.get(cols.get('vencimiento', '')), batch.excel_epoch)
            if not due or due >= today:
                continue
            ident = tuple(historical_match(cell(r, f)) for f in ('tribunal', 'rit', 'rut', 'programa'))
            if ident[1]:
                groups.setdefault(ident, []).append((r.record_id, due))
        accumulated = {}
        for items in groups.values():
            if len(items) > 1:
                oldest = min(due for _, due in items)
                accumulated.update({rid: (len(items), oldest) for rid, _ in items})
        return {}, accumulated

    def document_exception(self, reason):
        """Compatibilidad con sesiones antiguas; ya no es requisito para exportar."""
        if not reason.strip():
            raise ValueError('Indica el motivo de la excepción de cruce.')
        self.exception = reason.strip()

    def export(self, destination, *, backend='native', reduced_fidelity=False):
        """Exporta la propuesta o las ediciones humanas vigentes de Trabajo.

        Si existe al menos una edición humana, las filas realmente revisadas quedan
        REVISADAS y escriben sus campos de revisión. Las restantes se conservan como
        PENDIENTES, sin atribuirles una revisión inexistente.
        """
        if Path(destination).resolve() == Path(self.path).resolve():
            raise ValueError('El origen se conserva intacto. Elige otro nombre para la copia.')
        from .sync import require_current_copy

        require_current_copy(self)
        reviewed = {row.id: _human_reviewed(row) for row in self.rows}
        reviewed_stage = any(reviewed.values())
        package_input = bool(getattr(self, 'capture_digest', ''))
        notices = []
        base = _consolidated_base(self, notices) if package_input else None
        # Aviso visible (manual del Asistente): el consolidado existe pero la copia no parte de él.
        self.export_notice = ' '.join(notices)
        content, source_hash, primary_sheet, header_row = (
            (base[0], sha256(base[0]).hexdigest(), base[1], 1)
            if base
            else (self.content, self.source_hash, self.sheet, self.header)
        )
        if package_input and backend == 'native' and not base:
            backend = 'portable'
            reduced_fidelity = True
        batch = {
            'source_name': Path(self.path).stem + '.xlsx' if package_input else Path(self.path).name,
            'source_path': self.path,
            'source_hash': source_hash,
            'primary_sheet': primary_sheet,
            'header_row': header_row,
            'mode': self.mode,
            'export_stage': 'reviewed' if reviewed_stage else 'proposal',
        }
        from .record_edits import encode

        records = []
        for row in self.rows:
            is_reviewed = reviewed[row.id]
            records.append(
                {
                    'record_id': row.id,
                    'source_sheet': primary_sheet,
                    'source_row': base[2][row.id] if base else row.source_row,
                    'source_hash': source_hash,
                    'decision': 'excluded' if row.excluded else ('approved' if is_reviewed else 'pending'),
                    'evaluation_status': 'blocked' if row.warnings else 'reviewed',
                    'edit_reason': '; '.join(row.warnings),
                    'edited_observation': _review_value(row, 'OBSERVACION', row.observation),
                    'rule_ids_json': json.dumps(row.rules),
                    'decisions_json': encode(row),
                    'locally_validated': bool(is_reviewed and not row.excluded),
                    'rus_recorded': False,
                    'review_date': _review_value(row, 'FECHA_OBS'),
                    'tt_value': _review_value(row, 'TT'),
                    'workload_value': _review_value(row, 'CC'),
                    'resolution_value': _review_value(row, 'RES'),
                }
            )
        result = _export_preserved_payload(
            content,
            {'batch': batch, 'records': records, 'exceptions': [self.exception] if self.exception else []},
            destination,
            backend=backend,
            allow_reduced_fidelity=reduced_fidelity,
        )
        saved = (
            self.output,
            self.output_hash,
            self.content,
            self.source_hash,
            self.sheet,
            self.header,
            [row.source_row for row in self.rows],
        )
        from .sync import remember_export

        try:
            # Todo lo que puede fallar (lectura del archivo bloqueado, registro) queda dentro del try.
            self.output = str(result.path)
            self.output_hash = sha256(Path(self.output).read_bytes()).hexdigest()
            if base:
                # La copia parte del consolidado: el trabajo pasa a apuntar a su hoja y a sus filas, como un Excel tradicional.
                self.content, self.source_hash, self.sheet, self.header = (
                    content,
                    source_hash,
                    primary_sheet,
                    header_row,
                )
                for row in self.rows:
                    row.source_row = base[2][row.id]
            remember_export(self)
            from .review_metrics import record_export

            record_export(self)
        except Exception:
            (self.output, self.output_hash, self.content, self.source_hash, self.sheet, self.header, rows) = saved
            for row, number in zip(self.rows, rows):
                row.source_row = number
            Path(result.path).unlink(missing_ok=True)
            raise
        return self.output

    def refresh(self, resolutions=None):
        from .sync import refresh

        return refresh(self, resolutions)

    @classmethod
    def external(cls, path, config, mode='ESPERA', *, sheet=None):
        """Importa constancias tal como están, sin exigir un paso previo por el motor."""
        from .importing import read_external
        from .motor.utilidades import es_derivacion_sin_seg

        data = read_external(path, mode, sheet)
        obj = cls(config)
        obj.path = data['path']
        obj.output = obj.path
        obj.mode = data['mode']
        obj.as_of = date.today().isoformat()
        obj.sheet = data['sheet']
        obj.header = data['header']
        obj.mapping = data['mapping']
        obj.content = data['content']
        obj.source_hash = data['digest']
        obj.output_hash = data['digest']
        obj.needs_cross = False
        obj.cross_missing = False
        obj.external_input = True
        obj.warnings = []
        seen = {}
        used_ids = set()
        from .sitfa import source_index

        obj.sitfa_sources = source_index(data['content'])
        for number, values, review, rules in data['records']:
            identity = '|'.join(
                historical_match(values.get(obj.mapping.get(k, ''), ''))
                for k in ('rit', 'rut', 'nombre', 'tribunal', 'programa')
            )
            ordinal = seen.get(identity, 0)
            seen[identity] = ordinal + 1
            rid = (
                str(values.get('NURUS_ID_REGISTRO', '')).strip()
                or sha256((obj.sheet + '|' + identity + '|' + str(ordinal)).encode()).hexdigest()
            )
            if rid in used_ids:
                raise ValueError('El identificador de registro está duplicado en la planilla.')
            used_ids.add(rid)
            excluded = es_derivacion_sin_seg(str(values.get(obj.mapping.get('programa', ''), '')))
            events = list(dict.fromkeys(rules))
            actions = [] if excluded else actions_for(events)
            row = Row(rid, number, values, str(review.get('OBSERVACION', '')), events, actions, [], excluded, review)
            from .record_edits import restore

            restore(row, values.get('NURUS_DECISIONES', ''))
            row.sync_state = str(values.get('NURUS_DECISIONES', '') or '')
            row.sync_base = deepcopy(review)
            row.sync_local = deepcopy(review)
            _sync_resolution_warning(row)
            obj.rows.append(row)
        return obj

    def save_receipts(self):
        """Diario de recibos (EF-03): persiste solo los recibos, ligado a la trabajo.json vigente.

        Sustituye al guardado completo antes y después de cada efecto de Outlook dentro de un lote.
        Load solo lo aplica si `base` coincide con la trabajo.json en disco; save() completo lo retira.
        """
        directory = Path(self.storage_directory)
        if getattr(self, '_save_journal_blocked', ''):
            raise ValueError(journal_blocked_message())
        main = directory / 'trabajo.json'
        if not main.is_file():
            self.save(directory)
            return
        base = sha256(main.read_bytes()).hexdigest()
        atomic_json(directory / RECEIPT_JOURNAL, {'version': 1, 'base': base, 'receipts': self.receipts})

    def save(self, directory, *, snapshot=False):
        """snapshot=True guarda una copia de respaldo sin cambiar la sesión principal (storage_directory),
        donde los recibos de borradores y proyectos siguen persistiéndose."""
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        if not snapshot:
            self.storage_directory = str(directory)
        source = directory / (self.source_hash + '.bin')
        from .session import archive_bytes

        archive_bytes(source, self.content, self.source_hash)
        if self.original_hash and self.original_content:
            original = directory / (self.original_hash + '.bin')
            archive_bytes(original, self.original_content, self.original_hash)
        # La configuración (contactos incluidos) no se copia a la sesión: solo su huella.
        data = {
            k: v
            for k, v in self.__dict__.items()
            if k not in ('config', 'rows', 'content', 'original_content', 'activity_reviewed', 'signed_activity')
            and not k.startswith('_save_')
        }
        data['config_sha256'] = sha256(
            json.dumps(self.config, sort_keys=True, default=str, ensure_ascii=False).encode()
        ).hexdigest()
        data['rows'] = [asdict(r) for r in self.rows]
        previous = directory / 'trabajo.json'
        digest = sha256(json.dumps(data, sort_keys=True, default=str, ensure_ascii=False).encode()).hexdigest()
        saved = getattr(self, '_save_digests', {})
        stats = getattr(self, '_save_stats', {})
        key = str(directory.resolve())
        if previous.is_file():
            info = previous.stat()
            if saved.get(key) == digest and stats.get(key) == (info.st_mtime_ns, info.st_size):
                _discard_journal(directory, self)  # la principal ya contiene estos recibos (E3 WP-92 #2)
                return
        if previous.exists():
            try:
                old = json.loads(previous.read_text(encoding='utf-8'))
            except (ValueError, OSError):
                pass
            else:
                if isinstance(old, dict):
                    atomic_json(directory / 'trabajo.json.bak', old)
        atomic_json(previous, data)
        # La trabajo.json recién escrita ya contiene los recibos: el diario queda obsoleto.
        _discard_journal(directory, self)
        # Conserva también todas las copias que necesita la recuperación anterior.
        keep = {self.source_hash + '.bin'}
        if self.original_hash:
            keep.add(self.original_hash + '.bin')
        backup = directory / 'trabajo.json.bak'
        cleanup_safe = True
        if backup.is_file():
            try:
                old = json.loads(backup.read_text(encoding='utf-8'))
                if not isinstance(old, dict):
                    raise ValueError('Copia de sesión sin objeto JSON')
                for field in ('source_hash', 'original_hash'):
                    if old.get(field):
                        if not isinstance(old[field], str):
                            raise ValueError('Huella de copia inválida')
                        keep.add(old[field] + '.bin')
            except (ValueError, OSError):
                cleanup_safe = False
        for stale in directory.glob('*.bin') if cleanup_safe else ():
            if stale.name not in keep and not stale.is_symlink():
                stale.unlink(missing_ok=True)
        saved[str(directory.resolve())] = digest
        self._save_digests = saved
        info = previous.stat()
        stats[key] = (info.st_mtime_ns, info.st_size)
        self._save_stats = stats

    @classmethod
    def load(cls, directory, config=None):
        """Sesiones nuevas usan la configuración vigente; las antiguas conservan la que traían."""
        directory = Path(directory)
        raw = (directory / 'trabajo.json').read_bytes()
        data = json.loads(raw.decode('utf-8'))
        try:
            journal = _receipt_journal(directory, sha256(raw).hexdigest())
        except JournalUnreadable as exc:
            journal = None
            blocked = str(exc)
        else:
            blocked = ''
        if journal is not None:
            data['receipts'] = journal
        if 'config' not in data:
            from .config import defaults

            data['config'] = config if config is not None else defaults()
        data.pop('config_sha256', None)
        obj = cls(data['config'])
        obj.__dict__.update(data)
        # Se conserva el diario como evidencia y se bloquean nuevos borradores hasta resolverlo (H2).
        obj._save_journal_blocked = blocked
        obj.activity_reviewed = {}
        obj.signed_activity = {}
        if getattr(obj, 'needs_cross', False):
            obj.cross_missing = True
        obj.needs_cross = False
        if not hasattr(obj, 'cross_missing'):
            obj.cross_missing = False
        obj.rows = [Row(**r) for r in data['rows']]
        for row in obj.rows:
            _sync_resolution_warning(row)
        obj.content = (directory / (obj.source_hash + '.bin')).read_bytes()
        if obj.original_hash:
            obj.original_content = (directory / (obj.original_hash + '.bin')).read_bytes()
            if sha256(obj.original_content).hexdigest() != obj.original_hash:
                raise ValueError('La copia original archivada no coincide con su huella.')
        if sha256(obj.content).hexdigest() != obj.source_hash:
            raise ValueError('La copia de origen guardada no coincide con el trabajo.')
        if hasattr(obj, 'capture_context'):
            from .capture_findings import updated_contexts

            obj.capture_context = updated_contexts(obj, obj.rows)
        return obj
