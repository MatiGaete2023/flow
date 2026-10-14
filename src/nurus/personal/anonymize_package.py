"""Anonimización LOCAL de paquetes CSMP_CAPTURE_V1 (sin red ni originales).

Los datos libres del modo conservador requieren revisión humana. Las tablas, IDs
relacionales, procedencia y XLSX consolidado se reescriben coherentemente; los
documentos y originales binarios se rechazan antes de publicar la copia.
"""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import hmac
import json
from pathlib import Path
import re
import secrets
from tempfile import TemporaryDirectory

from nurus.capture import read_capture, write_capture
from nurus.capture.contract import MAPPING_VERSION, TABLES
from .anonymize import (
    OMITTED,
    Pseudonyms,
    Replacer,
    _identity_value,
    _protected_headers,
    column_kind,
    normalize,
    scrub_observation,
)

# Las claves técnicas no son textos libres y deben mantener sus uniones.
TECHNICAL = {
    'SITFA_CAUSA_RUS': 'cause_id',
    'SITFA_INGRESO_RUS': 'ingreso_id',
    'SITFA_PERSONA_RUS': 'person_id',
    'SITFA_CENTRO_RUS': 'program_id',
    'SITFA_SHA256': 'source_sha256',
    'SITFA_ARCHIVO': 'original_file',
    'SITFA_HOJA': 'original_sheet',
}
FIELDS_ID = {'query_id', 'record_id', 'event_id', 'incident_id', 'calendar_id', 'document_id', 'parent_event_id'}
REMOTE_ID = {'cause_id', 'ingreso_id', 'person_id', 'program_id'}
FREE_FIELDS = {'text', 'response', 'description', 'reference', 'message', 'error'}
CONTROL_FIELDS = {
    'capture_id',
    'source_system',
    'screen',
    'tab',
    'court_code',
    'court_name',
    'modality_code',
    'modality_raw',
    'record_type',
    'identity_status',
    'coverage',
    'event_type',
    'event_at',
    'response_at',
    'captured_at',
    'time_precision',
    'time_zone',
    'scope',
    'origin',
    'origin_verified',
    'status',
    'association',
    'phase',
    'code',
    'severity',
    'recoverable',
    'period_from',
    'period_to',
    'history_complete',
    'institutional_verified',
    'expected_pages',
    'saved_pages',
    'original_row',
    'download_state',
    'content_read',
    'history_complete',
}
_JSON_LIMIT = 30
_SENSITIVE_WORDS = re.compile(r'(?i)(rut|run|nombre|apellido|curador|litigante|domicilio|direccion|telefono|correo)')


class UnsafePackage(ValueError):
    """Contenido no controlado: ningún archivo de salida se debe publicar."""


class Mapper:
    def __init__(self, mode):
        if mode not in ('conservador', 'estricto'):
            raise ValueError('Modo desconocido.')
        self.mode = mode
        self.secret = secrets.token_bytes(32)
        self.pseudonyms = Pseudonyms()
        self.replacer = Replacer(self.pseudonyms)
        self.maps = {}
        self.protected = _protected_headers()
        self.capture_id = 'anon-' + secrets.token_hex(12)

    def token(self, kind, source):
        if source in ('', None):
            return ''
        key = str(source)
        bucket = self.maps.setdefault(kind, {})
        if key not in bucket:
            if kind in ('source_sha256', 'sha256'):
                bucket[key] = hmac.new(self.secret, (kind + ':' + key).encode(), sha256).hexdigest()
            elif kind == 'original_file':
                ext = Path(key).suffix.lower()
                if ext not in ('.xls', '.xlsx', '.html', '.htm', '.csv'):
                    ext = '.xls'
                bucket[key] = f'fuente_{len(bucket) + 1}{ext}'
            elif kind == 'original_sheet':
                bucket[key] = f'Hoja{len(bucket) + 1}'
            else:
                bucket[key] = f'{kind[:3]}_anon_{len(bucket) + 1}'
        return bucket[key]

    def free(self, value):
        if not value:
            return value
        if self.mode == 'estricto':
            return OMITTED
        return scrub_observation(str(value), self.replacer)

    def value(self, key, value):
        """Convierte una celda/entrada JSON; preserva fechas, umbrales y enums."""
        if value is None or value == '':
            return value
        label = str(key)
        name = normalize(label)
        if label in TECHNICAL:
            return self.value(TECHNICAL[label], value)
        if label in ('RIT', 'rit_original') or name == 'rit':
            return self.pseudonyms.get('rit', str(value))
        if label in FIELDS_ID or label in REMOTE_ID:
            return self.token(label, value)
        if label in ('source_sha256', 'sha256', 'original_file', 'original_sheet'):
            return self.token(label, value)
        if label == 'capture_id':
            return self.capture_id
        if label == 'local_file':
            # No hay binarios en un paquete aceptado para anonimización.
            return ''
        if label == 'author':
            if normalize(value) in ('desconocido', 'sin identificar'):
                return value
            self.replacer.add_person(str(value))
            return self.replacer.text(str(value))
        if label in FREE_FIELDS:
            return self.free(value)
        if label in CONTROL_FIELDS or name in self.protected:
            if isinstance(value, str):
                # Se mantienen términos institucionales y columnas del motor,
                # pero jamás se permite copiar identificadores literales.
                return self.replacer.text(value, free=False)
            return value
        kind = column_kind(label)
        if kind in ('persona', 'curador'):
            if isinstance(value, str):
                from .anonymize import curator_name

                self.replacer.add_person(curator_name(value) if kind == 'curador' else value)
                return self.replacer.text(value)
            return self.token('persona', value)
        if kind in ('rut', 'rit', 'mail', 'telefono', 'direccion'):
            return _identity_value(kind, value, self.pseudonyms, self.replacer)
        if kind == 'texto':
            return self.free(value)
        if isinstance(value, bool) or isinstance(value, (int, float)):
            if isinstance(value, (int, float)) and not isinstance(value, bool) and len(str(abs(int(value)))) >= 7:
                return self.token('numero', value)
            return value
        if isinstance(value, str):
            # Solo permiten salir datos ajenos a las columnas reconocidas como
            # texto libre depurado, nunca copiados literalmente.
            return self.free(value)
        raise UnsafePackage('Valor JSON no reconocido.')

    def json_field(self, raw, *, filters=False):
        from nurus.capture.contract import json_object

        obj = json_object(raw)

        def visit(value, key='', depth=0):
            if depth > _JSON_LIMIT:
                raise UnsafePackage('Estructura JSON demasiado profunda.')
            if isinstance(value, dict):
                result = {}
                for label, child in value.items():
                    if not isinstance(label, str) or not label or len(label) > 128:
                        raise UnsafePackage('Clave JSON incompatible.')
                    if _SENSITIVE_WORDS.search(label) and len(label) > 100:
                        raise UnsafePackage('Clave JSON sospechosa.')
                    result[label] = visit(child, label, depth + 1)
                return result
            if isinstance(value, list):
                if len(value) > 10000:
                    raise UnsafePackage('Lista JSON excesiva.')
                return [visit(child, key, depth + 1) for child in value]
            if isinstance(value, (dict, list)):
                raise UnsafePackage('Estructura JSON incompatible.')
            if filters and key in (
                'tribunal',
                'tribunal_nombre',
                'modality',
                'screen',
                'tab',
                'source_system',
                'court_code',
                'capture_tables',
            ):
                # Filtros de operación: no son identidad personal.
                return value
            return self.value(key, value)

        return json.dumps(visit(obj), ensure_ascii=False, separators=(',', ':'))

    def excel_value(self, label, value):
        name = str(label).strip()
        if name in TECHNICAL:
            return self.value(name, value)
        if name == 'SITFA_FILA':
            return value
        return None


def _seed_names(tables, mapper):
    """Sembrar identidades antes de escribir, incluso las de la segunda hoja."""
    for table, rows in tables.items():
        for row in rows:
            for key in ('author',):
                if key in row and row[key] and normalize(row[key]) not in ('desconocido', 'sin identificar'):
                    mapper.replacer.add_person(row[key])
            for field in ('raw_json', 'normalized_json'):
                if field not in row:
                    continue
                from nurus.capture.contract import json_object

                obj = json_object(row[field])
                for key, value in obj.items():
                    if not isinstance(value, str):
                        continue
                    kind = column_kind(key)
                    if kind == 'persona':
                        mapper.replacer.add_person(value)
                    elif kind == 'curador':
                        from .anonymize import curator_name

                        mapper.replacer.add_person(curator_name(value))


def _new_tables(tables, mapper):
    result = {}
    for table, rows in tables.items():
        if table not in TABLES:
            raise UnsafePackage('Tabla desconocida; actualizar el contrato antes de anonimizar.')
        result[table] = []
        for row in rows:
            if set(row) != set(TABLES[table]):
                # El lector acepta columnas futuras opcionales; esta versión no
                # las puede anonimizar con certeza y no debe descartarlas.
                raise UnsafePackage('El paquete contiene columnas nuevas no reconocidas.')
            out = {}
            for key, val in row.items():
                if key in ('raw_json', 'normalized_json', 'filters_json'):
                    out[key] = mapper.json_field(val, filters=key == 'filters_json')
                elif key in (
                    'document_id',
                    'parent_event_id',
                    'record_id',
                    'query_id',
                    'event_id',
                    'incident_id',
                    'calendar_id',
                ):
                    out[key] = mapper.token(key, val)
                else:
                    out[key] = mapper.value(key, val)
                if not isinstance(out[key], str):
                    raise UnsafePackage('Una celda del contrato dejó de ser texto.')
            result[table].append(out)
    return result


def _consolidated(raw, mapper):
    """Reconstituye el XLSX sin propiedades, enlaces, imágenes ni objetos."""
    from .anonymize import anonymize

    with TemporaryDirectory(prefix='csmp-anon-') as folder:
        original = Path(folder) / 'origen.xlsx'
        target = Path(folder) / 'salida.xlsx'
        original.write_bytes(raw)
        anonymize(
            original,
            target,
            text_mode=mapper.mode,
            pseudonyms=mapper.pseudonyms,
            replacer=mapper.replacer,
            field_transform=mapper.excel_value,
        )
        return target.read_bytes()


def _reserve_original_rits(capture, mapper):
    """Reserva todos los RIT originales para evitar colisiones con causas ficticias."""
    from io import BytesIO

    from openpyxl import load_workbook
    from .anonymize import RIT_FREE_RE

    def collect(value):
        if not isinstance(value, str):
            return
        for match in RIT_FREE_RE.finditer(value):
            mapper.pseudonyms.reserved_rits.add(
                f"{(match.group(1) or 'X').upper()}-{int(match.group(2))}-{match.group(3)}"
            )

    for rows in capture.tables.values():
        for row in rows:
            for value in row.values():
                collect(value)
    if capture.consolidated:
        book = load_workbook(BytesIO(capture.consolidated), data_only=True, read_only=True, keep_links=False)
        try:
            for sheet in book.worksheets:
                collect(sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    for value in row:
                        collect(value)
        finally:
            book.close()


def anonymize_capture(source, target, *, mode='conservador'):
    """Publica copia ZIP validada, conservando vínculos y sin filtrar binarios."""
    source, target = Path(source), Path(target)
    if source.suffix.lower() != '.zip' or target.suffix.lower() != '.zip':
        raise ValueError('Se requiere ZIP de captura de entrada y salida.')
    if source.resolve() == target.resolve() or target.exists():
        raise FileExistsError('No se permite sobrescribir el origen ni un destino existente.')
    capture = read_capture(source)  # valida hashes, tamaños, filas y todos los miembros
    if any(f['type'] == 'original' for f in capture.manifest['files']):
        raise UnsafePackage(
            'El ZIP incorpora originales/documentos binarios; no se crea copia. '
            'Exporta un paquete sin originales antes de anonimizar.'
        )
    if capture.manifest['mapping_version'] != MAPPING_VERSION:
        raise UnsafePackage('Versión de mapeo no controlada; se rechazó el paquete.')
    if capture.manifest.get('required_features'):
        raise UnsafePackage('El paquete requiere funcionalidades desconocidas.')
    mapper = Mapper(mode)
    _reserve_original_rits(capture, mapper)
    _seed_names(capture.tables, mapper)
    tables = _new_tables(capture.tables, mapper)
    if tables.get('consultas') is None:
        raise UnsafePackage('Falta tabla de consultas.')
    manifest = capture.manifest
    # Los campos no declarados del manifiesto NO se arrastran a la copia.
    metadata = {
        'capture_id': mapper.capture_id,
        'captured_at': manifest['captured_at'],
        'downloader': {'version': 'anonimizador-local', 'commit': 'COPIA_LOCAL'},
        'contains_sensitive_data': True,  # Nunca certificamos inocuidad del texto libre.
        'originals_authorized': False,
        'queries': deepcopy(tables['consultas']),
        'table_states': deepcopy(manifest['table_states']),
        'mapping_version': MAPPING_VERSION,
        'required_features': [],
    }
    consolidated = _consolidated(capture.consolidated, mapper) if capture.consolidated is not None else None
    path = write_capture(target, metadata, tables, consolidated=consolidated)
    # El contrato verifica de nuevo integridad e importación en write_capture.
    return path
