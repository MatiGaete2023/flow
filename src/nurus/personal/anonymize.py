"""Copia anónima de un Excel RUS/SITFA para compartir un caso de prueba.

Uso:
    python -m nurus.personal.anonymize original.xls copia_anonima.xlsx [--desplazar-fechas DIAS]

La copia se construye desde cero y solo con valores: no lleva fórmulas, comentarios,
hipervínculos, gráficos, tablas, filtros, formato condicional, nombres definidos,
vínculos externos, imágenes ni propiedades del original. Nombres, RUT, RIT, curadores,
domicilios, teléfonos y correos se reemplazan por valores ficticios estables (el mismo
dato recibe siempre el mismo reemplazo, así se conservan hermanos y causas). Los RUT
ficticios tienen dígito verificador inválido y el mismo formato que el original.
Tribunal, programa, fechas y números se conservan para que las reglas den el mismo
resultado.

Falla cerrado: antes de entregar la copia se revisa completa. Si queda un dato original,
un RUT con dígito válido, un teléfono o un correo real, la copia se borra y se informa.
Nunca sobrescribe archivos ni modifica el original.
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from itertools import chain
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.packaging.core import DocumentProperties

SENSITIVE = (
    "nombre",
    "apellido",
    "nna",
    "menor",
    "nino",
    "joven",
    "usuario",
    "beneficiario",
    "victima",
    "ruc",
    "madre",
    "padre",
    "litigante",
    "abogado",
    "apoderado",
    "tutor",
    "hermano",
    "familiar",
    "responsable",
    "profesional",
    "dupla",
    "contacto",
    "referente",
    "cuidador",
    "adulto",
    "educador",
    "psicolog",
    "trabajador",
    "asistente",
    "denunciante",
    "imputado",
    "demandante",
    "demandado",
    "solicitante",
    "interviniente",
)
KEEP = ("nombre centro", "nombre programa", "nombre tribunal", "nombre proyecto", "nombre institucion")
FREE_TEXT = ("observacion", "obs", "glosa", "detalle", "texto", "comentario", "nota", "motivo")
PLACEHOLDERS = {
    "",
    "no",
    "no posee",
    "no tiene",
    "sin designar",
    "sin curador",
    "sin rut",
    "n a",
    "na",
    "no aplica",
    "s i",
    "si",
    "sin informacion",
    "sin dato",
    "sd",
    "pendiente",
}
STOPWORDS = {
    "de",
    "del",
    "la",
    "las",
    "los",
    "san",
    "santa",
    "y",
    "e",
    "van",
    "von",
    "da",
    "dos",
    "abog",
    "sr",
    "sra",
    "srta",
    "don",
    "dona",
    "fono",
    "rut",
    "run",
}
DASH = "\\-\u2010\u2011\u2012\u2013\u2014\u2212"  # guiones; el ASCII va escapado para usarlo en clases
RUT_RE = re.compile(
    r"(?<![\w.])(\d{1,2}(?:\.\d{3}){2}|\d{1,2}(?: \d{3}){2}|\d{6,8})(\s*[" + DASH + r"]\s*|\s+)([\dkK])(?![\w])"
)
RUT_NO_DV_RE = re.compile(r"(?<![\w.])\d{1,2}(?:\.\d{3}){2}(?![\w.])")
RIT_STRICT_RE = re.compile(r"(?<!\w)([A-Za-z])[" + DASH + r"](\d{1,6})[" + DASH + r"/]((?:19|20)\d{2})(?!\w)")
RIT_FREE_RE = re.compile(
    r"(?<!\w)(?:([A-Za-z])[\s" + DASH + r"]+)?(\d{1,6})[\s" + DASH + r"/]+((?:19|20)\d{2})(?![\w])"
)
PHONE_RE = re.compile(
    r"(?<![\w+])(?:\+?56[\s-]?)?(?:9[\s-]?\d{4}[\s-]?\d{4}|\d{3}[\s-]\d{3}[\s-]\d{3}"
    r"|\d{2}[\s-]?\d{3}[\s-]?\d{4}|\d{2}[\s-]?\d{7})(?!\d)"
)
GROUPED_DIGITS_RE = re.compile(r"(?<![\w.])\d(?:[\s" + DASH + r"]?\d){7,}(?![\w])")
DATE_LIKE_RE = re.compile(
    r"\d{4}[\s" + DASH + r"/]\d{1,2}[\s" + DASH + r"/]\d{1,2}|\d{1,2}[\s" + DASH + r"/]\d{1,2}[\s" + DASH + r"/]\d{4}"
)
MAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
ADDRESS_RE = re.compile(
    r"(?i)\b(?:av(?:enida)?|calle|pasaje|pje|villa|pobl(?:aci[oó]n)?|block|depto|camino)\.?\s+"
    r"[\wÁÉÍÓÚÑáéíóúñ .'-]{2,40}?\d{1,5}"
)
DIGITS_RE = re.compile(r"(?<![\d.\-])\d{7,9}(?![\d\-])")
FAKE_DOMAIN = "@example.invalid"
SUFFIXES = {".xls", ".xlsx", ".xlsm"}
HEADER_LIMIT = 60
_FOLD = {"a": "aáàäâ", "e": "eéèëê", "i": "iíìïî", "o": "oóòöô", "u": "uúùüû", "n": "nñ", "c": "cç"}


@dataclass
class Report:
    target: Path
    changed: int = 0
    sheets: int = 0
    removed: dict[str, int] = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)


def nfc(value):
    return unicodedata.normalize("NFC", value) if isinstance(value, str) else value


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[-().:_/°º]+", " ", text.lower())).strip()


def check_digit(number):
    total, factor = 0, 2
    for char in reversed(str(number)):
        total += int(char) * factor
        factor = 2 if factor == 7 else factor + 1
    rest = 11 - total % 11
    return {11: "0", 10: "K"}.get(rest, str(rest))


def folded_pattern(text):
    """Patrón que encuentra `text` sin importar mayúsculas, tildes ni espacios repetidos."""
    base = normalize(text)
    parts = []
    for char in base:
        if char == " ":
            parts.append(r"[\s\-.,]+")
        else:
            parts.append(
                "[" + re.escape(_FOLD.get(char, char) + _FOLD.get(char, char).upper()) + "]"
                if char in _FOLD
                else re.escape(char)
            )
    return re.compile(r"(?<!\w)" + "".join(parts) + r"(?!\w)", re.IGNORECASE)


def _rut_key(digits, dv):
    return re.sub(r"\D", "", digits) + "-" + dv.upper()


class Pseudonyms:
    def __init__(self, reserved_rits=()):
        self.values = {}
        self.counters = {}
        self.reserved_rits = set(reserved_rits)

    def get(self, kind, original):
        original = str(original).strip()
        key = (kind, normalize(original) if kind in ("persona", "direccion") else original.upper())
        if kind == "rit":
            match = RIT_FREE_RE.search(original)
            if match:
                key = (kind, f"{(match.group(1) or '').upper()}-{int(match.group(2))}-{match.group(3)}")
        if key not in self.values:
            self.counters[kind] = number = self.counters.get(kind, 0) + 1
            if kind == "rut":
                base = 90_000_000 + number
                self.values[key] = f"{base}-{'1' if check_digit(base) != '1' else '2'}"
            elif kind == "rit":
                match = RIT_FREE_RE.search(original)
                letter = (match.group(1) or "X").upper() if match else "X"
                year = match.group(3) if match else "2000"
                while f"{letter}-{number}-{year}" in self.reserved_rits:
                    number += 1
                self.counters[kind] = number
                self.values[key] = f"{letter}-{number}-{year}"
            elif kind == "mail":
                self.values[key] = f"contacto{number}{FAKE_DOMAIN}"
            elif kind == "telefono":
                self.values[key] = f"900000{number:03d}"
            elif kind == "direccion":
                self.values[key] = f"Calle Ficticia {number}"
            else:
                self.values[key] = f"Persona Ficticia {number}"
        return self.values[key]

    def fakes(self, kind):
        return {value for (k, _), value in self.values.items() if k == kind}


def fake_rut(pseudonyms, match):
    """RUT ficticio con el mismo formato (puntos y separador) que el original."""
    digits, separator, dv = match.group(1), match.group(2), match.group(3)
    fake = pseudonyms.get("rut", _rut_key(digits, dv))
    number, _, fake_dv = fake.partition("-")
    thousands = "." if "." in digits else " " if " " in digits else ""
    if thousands:
        number = f"{int(number):,}".replace(",", thousands)
    return number + separator + fake_dv


def column_kind(header):
    name = normalize(header)
    if not name or any(name.startswith(keep) for keep in KEEP):
        return None
    compact = name.replace(" ", "")
    if name.startswith("rit") or re.search(r"\brit\b", name):
        return "rit"
    if "curador" in name or re.search(r"\bcur\b", name):
        return "curador"
    if compact in ("rut", "run") or name.startswith(("rut", "run")) or re.search(r"\b(rut|run)\b", name):
        return "rut"
    if any(word in name for word in ("correo", "email", "mail")):
        return "mail"
    if any(word in name for word in ("telefono", "fono", "celular")) or re.search(r"\btel\b", name):
        return "telefono"
    if any(word in name for word in ("direccion", "domicilio")):
        return "direccion"
    if any(re.search(r"\b" + word, name) for word in SENSITIVE):
        return "persona"
    if any(name.startswith(word) for word in FREE_TEXT):
        return "texto"
    return None


def _protected_headers():
    """Columnas que lee el motor y no son identidad: se conservan tal cual."""
    from nurus.rus.columns import COMMON, H2_COLUMNS, MODE_COLUMNS, VENCIMIENTO_ALIASES

    names = set(VENCIMIENTO_ALIASES)
    for key, aliases in {**COMMON, **H2_COLUMNS}.items():
        if key not in ("nombre", "rut", "rit", "curador"):
            names.update(aliases)
    for columns in MODE_COLUMNS.values():
        for aliases in columns.values():
            names.update(aliases)
    return {normalize(name) for name in names}


def header_row(rows, limit=HEADER_LIMIT):
    """Primera fila con al menos dos títulos y una columna de identidad (RIT, RUT/RUN, nombre o curador)."""
    for index, row in enumerate(rows[:limit]):
        texts = [value for value in row if isinstance(value, str) and value.strip()]
        names = [normalize(value) for value in texts]
        if len(names) < 2 or any(len(n) > 40 for n in names):
            continue
        kinds = [column_kind(value) for value in texts]
        identity = sum(kind in ("rit", "rut", "curador", "persona") for kind in kinds)
        if identity >= 2 or any(
            n in ("rit", "rut", "run") or n.startswith(("nombre", "rut ", "rit ", "run ")) for n in names
        ):
            return index
    return None


def _xls_sheets(source):
    import xlrd

    legacy = xlrd.open_workbook(str(source))
    for index in range(legacy.nsheets):
        old = legacy.sheet_by_index(index)
        rows = []
        for r in range(old.nrows):
            row = []
            for c in range(old.ncols):
                cell = old.cell(r, c)
                value = cell.value
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR) or value == "":
                    value = None
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, legacy.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif isinstance(value, float) and value.is_integer():
                    value = int(value)
                row.append(value)
            rows.append(row)
        yield old.name, rows, []


def _xlsx_sheets(source):
    # data_only: el valor calculado que lee el motor, nunca el texto de la fórmula.
    book = load_workbook(source, data_only=True, read_only=False)
    try:
        for sheet in book.worksheets:
            rows = [[None if cell.data_type == "e" else cell.value for cell in row] for row in sheet.iter_rows()]
            yield sheet.title, rows, [str(r) for r in sheet.merged_cells.ranges]
    finally:
        book.close()


def _load(source):
    loader = _xls_sheets if source.suffix.lower() == ".xls" else _xlsx_sheets
    return [(nfc(title), [[nfc(v) for v in row] for row in rows], merged) for title, rows, merged in loader(source)]


def _tokens(full_name, fake):
    parts = fake.split()
    words = [w.strip(",.;:()") for w in re.split(r"\s+", full_name.strip())]
    for index, token in enumerate(words):
        if len(token) >= 3 and normalize(token) not in STOPWORDS and re.search(r"[^\W\d_]", token):
            yield token, ("Apellido" if index else "Nombre") + parts[-1]


def is_placeholder(text):
    """Texto sin dato: marcadores conocidos o signos sueltos. Un texto con dígitos nunca es marcador."""
    norm = normalize(text)
    return norm in PLACEHOLDERS or not re.search(r"[a-z0-9]", norm)


def curator_name(text):
    """Nombre de un curador individual; vacío para placeholders y curadores institucionales."""
    if is_placeholder(text):
        return ""
    if "Institución:" in text:
        # Solo la persona que acompaña a la institución («abog.», «Sr.», entre paréntesis); la entidad se conserva.
        match = re.search(r"(?i)\b(?:abog(?:ad[oa])?|sr|sra|srta)\.?\s+([^()\[\],;:0-9]+)", text)
        return curator_name(match.group(1)) if match else ""
    rest = RUT_RE.sub(" ", text)
    rest = PHONE_RE.sub(" ", rest)
    rest = MAIL_RE.sub(" ", rest)
    rest = re.sub(r"(?i)\b(rut|run|fono|tel[eé]fono|abog|abogad[oa]|sr|sra|srta|curador[a]?)\b\.?", " ", rest)
    rest = re.sub(r"[()\[\]:;,/|0-9]", " ", rest)
    return re.sub(r"\s+", " ", rest).strip(" -.")


class Replacer:
    def __init__(self, pseudonyms):
        self.pseudonyms = pseudonyms
        self.names = {}  # nombre completo -> ficticio
        self.tokens = {}  # palabra de nombre -> ficticio
        self._patterns = None

    def add_person(self, full):
        full = re.sub(r"\s+", " ", full).strip()
        if len(full) < 3 or is_placeholder(full):
            return
        fake = self.pseudonyms.get("persona", full)
        self.names.setdefault(full, fake)
        for token, short in _tokens(full, fake):
            self.tokens.setdefault(token, short)
        self._patterns = None

    def patterns(self):
        if self._patterns is None:
            items = sorted({**self.tokens, **self.names}.items(), key=lambda item: -len(normalize(item[0])))
            self._patterns = [(folded_pattern(original), fake) for original, fake in items]
        return self._patterns

    def _mask_dotted(self, match):
        digits = re.sub(r"\D", "", match.group(0))
        fakes = {f.split("-")[0] for f in self.pseudonyms.fakes("rut")}
        return match.group(0) if digits in fakes else "#" * len(match.group(0))

    def _phone(self, match):
        return self.pseudonyms.get("telefono", re.sub(r"\D", "", match.group(0)))

    def identifiers(self, text, *, free, phones=None):
        """RUT, RIT, correos y RUT sin dígito en `text`. Con `free`, además teléfonos, domicilios y números largos.

        `phones=True` reemplaza teléfonos aunque `free` sea falso (columnas que lee el motor)."""
        result = RUT_RE.sub(lambda m: fake_rut(self.pseudonyms, m), text)
        result = RUT_NO_DV_RE.sub(self._mask_dotted, result)
        rit = RIT_FREE_RE if free else RIT_STRICT_RE
        result = rit.sub(lambda m: self.pseudonyms.get("rit", m.group(0)), result)
        result = MAIL_RE.sub(
            lambda m: m.group(0) if m.group(0).endswith(FAKE_DOMAIN) else self.pseudonyms.get("mail", m.group(0)),
            result,
        )
        if free or phones:
            result = PHONE_RE.sub(self._phone, result)
        if free:
            result = ADDRESS_RE.sub(lambda m: self.pseudonyms.get("direccion", m.group(0)), result)
            result = GROUPED_DIGITS_RE.sub(
                lambda m: (
                    m.group(0)
                    if DATE_LIKE_RE.fullmatch(m.group(0)) or re.sub(r"\D", "", m.group(0)).startswith("900000")
                    else "#" * len(m.group(0))
                ),
                result,
            )
            result = DIGITS_RE.sub(
                lambda m: m.group(0) if m.group(0).startswith("900000") else "#" * len(m.group(0)), result
            )
        return result

    def text(self, value, *, free=True):
        result = self.identifiers(value, free=free)
        for pattern, fake in self.patterns():
            result = pattern.sub(fake, result)
        return result


def _identity_value(kind, value, pseudonyms, replacer):
    if isinstance(value, (datetime, date, bool)):
        return value
    text = str(value).strip()
    if kind == "rut":
        if isinstance(value, (int, float)) or re.fullmatch(r"\d{6,9}", text):
            digits = re.sub(r"\D", "", text)
            fake = pseudonyms.get("rut", digits)
            return int(fake.split("-")[0]) if isinstance(value, (int, float)) else fake.split("-")[0]
        if not re.search(r"\d", text):
            return value
        return RUT_RE.sub(lambda m: fake_rut(pseudonyms, m), text) if RUT_RE.search(text) else "#" * len(text)
    if is_placeholder(text):
        return value
    if kind == "rit":
        return pseudonyms.get("rit", text)
    if kind == "persona":
        return replacer.text(text)
    return pseudonyms.get(kind, text)


def _long_number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return False
    return float(value).is_integer() and 7 <= len(str(abs(int(value)))) <= 9


def _sheet_plan(rows):
    index = header_row(rows)
    if index is None:
        return None, {}
    return index, {col: column_kind(value) for col, value in enumerate(rows[index])}


def _collect_names(sheets, replacer):
    for _, rows, _ in sheets:
        index, kinds = _sheet_plan(rows)
        if index is None:
            continue
        for row in rows[index + 1 :]:
            for col, value in enumerate(row):
                if not isinstance(value, str):
                    continue
                kind = kinds.get(col)
                if kind == "persona":
                    replacer.add_person(
                        re.sub(
                            r"(?i)^\s*(dupla|profesional|contacto|referente)\s*:\s*",
                            "",
                            MAIL_RE.sub(" ", PHONE_RE.sub(" ", value)),
                        )
                    )
                elif kind == "curador":
                    replacer.add_person(curator_name(value))


def _original_rits(sheets):
    """Reserva causas de todas las hojas y títulos antes de generar reemplazos."""
    result = set()
    for title, rows, _ in sheets:
        texts = chain((title,), (value for row in rows for value in row if isinstance(value, str)))
        for text in texts:
            for match in RIT_FREE_RE.finditer(text):
                result.add(f"{(match.group(1) or 'X').upper()}-{int(match.group(2))}-{match.group(3)}")
    return result


def _scan(book, pseudonyms, replacer, protected_cells):
    """Revisión final de la copia ya construida. Devuelve (filtraciones, residuos permitidos)."""
    found, residual = set(), 0
    fake_ruts = {f.split("-")[0] for f in pseudonyms.fakes("rut")}
    fake_phones = pseudonyms.fakes("telefono")
    names = [(folded_pattern(name), "persona") for name in replacer.names]
    tokens = [(folded_pattern(token), "palabra de nombre") for token in replacer.tokens if len(normalize(token)) >= 4]
    fake_rits = pseudonyms.fakes("rit")
    texts = [(sheet.title, False) for sheet in book.worksheets]
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                protected = (sheet.title, cell.row, cell.column) in protected_cells
                if isinstance(cell.value, str):
                    texts.append((cell.value, protected))
                elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and not protected:
                    digits = str(int(cell.value)) if float(cell.value).is_integer() else ""
                    if 7 <= len(digits) <= 9 and digits not in fake_ruts and digits not in fake_phones:
                        found.add("número de 7 a 9 dígitos")
    for text, protected in texts:
        text = nfc(text)
        for match in RIT_STRICT_RE.finditer(text):
            if f"{match.group(1).upper()}-{int(match.group(2))}-{match.group(3)}" not in fake_rits:
                found.add("RIT")
        for match in RUT_RE.finditer(text):
            digits = re.sub(r"\D", "", match.group(1))
            if digits not in fake_ruts and check_digit(digits) == match.group(3).upper():
                found.add("RUT con dígito válido")
        if any(not m.group(0).endswith(FAKE_DOMAIN) for m in MAIL_RE.finditer(text)):
            found.add("correo")
        for match in PHONE_RE.finditer(text):
            if re.sub(r"\D", "", match.group(0)) not in fake_phones:
                found.add("teléfono")
        for pattern, label in names + tokens:
            if pattern.search(text):
                if protected and label == "palabra de nombre":
                    residual += 1
                else:
                    found.add(label)
    return sorted(found), residual


def anonymize(source, target, shift_days=0):
    source, target = Path(source), Path(target)
    if source.suffix.lower() not in SUFFIXES:
        raise ValueError("El origen debe ser un Excel (.xls, .xlsx o .xlsm).")
    if target.suffix.lower() != ".xlsx":
        raise ValueError("La copia anónima se guarda como .xlsx.")
    if target.exists():
        raise FileExistsError(f"El destino ya existe; elige otro nombre: {target.name}")
    if source.resolve() == target.resolve():
        raise ValueError("El destino debe ser distinto del original.")
    sheets = _load(source)
    report = Report(target)
    pseudonyms = Pseudonyms(_original_rits(sheets))
    replacer = Replacer(pseudonyms)
    _collect_names(sheets, replacer)
    protected_headers = _protected_headers()
    book = Workbook()
    book.remove(book.active)
    protected_cells = set()
    used_titles = set()
    for number, (title, rows, merged) in enumerate(sheets, 1):
        index, kinds = _sheet_plan(rows)
        others = {t for t, _, _ in sheets if t != title}
        new_title = replacer.text(title)[:31] or f"Hoja{number}"
        if new_title in used_titles or (new_title != title and new_title in others):
            new_title = f"Hoja{number}"
        used_titles.add(new_title)
        sheet = book.create_sheet(new_title)
        protected_cols = (
            {col for col, value in enumerate(rows[index]) if normalize(value) in protected_headers}
            if index is not None
            else set()
        )
        for r, row in enumerate(rows):
            body = index is not None and r > index
            for c, value in enumerate(row):
                if value is None or value == "":
                    continue
                kind = kinds.get(c) if body else None
                new = value
                if shift_days and isinstance(value, (datetime, date)):
                    new = value + timedelta(days=shift_days)
                elif index is not None and r == index:
                    new = value
                elif body and c in protected_cols:
                    new = replacer.identifiers(value, free=False, phones=True) if isinstance(value, str) else value
                    protected_cells.add((new_title, r + 1, c + 1))
                elif kind in ("rut", "rit", "mail", "telefono", "direccion", "persona"):
                    new = _identity_value(kind, value, pseudonyms, replacer)
                elif _long_number(value):
                    # Fuera de columnas que lee el motor, un entero de 7 a 9 dígitos puede ser RUT o teléfono.
                    new = "#" * len(str(int(value)))
                elif isinstance(value, str):
                    new = replacer.text(value, free=kind != "rit")
                if new != value:
                    report.changed += 1
                cell = sheet.cell(r + 1, c + 1, new)
                if isinstance(new, str):
                    # openpyxl interpreta «=» como fórmula incluso si el origen era texto literal.
                    cell.data_type = "s"
        for cells in merged:
            sheet.merge_cells(cells)
        report.sheets += index is not None
    book.properties = DocumentProperties(creator="anonimizado", lastModifiedBy="anonimizado")
    found, residual = _scan(book, pseudonyms, replacer, protected_cells)
    if found:
        raise ValueError(
            "No se creó la copia: quedarían datos personales ("
            + ", ".join(found)
            + "). Revisa el original o informa el caso sin adjuntarlo."
        )
    handle = tempfile.NamedTemporaryFile(prefix=".nurus-anon-", suffix=".xlsx", dir=target.parent, delete=False)
    handle.close()
    temp = Path(handle.name)
    try:
        book.save(temp)
        try:
            os.link(temp, target)  # falla si el destino apareció mientras tanto; nunca sobrescribe
        except FileExistsError:
            raise FileExistsError(f"El destino ya existe; elige otro nombre: {target.name}") from None
        except OSError:
            created = False
            try:
                with target.open("xb") as output:
                    created = True
                    output.write(temp.read_bytes())
            except Exception:
                if created:
                    target.unlink(missing_ok=True)
                raise
    finally:
        temp.unlink(missing_ok=True)
        book.close()
    if residual:
        report.warnings.append(
            f"{residual} celda(s) de programa o tribunal contienen palabras iguales a "
            "nombres; se conservaron porque las lee el motor. Revisa la copia."
        )
    if shift_days:
        report.warnings.append(
            "Las fechas se desplazaron: usa la fecha de corte desplazada; las reglas pueden diferir del original."
        )
    report.warnings.append("Revisa la copia antes de compartirla; no conserva formato, fórmulas ni objetos.")
    return report


def main(argv=None):
    parser = argparse.ArgumentParser(description="Copia anónima de un Excel RUS/SITFA.")
    parser.add_argument("origen", type=Path)
    parser.add_argument("destino", type=Path)
    parser.add_argument("--desplazar-fechas", type=int, default=0, metavar="DIAS")
    args = parser.parse_args(argv)
    try:
        report = anonymize(args.origen, args.destino, args.desplazar_fechas)
    except (OSError, ValueError) as exc:
        print(exc)
        return 1
    print(f"Copia anónima creada: {report.target} ({report.changed} celdas reemplazadas en {report.sheets} hojas).")
    for warning in report.warnings:
        print("Advertencia:", warning)
    return 0


if __name__ == "__main__":
    sys.exit(main())
