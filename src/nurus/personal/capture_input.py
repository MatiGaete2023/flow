"""Importador transaccional del paquete hacia Work; motor tradicional conservado."""

from collections import defaultdict
from copy import deepcopy
from datetime import date, timedelta
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from openpyxl import Workbook
from nurus.capture import read_capture
from nurus.capture.contract import COMMON as SOURCE_FIELDS
from nurus.capture.temporal import (
    chronology,
    attach_displays,
    identity,
    DISPLAY_VERSION,
    RULESET_VERSION,
)
from nurus.rus.columns import COMMON, MODE_COLUMNS, map_columns
from nurus.rus.models import ReadBatch, SourceRecord, SourceReference, Mode
from nurus.rus.util import as_date
from .motor.historia import construir, dia_iso, observaciones
from .work import Work
from nurus.services.excel_cells import append_literal_row

CANONICAL = {key: names[0] for key, names in COMMON.items()}
for columns in MODE_COLUMNS.values():
    CANONICAL.update({key: names[0] for key, names in columns.items()})
MAX_ANALYSIS_RECORDS = 20000  # Techo operativo ensayado; no altera el límite agregado del contrato.
VENTANA_INFORMES_DIAS = 60  # Calendario de informes por vencer: mes actual y siguiente


# Tipos de registro que evalúa el motor tradicional; EGRESADOS/OTRO quedan sin evaluar.
ENGINE_TYPES = ('ESPERA', 'CUMPLIMIENTO', 'INFORMES')


STATE_TYPES = ('ESPERA', 'CUMPLIMIENTO', 'EGRESADOS')
STATE_LABELS = {'ESPERA': 'Espera', 'CUMPLIMIENTO': 'Cumplimiento', 'EGRESADOS': 'Egresados'}


def mark_inconsistent_states(rows, provenance):
    """Q-03: un mismo ingreso verificado en dos pestañas de estado del lote es un estado RUS inconsistente.

    Informes no cuenta: un informe pendiente convive con su ingreso en Cumplimiento. Solo avisa; no corrige.
    """
    states = defaultdict(set)
    for row in rows:
        record = provenance[row.id]
        if record['identity_status'] == 'VERIFICADA' and record['record_type'] in STATE_TYPES:
            states[identity(record)].add(record['record_type'])
    for row in rows:
        record = provenance[row.id]
        found = states.get(identity(record), set()) if record['identity_status'] == 'VERIFICADA' else set()
        if record['record_type'] in STATE_TYPES and len(found) > 1:
            names = ' y '.join(STATE_LABELS[t] for t in STATE_TYPES if t in found)
            row.warnings.append(f'Q-03: el mismo ingreso figura en {names} en este lote; revisar su estado en RUS.')


def mark_search_orders(rows, provenance, capture, as_of):
    """OB-01..OB-05 solo mostrar para filas de «Búsqueda Lit» (pantalla litigantes, menú 94); sin cruces."""
    from .ordenes_busqueda import avisos_orden

    screens = {query['query_id']: query.get('screen', '') for query in capture.rows('consultas')}
    for row in rows:
        record = provenance[row.id]
        if screens.get(record.get('query_id')) == 'litigantes':
            row.warnings.extend(avisos_orden(json.loads(record['normalized_json']), as_of))


def _historias_por_causa(capture):
    """{(source_system, court_code, cause_id): Historia}; vacío si la fuente no se solicitó."""
    rango = {'COMPLETA': 3, 'PARCIAL': 2, 'VACIA_COMPROBADA': 1}
    filas = defaultdict(list)
    for fila in capture.rows('historia_tramites'):
        filas[fila['query_id']].append(fila)
    resultado, mejor = {}, {}
    for consulta in capture.manifest['queries']:
        if consulta['screen'] != 'historia_tramites':
            continue
        filtros = json.loads(consulta['filters_json'])
        causa = filtros.get('historia_causa')
        if not causa:
            continue
        clave = (causa['source_system'], causa['court_code'], causa['cause_id'])
        puntaje = rango.get(consulta['coverage'], 0)
        if clave in mejor and mejor[clave] >= puntaje:
            continue
        ventana = filtros.get('ventana') or {}
        desde, hasta = (dia_iso(ventana.get(k)) for k in ('desde', 'hasta'))
        resultado[clave] = construir(filas.get(consulta['query_id'], ()), consulta['coverage'], desde, hasta)
        mejor[clave] = puntaje
    return resultado


def captured_rule_context(work, source):
    """Fuentes de la sesión para recalcular una fila, sin reabrir el ZIP ni elevar identidad."""
    if not source or source['identity_status'] != 'VERIFICADA':
        return None, ()
    tables = getattr(work, 'capture_rule_sources', None)
    if tables is None:
        raise ValueError(
            'La sesión anterior no conserva fuentes del motor. Reimporta el paquete antes de editar datos.'
        )
    cached = getattr(work, '_save_rule_source_index', None)
    if cached is None or cached[0] != work.capture_digest:
        from nurus.capture.contract import Capture

        capture = Capture(work.capture_manifest, tables, work.capture_digest)
        diaries = defaultdict(list)
        for event in capture.rows('bitacoras'):
            if event.get('identity_status') == 'VERIFICADA' and event.get('scope', 'INGRESO') == 'INGRESO':
                diaries[identity(event)].append(event)
        cached = (work.capture_digest, _historias_por_causa(capture), {k: observaciones(v) for k, v in diaries.items()})
        work._save_rule_source_index = cached  # índice derivado; Work.save excluye los campos _save_*.
    return cached[1].get((source['source_system'], source['court_code'], source['cause_id'])), cached[2].get(
        identity(source), ()
    )


def import_capture(path, config, *, as_of=None, previous=None, resolutions=None):
    from .operations import checkpoint, progress

    checkpoint()
    progress('Validando paquete y cobertura')
    capture = read_capture(path)
    as_of = as_of or date.today()
    config_hash = sha256(json.dumps(config, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()
    if (
        previous
        and getattr(previous, 'capture_digest', '') == capture.digest
        and previous.as_of == as_of.isoformat()
        and getattr(previous, 'capture_config_hash', '') == config_hash
        and getattr(previous, 'ruleset_version', '') == RULESET_VERSION
        and getattr(previous, 'mapping_version', '') == capture.manifest['mapping_version']
    ):
        if resolutions:
            raise ValueError('Respuesta de conciliación ajena o innecesaria; trabajo anterior intacto.')
        cached = deepcopy(previous)
        # Los mismos bytes pueden reimportarse después de mover el paquete.
        # La procedencia debe apuntar al ZIP validado en esta importación.
        cached.path = str(Path(path).resolve())
        from .capture_findings import updated_contexts

        cached.capture_context = updated_contexts(cached, cached.rows)
        if getattr(previous, 'display_version', '') != DISPLAY_VERSION:
            # Una versión nueva de la vista solo se recalcula: no invalida ediciones ni pide conciliación.
            attach_displays(capture, cached.capture_context, cached.capture_sources, date.fromisoformat(cached.as_of))
            cached.display_version = DISPLAY_VERSION
        from nurus.capture.coverage import record_coverages

        cached.capture_coverage = record_coverages(capture.manifest, capture.tables, cached.capture_sources)
        cached.capture_incidents = deepcopy(capture.rows('incidencias'))
        from .capture_changes import compare

        cached.capture_changes = compare(
            cached,
            previous,
            {row.id: row.id for row in cached.rows},
            {row.id: () for row in cached.rows},
        )
        return cached
    records = list(capture.rows('registros'))
    todos = list(records)
    # Con registros de Seguimiento en la captura, el calendario de informes por vencer solo aporta fechas al cruce
    # (C-10 y correo preventivo): sus filas no son registros propios, para no duplicar a la misma persona. Un paquete
    # hecho solo del calendario sigue mostrándolo como filas.
    pantallas = {q['query_id']: q['screen'] for q in capture.rows('consultas')}
    solo_cruce = any(pantallas.get(r['query_id']) != 'calendario_informes' for r in records)
    if solo_cruce:
        records = [r for r in records if pantallas.get(r['query_id']) != 'calendario_informes']

    # Un productor independiente puede aportar calendarios sin duplicarlos en
    # registros. Solo se omite una copia de la misma fila física, nunca por RIT.
    def physical(row):
        return tuple(row[k] for k in ('query_id', 'source_sha256', 'original_file', 'original_sheet', 'original_row'))

    represented = {physical(row) for row in todos}
    record_ids = {row['record_id'] for row in todos}
    for calendar in capture.rows('calendarios'):
        # WD-1: las proyecciones por ingreso complementan su registro; no son registros propios.
        if physical(calendar) in represented or solo_cruce or calendar['event_type'] == 'INFORME_PROYECTADO':
            continue
        record = {k: calendar[k] for k in SOURCE_FIELDS}
        is_report = calendar['event_type'].strip().upper() in ('INFORME', 'INFORMES')
        record.update(
            record_id='calendar:' + calendar['calendar_id'],
            record_type='INFORMES' if is_report else 'OTRO',
            identity_status='VERIFICADA' if calendar['association'] == 'VERIFICADA' else 'INCIERTA',
            raw_json=calendar['raw_json'],
            normalized_json=calendar['normalized_json'],
            source_table='calendarios',
            calendar_id=calendar['calendar_id'],
        )
        if record['record_id'] in record_ids:
            raise ValueError('Identificador de calendario colisiona con registro; trabajo anterior intacto.')
        records.append(record)
        record_ids.add(record['record_id'])
    if len(records) > MAX_ANALYSIS_RECORDS:
        raise ValueError(
            'La captura excede el límite operativo de 20.000 registros; '
            'divida la exportación en lotes. El trabajo anterior sigue intacto.'
        )
    canonical = []
    groups = defaultdict(list)
    provenance = {}
    contexts = {}
    row_modes = {}
    calendars = defaultdict(list)
    for row in capture.rows('calendarios'):
        if row['association'] == 'VERIFICADA' and row['event_type'].strip().upper() in ('INFORME', 'INFORMES'):
            calendars[identity(row)].append(row)
    cross = {}
    from .calendario_ingreso import avisos_cotejo, proxima_pendiente, proyecciones

    projected = proyecciones(capture, identity)
    cotejo = {}
    event_index = defaultdict(list)
    bitacora_index = defaultdict(list)
    historias = _historias_por_causa(capture)
    historia_by_record = {}
    bitacora_by_record = {}
    for table in ('bitacoras', 'historia_tramites'):
        for event in capture.rows(table):
            if event.get('identity_status') == 'VERIFICADA' and event.get('scope', 'INGRESO') == 'INGRESO':
                event_index[identity(event)].append(event)
                if table == 'bitacoras':
                    bitacora_index[identity(event)].append(event)
    for number, record in enumerate(records, 2):
        if number % 100 == 0:
            checkpoint()
            progress('Relacionando fuentes por ingreso', number - 1, len(records))
        raw = json.loads(record['normalized_json'])
        mode = record['record_type']
        if mode not in ENGINE_TYPES:
            mode = 'ESPERA'
        mapping = map_columns(raw, mode)
        values = {CANONICAL[key]: raw[column] for key, column in mapping.items()}
        # Solo campos mapeados. Raw/normalizado completos quedan en procedencia.
        values.update(
            NURUS_TIPO_REGISTRO=record['record_type'],
            MODALIDAD=record['modality_raw'],
            SITFA_CAUSA_RUS=record['cause_id'],
            SITFA_INGRESO_RUS=record['ingreso_id'],
            SITFA_TRIBUNAL_CODIGO=record['court_code'],
            SITFA_SHA256=record['source_sha256'],
            SITFA_ARCHIVO=record['original_file'],
        )
        canonical.append(values)
        groups[(mode, tuple(sorted(mapping)))].append((record, number, values, mapping))
        provenance[record['record_id']] = record
        row_modes[record['record_id']] = mode
        contexts[record['record_id']] = chronology(capture, record, as_of, event_index=event_index)
        if record['identity_status'] == 'VERIFICADA':
            historia_by_record[record['record_id']] = historias.get(
                (record['source_system'], record['court_code'], record['cause_id'])
            )
            bitacora_by_record[record['record_id']] = observaciones(bitacora_index.get(identity(record), ()))
        matching = calendars.get(identity(record), []) if record['identity_status'] == 'VERIFICADA' else []
        dates = []
        for cal in matching:
            source = json.loads(cal['normalized_json'])
            cm = map_columns(source, 'INFORMES')
            due = as_date(source.get(cm.get('vencimiento', ''), ''))
            if due and as_of < due <= as_of + timedelta(days=VENTANA_INFORMES_DIAS):
                dates.append(due)
        if len(set(dates)) == 1:
            cross[record['record_id']] = dates[0]
        items = projected.get(identity(record), []) if record['identity_status'] == 'VERIFICADA' else []
        if items and mode in ('CUMPLIMIENTO', 'INFORMES'):
            due = None
            if mode == 'INFORMES':
                column = map_columns(raw, 'INFORMES').get('vencimiento')
                due = as_date(raw.get(column, '')) if column else None
            cotejo[record['record_id']] = avisos_cotejo(
                items, as_of, vencimiento_informes=due, fecha_hoja2=cross.get(record['record_id'])
            )
            # Calendario como fuente (decisión 19-oct-2026): suple a Hoja2 solo cuando esta no aporta fecha.
            if record['record_id'] not in cross and proxima_pendiente(items, as_of):
                cross[record['record_id']] = proxima_pendiente(items, as_of)
    headers = list(dict.fromkeys(key for row in canonical for key in row))
    if not headers:
        headers = ['RIT', 'TRIBUNAL', 'NOMBRE', 'DERIVACION']
    book = Workbook()
    sheet = book.active
    sheet.title = 'Trabajo'
    sheet.append(headers)
    for number, row in enumerate(canonical, 2):
        if number % 100 == 0:
            checkpoint()
            progress('Construyendo copia Excel', number - 1, len(canonical))
        append_literal_row(sheet, (row.get(k, '') for k in headers))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    out = BytesIO()
    book.save(out)
    book.close()
    content = out.getvalue()
    digest = sha256(content).hexdigest()
    candidate = Work(config)
    candidate.rows = []
    candidate.mapping = {key: title for key, title in CANONICAL.items() if title in headers}
    candidate.mode = next(iter(groups), ('ESPERA', ()))[0]
    candidate.sheet = 'Trabajo'
    candidate.header = 1
    candidate.as_of = as_of.isoformat()
    candidate.path = str(Path(path).resolve())
    candidate.content = content
    candidate.source_hash = digest
    candidate.sitfa_sources = {}
    candidate.cross_dates = {}
    for (mode, _), items in groups.items():
        checkpoint()
        progress('Analizando reglas tradicionales: ' + mode)
        keys = set(key for _, _, _, mapping in items for key in mapping)
        mapping = {key: CANONICAL[key] for key in keys}
        batch = ReadBatch(
            mode=Mode(mode),
            records=tuple(
                SourceRecord(values, SourceReference(Path(path).name, digest, 'Trabajo', number), record['record_id'])
                for record, number, values, _ in items
            ),
            source_path=str(path),
            workbook_name=Path(path).name,
            workbook_sha256=digest,
            primary_sheet='Trabajo',
            header_row=1,
            column_mapping=mapping,
            source_bytes=content,
        )
        analyzed = Work(config).analyze(
            path,
            mode,
            as_of=as_of,
            batch=batch,
            cross_by_record=cross,
            historia_by_record=historia_by_record,
            bitacora_by_record=bitacora_by_record,
        )
        # Un grupo solo de EGRESADOS/OTRO no pasa por el motor: sus avisos de columnas serían falsos.
        if any(record['record_type'] in ENGINE_TYPES for record, _, _, _ in items):
            candidate.warnings.extend(analyzed.warnings)
        candidate.cross_dates.update(analyzed.cross_dates)
        for row in analyzed.rows:
            record = provenance[row.id]
            if record['record_type'] not in ENGINE_TYPES:
                row.observation = ''
                row.rules = []
                row.actions = []
                row.warnings.append('Tipo de fuente fuera del motor tradicional; no evaluado.')
            row.warnings.append(contexts[row.id]['reason'])
            candidate.rows.append(row)
    attach_displays(capture, contexts, provenance, as_of)
    candidate.rows.sort(key=lambda row: row.source_row)
    mark_inconsistent_states(candidate.rows, provenance)
    for row in candidate.rows:
        row.warnings.extend(cotejo.get(row.id, ()))
    mark_search_orders(candidate.rows, provenance, capture, as_of)
    from .contexto_excel import marcar_paquete

    marcar_paquete(candidate.rows, provenance, identity)
    from .bitacora_avisos import aplicar_reiteracion

    aplicar_reiteracion(candidate.rows, contexts, config)
    candidate.capture_digest = capture.digest
    candidate.capture_id = capture.manifest['capture_id']
    candidate.capture_config_hash = config_hash
    candidate.capture_manifest = capture.manifest
    candidate.capture_rule_sources = {table: capture.rows(table) for table in ('historia_tramites', 'bitacoras')}
    candidate.capture_sources = provenance
    candidate.capture_context = contexts
    candidate.capture_modes = row_modes
    candidate.ruleset_version = RULESET_VERSION
    candidate.display_version = DISPLAY_VERSION
    candidate.mapping_version = capture.manifest['mapping_version']
    candidate.source_scope = 'PAQUETE'
    candidate.signed_activity = {}
    candidate.activity_reviewed = {}
    from nurus.capture.coverage import record_coverages

    candidate.capture_coverage = record_coverages(capture.manifest, capture.tables, provenance)
    candidate.capture_incidents = deepcopy(capture.rows('incidencias'))
    from .capture_reconciliation import reconcile

    reconcile(candidate, previous, resolutions=resolutions)
    from .capture_findings import updated_contexts

    candidate.capture_context = updated_contexts(candidate, candidate.rows)
    # NO_SOLICITADA no falta: el usuario no pidió esa fuente (p. ej. bitácoras de un lote sin lectura).
    missing = [
        k + ': ' + v
        for k, v in capture.manifest['table_states'].items()
        if v not in ('COMPLETA', 'VACIA_COMPROBADA', 'NO_SOLICITADA')
    ]
    if missing:
        candidate.warnings.append('Fuentes faltantes: ' + ', '.join(missing))
    progress('Paquete validado; trabajo candidato disponible', len(records), len(records))
    return candidate
