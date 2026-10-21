"""Detección por contenido y lectura/escritura de tablas: xlsx, xls, XML 2003, HTML, CSV."""

from __future__ import annotations

import csv
import re
import zipfile
from datetime import date, datetime
from html import escape
from html.parser import HTMLParser
from io import BytesIO, StringIO
from xml.etree import ElementTree

from .grid import MAX_CELLS, Sheet
from .textengine import nfc

OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
TEXT_SUFFIXES = {".txt", ".md", ".log", ".text"}
SS = "urn:schemas-microsoft-com:office:spreadsheet"
MAX_ROWS = 1_048_576


class Unsupported(ValueError):
    """Formato que esta herramienta no puede anonimizar con garantías."""


def decode_text(data):
    """Texto sin reemplazos: UTF-8 (con o sin BOM), declarada en HTML/XML o Windows-1252."""
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16")
    head = data[:4096]
    declared = re.search(rb"charset\s*=\s*[\"']?\s*([A-Za-z0-9_-]+)", head, re.I) or re.search(
        rb"encoding\s*=\s*[\"']([A-Za-z0-9_-]+)", head, re.I
    )
    names = []
    if declared:
        label = declared.group(1).decode("ascii").lower()
        names.append({"iso-8859-1": "cp1252", "latin1": "cp1252", "latin-1": "cp1252"}.get(label, label))
    names += ["utf-8-sig", "cp1252"]
    for name in names:
        try:
            return data.decode(name)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("latin-1")


def sniff(data, name=""):
    """Tipo lógico por contenido: xlsx, xls, xml2003, html, json, ics, csv, txt, docx, zip, pdf o None."""
    suffix = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    head = data[:4096].lstrip(b"\xef\xbb\xbf \t\r\n")
    if data[:4] == b"PK\x03\x04" or data[:4] == b"PK\x05\x06":
        try:
            names = set(zipfile.ZipFile(BytesIO(data)).namelist())
        except zipfile.BadZipFile:
            return None
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "word/document.xml" in names:
            return "docx"
        if any(n.startswith(("ppt/", "visio/")) for n in names):
            return None
        return "zip"
    if data[:8] == OLE:
        return "xls" if suffix in (".xls", ".xlt", "") else None
    if data[:5] == b"%PDF-":
        return "pdf"
    if re.match(rb"(?i)<\?xml", head) and SS.encode() in data[:8192]:
        return "xml2003"
    if head[:1] in (b"<",) and re.search(rb"(?i)<(!doctype html|html|table|body|head|div|p)\b", head):
        return "html"
    if data[:2] in (b"\xff\xd8", b"\x89P", b"GI", b"BM") or data[:4] in (b"RIFF", b"MZ\x90\x00"):
        return None
    if b"\x00" in data[:2048] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return None
    if re.match(rb"\s*[\[{]", head) and suffix in (".json", ".jsonl", ".ndjson", ""):
        return "json"
    if suffix in (".json", ".jsonl", ".ndjson"):
        return "json"
    if head.upper().startswith(b"BEGIN:VCALENDAR") or suffix == ".ics":
        return "ics"
    if suffix in (".csv", ".tsv"):
        return "csv"
    if suffix in TEXT_SUFFIXES or suffix == "":
        return "txt"
    if suffix in (".xls", ".xlsx", ".xlsm", ".html", ".htm", ".xml"):
        return None
    return None


def _check_dimensions(rows, cells):
    cells += sum(len(row) for row in rows)
    if len(rows) > MAX_ROWS or cells > MAX_CELLS:
        raise Unsupported("El archivo supera el límite de celdas admitido.")
    return cells


def has_formulas(data):
    """¿Algún hoja del libro guarda fórmulas? Solo se conservan sus valores calculados."""
    from . import safezip

    safezip.validate(data)  # límites de tamaño y compresión antes de descomprimir hojas
    with zipfile.ZipFile(BytesIO(data)) as book:
        for name in book.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                head = book.read(name)
                if re.search(rb"<f[ >/]", head):
                    return True
    return False


def read_xlsx(data):
    from openpyxl import load_workbook

    from .safezip import validate

    validate(data)
    book = load_workbook(BytesIO(data), data_only=True, read_only=False, keep_links=False)
    sheets, cells = [], 0
    try:
        for sheet in book.worksheets:
            rows = [[None if cell.data_type == "e" else cell.value for cell in row] for row in sheet.iter_rows()]
            cells = _check_dimensions(rows, cells)
            sheets.append(Sheet(nfc(sheet.title), _nfc_rows(rows), [str(r) for r in sheet.merged_cells.ranges]))
    finally:
        book.close()
    return sheets


def read_xls(data):
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover - xlrd se declara como dependencia
        raise Unsupported("Falta el componente xlrd para leer .xls antiguos.") from exc
    book = xlrd.open_workbook(file_contents=data)
    sheets, cells = [], 0
    for index in range(book.nsheets):
        old = book.sheet_by_index(index)
        rows = []
        for r in range(old.nrows):
            row = []
            for c in range(old.ncols):
                cell = old.cell(r, c)
                value = cell.value
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR) or value == "":
                    value = None
                elif cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, book.datemode)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    value = bool(value)
                elif isinstance(value, float) and value.is_integer():
                    value = int(value)
                row.append(value)
            rows.append(row)
        cells = _check_dimensions(rows, cells)
        sheets.append(Sheet(nfc(old.name), _nfc_rows(rows), []))
    return sheets


def read_xml2003(data):
    if re.search(rb"<!(?:DOCTYPE|ENTITY)", data[:65536], re.I):
        raise Unsupported("XML con declaraciones DOCTYPE/ENTITY rechazado por seguridad.")
    root = ElementTree.fromstring(data)
    ns = {"ss": SS}
    key_index, key_name = f"{{{SS}}}Index", f"{{{SS}}}Name"
    sheets, cells = [], 0
    for number, worksheet in enumerate(root.findall("ss:Worksheet", ns), 1):
        grid = []
        for table in worksheet.findall("ss:Table", ns):
            for row in table.findall("ss:Row", ns):
                if row.get(key_index):
                    target = int(row.get(key_index))
                    if not len(grid) < target <= MAX_ROWS:
                        raise Unsupported("Índice de fila XML inválido.")
                    grid.extend([[] for _ in range(target - 1 - len(grid))])
                values = []
                for cell in row.findall("ss:Cell", ns):
                    if cell.get(key_index):
                        position = int(cell.get(key_index))
                        if not 1 <= position <= 16384:
                            raise Unsupported("Índice de columna XML inválido.")
                        values.extend([None] * (position - 1 - len(values)))
                    node = cell.find("ss:Data", ns)
                    value = "".join(node.itertext()) if node is not None else None
                    if node is not None and value:
                        kind = node.get(f"{{{SS}}}Type")
                        if kind == "Number":
                            try:
                                number_value = float(value)
                                value = int(number_value) if number_value.is_integer() else number_value
                            except ValueError:
                                pass
                        elif kind == "DateTime":
                            try:
                                value = datetime.fromisoformat(value.rstrip("Z"))
                            except ValueError:
                                pass
                        elif kind == "Boolean":
                            value = value.strip() in ("1", "true", "True")
                    values.append(value if value != "" else None)
                grid.append(values)
        cells = _check_dimensions(grid, cells)
        sheets.append(Sheet(nfc(worksheet.get(key_name) or f"Hoja{number}"), _nfc_rows(grid), []))
    return sheets


def _nfc_rows(rows):
    return [[nfc(v) for v in row] for row in rows]


def read_csv(data):
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if text.count(";") > text.count(",") else ","
    csv.field_size_limit(16 * 1024 * 1024)
    rows = [
        [v if v != "" else None for v in row] for row in csv.reader(StringIO(text, newline=""), delimiter=delimiter)
    ]
    _check_dimensions(rows, 0)
    return [Sheet("Hoja1", _nfc_rows(rows), [])], delimiter, data[:3] == b"\xef\xbb\xbf"


SKIP_TAGS = {"script", "style", "noscript", "template", "svg", "head", "iframe", "object", "embed", "canvas"}
BLOCK_TAGS = {
    "p",
    "div",
    "br",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "li",
    "ul",
    "ol",
    "section",
    "article",
    "header",
    "footer",
    "pre",
    "blockquote",
    "hr",
    "form",
    "body",
    "tr",
    "dd",
    "dt",
    "dl",
    "fieldset",
    "legend",
    "label",
}


class _Blocks(HTMLParser):
    """Esqueleto del HTML: tablas y bloques de texto. Descarta atributos, enlaces, scripts, imágenes y comentarios."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.skip, self.depth, self.span = [], 0, 0, 1
        self.text, self.rows, self.row, self.cell, self.in_cell = [], [], None, [], False

    def flush(self):
        text = re.sub(r"\s+", " ", "".join(self.text)).strip()
        if text:
            self.blocks.append(("text", text))
        self.text = []

    def handle_starttag(self, tag, attrs):
        if tag in SKIP_TAGS:
            self.skip += 1
            return
        if self.skip:
            return
        if tag == "table":
            if self.depth == 0:
                self.flush()
                self.rows = []
            self.depth += 1
        elif self.depth == 1 and tag == "tr":
            self.row = []
        elif self.depth == 1 and tag in ("td", "th"):
            self.cell, self.in_cell = [], True
            colspan = str(dict(attrs).get("colspan") or "1")
            self.span = min(int(colspan), 50) if colspan.isdigit() and int(colspan) > 0 else 1
        elif self.depth <= 1 and tag in BLOCK_TAGS and self.depth == 0:
            self.flush()
        elif self.in_cell and tag in BLOCK_TAGS:
            self.cell.append(" ")

    def handle_endtag(self, tag):
        if tag in SKIP_TAGS:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if self.depth == 1 and tag in ("td", "th") and self.in_cell:
            text = re.sub(r"\s+", " ", "".join(self.cell)).strip()
            if self.row is not None:
                self.row.append(text or None)
                self.row.extend([None] * (self.span - 1))
            self.in_cell = False
        elif self.depth == 1 and tag == "tr" and self.row is not None:
            if any(self.row):
                self.rows.append(self.row)
            self.row = None
        elif tag == "table" and self.depth:
            self.depth -= 1
            if self.depth == 0:
                if self.rows:
                    self.blocks.append(("table", self.rows))
                self.rows = []
        elif self.depth == 0 and tag in BLOCK_TAGS:
            self.flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.in_cell:
            self.cell.append(data)
        elif self.depth == 0:
            self.text.append(data)

    def close(self):
        super().close()
        self.flush()


def read_html(data):
    """HTML -> [('text', str) | ('table', filas)] sin atributos ni scripts."""
    parser = _Blocks()
    parser.feed(decode_text(data))
    parser.close()
    cells = 0
    for kind, content in parser.blocks:
        if kind == "table":
            cells = _check_dimensions(content, cells)
    return [(kind, _nfc_rows(content) if kind == "table" else nfc(content)) for kind, content in parser.blocks]


def write_xlsx(sheets):
    from openpyxl import Workbook
    from openpyxl.packaging.core import DocumentProperties

    book = Workbook()
    book.remove(book.active)
    try:
        for sheet in sheets:
            target = book.create_sheet(sheet.title[:31])
            for r, row in enumerate(sheet.rows, 1):
                for c, value in enumerate(row, 1):
                    if value is None or value == "":
                        continue
                    cell = target.cell(r, c, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
            for cells in sheet.merged:
                target.merge_cells(cells)
        if not book.worksheets:
            book.create_sheet("Hoja1")
        book.properties = DocumentProperties(creator="anonimizado", lastModifiedBy="anonimizado")
        buffer = BytesIO()
        book.save(buffer)
        return buffer.getvalue()
    finally:
        book.close()


def write_csv(sheet, delimiter, bom):
    stream = StringIO(newline="")
    writer = csv.writer(stream, delimiter=delimiter, lineterminator="\r\n")
    for row in sheet.rows:
        writer.writerow(["" if v is None else _csv_value(v) for v in row])
    return (("﻿" if bom else "") + stream.getvalue()).encode("utf-8")


def _csv_value(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value[:1] in ("=", "+", "@", "\t", "\r"):
        return "'" + value  # evita que Excel ejecute una celda de texto como fórmula
    return value


def write_html(blocks):
    parts = ['<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Copia anónima</title></head><body>']
    for kind, content in blocks:
        if kind == "text":
            parts.append(f"<p>{escape(content)}</p>")
        else:
            parts.append('<table border="1">')
            for row in content:
                cells = "".join(f"<td>{escape('' if v is None else str(_csv_value(v)))}</td>" for v in row)
                parts.append(f"<tr>{cells}</tr>")
            parts.append("</table>")
    parts.append("</body></html>")
    return "\n".join(parts).encode("utf-8")
