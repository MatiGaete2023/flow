import hashlib
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from anonimizador_csmp.engine import Rejected, process_bytes, run
from anonimizador_csmp.textengine import OMITTED, Options, Session
from conftest import make_xlsx, valid_rut

HEAD = ["RIT", "RUT", "NOMBRE", "TRIBUNAL", "T ESPERA", "FEC. NACIMIENTO", "OBSERVACION"]


def rows_of(path):
    book = load_workbook(path)
    try:
        return [list(r) for r in book.active.iter_rows(values_only=True)]
    finally:
        book.close()


def anon(tmp_path, rows, **options):
    source = make_xlsx(tmp_path / "origen.xlsx", rows)
    result = run([source], Options(**options), tmp_path / "salida")
    assert result.ok, [i.detail for i in result.items]
    return rows_of(result.outputs[0])


def test_estricto_omite_observaciones_y_seudonimiza_identidad(tmp_path, rut):
    out = anon(
        tmp_path,
        [HEAD, ["P-77-2026", rut, "Juan Pérez Soto", "LAJA", 45, datetime(2015, 3, 4), "Habló con Ana Rojas"]],
        mode="estricto",
    )
    assert out[0] == HEAD
    rit, new_rut, name, court, wait, birth, obs = out[1]
    assert rit != "P-77-2026" and rit.startswith("P-") and rit.endswith("-2026")
    assert new_rut != rut and name == "Persona Ficticia 1"
    assert court == "LAJA" and wait == 45 and birth == datetime(2015, 3, 4)
    assert obs == OMITTED


def test_conservador_mantiene_contexto_sin_identidades(tmp_path):
    out = anon(
        tmp_path,
        [
            HEAD,
            ["P-1-2026", None, "Juan Pérez Soto", "LAJA", 3, None, "Audiencia 05-10-2026; juan perez soto no asistió"],
        ],
        mode="conservador",
    )
    obs = out[1][6]
    assert "05-10-2026" in obs and "Persona Ficticia 1" in obs and "Pérez" not in obs


def test_mismo_nombre_mismo_reemplazo_y_hermanos_distintos(tmp_path):
    out = anon(
        tmp_path,
        [
            HEAD,
            ["P-1-2026", None, "Ana Díaz", "LAJA", 1, None, None],
            ["P-2-2026", None, "Luis Díaz", "LAJA", 2, None, None],
            ["P-3-2026", None, "ANA DIAZ", "LAJA", 3, None, None],
        ],
    )
    assert out[1][2] == out[3][2] != out[2][2]


def test_numeros_largos_y_ids_se_vinculan_entre_hojas(tmp_path):
    source = make_xlsx(
        tmp_path / "libro.xlsx",
        [["RIT", "NOMBRE", "Causa RUS", "Folio"], ["P-1-2026", "Ana Díaz", 1234567890, 12345678]],
        extra_sheets=[("Resumen", [["RIT", "NOMBRE", "Causa RUS"], ["P-1-2026", "Ana Díaz", "1234567890"]])],
    )
    result = run([source], Options(), tmp_path / "o")
    book = load_workbook(result.outputs[0])
    first, second = (
        list(book.worksheets[0].iter_rows(values_only=True)),
        list(book.worksheets[1].iter_rows(values_only=True)),
    )
    assert first[1][2] == second[1][2] and str(first[1][2]).startswith("ID-")
    assert "12345678" not in str(first[1][3])


def test_columna_desconocida_se_omite_pero_estructura_se_conserva(tmp_path):
    out = anon(
        tmp_path,
        [
            HEAD + ["DEFENSOR", "ESTADO", "OTRA"],
            ["P-1-2026", None, "Ana Díaz", "LAJA", 1, None, None, "Carlos Muñoz", "VALIDADA", "dato libre"],
        ],
    )
    assert out[1][7] == "Persona Ficticia 2" and out[1][8] == "VALIDADA" and out[1][9] == OMITTED


def test_titulos_y_propiedades_no_filtran(tmp_path):
    source = make_xlsx(
        tmp_path / "x.xlsx", [HEAD, ["P-1-2026", None, "Ana Díaz", "LAJA", 1, None, None]], title="Ana Díaz"
    )
    from openpyxl import load_workbook as load

    book = load(source)
    book.properties.creator = "Ana Díaz"
    book.active["A5"] = "=1+1"
    book.save(source)
    result = run([source], Options(), tmp_path / "o")
    out = load(result.outputs[0])
    assert out.active.title == "Hoja1" and out.properties.creator == "anonimizado"
    assert "Díaz" not in out.active.title


def test_texto_con_igual_no_se_interpreta_como_formula(tmp_path):
    source = make_xlsx(tmp_path / "f.xlsx", [HEAD, ["P-1-2026", None, "Ana Díaz", "LAJA", 1, None, "x"]])
    book = load_workbook(source)
    book.active["G2"].value = '=HYPERLINK("x")'
    book.active["G2"].data_type = "s"
    book.save(source)
    result = run([source], Options(mode="conservador"), tmp_path / "salida")
    out = load_workbook(result.outputs[0])
    assert out.active["G2"].data_type == "s" and isinstance(out.active["G2"].value, str)


def test_desplazar_fechas_en_celdas_y_textos(tmp_path):
    out = anon(
        tmp_path,
        [
            HEAD,
            ["P-1-2026", None, "Ana Díaz", "LAJA", 1, date(2015, 3, 4), None],
            ["P-2-2026", None, "Luis", "LAJA", 1, "2015-03-04", None],
        ],
        shift_days=10,
    )
    assert out[1][5].date() == date(2015, 3, 14) or out[1][5] == datetime(2015, 3, 14)
    assert out[2][5] == "2015-03-14"


def test_generalizar_nacimiento(tmp_path):
    out = anon(
        tmp_path,
        [
            HEAD,
            ["P-1-2026", None, "Ana Díaz", "LAJA", 1, datetime(2015, 3, 4), None],
            ["P-2-2026", None, "Luis", "LAJA", 1, "04-03-2016", None],
        ],
        generalize_birth=True,
    )
    assert out[1][5] == datetime(2015, 1, 1) and out[2][5] == "01-01-2016"


def test_pares_etiqueta_valor_sin_encabezado(tmp_path):
    source = make_xlsx(
        tmp_path / "ficha.xlsx",
        [["Nombre:", "Ana Díaz"], ["RUT:", "11111111-1"], ["Tribunal", "LAJA"]],
        title="Resumen",
    )
    result = run([source], Options(), tmp_path / "o")
    out = rows_of(result.outputs[0])
    assert out[0][1] == "Persona Ficticia 1" and out[1][1] != "11111111-1" and out[2][1] == "LAJA"


def test_origen_intacto_y_nunca_sobrescribe(tmp_path):
    source = make_xlsx(tmp_path / "o.xlsx", [HEAD, ["P-1-2026", None, "Ana Díaz", "LAJA", 1, None, None]])
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    result = run([source], Options(), tmp_path)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    assert result.outputs[0] != source and result.outputs[0].exists()


@pytest.mark.parametrize("name", ["tabla.xls", "tabla.xlsx", "tabla.html"])
def test_html_con_extension_excel_se_lee_por_contenido(tmp_path, name):
    html = (
        "<html><body><table><tr><th>RIT</th><th>NOMBRE</th><th>OBSERVACION</th></tr>"
        "<tr><td>P-2-2026</td><td>Ana Díaz</td><td>Llamar a Pedro Soto</td></tr></table></body></html>"
    )
    path = tmp_path / name
    path.write_text(html, encoding="utf-8")
    result = run([path], Options(mode="conservador"), tmp_path / "o")
    assert result.ok
    out = result.outputs[0]
    if name.endswith(".html"):
        text = out.read_text(encoding="utf-8")
        assert "Ana" not in text and "Soto" not in text and "Persona Ficticia 1" in text
    else:
        assert out.suffix == ".xlsx"
        rows = rows_of(out)
        assert rows[1][1] == "Persona Ficticia 1" and "Soto" not in rows[1][2]


def test_xml_spreadsheet_2003(tmp_path):
    xml = (
        '<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
        'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="ESPERA"><Table>'
        '<Row><Cell><Data ss:Type="String">RIT</Data></Cell><Cell><Data ss:Type="String">NOMBRE</Data></Cell>'
        '<Cell><Data ss:Type="String">T ESPERA</Data></Cell></Row>'
        '<Row><Cell><Data ss:Type="String">P-5-2026</Data></Cell><Cell><Data ss:Type="String">Ana Díaz</Data></Cell>'
        '<Cell><Data ss:Type="Number">12</Data></Cell></Row></Table></Worksheet></Workbook>'
    )
    path = tmp_path / "SITFA.xls"
    path.write_text(xml, encoding="utf-8")
    result = run([path], Options(), tmp_path / "o")
    rows = rows_of(result.outputs[0])
    assert rows[1][1] == "Persona Ficticia 1" and rows[1][2] == 12


def test_xml_con_entidades_se_rechaza(tmp_path):
    path = tmp_path / "x.xls"
    path.write_text(
        '<?xml version="1.0"?><!DOCTYPE a [<!ENTITY e "x">]><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"/>'
    )
    result = run([path], Options(), tmp_path / "o")
    assert not result.ok and not result.outputs


def test_xlsx_dañado_y_pdf_se_rechazan_con_mensaje_generico(tmp_path):
    bad = tmp_path / "roto.xlsx"
    bad.write_bytes(b"PK\x03\x04basura")
    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF-1.4 contenido")
    result = run([bad, pdf], Options(), tmp_path / "o")
    assert [i.ok for i in result.items] == [False, False]
    assert all("roto" not in i.detail and "doc" not in i.detail for i in result.items)


def test_csv_con_punto_y_coma_y_bom(tmp_path):
    path = tmp_path / "datos.csv"
    path.write_bytes("﻿RIT;NOMBRE;OBSERVACION\r\nP-1-2026;Ana Díaz;Llamó Ana Díaz\r\n".encode("utf-8"))
    result = run([path], Options(mode="conservador"), tmp_path / "o")
    text = result.outputs[0].read_bytes().decode("utf-8-sig")
    assert text.startswith("RIT;NOMBRE;OBSERVACION") and "Díaz" not in text and text.count("Persona Ficticia 1") == 2


def test_filtracion_detectada_no_crea_copia():
    session = Session(Options())
    session.replacer.add_person("Pedro Soto")
    found, _ = session.leaks([("fue Pedro Soto", False)])
    assert "nombre" in found
    with pytest.raises(Rejected):
        from anonimizador_csmp.engine import _scan

        _scan([(valid_rut(), False)], session)
    _ = process_bytes


def test_csv_no_deja_formulas_activas(tmp_path):
    path = tmp_path / "f.csv"
    path.write_text('RIT,NOMBRE,OBSERVACION\nP-1-2026,Ana,"=HYPERLINK(""http://x"")"\n', encoding="utf-8")
    result = run([path], Options(mode="conservador"), tmp_path / "o")
    text = result.outputs[0].read_text(encoding="utf-8")
    assert ",=" not in text and "http://x" not in text


def test_xls_binario_real(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sheet = book.add_sheet("ESPERA")
    for c, value in enumerate(["RIT", "NOMBRE", "T ESPERA"]):
        sheet.write(0, c, value)
    for c, value in enumerate(["P-1-2026", "Ana Díaz", 12]):
        sheet.write(1, c, value)
    path = tmp_path / "sitfa.xls"
    book.save(str(path))
    result = run([path], Options(), tmp_path / "o")
    rows = rows_of(result.outputs[0])
    assert result.ok and rows[1][1] == "Persona Ficticia 1" and rows[1][2] == 12
