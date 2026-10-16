"""WP-16, revisión: casos que antes filtraban datos. Todos los datos son ficticios.

Cada caso construye un original con un dato personal escondido en un lugar distinto y
exige que la copia no lo contenga, o que el anonimizador se niegue a crearla.
"""

import html
import re
import unicodedata
import zipfile
from datetime import datetime

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from nurus.personal.anonymize import anonymize, check_digit

ENCABEZADOS = [
    'RIT',
    'TRIBUNAL',
    'NOMBRE',
    'RUT',
    'FEC. NACIMIENTO',
    'NOMBRE CENTRO',
    'T ESPERA',
    'CURADOR',
    'OBSERVACION',
    'PROFESIONAL',
    '',
]
NOMBRE = 'José Ñuñez Pérez'


def rut(number, formato='{n}-{dv}'):
    # Se calcula en ejecución: este archivo no contiene RUT con dígito válido.
    puntos = f'{number:,}'.replace(',', '.')
    return formato.format(n=number, p=puntos, dv=check_digit(number))


RUT_NNA = rut(15_284_731)


def plegar(texto):
    texto = unicodedata.normalize('NFKD', texto).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', texto)


def base(ws, filas_titulo=0, obs='', curador='NO POSEE', profesional='', libre=''):
    for _ in range(filas_titulo):
        ws.append(['Informe ficticio'])
    ws.append(ENCABEZADOS)
    ws.append(
        ['P-1534-2025', 'LAJA', NOMBRE, RUT_NNA, datetime(2010, 1, 1), 'PRM UNO', 45, curador, obs, profesional, libre]
    )


def contenido(ruta):
    with zipfile.ZipFile(ruta) as archivo:
        return html.unescape('\n'.join(archivo.read(n).decode('utf-8', 'ignore') for n in archivo.namelist()))


def objetos(wb, ws):
    base(ws)
    grafico = BarChart()
    grafico.title = 'Espera de Ñuñez'
    grafico.add_data(Reference(ws, min_col=7, min_row=1, max_row=2))
    ws.add_chart(grafico, 'M2')
    ws.auto_filter.ref = 'A1:K2'
    ws.auto_filter.add_filter_column(2, ['Ñuñez'])
    ws.defined_names['Nunez_hoja'] = DefinedName('Nunez_hoja', attr_text='ESPERA!$C$2')
    ws.conditional_formatting.add(
        'I2:I9',
        Rule(
            type='containsText',
            operator='containsText',
            text='Ñuñez',
            formula=['NOT(ISERROR(SEARCH("Ñuñez",I2)))'],
            dxf=DifferentialStyle(),
        ),
    )
    ws.oddHeader.center.text = 'Ñuñez'
    otra = wb.create_sheet('Tabla')
    otra.append(['RIT', 'NNA'])
    otra.append(['P-1-2026', 'x'])
    otra.add_table(Table(displayName='Tabla_Nunez', ref='A1:B2'))


def vinculo_externo(wb, ws):
    from openpyxl.packaging.relationship import Relationship
    from openpyxl.workbook.external_link.external import (
        ExternalBook,
        ExternalCell,
        ExternalLink,
        ExternalRow,
        ExternalSheetData,
        ExternalSheetDataSet,
        ExternalSheetNames,
    )

    base(ws)
    libro = ExternalBook(
        sheetNames=ExternalSheetNames(sheetName=['Hoja1']),
        id='rId1',
        sheetDataSet=ExternalSheetDataSet(
            sheetData=[
                ExternalSheetData(
                    sheetId=0, row=[ExternalRow(r=1, cell=[ExternalCell(r='A1', t='str', v='Persona Canario Externa')])]
                )
            ]
        ),
    )
    vinculo = ExternalLink(externalBook=libro)
    vinculo.file_link = Relationship(
        type='externalLinkPath', TargetMode='External', Target='file:///C:/Users/persona.canario/Causas/CanarioExterno.xlsx'
    )
    wb._external_links.append(vinculo)
    ws['M2'] = '=[1]Hoja1!A1'


def hoja_secundaria(wb, ws):
    base(ws)
    otra = wb.create_sheet('Listado')
    otra.append(['N° CAUSA', 'NNA', 'RUN'])
    otra.append(['C-1', 'Ñuñez Pérez José', int(rut(15_284_731, '{n}{dv}'))])


def hojas_con_nombre(wb, ws):
    base(ws)
    wb.create_sheet('Ñuñez')
    wb.create_sheet('Pérez')


def hoja_extra(encabezados, *filas):
    def construir(wb, ws):
        base(ws)
        otra = wb.create_sheet('Extra')
        otra.append(encabezados)
        for fila in filas:
            otra.append(fila)

    return construir


CASOS = {
    'encabezado_en_fila_20': (lambda wb, ws: base(ws, filas_titulo=19), [NOMBRE, 'Ñuñez']),
    'texto_libre_variantes': (
        lambda wb, ws: base(
            ws,
            obs=(
                'JOSE NUNEZ fono +56 9 8765 4321, vive en Av. Los Aromos 1234 Coronel; RUT '
                + rut(15_284_731, '{p} {dv}')
                + ' / '
                + rut(15_284_731, '{n}{dv}')
                + ' / '
                + rut(15_284_731, '{n} - {dv}')
                + '; causa 1534-2025 y P 1534 2025'
            ),
        ),
        [
            'JOSE NUNEZ',
            '8765 4321',
            'Los Aromos',
            rut(15_284_731, '{p} {dv}'),
            rut(15_284_731, '{n}{dv}'),
            rut(15_284_731, '{n} - {dv}'),
            '1534-2025',
            'P 1534 2025',
        ],
    ),
    'curador_sin_rut': (
        lambda wb, ws: base(ws, curador='Abog. Marcela Quiroga Lillo, fono 41-2223344'),
        ['Marcela Quiroga Lillo', 'Quiroga', '2223344'],
    ),
    'curador_rut_seis_digitos': (
        lambda wb, ws: base(ws, curador='Ramiro Tapia Ulloa ' + rut(912_345)),
        ['Ramiro Tapia Ulloa', rut(912_345), 'Tapia'],
    ),
    'columna_profesional_y_sin_encabezado': (
        lambda wb, ws: base(ws, profesional='Dupla: Ignacia Bustos Mella', libre='hablar con Ñuñez'),
        ['Ignacia Bustos Mella', 'Bustos', 'Ñuñez'],
    ),
    'objetos_del_libro': (objetos, [NOMBRE, 'Ñuñez', 'Nunez']),
    'formula_con_telefono': (
        lambda wb, ws: (
            base(ws),
            ws.__setitem__('M5', '=HYPERLINK("mailto:jose.nunez@example.invalid","José Ñuñez Pérez 987654321")'),
        ),
        [NOMBRE, '987654321', 'jose.nunez'],
    ),
    'texto_nfd': (lambda wb, ws: base(ws, obs=unicodedata.normalize('NFD', 'Visita a José Ñuñez Pérez')), [NOMBRE]),
    'vinculo_externo': (vinculo_externo, ['Persona Canario Externa', 'persona.canario', 'CanarioExterno']),
    'run_numerico_en_hoja_secundaria': (hoja_secundaria, ['Ñuñez Pérez José', rut(15_284_731, '{n}{dv}')]),
    'hojas_con_nombre': (hojas_con_nombre, ['Ñuñez', 'Pérez']),
    # Revisión del 11 de octubre de 2026: formatos y lugares que antes pasaban sin cambio.
    'rut_sin_dv_guion_largo_y_espacios': (
        lambda wb, ws: base(
            ws,
            obs=(
                'RUT 15.284.731; '
                + rut(16_345_678, '{p}\u2013{dv}')
                + '; '
                + rut(17_456_789, '{n}').replace('17456789', '17 456 789')
                + '-'
                + check_digit(17_456_789)
            ),
        ),
        ['15.284.731', rut(16_345_678, '{p}\u2013{dv}'), '17 456 789'],
    ),
    'telefono_en_grupos_de_tres': (lambda wb, ws: base(ws, obs='llamar al 912 345 678'), ['912 345 678']),
    'numeros_en_celdas_numericas': (
        hoja_extra(
            ['RIT', 'NOMBRE', 'TEL.', 'R.U.N.', 'CONTACTO'],
            ['P-2-2026', 'Persona Dos', 912345678, 15_284_731, 987654321],
        ),
        ['912345678', '15284731', '987654321'],
    ),
    'curador_ad_litem': (
        hoja_extra(['RIT', 'NOMBRE', 'CUR. AD LITEM'], ['P-3-2026', 'Persona Tres', 'Rosa Quintanilla Soto']),
        ['Quintanilla'],
    ),
    'encabezado_sin_nombre_ni_rut': (
        hoja_extra(['N° RIT', 'NNA', 'CURADOR'], ['P-4-2026', 'Rosa Quintanilla', 'Ignacio Valdebenito']),
        ['Quintanilla', 'Valdebenito'],
    ),
    'encabezado_run_nna': (
        hoja_extra(['CAUSA', 'RUN NNA', 'NNA'], ['C-5-2026', rut(15_284_731), 'Rosa Quintanilla']),
        ['Quintanilla', rut(15_284_731)],
    ),
    'telefono_en_columna_protegida': (
        hoja_extra(['RIT', 'NOMBRE', 'DERIVACION'], ['P-6-2026', 'Persona Seis', 'Derivó fono 9 1234 5678']),
        ['1234 5678'],
    ),
    'curador_institucional_con_abogado': (
        lambda wb, ws: base(ws, curador='Institución: Corp X (abog. Ignacio Valdebenito)'),
        ['Valdebenito'],
    ),
    'rit_con_barra': (lambda wb, ws: base(ws, obs='causa C-2345/2023'), ['2345/2023']),
    **{
        f'columna_identidad_{numero}': (
            hoja_extra(['RIT', encabezado], ['P-700-2026', 'Teresa Malibran Arrieta']),
            ['Malibran', 'Arrieta'],
        )
        for numero, encabezado in enumerate(
            ['APELLIDO PATERNO', 'APELLIDOS', 'DEMANDANTE', 'DEMANDADO', 'SOLICITANTE', 'INTERVINIENTE'], 1
        )
    },
}


@pytest.mark.parametrize('caso', list(CASOS))
def test_ningun_dato_personal_sobrevive(tmp_path, caso):
    construir, secretos = CASOS[caso]
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'ESPERA'
    construir(libro, hoja)
    origen, destino = tmp_path / 'original.xlsx', tmp_path / 'copia.xlsx'
    libro.save(origen)
    try:
        anonymize(origen, destino)
    except ValueError as error:
        assert 'No se creó la copia' in str(error) and not destino.exists()
        return
    texto = plegar(contenido(destino))
    assert [s for s in secretos if plegar(s) in texto] == []
    assert not list(tmp_path.glob('.*.tmp.xlsx'))


def test_placeholder_en_rut_no_inventa_un_rut(tmp_path):
    from openpyxl import load_workbook

    libro = Workbook()
    hoja = libro.active
    hoja.append(['RIT', 'NOMBRE', 'RUT'])
    hoja.append(['P-1-2026', 'Persona Uno', 'SIN RUT'])
    libro.save(tmp_path / 'o.xlsx')
    anonymize(tmp_path / 'o.xlsx', tmp_path / 'c.xlsx')
    assert load_workbook(tmp_path / 'c.xlsx').active['C2'].value == 'SIN RUT'


def test_columna_protegida_conserva_el_texto_pero_no_el_rut(tmp_path):
    from openpyxl import load_workbook

    libro = Workbook()
    hoja = libro.active
    hoja.append(['RIT', 'NOMBRE', 'NOMBRE CENTRO'])
    hoja.append(['P-1-2026', 'Persona Uno', 'PRM ' + rut(15_284_731)])
    libro.save(tmp_path / 'o.xlsx')
    anonymize(tmp_path / 'o.xlsx', tmp_path / 'c.xlsx')
    valor = load_workbook(tmp_path / 'c.xlsx').active['C2'].value
    assert valor.startswith('PRM ') and rut(15_284_731) not in valor
