from datetime import date
import json
import pytest
from openpyxl import load_workbook
from nurus.capture import write_capture
from nurus.capture.diaries import export_diaries
from nurus.capture.requests import read_request, write_request
from nurus.capture.temporal import chronology
from nurus.personal.capture_input import import_capture
from nurus.personal.capture_products import export_requests, evidence_text
from nurus.personal.config import defaults
from nurus.personal.product_state import stamp, stale
from nurus.personal.outputs import Draft
from nurus.personal.work import Work
from test_capture_contract import fixture
from bitacora_fixtures import manifest


def test_dates_ties_origin_and_ingreso_are_not_inferred(tmp_path):
    m, t = fixture()
    entry = t['bitacoras'][0]
    # Mismo día sin hora: orden incierto. Un ingreso distinto nunca se cruza.
    t['bitacoras'].append(
        {
            **entry,
            'event_id': 'nuevo',
            'origin': 'OTRO',
            'event_at': '2026-09-10',
            'text': 'Respuesta informada; no acredita recepción.',
        }
    )
    t['bitacoras'].append({**entry, 'event_id': 'posterior', 'origin': 'OTRO', 'event_at': '2026-09-11'})
    path = write_capture(tmp_path / 'capture.zip', m, t)
    from nurus.capture import read_capture

    capture = read_capture(path)
    context = chronology(capture, t['registros'][0], date(2026, 10, 8), fixture_preview=True)
    assert [e['event_id'] for e in context['last_center']] == ['e0']
    assert [e['event_id'] for e in context['later']] == ['posterior']
    assert [e['event_id'] for e in context['uncertain']] == ['nuevo']
    assert chronology(capture, t['registros'][0], date(2026, 10, 8))['class'] == 'NO_EVALUABLE'
    capture.manifest['contains_sensitive_data'] = True
    assert chronology(capture, t['registros'][0], date(2026, 10, 8), fixture_preview=True)['class'] == 'NO_EVALUABLE'


def test_no_center_origin_no_fictional_last_date(tmp_path):
    m, t = fixture()
    t['bitacoras'][0]['origin'] = 'DESCONOCIDO'
    from nurus.capture import read_capture

    capture = read_capture(write_capture(tmp_path / 'capture.zip', m, t))
    context = chronology(capture, t['registros'][0], date(2026, 10, 8), fixture_preview=True)
    assert not context['last_center'] and context['class'] == 'NO_EVALUABLE'


def test_timezone_label_without_offset_does_not_prove_order(tmp_path):
    from nurus.capture import read_capture

    m, t = fixture()
    entry = t['bitacoras'][0]
    entry.update(event_at='2026-09-10T10:00:00', time_precision='SEGUNDO', time_zone='America/Santiago')
    t['bitacoras'].append({**entry, 'event_id': 'reply', 'origin': 'OTRO', 'event_at': '2026-09-10T11:00:00-03:00'})
    capture = read_capture(write_capture(tmp_path / 'capture.zip', m, t))
    result = chronology(capture, t['registros'][0], date(2026, 10, 8), fixture_preview=True)
    assert not result['later'] and [e['event_id'] for e in result['uncertain']] == ['reply']


def test_temporal_window_does_not_seek_older_observation(tmp_path):
    from nurus.capture import read_capture
    from nurus.capture.temporal import earliest

    assert earliest(date(2026, 6, 30)) == date(2026, 2, 28)
    m, t = fixture()
    t['bitacoras'][0]['event_at'] = '2026-06-07'
    capture = read_capture(write_capture(tmp_path / 'capture.zip', m, t))
    result = chronology(capture, t['registros'][0], date(2026, 10, 8), fixture_preview=True)
    assert result['period_from'] == '2026-06-08'
    assert not result['last_center'] and not result['recent']
    assert result['class'] == 'NO_EVALUABLE'
    assert len(capture.rows('bitacoras')) == 4  # Copia capturada no recortada.


def test_long_diary_text_is_split_exactly_and_formulas_are_text(tmp_path):
    source = manifest(tmp_path)
    data = json.loads(source.read_text(encoding='utf-8'))
    data['registros'] = data['registros'][:1]
    html = tmp_path / 'captura.html'
    body = html.read_text(encoding='utf-8').replace(
        'Texto completo que la tabla oculta.', '=1+2 ' + ('😀texto\n' * 6000)
    )
    html.write_text(body, encoding='utf-8')
    from hashlib import sha256

    data['registros'][0]['sha256'] = sha256(html.read_bytes()).hexdigest()
    source.write_text(json.dumps(data), encoding='utf-8')
    path = export_diaries(source, tmp_path / 'diaries.xlsx')
    book = load_workbook(path)
    parts = [row[-1] for row in book['Textos íntegros'].iter_rows(min_row=2, values_only=True) if row[1] == 'texto']
    # Parser conserva el atributo completo (saltos HTML pueden normalizarse).
    from nurus.capture.diaries import import_lote

    original = import_lote(source)['queries'][0]['entries'][0]['texto']
    assert ''.join(parts) == original
    assert all(c.data_type == 's' for row in book['Textos íntegros'].iter_rows(min_row=2) for c in (row[-1],))
    assert book['Lecturas'].max_row == 2 and book['Consultas'].max_row == 5
    book.close()


def test_request_roundtrip_and_code_url_rejection(tmp_path):
    m, t = fixture(partial=True)
    work = import_capture(write_capture(tmp_path / 'capture.zip', m, t), defaults())
    path = export_requests(work, tmp_path / 'request.json')
    data = read_request(path)
    assert len(data['requests']) == 8
    data['requests'][0]['url'] = 'https://example.com'
    with pytest.raises(ValueError):
        write_request(tmp_path / 'bad.json', data)
    assert not (tmp_path / 'bad.json').exists()
    assert 'NO_CONSULTADA' in evidence_text(work, work.rows[0])


def test_product_source_rules_and_cut_invalidation_and_marks_not_saved(tmp_path):
    m, t = fixture()
    work = import_capture(write_capture(tmp_path / 'capture.zip', m, t), defaults())
    draft = Draft('Ficticio', 'Ficticio', record_ids=[work.rows[0].id], kind='programa_espera')
    stamp(work, draft)
    assert not stale(work, draft)
    work.ruleset_version = 'new'
    assert stale(work, draft)
    work.activity_reviewed = {'old': 'local'}
    work.signed_activity = {'old': 'local'}
    work.rows[0].review['OBSERVACION'] = 'Humano'
    work.save(tmp_path / 'session')
    data = json.loads((tmp_path / 'session/trabajo.json').read_text(encoding='utf-8'))
    assert 'activity_reviewed' not in data and 'signed_activity' not in data
    loaded = Work.load(tmp_path / 'session')
    assert loaded.rows[0].review['OBSERVACION'] == 'Humano' and not loaded.activity_reviewed


def test_engine_upgrade_invalidates_prepared_products_without_changing_session(tmp_path, monkeypatch):
    from nurus.capture import temporal

    m, t = fixture()
    work = import_capture(write_capture(tmp_path / 'capture.zip', m, t), defaults())
    draft = Draft('Ficticio', 'Ficticio', record_ids=[work.rows[0].id], kind='programa_espera')
    stamp(work, draft)
    stored = work.ruleset_version
    monkeypatch.setattr(temporal, 'RULESET_VERSION', 'actualizacion-ficticia')
    assert stale(work, draft) and work.ruleset_version == stored


def test_inherited_mail_statement_needs_correction_for_new_package(tmp_path):
    from nurus.personal.capture_products import institutional_statement_issue
    from nurus.personal.outputs import create_draft
    from nurus.personal.product_checks import check_products

    m, t = fixture()
    work = import_capture(write_capture(tmp_path / 'capture.zip', m, t), defaults())
    draft = Draft(
        'Ficticio',
        'Se registraron observaciones en bitácora.',
        to='ficticio@example.test',
        record_ids=[work.rows[0].id],
    )
    stamp(work, draft)
    assert institutional_statement_issue(work, draft)
    assert any('no acredita registro RUS' in issue for c in check_products(work, [draft], []) for issue in c.issues)
    with pytest.raises(ValueError, match='no acredita'):
        create_draft(work, draft, confirmed=True)
    draft.body = 'Se prepararon propuestas para revisión humana.'
    assert not institutional_statement_issue(work, draft)
