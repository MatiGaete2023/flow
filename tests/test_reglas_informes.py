"""R-T16 a R-T32: Informes, contexto ficticio y decisiones humanas con corte fijo."""

from copy import deepcopy
from datetime import date
import json
import sys
from types import ModuleType, SimpleNamespace

import pytest
from openpyxl import load_workbook
from nurus.capture import read_capture, write_capture
from nurus.capture.contract import COMMON
from nurus.capture.fixtures import fixture
from nurus.capture.temporal import chronology
from nurus.personal.capture_input import import_capture
from nurus.personal.config import defaults
from nurus.personal.outputs import create_draft, prepare_drafts
from nurus.personal.record_edits import apply
from nurus.personal.resolutions import automatic_project_selections

from test_mayoria_edad import CORTE, analizar, persona

BASE = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'FECHA VENCIMIENTO']


def test_rt16_informe_vence_ayer_hoy_30_31(tmp_path):
    casos = ['2026-10-08', '2026-10-09', '2026-11-08', '2026-11-09']
    filas = analizar(tmp_path, 'INFORMES', BASE, [persona(n) + [fecha] for n, fecha in enumerate(casos, 1)]).rows
    assert [f.rules for f in filas] == [
        ['INFORMES.I01_VENCIDO_GENERAL'],
        ['INFORMES.I02_POR_VENCER_GENERAL'],
        ['INFORMES.I02_POR_VENCER_GENERAL'],
        [],
    ]
    assert filas[0].actions == ['programa_vencido', 'PC_INFO']
    assert filas[1].actions == ['programa_por_vencer']
    assert filas[3].observation == '' and filas[3].actions == []


def test_rt16_dce_distingue_informe_diagnostico(tmp_path):
    # WA-7 (19-oct-2026): DCE se pide en cuenta pasados 10 días del vencimiento RUS; antes, correo I-02 DCE.
    filas = analizar(
        tmp_path, 'INFORMES', BASE, [persona(1, programa='DCE UNO') + [f] for f in ('2026-09-28', '2026-10-08')]
    ).rows
    assert [f.rules for f in filas] == [['INFORMES.I01_VENCIDO_DCE'], ['INFORMES.I02_POR_VENCER_DCE']]


def test_rt16_sin_fecha_no_emite_frase_e_informa_incidencia(tmp_path):
    fila = analizar(tmp_path, 'INFORMES', BASE, [persona(1) + ['']]).rows[0]
    assert fila.rules == [] and fila.observation == ''
    assert any('sin fecha de vencimiento' in aviso for aviso in fila.warnings)


def test_rt27_regla_valida_con_dato_faltante_de_otra(tmp_path):
    encabezados = [
        'RIT',
        'TRIBUNAL',
        'NOMBRE',
        'RUT',
        'NOMBRE CENTRO',
        'DIAS DE CUMPLIMIENTO',
        'DIAS PARA EGRESAR',
        'FEC. INGRESO EFECTIVO',
        'FEC. EGRESO PROYECTADO',
    ]
    fila = analizar(
        tmp_path, 'CUMPLIMIENTO', encabezados, [persona(1) + ['sin dato', 10, '2026-01-01', '2026-10-19']]
    ).rows[0]
    assert fila.rules == ['CUMPLIMIENTO.C05_POR_VENCER']
    assert '19 de octubre de 2026' in fila.observation
    assert any('días de cumplimiento ausentes o inválidos' in aviso for aviso in fila.warnings)


def test_rt28_tribunal_no_soportado_sin_textos_inventados(tmp_path):
    encabezados = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'T ESPERA']
    fila = analizar(tmp_path, 'ESPERA', encabezados, [persona(1, tribunal='TRIBUNAL FICTICIO') + [90]]).rows[0]
    assert fila.rules == [] and fila.observation == '' and fila.actions == []
    assert any('Tribunal fuera de Laja, Mulchén y Tomé' in aviso for aviso in fila.warnings)
    assert fila.values['TRIBUNAL'] == 'TRIBUNAL FICTICIO'


ESPERA = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'T ESPERA', 'CURADOR']


def test_rt29_recalculo_no_sobrescribe_texto_humano(tmp_path):
    trabajo = analizar(tmp_path, 'ESPERA', ESPERA, [persona(1) + [10, 'SI']])
    fila = trabajo.rows[0]
    decisiones = {'resolution': 'none', 'mail': {'programa_espera': 'omit'}}
    apply(trabajo, [fila.id], review={'OBSERVACION': 'Texto revisado por la persona usuaria.'}, decisions=decisiones)
    apply(trabajo, [trabajo.rows[0].id], overrides={'espera': 45})
    fila = trabajo.rows[0]
    assert 'ESPERA.E05_PROYECTO_Y_CORREO' in fila.rules and 'ESPERA.E06_SIN_RESOLUCION' not in fila.rules
    assert fila.review['OBSERVACION'] == 'Texto revisado por la persona usuaria.'
    assert fila.decisions['mail'] == decisiones['mail']
    assert fila.decisions['resolution'] == 'none' and fila.review['RES'] == ''
    assert prepare_drafts(trabajo, 'programa_espera', directory=tmp_path / 'salida') == []
    assert automatic_project_selections(trabajo) == []


def test_rt30_nominas_exactas_por_programa(tmp_path):
    filas = [
        persona(1) + [45, 'SI'],
        persona(2) + [60, 'SI'],
        persona(3) + [90, 'SI'],
        persona(4, programa='PRM DOS') + [45, 'SI'],
        persona(5) + [10, 'SI'],
    ]
    trabajo = analizar(tmp_path, 'ESPERA', ESPERA, filas)
    borradores = prepare_drafts(trabajo, 'programa_espera', directory=tmp_path / 'salida')
    rit = {fila.id: fila.values['RIT'] for fila in trabajo.rows}
    por_programa = {b.program: b for b in borradores}
    assert len(borradores) == 2 and set(por_programa) == {'PRM UNO', 'PRM DOS'}
    for programa, esperadas in (('PRM UNO', ['X-1-2026', 'X-2-2026', 'X-3-2026']), ('PRM DOS', ['X-4-2026'])):
        ids = por_programa[programa].record_ids
        assert len(ids) == len(set(ids)) == len(esperadas)
        assert sorted(rit[i] for i in ids) == esperadas
        fuente = {fila.id: fila.source_row for fila in trabajo.rows}
        assert [fuente[i] for i in ids] == ([2, 3, 4] if programa == 'PRM UNO' else [5])
        hoja = load_workbook(por_programa[programa].attachments[0]).active
        assert [celda.value for celda in hoja[1]] == ['RIT', 'TRIBUNAL', 'NOMBRE', 'DERIVACION', 'T ESPERA']
        assert [hoja.cell(n, 1).value for n in range(2, hoja.max_row + 1)] == esperadas
        assert hoja.max_row == len(esperadas) + 1


def test_rt32_regla_desactivada_se_omite_y_se_conserva_la_configuracion(tmp_path):
    configuracion = defaults()
    configuracion['desactivadas'] = ['COMUN.CURADOR']
    umbrales = dict(configuracion['umbrales'])
    trabajo = analizar(tmp_path, 'ESPERA', ESPERA, [persona(1) + [45, '']], configuracion=configuracion)
    assert 'COMUN.CURADOR' not in trabajo.rows[0].rules
    assert trabajo.rows[0].rules == ['ESPERA.E05_PROYECTO_Y_CORREO']
    assert configuracion['umbrales'] == umbrales and configuracion['desactivadas'] == ['COMUN.CURADOR']
    activo = analizar(tmp_path, 'ESPERA', ESPERA, [persona(1) + [45, '']])
    assert 'COMUN.CURADOR' in activo.rows[0].rules


def _captura(tmp_path, modificar=None):
    manifiesto, tablas = fixture(count=2)
    if modificar:
        modificar(manifiesto, tablas)
    ruta = write_capture(tmp_path / 'captura.zip', manifiesto, tablas)
    return read_capture(ruta), ruta


def _contexto(captura, numero=0):
    return chronology(captura, captura.rows('registros')[numero], CORTE, fixture_preview=True)


def _ids(eventos):
    return [evento['event_id'] for evento in eventos]


def _tramite(base, **cambios):
    evento = {k: base[k] for k in COMMON}
    evento.update(
        event_id='tramite-ficticio',
        event_at='2026-09-20',
        time_precision='DIA',
        time_zone='DESCONOCIDA',
        scope='INGRESO',
        event_type='Trámite ficticio',
        description='Título ficticio; contenido no cotejado.',
        reference='referencia-ficticia',
        status='CAPTURADO',
        document_id='',
        identity_status='VERIFICADA',
    )
    evento.update(cambios)
    return evento


def test_rt17_dos_ingresos_misma_causa_sin_contaminacion(tmp_path):
    def modificar(manifiesto, tablas):
        for numero in (0, 1):
            tablas['bitacoras'].append(
                {
                    **tablas['bitacoras'][numero],
                    'event_id': f'novedad-{numero}',
                    'origin': 'OTRO',
                    'event_at': '2026-09-20',
                }
            )

    captura, _ = _captura(tmp_path, modificar)
    assert captura.rows('registros')[0]['cause_id'] == captura.rows('registros')[1]['cause_id']
    for numero in (0, 1):
        resultado = _contexto(captura, numero)
        assert _ids(resultado['last_center']) == [f'e{numero}']
        assert _ids(resultado['later']) == [f'novedad-{numero}']
        assert set(_ids(resultado['recent'])) == {f'e{numero}', f'novedad-{numero}'}


def test_rt18_bitacora_totalmente_fallida_es_no_evaluable(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['bitacoras'] = []
        manifiesto['table_states']['bitacoras'] = 'FALLIDA'

    captura, _ = _captura(tmp_path, modificar)
    assert captura.manifest['table_states']['bitacoras'] == 'FALLIDA'
    resultado = _contexto(captura)
    assert resultado['class'] == 'NO_EVALUABLE' and not resultado['last_center']
    assert 'NO_ACREDITADA' in resultado['reason']
    assert 'sin observaciones' not in resultado['reason'].lower()
    assert resultado['findings'][0]['status'] == 'NO_EVALUABLE'


def _discrepancia_visual(manifiesto, tablas):
    tablas['bitacoras'][0]['rit_original'] = 'VISUAL-FICTICIO-DISTINTO'


def test_rt19_ids_iguales_admiten_rit_visual_distinto(tmp_path):
    captura, _ = _captura(tmp_path, _discrepancia_visual)
    resultado = _contexto(captura)
    assert _ids(resultado['last_center']) == ['e0']
    assert resultado['last_center'][0]['rit_original'] == 'VISUAL-FICTICIO-DISTINTO'


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason='[CONFLICTO_ABIERTO] R-T19: se admite el RIT visual distinto por IDs iguales, '
    'pero falta advertencia trazable de discrepancia visual; 03 §12 exige ambas condiciones.',
)
def test_rt19_discrepancia_visual_deja_advertencia_trazable(tmp_path):
    captura, _ = _captura(tmp_path, _discrepancia_visual)
    resultado = _contexto(captura)
    assert any('RIT' in hallazgo['reason'] and 'e0' in hallazgo['source_ids'] for hallazgo in resultado['findings'])


def test_rt20_nombre_igual_ids_ingreso_distintos_no_se_cruzan(tmp_path):
    def modificar(manifiesto, tablas):
        primero, segundo = tablas['registros']
        datos = json.loads(segundo['normalized_json'])
        datos['NOMBRE'] = json.loads(primero['normalized_json'])['NOMBRE']
        segundo.update(raw_json=json.dumps(datos), normalized_json=json.dumps(datos))
        tablas['bitacoras'][1]['author'] = tablas['bitacoras'][0]['author']
        tablas['bitacoras'][1]['event_at'] = '2026-09-25'

    captura, _ = _captura(tmp_path, modificar)
    for numero in (0, 1):
        resultado = _contexto(captura, numero)
        assert _ids(resultado['last_center']) == [f'e{numero}']
        assert _ids(resultado['recent']) == [f'e{numero}']


def test_rt21_observaciones_del_centro_empatadas_se_conservan(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['bitacoras'].append(
            {**tablas['bitacoras'][0], 'event_id': 'empate', 'text': 'Segunda observación ficticia empatada.'}
        )

    captura, _ = _captura(tmp_path, modificar)
    resultado = _contexto(captura)
    assert sorted(_ids(resultado['last_center'])) == ['e0', 'empate']
    assert len({evento['text'] for evento in resultado['last_center']}) == 2


def test_rt22_observacion_fuera_de_ventana_no_inventa_ultima(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['bitacoras'][0]['event_at'] = '2026-06-08'  # Un día antes de la ventana.
        tablas['bitacoras'].append(
            {**tablas['bitacoras'][0], 'event_id': 'reciente-otro', 'origin': 'OTRO', 'event_at': '2026-09-20'}
        )

    captura, _ = _captura(tmp_path, modificar)
    resultado = _contexto(captura)
    assert resultado['period_from'] == '2026-06-09'
    assert not resultado['last_center'] and not resultado['later']
    assert _ids(resultado['recent']) == ['reciente-otro']
    assert 'NO_ACREDITADA' in resultado['reason']


def test_rt23_mismo_dia_sin_hora_es_incierto_y_focalizado(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['bitacoras'].append({**tablas['bitacoras'][0], 'event_id': 'mismo-dia', 'origin': 'OTRO'})

    captura, _ = _captura(tmp_path, modificar)
    resultado = _contexto(captura)
    assert _ids(resultado['uncertain']) == ['mismo-dia']
    assert not resultado['later'] and resultado['class'] == 'FOCALIZADA'


def debe(condicion, motivo):
    """Invariante previa al conflicto (revisión externa H6).

    pytest.fail no es AssertionError: un xfail estricto con raises=AssertionError no puede absorber su falla.
    """
    if not condicion:
        pytest.fail('Invariante previa al conflicto R-T: ' + motivo)


FUERA_DE_ALCANCE = (
    '[CONFLICTO_ABIERTO] {caso}: falta el hallazgo {regla}; historia SITFA/PDF y CT-03..CT-10 '
    'siguen fuera del alcance de PLAN_VIGENTE; se prueba la expectativa, sin activar reglas.'
)


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=FUERA_DE_ALCANCE.format(caso='R-T24', regla='CT-05 con cotejo de documento'),
)
def test_rt24_informe_vencido_y_tramite_exigen_cotejo(tmp_path):
    def modificar(manifiesto, tablas):
        registro = tablas['registros'][0]
        datos = json.loads(registro['normalized_json'])
        datos['FECHA VENCIMIENTO'] = '2026-10-08'
        registro.update(record_type='INFORMES', raw_json=json.dumps(datos), normalized_json=json.dumps(datos))
        tablas['historia_tramites'] = [
            _tramite(
                registro,
                event_id='tramite-informe',
                event_type='Informe',
                description='Ingreso de informe; contenido no cotejado.',
            )
        ]
        manifiesto['table_states']['historia_tramites'] = 'COMPLETA'

    captura, ruta = _captura(tmp_path, modificar)
    trabajo = import_capture(ruta, defaults(), as_of=CORTE)
    debe(
        'INFORMES.I01_VENCIDO_GENERAL' in trabajo.rows[0].rules,
        "'INFORMES.I01_VENCIDO_GENERAL' in trabajo.rows[0].rules",
    )
    resultado = _contexto(captura)
    debe(resultado['class'] == 'FOCALIZADA', "resultado['class'] == 'FOCALIZADA'")
    debe('tramite-informe' in _ids(resultado['later']), "'tramite-informe' in _ids(resultado['later'])")
    debe('subsanado' not in resultado['reason'].lower(), "'subsanado' not in resultado['reason'].lower()")
    assert any(
        h['rule_id'] == 'CT-05' and 'tramite-informe' in h['source_ids'] and 'cotej' in h['reason'].lower()
        for h in resultado['findings']
    )


@pytest.mark.xfail(
    strict=True, raises=AssertionError, reason=FUERA_DE_ALCANCE.format(caso='R-T25', regla='CT-08 y cobertura parcial')
)
def test_rt25_firma_parcial_no_acredita_ausencia_de_actividad(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['historia_tramites'] = [
            _tramite(
                tablas['registros'][0], event_id='firma-parcial', event_type='Resolución firmada', coverage='PARCIAL'
            )
        ]
        manifiesto['table_states']['historia_tramites'] = 'PARCIAL'

    captura, _ = _captura(tmp_path, modificar)
    resultado = _contexto(captura)
    debe(_ids(resultado['later']) == ['firma-parcial'], "_ids(resultado['later']) == ['firma-parcial']")
    debe(resultado['later'][0]['coverage'] == 'PARCIAL', "resultado['later'][0]['coverage'] == 'PARCIAL'")
    debe(
        'historia completa' in resultado['reason'] and 'no acredita' in resultado['reason'],
        "'historia completa' in resultado['reason'] and 'no acredita' in resultado['reason']",
    )
    assert any(
        h['rule_id'] == 'CT-08' and h['status'] == 'NO_EVALUABLE' and 'firma-parcial' in h['source_ids']
        for h in resultado['findings']
    )


@pytest.mark.xfail(
    strict=True, raises=AssertionError, reason=FUERA_DE_ALCANCE.format(caso='R-T26', regla='CT-07 posible reiteración')
)
def test_rt26_reiteracion_senalada_sin_suprimir_propuesta(tmp_path):
    def modificar(manifiesto, tablas):
        tablas['bitacoras'].append(
            {
                **tablas['bitacoras'][0],
                'event_id': 'reiterada',
                'event_at': '2026-09-20',
                'text': tablas['bitacoras'][0]['text'] + ' ',
            }
        )

    captura, ruta = _captura(tmp_path, modificar)
    trabajo = import_capture(ruta, defaults(), as_of=CORTE)
    debe(
        trabajo.rows[0].observation and trabajo.rows[0].actions,
        'trabajo.rows[0].observation and trabajo.rows[0].actions',
    )
    resultado = _contexto(captura)
    debe(set(_ids(resultado['recent'])) == {'e0', 'reiterada'}, "set(_ids(resultado['recent'])) == {'e0', 'reiterada'}")
    debe(_ids(resultado['last_center']) == ['reiterada'], "_ids(resultado['last_center']) == ['reiterada']")
    assert any(
        h['rule_id'] == 'CT-07'
        and set(h['source_ids']) == {'e0', 'reiterada'}
        and 'posible' in h['reason'].lower()
        and not h['proposed_actions']
        for h in resultado['findings']
    )


def test_rt29_reimportacion_conserva_texto_y_decisiones_humanas(tmp_path):
    _, ruta = _captura(tmp_path)
    trabajo = import_capture(ruta, defaults(), as_of=CORTE)
    id_fila = trabajo.rows[0].id
    decisiones = {'resolution': 'none', 'mail': {'programa_espera': 'omit'}}
    apply(trabajo, [id_fila], review={'OBSERVACION': 'Texto humano que permanece.'}, decisions=decisiones)
    anterior = deepcopy(trabajo.rows[0])
    # Un nuevo corte evita la ruta de caché y vuelve a ejecutar el motor/conciliación.
    nueva = import_capture(ruta, trabajo.config, as_of=date(2026, 10, 10), previous=trabajo)
    assert nueva.as_of == '2026-10-10'
    assert nueva.rows[0].review == anterior.review
    assert nueva.rows[0].decisions == anterior.decisions
    assert nueva.rows[0].review['OBSERVACION'] == 'Texto humano que permanece.'
    assert nueva.rows[0].decisions['resolution'] == 'none'
    assert nueva.rows[0].decisions['mail']['programa_espera'] == 'omit'
    assert all(id_fila != nueva.rows[0].id for id_fila, tipo in automatic_project_selections(nueva))
    borradores = prepare_drafts(nueva, 'programa_espera', directory=tmp_path / 'salida-reimportada')
    assert all(nueva.rows[0].id not in borrador.record_ids for borrador in borradores)


def test_rt31_borrador_guardado_nunca_se_registra_como_envio(tmp_path, monkeypatch):
    from nurus.adapters import outlook

    trabajo = analizar(tmp_path, 'ESPERA', ESPERA, [persona(1) + [45, 'SI']])
    borrador = prepare_drafts(trabajo, 'programa_espera', directory=tmp_path / 'salida')[0]
    llamadas = []

    class CorreoFicticio:
        EntryID = 'borrador-ficticio'
        HTMLBody = '<html><body>Firma ficticia</body></html>'

        def Display(self, modal):
            assert modal is False

        def Save(self):
            llamadas.append('Save')

        def Send(self):
            pytest.fail('El guardado de borradores no puede enviar correos.')

        def Close(self, modo):
            assert modo == 0

    correo = CorreoFicticio()
    correo.Attachments = SimpleNamespace(Add=lambda ruta: llamadas.append('adjunto'))
    carpeta = SimpleNamespace(
        Items=SimpleNamespace(Add=lambda tipo: correo),
        Store=SimpleNamespace(StoreID='almacen-ficticio'),
        FolderPath='Borradores ficticios',
    )

    def carpeta_predeterminada(numero):
        assert numero == 16
        return carpeta

    cliente = ModuleType('win32com.client')
    cliente.Dispatch = lambda nombre: SimpleNamespace(
        Session=SimpleNamespace(Accounts=(), GetDefaultFolder=carpeta_predeterminada)
    )
    win32 = ModuleType('win32com')
    win32.client = cliente
    pythoncom = ModuleType('pythoncom')
    pythoncom.CoInitialize = lambda: None
    pythoncom.CoUninitialize = lambda: None
    for nombre, modulo in (('win32com', win32), ('win32com.client', cliente), ('pythoncom', pythoncom)):
        monkeypatch.setitem(sys.modules, nombre, modulo)
    monkeypatch.setattr(outlook.platform, 'system', lambda: 'Windows')
    comprobante = create_draft(trabajo, borrador, confirmed=True)
    assert comprobante.entry_id == 'borrador-ficticio'
    assert llamadas.count('Save') == 1 and llamadas.count('adjunto') == 1
    assert list(trabajo.receipts.values())[0]['kind'] == 'draft'
    assert list(trabajo.receipts.values())[0]['state'] == 'created'
    assert list(trabajo.receipts.values())[0]['record_ids'] == [trabajo.rows[0].id]
    assert correo.CC == 'ucc_concepcion@pjud.cl'
    assert not any(comprobante.get('state') in ('sent', 'enviado') for comprobante in trabajo.receipts.values())


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason='[CONFLICTO_ABIERTO] R-T32: CT-01/CT-02 desactivadas no producen '
    'explicación específica de exclusión en los hallazgos; sólo el aviso general de no aceptación.',
)
def test_rt32_reglas_contextuales_desactivadas_explican_exclusion(tmp_path):
    _, ruta = _captura(tmp_path)
    configuracion = defaults()
    configuracion['desactivadas'] = ['CT-01', 'CT-02']
    configuracion['umbrales']['mayoria'] = 61
    trabajo = import_capture(ruta, configuracion, as_of=CORTE)
    apply(trabajo, [trabajo.rows[0].id], decisions={'resolution': 'none', 'mail': {'programa_espera': 'omit'}})
    nueva = import_capture(ruta, configuracion, as_of=CORTE, previous=trabajo)
    debe(nueva.config['desactivadas'] == ['CT-01', 'CT-02'], "nueva.config['desactivadas'] == ['CT-01', 'CT-02']")
    debe(nueva.config['umbrales']['mayoria'] == 61, "nueva.config['umbrales']['mayoria'] == 61")
    debe(nueva.rows[0].decisions == trabajo.rows[0].decisions, 'nueva.rows[0].decisions == trabajo.rows[0].decisions')
    hallazgos = nueva.capture_context[nueva.rows[0].id]['findings']
    debe(
        not any(h['rule_id'] in ('CT-01', 'CT-02') and h['status'] == 'APLICA' for h in hallazgos),
        "not any(h['rule_id'] in ('CT-01', 'CT-02') and h['status'] == 'APLICA' for h in hallazgos)",
    )
    assert all(
        any(h['rule_id'] == regla and 'desactiv' in h['reason'].lower() for h in hallazgos)
        for regla in ('CT-01', 'CT-02')
    )
