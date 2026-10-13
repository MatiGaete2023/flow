"""Historia_SITFA.xlsx: hojas exactas, orden, texto literal y verificación del respaldo."""

import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from nurus.capture.historias import HOJAS, LEYENDAS, export_historias

FIXTURE = (
    Path(__file__).resolve().parents[1] / 'descargador' / 'fixtures' / 'historia_ficticia' / 'ventana_historia.html'
)


def manifiesto(tmp_path, html=None, fallida=True):
    carpeta = tmp_path / 'historia'
    carpeta.mkdir()
    data = html if html is not None else FIXTURE.read_bytes()
    sha = hashlib.sha256(data).hexdigest()
    (carpeta / 'h.html').write_bytes(data)
    causas = {
        'RUS|49|100': {
            'tribunal': '49',
            'tribunal_nombre': 'Juzgado ficticio',
            'rit': 'C-99999-2099',
            'causa_id': '100',
            'estado': 'LEIDA',
            'archivo': 'h.html',
            'sha256': sha,
            'leida': '2099-10-13T12:00:00+00:00',
            'desde': '2099-06-13',
            'hasta': '2099-10-13',
            'en_ventana': 3,
            'anteriores': 0,
            'incompletas': 1,
            'advertencias': ['HISTORIA_FILA_INCOMPLETA'],
        }
    }
    if fallida:
        causas['RUS|49|101'] = {
            'tribunal': '49',
            'rit': 'C-99998-2099',
            'causa_id': '101',
            'estado': 'FALLIDA',
            'error': 'HISTORIA_SESION',
        }
    vinculos = [
        {
            'rit': 'C-99999-2099',
            'nombre': 'NOMBRECANARIO',
            'rut': '11.111.111-1',
            'programa': 'Centro',
            'tribunal_codigo': '49',
            'causa_id': '100',
            'ingreso_id': '201',
            'pestana': 'Cumplimiento',
            'modalidad': '2',
            'clave': 'RUS|49|100',
        }
    ]
    estado = {
        'version': 1,
        'tipo': 'HISTORIA_LECTURA',
        'estado': 'COMPLETA',
        'criterio_temporal': 'ventana_4_meses',
        'causas': causas,
        'vinculos': vinculos,
        'descargador': {'version': '2.8.0', 'commit': 'abc'},
        'extension_version': '2.8.0',
    }
    ruta = carpeta / 'historias.json'
    ruta.write_text(json.dumps(estado), encoding='utf-8')
    return ruta


def libro(tmp_path, **kw):
    destino = export_historias(manifiesto(tmp_path, **kw), tmp_path / 'Historia_SITFA.xlsx')
    return load_workbook(destino)


def test_hojas_columnas_y_orden(tmp_path):
    wb = libro(tmp_path)
    assert wb.sheetnames == ['Trámites', 'Causas', 'Vínculos', 'Control']
    for nombre, columnas in HOJAS.items():
        assert tuple(c.value for c in wb[nombre][1]) == columnas
    filas = list(wb['Trámites'].iter_rows(min_row=2, values_only=True))
    assert [f[2] for f in filas] == ['9120', '9119', '9117']
    assert [f[3] for f in filas] == ['13/10/2099', '09/10/2099', '02/10/2099']
    assert all(f[8] in ('Sí', 'No') and f[9] == '100' and f[0] == 'Juzgado ficticio' for f in filas)
    assert any(f[7] == 'Sin referencia.' for f in filas)
    assert len(filas) == 3


def test_causa_fallida_solo_en_causas(tmp_path):
    wb = libro(tmp_path)
    causas = {r[1]: r for r in wb['Causas'].iter_rows(min_row=2, values_only=True)}
    assert causas['C-99998-2099'][2] == 'FALLIDA' and causas['C-99998-2099'][8] == 'HISTORIA_SESION'
    assert causas['C-99999-2099'][3] == 3 and causas['C-99999-2099'][5] == 1
    assert 'C-99998-2099' not in {r[1] for r in wb['Trámites'].iter_rows(min_row=2, values_only=True)}
    vinculo = next(wb['Vínculos'].iter_rows(min_row=2, values_only=True))
    assert vinculo[8] == '100' and vinculo[9] == 'LEIDA'


def test_control_con_leyendas(tmp_path):
    wb = libro(tmp_path)
    control = {r[0]: r[1] for r in wb['Control'].iter_rows(min_row=2, values_only=True) if r[0] != 'Leyenda'}
    assert control['Criterio'] == 'ventana_4_meses' and control['Tope alcanzado'] == 'No'
    assert control['Desde'] == '2099-06-13' and 'FALLIDA: 1' in control['Causas por estado']
    leyendas = [r[1] for r in wb['Control'].iter_rows(min_row=2, values_only=True) if r[0] == 'Leyenda']
    assert tuple(leyendas) == LEYENDAS


def test_formula_queda_como_texto(tmp_path):
    html = FIXTURE.read_bytes().replace('Sin referencia.'.encode(), b'=1+1', 1)
    wb = libro(tmp_path, html=html)
    celdas = [c for fila in wb['Trámites'].iter_rows(min_row=2) for c in fila if c.value == '=1+1']
    assert celdas and all(c.data_type == 's' for c in celdas)


def test_html_alterado_y_destino_existente(tmp_path):
    ruta = manifiesto(tmp_path)
    (ruta.parent / 'h.html').write_bytes(FIXTURE.read_bytes() + b' ')
    with pytest.raises(ValueError):
        export_historias(ruta, tmp_path / 'x.xlsx')
    (ruta.parent / 'h.html').write_bytes(FIXTURE.read_bytes())
    export_historias(ruta, tmp_path / 'x.xlsx')
    with pytest.raises(ValueError):
        export_historias(ruta, tmp_path / 'x.xlsx')


def test_e3_texto_largo_se_conserva_integro_en_cuatro_hojas(tmp_path):
    texto = ('=' + 'Texto ficticio 😀 ' * 5000).rstrip()
    html = FIXTURE.read_bytes().replace(b'windows-1252', b'utf-8').replace(b'Tengase presente canario', texto.encode())
    wb = libro(tmp_path, html=html)
    assert wb.sheetnames == ['Trámites', 'Causas', 'Vínculos', 'Control']
    hoja = wb['Trámites']
    encabezados = [c.value for c in hoja[1]]
    indices = [encabezados.index('Desc. Trámite')] + [
        i for i, e in enumerate(encabezados) if e.startswith('Desc. Trámite (parte ')
    ]
    fila = next(r for r in hoja.iter_rows(min_row=2) if r[2].value == '9120')
    partes = [fila[i].value for i in indices]
    assert ''.join(partes) == texto
    assert all(len(p.encode('utf-16-le')) // 2 <= 30000 for p in partes)
    assert all(fila[i].data_type == 's' for i in indices)
