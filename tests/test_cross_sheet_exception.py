from __future__ import annotations


from openpyxl import Workbook

from nurus.rus import Mode


def _write_compliance(path, *, extra_sheets=()):
    book = Workbook()
    sheet = book.active
    sheet.title = "Cumplimiento"
    sheet.append([
        "DERIVACION", "TRIBUNAL", "NOMBRE", "RIT", "RUT",
        "DIAS DE CUMPLIMIENTO", "DIAS PARA EGRESAR",
    ])
    sheet.append(["PRM Centro", "Juzgado de Laja", "Ana", "C-1", "1-9", 100, 100])
    for name in extra_sheets:
        cross = book.create_sheet(name)
        cross.append([
            "NOMBRE CENTRO", "TRIBUNAL", "NOMBRE MENOR", "RIT", "RUT MENOR",
            "FECHA VENCIMIENTO",
        ])
        cross.append(["PRM Centro", "Juzgado de Laja", "Ana", "C-1", "1-9", "20/09/2026"])
    book.save(path)



def test_hoja2_name_does_not_override_two_valid_cross_sheets(tmp_path):
    from nurus.rus.reader import read_workbook
    path = tmp_path / "ambiguous.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2", "Otro cruce"))
    batch = read_workbook(path, Mode.CUMPLIMIENTO)
    assert not batch.cross_sheet
    assert any(item.startswith("CROSS_SHEET_AMBIGUOUS") for item in batch.warnings)
    explicit = read_workbook(path, Mode.CUMPLIMIENTO, cross_sheet_name="Otro cruce")
    assert explicit.cross_sheet == "Otro cruce"


def test_cross_and_primary_headers_can_be_on_different_rows(tmp_path):
    from openpyxl import load_workbook
    from nurus.rus.reader import read_workbook
    path = tmp_path / "displaced.xlsx"
    _write_compliance(path, extra_sheets=("Hoja2",))
    book = load_workbook(path)
    book["Cumplimiento"].insert_rows(1, 4)
    book["Hoja2"].insert_rows(1, 2)
    book.save(path)
    result = read_workbook(path, Mode.CUMPLIMIENTO, header_row=5)
    assert result.records[0].source.row_number == 6
    assert result.cross_records[0].source.row_number == 4


def test_invalid_hoja2_does_not_hide_a_valid_cross(tmp_path):
    from openpyxl import load_workbook
    from nurus.rus.reader import read_workbook
    path = tmp_path / "cross.xlsx"
    _write_compliance(path, extra_sheets=("Cruce",))
    book = load_workbook(path)
    book.create_sheet("Hoja2").append(["Solo notas"])
    book.save(path)
    assert read_workbook(path, Mode.CUMPLIMIENTO).cross_sheet == "Cruce"


