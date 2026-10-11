"""Exportación: sin archivo residual ante error y sin fórmulas en la copia."""

from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from nurus.personal.config import defaults
from nurus.personal.work import Work

ENCABEZADOS = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'T ESPERA', 'OBSERVACION']


def _trabajo(tmp_path, filas):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'ESPERA'
    hoja.append(ENCABEZADOS)
    for fila in filas:
        hoja.append(fila)
    ruta = tmp_path / 'fuente.xlsx'
    libro.save(ruta)
    trabajo = Work(defaults())
    trabajo.analyze(ruta, 'ESPERA', as_of=date(2026, 10, 9))
    return trabajo


def test_exportacion_fallida_no_deja_archivo_y_restaura_output(tmp_path):
    trabajo = _trabajo(
        tmp_path,
        [
            ['X-1-2026', 'LAJA', '=HYPERLINK("https://ejemplo.invalid","Persona")', '11111111-1', 'PRM UNO', 40, ''],
            ['X-2-2026', 'LAJA', 'Persona Ficticia', '22222222-2', 'PRM UNO', 40, ''],
        ],
    )
    antes = (trabajo.output, trabajo.output_hash)
    destino = tmp_path / 'salida.xlsx'
    with pytest.raises(Exception):
        trabajo.export(destino, backend='portable', reduced_fidelity=True)
    assert not destino.exists()
    assert (trabajo.output, trabajo.output_hash) == antes


def test_observacion_formula_se_exporta_como_texto(tmp_path):
    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '=1+1']])
    destino = tmp_path / 'salida.xlsx'
    trabajo.export(destino, backend='portable', reduced_fidelity=True)
    hoja = load_workbook(destino)['ESPERA']
    columnas = {celda.value: celda.column for celda in hoja[1]}
    celda = hoja.cell(2, columnas['OBSERVACION'])
    assert celda.value == '=1+1' and celda.data_type == 's'
    assert not any(c.data_type == 'f' for fila in hoja.iter_rows() for c in fila)


def test_lectura_fallida_de_la_copia_restaura_output_y_borra(tmp_path, monkeypatch):
    import nurus.personal.work as modulo

    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '']])
    antes = (trabajo.output, trabajo.output_hash)
    destino = tmp_path / 'salida.xlsx'
    original = modulo.Path.read_bytes

    def bloqueado(self):
        if self == destino:
            raise PermissionError('archivo bloqueado')
        return original(self)

    monkeypatch.setattr(modulo.Path, 'read_bytes', bloqueado)
    with pytest.raises(PermissionError):
        trabajo.export(destino, backend='portable', reduced_fidelity=True)
    assert (trabajo.output, trabajo.output_hash) == antes
    assert not destino.exists()


def test_fallo_al_calcular_huella_no_publica_copia(tmp_path, monkeypatch):
    import nurus.services.exports as modulo

    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '']])
    antes = (trabajo.output, trabajo.output_hash)
    destino = tmp_path / 'salida.xlsx'

    def huella_bloqueada(path):
        raise PermissionError('archivo bloqueado al calcular huella')

    monkeypatch.setattr(modulo, '_hash', huella_bloqueada)
    with pytest.raises((PermissionError, modulo.ExportError), match='huella'):
        trabajo.export(destino, backend='portable', reduced_fidelity=True)
    assert (trabajo.output, trabajo.output_hash) == antes
    assert not destino.exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == ['fuente.xlsx']


def test_fallo_al_abrir_temporal_para_copiar_no_deja_archivo_vacio(tmp_path, monkeypatch):
    import nurus.services.exports as modulo

    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '']])
    destino = tmp_path / 'salida.xlsx'
    abrir = Path.open

    def temporal_bloqueado(self, mode='r', *args, **kwargs):
        if self.name.startswith('nurus-salida-') and mode == 'rb' and destino.exists():
            raise PermissionError('temporal bloqueado al copiar')
        return abrir(self, mode, *args, **kwargs)

    monkeypatch.setattr(Path, 'open', temporal_bloqueado)
    with pytest.raises(modulo.ExportError, match='bloqueado'):
        trabajo.export(destino, backend='portable', reduced_fidelity=True)
    assert not destino.exists()
    assert not trabajo.output and not trabajo.output_hash
    assert sorted(p.name for p in tmp_path.iterdir()) == ['fuente.xlsx']


def test_copia_portable_no_muestra_apostrofo_en_la_propuesta(tmp_path):
    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '']])
    trabajo.rows[0].observation = '- revisar =1+1'
    destino = tmp_path / 'salida.xlsx'
    trabajo.export(destino, backend='portable', reduced_fidelity=True)
    hoja = load_workbook(destino)['ESPERA']
    columnas = {celda.value: celda.column for celda in hoja[1]}
    for titulo in ('NURUS_PROPUESTA', 'OBSERVACION'):
        celda = hoja.cell(2, columnas[titulo])
        assert celda.value == '- revisar =1+1' and celda.data_type == 's', (titulo, celda.value)


def test_copia_portable_revisada_conserva_observacion_literal_sin_apostrofo(tmp_path):
    from nurus.personal.record_edits import apply

    trabajo = _trabajo(tmp_path, [['X-1-2026', 'LAJA', 'Persona Ficticia', '11111111-1', 'PRM UNO', 40, '']])
    apply(trabajo, [trabajo.rows[0].id], review={'OBSERVACION': '=texto revisado'})
    destino = tmp_path / 'salida.xlsx'
    trabajo.export(destino, backend='portable', reduced_fidelity=True)
    copia = load_workbook(destino)
    try:
        hoja = copia['ESPERA']
        columnas = {celda.value: celda.column for celda in hoja[1]}
        for titulo in ('NURUS_OBSERVACION_FINAL', 'OBSERVACION'):
            celda = hoja.cell(2, columnas[titulo])
            assert celda.value == '=texto revisado' and celda.data_type == 's'
    finally:
        copia.close()
