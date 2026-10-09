"""Regresiones de H01–H08 y cancelación: fuentes exclusivamente sintéticas."""
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
import random
import string

from openpyxl import Workbook, load_workbook
import pytest

from nurus.capture import CaptureError, read_capture, write_capture
from nurus.capture.acceptance import validate_acta
from nurus.capture.coverage import record_coverages
from nurus.capture.fixtures import fixture
from nurus.capture.requests import MAX_BYTES, read_request, write_request
from nurus.personal.capture_input import import_capture
from nurus.personal.capture_products import evidence_text, export_evidence, request_document
from nurus.personal.capture_reconciliation import ReconciliationRequired
from nurus.personal.config import defaults
from nurus.personal.operations import Operation, OperationCancelled, operation_scope
from nurus.personal.outputs import Draft
from nurus.personal.product_state import stamp, stale
from nurus.personal.record_edits import apply, undo
from nurus.personal.work import Work
from nurus.services.excel_cells import append_literal_row
from test_capture_acceptance import executed

CUT = date(2026, 10, 8)


def work_from(tmp_path, *, partial=False):
    metadata, tables = fixture(partial=partial)
    path = write_capture(tmp_path / 'base.zip', metadata, tables)
    return import_capture(path, defaults(), as_of=CUT), metadata, tables


@pytest.mark.parametrize('text,valid', [('x'*32767, True), ('x'*32768, False),
                                     ('😀'*16383+'x', True), ('😀'*16384, False)])
def test_excel_limit_checked_before_cell_assignment(text, valid):
    book = Workbook(); sheet = book.active
    if valid:
        append_literal_row(sheet, ['=1+1', text], row_number=1)
        assert sheet['B1'].value == text and sheet['A1'].data_type == 's'
    else:
        with pytest.raises(ValueError, match='UTF-16'):
            append_literal_row(sheet, [text], row_number=1)
        assert sheet._cells == {}
    book.close()


def test_oversized_capture_does_not_replace_previous(tmp_path):
    old, metadata, tables = work_from(tmp_path)
    before = deepcopy(old.__dict__)
    raw = json.loads(tables['registros'][0]['normalized_json'])
    raw['NOMBRE'] = ''.join(random.Random(1).choices(string.ascii_letters, k=33000))
    tables['registros'][0].update(raw_json=json.dumps(raw), normalized_json=json.dumps(raw))
    path = write_capture(tmp_path/'long.zip', metadata, tables)
    with pytest.raises(ValueError, match='límite Excel'):
        import_capture(path, defaults(), as_of=CUT, previous=old)
    assert old.__dict__ == before


def traditional(work):
    return [f['rule_id'] for f in work.capture_context[work.rows[0].id]['findings']
            if not f['rule_id'].startswith('CT-')]


def test_findings_edit_undo_recovery_and_failed_edit_are_transactional(tmp_path):
    work, _, _ = work_from(tmp_path)
    original = traditional(work)
    old = apply(work, [work.rows[0].id], review={'OBSERVACION':'Texto humano'}, overrides={'espera':'0'})
    assert traditional(work) == work.rows[0].rules != original
    assert work.rows[0].review['OBSERVACION'] == 'Texto humano'
    work.save(tmp_path/'session'); restored = Work.load(tmp_path/'session')
    assert traditional(restored) == restored.rows[0].rules
    undo(work, old)
    assert traditional(work) == work.rows[0].rules == original
    before = deepcopy(work.__dict__)
    with pytest.raises(ValueError):
        apply(work, [work.rows[0].id], overrides={'no_column':'x'})
    assert work.__dict__ == before


def test_excel_source_edit_refreshes_findings_without_erasing_review(tmp_path):
    work, _, _ = work_from(tmp_path)
    work.rows[0].review['OBSERVACION'] = 'Texto humano'
    work.export(tmp_path/'output.xlsx', backend='portable', reduced_fidelity=True)
    book = load_workbook(work.output); sheet = book['Trabajo']
    headers = {c.value:c.column for c in sheet[1]}
    sheet.cell(2, headers[work.mapping['espera']], '0'); book.save(work.output); book.close()
    assert work.refresh()
    assert traditional(work) == work.rows[0].rules
    assert work.rows[0].review['OBSERVACION'] == 'Texto humano'


def test_equivalent_capture_and_changed_ids_preserve_decisions_products_and_two_ingresos(tmp_path):
    old, metadata, tables = work_from(tmp_path)
    old.rows[0].review['OBSERVACION'] = 'Humano 0'; old.rows[0].excluded = True
    old.rows[1].review['OBSERVACION'] = 'Humano 1'
    draft = Draft('Ficticio','Texto',kind='programa_espera',record_ids=['r0']); stamp(old,draft)
    from dataclasses import asdict
    old.session = {'drafts':[asdict(draft)],'projects':[],'prepared_selection':['r0']}
    metadata['downloader']['commit'] = 'BUILD_COMPLEMENTARIO'
    for row in tables['registros']: row['record_id'] = 'new-' + row['record_id']
    path = write_capture(tmp_path/'new.zip', metadata, tables)
    new = import_capture(path, defaults(), as_of=CUT, previous=old)
    assert new.rows[0].review['OBSERVACION'] == 'Humano 0' and new.rows[0].excluded
    assert new.rows[1].review['OBSERVACION'] == 'Humano 1'
    assert new.session['drafts'][0]['record_ids'] == ['new-r0']
    assert new.session['prepared_selection'] == ['new-r0']
    assert stale(new, Draft(**new.session['drafts'][0])) and not new.output
    assert old.rows[0].id == 'r0' and old.session['prepared_selection'] == ['r0']


def test_complementary_events_retain_review_and_changed_data_requires_choice(tmp_path):
    old, metadata, tables = work_from(tmp_path)
    old.rows[0].review['OBSERVACION'] = 'Humano'; old.rows[0].excluded = True
    before = deepcopy(old.__dict__)
    tables['bitacoras'].append({**tables['bitacoras'][0], 'event_id':'new-event','event_at':'2026-10-07'})
    complementary = import_capture(write_capture(tmp_path/'complement.zip',metadata,tables),defaults(),as_of=CUT,previous=old)
    assert complementary.rows[0].excluded and complementary.rows[0].review['OBSERVACION']=='Humano'
    raw = json.loads(tables['registros'][0]['normalized_json']); raw['T ESPERA']='0'
    tables['registros'][0]['normalized_json']=json.dumps(raw)
    path = write_capture(tmp_path/'conflict.zip',metadata,tables)
    with pytest.raises(ReconciliationRequired) as caught:
        import_capture(path,defaults(),as_of=CUT,previous=old)
    assert caught.value.conflicts[0].can_keep and old.__dict__==before
    keep=import_capture(path,defaults(),as_of=CUT,previous=old,resolutions={'r0':'KEEP'})
    reset=import_capture(path,defaults(),as_of=CUT,previous=old,resolutions={'r0':'RESET'})
    assert keep.rows[0].excluded and keep.rows[0].review['OBSERVACION']=='Humano'
    assert not reset.rows[0].excluded and reset.rows[0].review.get('OBSERVACION')!='Humano'


def test_ambiguous_rows_or_missing_record_never_auto_copy_decisions(tmp_path):
    old, metadata, tables = work_from(tmp_path); old.rows[0].excluded=True
    tables['registros'].append({**tables['registros'][0], 'record_id':'duplicate'})
    path=write_capture(tmp_path/'ambiguous.zip',metadata,tables)
    with pytest.raises(ReconciliationRequired) as caught:
        import_capture(path,defaults(),as_of=CUT,previous=old)
    assert all(not c.can_keep for c in caught.value.conflicts)
    answers={c.record_id:'RESET' for c in caught.value.conflicts}
    resolved=import_capture(path,defaults(),as_of=CUT,previous=old,resolutions=answers)
    assert not any(r.excluded for r in resolved.rows)
    tables['registros']=tables['registros'][1:4]
    path=write_capture(tmp_path/'missing.zip',metadata,tables)
    with pytest.raises(ReconciliationRequired) as caught:
        import_capture(path,defaults(),as_of=CUT,previous=old)
    assert caught.value.conflicts[0].record_id=='missing:r0'


def test_query_row_and_table_cannot_overclaim_coverage(tmp_path):
    metadata,tables=fixture()
    tables['consultas'][0].update(coverage='PARCIAL',expected_pages='2',saved_pages='1')
    with pytest.raises(CaptureError,match='cobertura'):
        write_capture(tmp_path/'bad-row.zip',metadata,tables)
    for name in ('registros','bitacoras'):
        for row in tables[name]:
            if row['query_id']=='q1':row['coverage']='PARCIAL'
    with pytest.raises(CaptureError,match='tabla'):
        write_capture(tmp_path/'bad-table.zip',metadata,tables)
    metadata['table_states'].update(registros='PARCIAL',bitacoras='PARCIAL')
    work=import_capture(write_capture(tmp_path/'partial.zip',metadata,tables),defaults(),as_of=CUT)
    request=request_document(work)['requests']
    assert [r['ingreso_id'] for r in request if r['source']=='bitacoras']==['ingreso-ficticio-0']


def test_complete_empty_failed_and_partial_scopes_are_distinguished(tmp_path):
    metadata,tables=fixture()
    tables['bitacoras']=[]; metadata['table_states']['bitacoras']='PARCIAL'
    for i,state in enumerate(('COMPLETA','VACIA_COMPROBADA','FALLIDA','PARCIAL')):
        record=tables['registros'][i]
        query={**tables['consultas'][i], 'query_id':f'd{i}','screen':'bitacoras','coverage':state}
        query['filters_json']=json.dumps({'capture_tables':['bitacoras'],'remote_identity':
                            {k:record[k] for k in ('source_system','court_code','cause_id','ingreso_id')}})
        if state!='COMPLETA':query.update(saved_pages='0')
        tables['consultas'].append(query)
        if state=='COMPLETA':tables['bitacoras'].append({**record,
            **{k:'' for k in ('response','response_at')}, 'event_id':'complete','event_at':'2026-10-01',
            'time_precision':'DIA','time_zone':'DESCONOCIDA','origin':'CENTRO','origin_verified':'true',
            'author':'Ficticio','event_type':'Observación','text':'Ficticio','query_id':f'd{i}'})
    # Only the diary columns, not the extra record fields.
    from nurus.capture.contract import TABLES
    tables['bitacoras']=[{k:r[k] for k in TABLES['bitacoras']} for r in tables['bitacoras']]
    metadata['table_states']['consultas']='PARCIAL'
    path=write_capture(tmp_path/'mixed.zip',metadata,tables)
    work=import_capture(path,defaults(),as_of=CUT)
    assert [r['ingreso_id'] for r in request_document(work)['requests'] if r['source']=='bitacoras']==['ingreso-ficticio-2','ingreso-ficticio-3']
    assert work.capture_coverage['r1']['bitacoras']=='VACIA_COMPROBADA'


@pytest.mark.parametrize('counts',[{'observed':1,'matched':2},
    {'observed':2,'matched':1,'failed':0,'unknown':0}, {'expected':2,'observed':1},
    {'observed':1,'unknown':1}])
def test_impossible_acceptance_counts_cannot_pass(counts):
    data=executed();data['records'][0]['counts']=counts
    with pytest.raises(CaptureError):validate_acta(data)


def test_count_partition_is_optional_but_coherent():
    data=executed(); data['records'][0]['counts']={'expected':2,'observed':2,'matched':2,'failed':0,'unknown':0}
    assert validate_acta(data)==data
    data['records'][0]['counts']={}; assert validate_acta(data)==data


def request(reason='ficticio'):
    return {'schema_id':'CSMP_REQUEST','schema_version':'1.0.0','capture_sha256':'a'*64,
            'as_of':'2026-10-08','requests':[{'record_id':'r0','source_system':'RUS','court_code':'FICTICIO',
             'cause_id':'c0','ingreso_id':'i0','source':'bitacoras','reason':reason}]}


def test_request_writer_and_reader_share_byte_limit(tmp_path,monkeypatch):
    from nurus.capture import requests
    data=request('😀'*50)
    size=len(json.dumps(data,ensure_ascii=False,indent=2).encode())
    monkeypatch.setattr(requests,'MAX_BYTES',size)
    path=write_request(tmp_path/'exact.json',data)
    assert read_request(path)==data
    monkeypatch.setattr(requests,'MAX_BYTES',size-1)
    with pytest.raises(CaptureError):write_request(tmp_path/'too-large.json',data)
    assert not (tmp_path/'too-large.json').exists()
    with pytest.raises(CaptureError):read_request(path)


def test_reported_large_request_is_rejected_before_publication(tmp_path):
    data=request('x'*2000)
    data['requests']=[{**data['requests'][0],'record_id':f'r{i}','ingreso_id':f'i{i}'} for i in range(2200)]
    assert len(json.dumps(data,ensure_ascii=False,indent=2).encode())>MAX_BYTES
    with pytest.raises(CaptureError):write_request(tmp_path/'oversized.json',data)
    assert not list(tmp_path.iterdir())


def test_incident_survives_import_recovery_view_and_excel(tmp_path):
    metadata,tables=fixture()
    incident={'capture_id':metadata['capture_id'],'query_id':'q1','record_id':'r0','incident_id':'i0',
              'phase':'LECTURA','code':'COTEJO','severity':'ADVERTENCIA','message':'=Advertencia ficticia','recoverable':'true'}
    tables['incidencias']=[incident];metadata['table_states']['incidencias']='COMPLETA'
    work=import_capture(write_capture(tmp_path/'incident.zip',metadata,tables),defaults(),as_of=CUT)
    work.save(tmp_path/'session');work=Work.load(tmp_path/'session')
    assert work.capture_incidents==[incident]
    assert incident['message'] in evidence_text(work,work.rows[0])
    assert incident['message'] not in evidence_text(work,work.rows[1])
    export_evidence(work,tmp_path/'evidence.xlsx')
    book=load_workbook(tmp_path/'evidence.xlsx');cell=book['Incidencias'].cell(2,8)
    assert cell.value==incident['message'] and cell.data_type=='s';book.close()


def test_workbook_stage_cancels_without_touching_previous(tmp_path):
    old,_,_=work_from(tmp_path);before=deepcopy(old.__dict__)
    metadata,tables=fixture(250)
    path=write_capture(tmp_path/'large.zip',metadata,tables)
    op=Operation(lambda event:op.cancelled.set() if event[1]=='Construyendo copia Excel' else None)
    with operation_scope(op),pytest.raises(OperationCancelled):
        import_capture(path,defaults(),as_of=CUT,previous=old)
    assert old.__dict__==before


def test_operational_limit_is_independent_of_contract_limit(tmp_path,monkeypatch):
    from nurus.personal import capture_input
    old,_,_=work_from(tmp_path);before=deepcopy(old.__dict__)
    metadata,tables=fixture(5);path=write_capture(tmp_path/'over-limit.zip',metadata,tables)
    assert len(read_capture(path).rows('registros'))==5
    monkeypatch.setattr(capture_input,'MAX_ANALYSIS_RECORDS',4)
    with pytest.raises(ValueError,match='límite operativo'):
        import_capture(path,defaults(),as_of=CUT,previous=old)
    assert old.__dict__==before


def test_primary_order_and_unverified_physical_copy_keep_only_matching_edits(tmp_path):
    old,metadata,tables=work_from(tmp_path)
    old.rows[0].review['OBSERVACION']='Humano 0';old.rows[1].review['OBSERVACION']='Humano 1'
    metadata['capture_id']='capture-new'
    for name,rows in tables.items():
        for row in rows:
            if 'capture_id' in row:row['capture_id']=metadata['capture_id']
    tables['registros'].reverse()
    new=import_capture(write_capture(tmp_path/'reordered.zip',metadata,tables),defaults(),as_of=CUT,previous=old)
    assert {r.id:r.review.get('OBSERVACION') for r in new.rows}['r0']=='Humano 0'
    # Sin ingreso probado solo se admite la misma fila física íntegra, no RIT/nombre.
    metadata,tables=fixture(1)
    tables['registros'][0]['identity_status']='INCIERTA'
    old=import_capture(write_capture(tmp_path/'uncertain.zip',metadata,tables),defaults(),as_of=CUT)
    old.rows[0].excluded=True;tables['registros'][0]['record_id']='new-physical'
    new=import_capture(write_capture(tmp_path/'same-physical.zip',metadata,tables),defaults(),as_of=CUT,previous=old)
    assert new.rows[0].excluded
    tables['registros'][0]['original_row']='999'
    with pytest.raises(ReconciliationRequired):
        import_capture(write_capture(tmp_path/'different-physical.zip',metadata,tables),defaults(),as_of=CUT,previous=old)
