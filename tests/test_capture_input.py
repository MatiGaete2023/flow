from copy import deepcopy
from datetime import date
from pathlib import Path
import sys
import json
import pytest
from openpyxl import load_workbook
from nurus.capture import write_capture, CaptureError
from nurus.personal.capture_input import import_capture
from nurus.personal.config import defaults
from nurus.personal.work import Work
from test_capture_contract import fixture, mutate


def test_four_modalities_edits_reimport_and_recovery(tmp_path):
    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    before_modules = set(sys.modules)
    work = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    assert len(work.rows) == 4 and all(r.observation for r in work.rows)
    assert not work.activity_reviewed and not work.signed_activity
    work.rows[0].review = {'OBSERVACION': 'Texto humano', 'TT': 'T', 'CC': 1, 'RES': 'PC_IE'}
    work.rows[0].decisions['resolution'] = 'PC_IE'
    restored = import_capture(path, defaults(), as_of=date(2026, 10, 8), previous=work)
    assert restored.rows[0].review == work.rows[0].review
    assert all(ctx['class'] == 'NO_EVALUABLE' for ctx in restored.capture_context.values())
    restored.save(tmp_path / 'session')
    loaded = Work.load(tmp_path / 'session')
    assert loaded.capture_digest == work.capture_digest and loaded.rows[0].review['OBSERVACION'] == 'Texto humano'
    assert not any(name in set(sys.modules) - before_modules for name in ('nurus.personal.download_link',))


def test_corrupt_import_preserves_previous_and_separates_ingresos(tmp_path):
    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    previous = import_capture(path, defaults())
    previous.rows[0].review['OBSERVACION'] = 'Edición'
    before = deepcopy(previous.__dict__)
    mutate(path, lambda items: items.update({'escape': b'bad'}))
    with pytest.raises(CaptureError):
        import_capture(path, defaults(), previous=previous)
    assert previous.__dict__ == before
    for rid, record in previous.capture_sources.items():
        assert all(e['ingreso_id'] == record['ingreso_id'] for e in previous.capture_context[rid]['recent'])


def test_export_and_refresh_portable_preserve_human_fields(tmp_path):
    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    work = import_capture(path, defaults())
    work.rows[0].review = {'OBSERVACION': 'Humano', 'TT': '1', 'CC': '2', 'RES': 'PC_IE'}
    work.export(tmp_path / 'out.xlsx', backend='portable', reduced_fidelity=True)
    book = load_workbook(work.output)
    sheet = book['Trabajo']
    headers = {c.value: c.column for c in sheet[1]}
    assert sheet.cell(2, headers['OBSERVACION']).value == 'Humano'
    assert sheet.cell(2, headers['CC']).value == '2'
    book.close()
    assert work.refresh() is False


def test_reanalysis_new_cut_preserves_human_decisions_and_invalidates_product(tmp_path):
    from nurus.personal.outputs import Draft
    from nurus.personal.product_state import stamp, stale

    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    old = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    old.rows[0].review['OBSERVACION'] = 'Humano'
    draft = Draft('Ficticio', 'Ficticio', kind='programa_espera', record_ids=[old.rows[0].id])
    stamp(old, draft)
    new = import_capture(path, defaults(), as_of=date(2026, 10, 9), previous=old)
    assert new.rows[0].review['OBSERVACION'] == 'Humano' and stale(new, draft)
    assert old.as_of == '2026-10-08'


def test_old_rules_version_is_not_reused_from_session_cache(tmp_path):
    from nurus.capture.temporal import RULESET_VERSION

    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    old = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    old.ruleset_version = 'anterior'
    old.rows[0].review['OBSERVACION'] = 'Humano'
    old.rows[0].excluded = True
    new = import_capture(path, defaults(), as_of=date(2026, 10, 8), previous=old)
    assert (
        new.ruleset_version == RULESET_VERSION
        and new.rows[0].review['OBSERVACION'] == 'Humano'
        and new.rows[0].excluded
    )


def test_mixed_modes_export_refresh_and_individual_recalculation(tmp_path):
    m, t = fixture()
    for index, mode in enumerate(('ESPERA', 'CUMPLIMIENTO', 'INFORMES', 'ESPERA')):
        r = t['registros'][index]
        raw = json.loads(r['normalized_json'])
        if mode == 'CUMPLIMIENTO':
            raw.update({'DIAS DE CUMPLIMIENTO': '50', 'DIAS PARA EGRESAR': '10', 'FEC.EGRESO PROYECTADO': '2026-10-18'})
        if mode == 'INFORMES':
            raw['FECHA VENCIMIENTO'] = '2026-10-15'
        r.update(record_type=mode, raw_json=json.dumps(raw), normalized_json=json.dumps(raw))
    work = import_capture(write_capture(tmp_path / 'mix.zip', m, t), defaults(), as_of=date(2026, 10, 8))
    assert any(rule.startswith('CUMPLIMIENTO.') for rule in work.rows[1].rules)
    assert any(rule.startswith('INFORMES.') for rule in work.rows[2].rules)
    work.export(tmp_path / 'mix.xlsx', backend='portable', reduced_fidelity=True)
    book = load_workbook(work.output)
    sheet = book['Trabajo']
    headers = {c.value: c.column for c in sheet[1]}
    sheet.cell(work.rows[2].source_row, headers['OBSERVACION'], 'Edición informe')
    book.save(work.output)
    book.close()
    assert work.refresh() is True
    assert 'vencimiento' in work.mapping and 'dias_egresar' in work.mapping
    from nurus.personal.record_edits import recalculate

    recalculate(work, work.rows[2])
    assert any(rule.startswith('INFORMES.') for rule in work.rows[2].rules)
    assert work.rows[2].review['OBSERVACION'] == 'Edición informe'


def test_calendar_only_and_mirrored_physical_source(tmp_path):
    from nurus.capture.contract import COMMON

    m, t = fixture()
    source = t['registros'][0]
    calendar = {k: source[k] for k in COMMON}
    raw = {
        'RIT': 'FICT-0-2026',
        'TRIBUNAL': 'Laja',
        'NOMBRE': 'Persona Ficticia',
        'DERIVACION': 'PRM FICTICIO',
        'FECHA VENCIMIENTO': '2026-10-15',
    }
    calendar.update(
        calendar_id='cal0',
        event_at='2026-10-15',
        event_type='INFORME',
        status='PENDIENTE',
        association='NO_ASOCIADO',
        raw_json=json.dumps(raw),
        normalized_json=json.dumps(raw),
    )
    t['calendarios'] = [calendar]
    m['table_states']['calendarios'] = 'COMPLETA'
    mirrored = import_capture(write_capture(tmp_path / 'mirrored.zip', m, t), defaults(), as_of=date(2026, 10, 8))
    assert len(mirrored.rows) == 4
    t['registros'] = []
    m['table_states']['registros'] = 'VACIA_COMPROBADA'
    work = import_capture(write_capture(tmp_path / 'calendar.zip', m, t), defaults(), as_of=date(2026, 10, 8))
    assert len(work.rows) == 1 and work.rows[0].id == 'calendar:cal0'
    record = work.capture_sources[work.rows[0].id]
    assert record['source_table'] == 'calendarios' and record['identity_status'] == 'INCIERTA'
    assert work.capture_context[work.rows[0].id]['class'] == 'NO_EVALUABLE'
    work.export(tmp_path / 'calendar.xlsx', backend='portable', reduced_fidelity=True)
    assert Path(work.output).is_file()
    calendar.update(event_type='MEDIDA', association='VERIFICADA')
    other = import_capture(write_capture(tmp_path / 'measure.zip', m, t), defaults(), as_of=date(2026, 10, 8))
    assert other.capture_sources[other.rows[0].id]['record_type'] == 'OTRO'
    assert not other.rows[0].observation and not other.rows[0].rules and not other.rows[0].actions
    assert all(due is None for due in other.cross_dates.values())
    from nurus.personal.record_edits import recalculate

    other.rows[0].review['OBSERVACION'] = 'Decisión manual'
    recalculate(other, other.rows[0])
    assert not other.rows[0].observation and not other.rows[0].actions
    assert other.rows[0].review['OBSERVACION'] == 'Decisión manual'


def test_cumplimiento_package_has_no_spurious_cross_or_source_warnings(tmp_path):
    m, t = fixture()
    for record in t['registros']:
        raw = json.loads(record['normalized_json'])
        raw.update({'DIAS DE CUMPLIMIENTO': '50', 'DIAS PARA EGRESAR': '10', 'FEC.EGRESO PROYECTADO': '2026-10-18'})
        record.update(record_type='CUMPLIMIENTO', raw_json=json.dumps(raw), normalized_json=json.dumps(raw))
    work = import_capture(write_capture(tmp_path / 'c.zip', m, t), defaults(), as_of=date(2026, 10, 8))
    # El cruce de un paquete llega por registro: no se avisa «Sin cruce utilizable».
    assert not any('Sin cruce utilizable' in w for w in work.warnings)
    assert 'Fuentes faltantes: ' not in work.warnings


def test_reimportar_sin_conflictos_rechaza_respuestas_ajenas_incluso_en_cache(tmp_path):
    m, t = fixture()
    path = write_capture(tmp_path / 'full.zip', m, t)
    previous = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    before = deepcopy(previous.__dict__)
    with pytest.raises(ValueError, match='ajena o innecesaria'):
        import_capture(path, defaults(), as_of=date(2026, 10, 8), previous=previous, resolutions={'r0': 'RESET'})
    assert previous.__dict__ == before
