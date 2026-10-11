"""Reglas de mayoría de edad en Espera y Cumplimiento, con fecha de corte fija."""

from datetime import date

import pytest
from openpyxl import Workbook

from nurus.personal.config import defaults
from nurus.personal.work import Work

CORTE = date(2026, 10, 9)
CASOS = [
    ("2008-10-09", "MAYORIA", "9 de octubre de 2026"),
    ("09-10-2008", "MAYORIA", "9 de octubre de 2026"),
    ("2008-10-10", "PROXIMA", "10 de octubre de 2026"),
    ("10/10/2008", "PROXIMA", "10 de octubre de 2026"),
    ("2008-12-08", "PROXIMA", "8 de diciembre de 2026"),  # 60 días.
    ("2008-12-09", "NINGUNA", ""),  # 61 días.
    ("2008-02-29", "MAYORIA", "28 de febrero de 2026"),
    ("2009-01-15", "NINGUNA", ""),
    ("", "NINGUNA", ""),
]
ENCABEZADOS = {
    "ESPERA": ["RIT", "TRIBUNAL", "NOMBRE", "RUT", "NOMBRE CENTRO", "T ESPERA", "FEC. NACIMIENTO"],
    "CUMPLIMIENTO": ["RIT", "TRIBUNAL", "NOMBRE", "RUT", "NOMBRE CENTRO", "FEC. NACIMIENTO"],
}


def _analizar(tmp_path, modo):
    libro = Workbook()
    hoja = libro.active
    hoja.title = modo
    hoja.append(ENCABEZADOS[modo])
    for numero, (nacimiento, _, _) in enumerate(CASOS, 1):
        fila = [f"X-{numero}-2026", "LAJA", f"Persona Ficticia {numero}", "11111111-0", "PRM uno"]
        if modo == "ESPERA":
            fila.append(40)
        fila.append(nacimiento)
        hoja.append(fila)
    ruta = tmp_path / f"edad_{modo}.xlsx"
    libro.save(ruta)
    contenido = ruta.read_bytes()
    trabajo = Work(defaults())
    trabajo.analyze(ruta, modo, as_of=CORTE)
    assert trabajo.as_of == CORTE.isoformat()
    assert trabajo.receipts == {}  # El análisis sólo propone; no registra egresos.
    assert ruta.read_bytes() == contenido
    return trabajo.rows


def _clase(observacion):
    if "alcanzó la mayoría de edad" in observacion:
        return "MAYORIA"
    if "alcanzará la mayoría de edad" in observacion:
        return "PROXIMA"
    return "NINGUNA"


@pytest.mark.parametrize("modo", ["ESPERA", "CUMPLIMIENTO"])
def test_rt02_limites_1_60_61_dias_y_formatos_sitfa(tmp_path, modo):
    filas = _analizar(tmp_path, modo)
    assert len(filas) == len(CASOS)
    for fila, (nacimiento, esperado, fecha_texto) in zip(filas, CASOS):
        assert _clase(fila.observation) == esperado, (modo, nacimiento, fila.observation)
        assert ("COMUN.MAYORIA_EDAD" in fila.rules) == (esperado == "MAYORIA")
        assert ("COMUN.PROXIMA_MAYORIA" in fila.rules) == (esperado == "PROXIMA")
        if esperado != "MAYORIA":
            assert "Se sugiere egresar la medida." not in fila.observation
        if fecha_texto:
            assert fecha_texto in fila.observation


def test_rt03_mayoria_cumplida_sugiere_sin_ejecutar_egreso(tmp_path):
    filas = _analizar(tmp_path, "ESPERA")
    assert "Se sugiere egresar la medida." in filas[0].observation
    assert "Se sugiere egresar la medida." not in filas[2].observation


def test_umbral_de_proxima_mayoria_respeta_configuracion(tmp_path):
    configuracion = defaults()
    configuracion["umbrales"]["mayoria"] = 120
    libro = Workbook()
    hoja = libro.active
    hoja.title = "CUMPLIMIENTO"
    hoja.append(ENCABEZADOS["CUMPLIMIENTO"])
    hoja.append(["X-1-2026", "LAJA", "Persona Ficticia", "11111111-0", "PRM uno", "2009-01-15"])
    ruta = tmp_path / "umbral.xlsx"
    libro.save(ruta)
    trabajo = Work(configuracion)
    trabajo.analyze(ruta, "CUMPLIMIENTO", as_of=CORTE)
    assert _clase(trabajo.rows[0].observation) == "PROXIMA"


def analizar(tmp_path, modo, encabezados, filas, cruce=None, configuracion=None, corte=CORTE):
    libro = Workbook()
    hoja = libro.active
    hoja.title = modo
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    if cruce:
        hoja2 = libro.create_sheet('Hoja2')
        for fila in cruce:
            hoja2.append(fila)
    ruta = tmp_path / f'rt_{modo.lower()}.xlsx'
    libro.save(ruta)
    trabajo = Work(configuracion or defaults())
    trabajo.analyze(ruta, modo, as_of=corte)
    assert trabajo.as_of == corte.isoformat()
    return trabajo


def persona(numero, tribunal='LAJA', programa='PRM UNO'):
    return [f'X-{numero}-2026', tribunal, f'Persona Ficticia {numero}', '11111111-0', programa]
