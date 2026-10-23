"""D6: con consolidado en el paquete, la copia del Asistente parte de ese libro; sin él, queda como antes."""

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import json
import pytest
from openpyxl import Workbook, load_workbook
from nurus.capture import write_capture
from nurus.personal.capture_input import import_capture
from nurus.personal.config import defaults
from nurus.personal.work import _consolidated_base
from test_capture_contract import fixture

ORIGIN = (
    'SITFA_LOTE',
    'SITFA_CONSULTA',
    'SITFA_ARCHIVO',
    'SITFA_PAGINA',
    'SITFA_FILA',
    'SITFA_SHA256',
    'SITFA_FECHA_DESCARGA',
)


def consolidado(
    registros,
    *,
    hojas=1,
    invertir=True,
    duplicar=False,
    incompleto=False,
    identidad_alterada=False,
    transformar=None,
    quitar_columnas=(),
):
    book = Workbook()
    resumen = book.active
    resumen.title = 'Resumen'
    resumen.append(['Resultado'])
    destinos = [book.create_sheet(f'Datos{i}') for i in range(hojas)]
    encabezados = (
        ['RIT']
        if incompleto
        else list(dict.fromkeys(key for registro in registros for key in json.loads(registro['normalized_json'])))
    )
    encabezados = [key for key in encabezados if key not in quitar_columnas]
    for sheet in destinos:
        sheet.append([*encabezados, *ORIGIN])
    lista = list(reversed(registros)) if invertir else list(registros)
    for i, r in enumerate(lista):
        datos = json.loads(r['normalized_json'])
        if identidad_alterada:
            datos['NOMBRE'] = 'Persona Diferente'
        if transformar:
            transformar(datos)
        destinos[i % hojas].append(
            [
                *[datos.get(h, '') for h in encabezados],
                '',
                '',
                r['original_file'],
                '1',
                r['original_row'],
                r['source_sha256'],
                '',
            ]
        )
    if duplicar:
        r = lista[0]
        destinos[0].append(
            [
                *[json.loads(r['normalized_json']).get(h, '') for h in encabezados],
                '',
                '',
                r['original_file'],
                '1',
                r['original_row'],
                r['source_sha256'],
                '',
            ]
        )
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def importar(tmp_path, **opciones):
    metadata, tables = fixture()
    raw = consolidado(tables['registros'], **opciones)
    path = write_capture(tmp_path / 'con.zip', metadata, tables, consolidated=raw)
    return import_capture(path, defaults(), as_of=date(2026, 10, 8)), tables, raw


def test_la_base_apunta_a_la_hoja_y_a_las_filas_del_consolidado(tmp_path):
    work, tables, raw = importar(tmp_path)
    base = _consolidated_base(work)
    assert base and base[0] == raw and base[1] == 'Datos0'
    by_id = {r['record_id']: r for r in tables['registros']}
    book = load_workbook(BytesIO(raw))
    sheet = book['Datos0']
    encabezado = [c.value for c in sheet[1]]
    col_fila = encabezado.index('SITFA_FILA')
    for record_id, numero in base[2].items():
        assert str(sheet.cell(numero, col_fila + 1).value) == by_id[record_id]['original_row']
    book.close()
    destino = tmp_path / 'copia.xlsx'
    work.export(destino, backend='portable', reduced_fidelity=True)
    copia = load_workbook(destino)
    try:
        assert copia.sheetnames[0] == 'Resumen'
        assert work.sheet == 'Datos0'
    finally:
        copia.close()


@pytest.mark.parametrize('opcion', ['incompleto', 'identidad_alterada'])
def test_consolidado_incompatible_con_el_trabajo_utiliza_copia_portable(tmp_path, opcion):
    work, _, _ = importar(tmp_path, **{opcion: True})
    assert _consolidated_base(work) is None
    destino = tmp_path / 'copia.xlsx'
    work.export(destino)
    copia = load_workbook(destino)
    try:
        assert copia.sheetnames[0] == 'Trabajo'
    finally:
        copia.close()


def test_con_registros_en_varias_hojas_o_duplicados_no_hay_base_y_queda_portable(tmp_path):
    for nombre, opciones in (('hojas', {'hojas': 2}), ('duplicados', {'duplicar': True})):
        carpeta = tmp_path / nombre
        carpeta.mkdir()
        work, _, _ = importar(carpeta, **opciones)
        assert _consolidated_base(work) is None
        destino = carpeta / 'copia.xlsx'
        work.export(destino)
        assert load_workbook(destino).sheetnames[0] == 'Trabajo'
        # F1 (revisión integral): el descarte del consolidado se informa, como promete el manual.
        assert 'fidelidad reducida' in work.export_notice


def test_paquete_sin_consolidado_sigue_igual(tmp_path):
    metadata, tables = fixture()
    work = import_capture(write_capture(tmp_path / 'sin.zip', metadata, tables), defaults(), as_of=date(2026, 10, 8))
    assert _consolidated_base(work) is None
    destino = tmp_path / 'copia.xlsx'
    work.export(destino)
    assert load_workbook(destino).sheetnames[0] == 'Trabajo'
    # Sin consolidado no hay nada descartado: no se avisa.
    assert work.export_notice == ''


def test_reimportar_el_mismo_paquete_movido_conserva_la_base_y_ediciones(tmp_path):
    previous, _, raw = importar(tmp_path)
    old_path = Path(previous.path)
    new_path = old_path.rename(tmp_path / 'movido.zip')
    previous.rows[0].review['OBSERVACION'] = 'Edición humana ficticia'
    restored = import_capture(new_path, defaults(), as_of=date(2026, 10, 8), previous=previous)
    assert restored.path == str(new_path.resolve())
    assert previous.path == str(old_path.resolve())
    assert restored.rows[0].review == previous.rows[0].review
    base = _consolidated_base(restored)
    assert base and base[0] == raw


@pytest.mark.parametrize('campo,valor', [('T ESPERA', 9999), ('FEC. RESOLUCIÓN', '2026-08-02')])
def test_diferencias_entre_los_datos_analizados_y_consolidado_utilizan_portable(tmp_path, campo, valor):
    work, _, _ = importar(tmp_path, transformar=lambda datos: datos.update({campo: valor}))
    assert _consolidated_base(work) is None
    destino = tmp_path / 'copia.xlsx'
    work.export(destino)
    copia = load_workbook(destino)
    try:
        assert copia.sheetnames[0] == 'Trabajo'
        columnas = {celda.value: celda.column for celda in copia['Trabajo'][1]}
        assert str(copia['Trabajo'].cell(2, columnas[campo]).value) != str(valor)
    finally:
        copia.close()


def test_columna_analizada_omitida_en_consolidado_utiliza_portable(tmp_path):
    work, _, _ = importar(tmp_path, quitar_columnas=['T ESPERA'])
    assert _consolidated_base(work) is None
    work.export(tmp_path / 'copia.xlsx')


@pytest.mark.parametrize('fecha', [datetime(2026, 8, 1), '01/08/2026', 46235])
def test_numeros_y_fechas_equivalentes_conservan_la_base_consolidada(tmp_path, fecha):
    def convertir(datos):
        datos['T ESPERA'] = 40.0
        datos['FEC. RESOLUCIÓN'] = fecha

    work, _, _ = importar(tmp_path, transformar=convertir)
    assert _consolidated_base(work) is not None
    work.export(tmp_path / 'copia.xlsx', backend='portable', reduced_fidelity=True)


def test_modalidades_mixtas_comparan_las_columnas_de_cada_fila(tmp_path):
    metadata, tables = fixture()
    registro = tables['registros'][1]
    registro['record_type'] = 'CUMPLIMIENTO'
    datos = json.loads(registro['normalized_json'])
    datos.pop('T ESPERA')
    datos['DIAS DE CUMPLIMIENTO'] = '45'
    registro['normalized_json'] = json.dumps(datos)
    raw = consolidado(tables['registros'])
    path = write_capture(tmp_path / 'valido.zip', metadata, tables, consolidated=raw)
    work = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    assert _consolidated_base(work) is not None
    work.export(tmp_path / 'copia.xlsx', backend='portable', reduced_fidelity=True)

    raw = consolidado(tables['registros'], transformar=lambda datos: datos.update({'DIAS DE CUMPLIMIENTO': 9999}))
    path = write_capture(tmp_path / 'distinto.zip', metadata, tables, consolidated=raw)
    work = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    assert _consolidated_base(work) is None


def test_valor_nuevo_en_una_celda_analizada_vacia_utiliza_portable(tmp_path):
    metadata, tables = fixture()
    for registro in tables['registros']:
        datos = json.loads(registro['normalized_json'])
        datos['FEC. RESOLUCIÓN'] = ''
        registro['normalized_json'] = json.dumps(datos)
    raw = consolidado(tables['registros'], transformar=lambda datos: datos.update({'FEC. RESOLUCIÓN': '2026-08-01'}))
    path = write_capture(tmp_path / 'distinto.zip', metadata, tables, consolidated=raw)
    work = import_capture(path, defaults(), as_of=date(2026, 10, 8))
    assert _consolidated_base(work) is None


def test_ediciones_humanas_no_se_confunden_con_el_origen_consolidado(tmp_path):
    from nurus.personal.record_edits import apply

    work, _, _ = importar(tmp_path)
    apply(work, [work.rows[0].id], overrides={'espera': 9999}, review={'OBSERVACION': 'Revisión humana'})
    assert _consolidated_base(work) is not None
