"""Decisión x (15-oct): orden de las observaciones, una sola vez cada dato y reglas nuevas de Cumplimiento."""

import json
import re
from copy import deepcopy

import pytest
from types import SimpleNamespace
from historia_fixtures import CUMPLIMIENTO, ESPERA, bita, correr, cum, cumplimiento, esp, fila, hist

from nurus.personal.config import (
    LEGACY_TEXTS_REV4,
    LEGACY_TEXTS_REV5,
    LEGACY_TEXTS_REV6,
    TEXT_REVISION,
    Configuration,
    defaults,
)
from nurus.personal.motor.textos import ESTRUCTURA

ACLARACION = 'La historia es de la causa y no permite atribuirlo a un ingreso determinado.'


def _cuerpo(fila_):
    return fila_.observation.split(': ', 1)[1]


def test_estructura_coincide_con_las_oraciones_de_cada_texto_base():
    textos = defaults()['textos']
    for clave, letras in ESTRUCTURA.items():
        ambito, nombre = clave.split('.', 1)
        oraciones = re.split(r'(?<=\.)\s+', textos[ambito][nombre]['texto'].strip())
        assert len(oraciones) == len(letras), clave
        assert set(letras) <= set('CATF'), clave


def test_espera_ordena_constancias_luego_acciones_del_centro_y_al_final_el_tribunal(tmp_path):
    f = esp(tmp_path, fila(1) + [45], extra=[''], encabezados=['CURADOR'])
    assert _cuerpo(f) == (
        'Medida revisada, a la espera de ingreso efectivo. Al 14 de octubre de 2026, RUS registra 45 días de espera. '
        'No registra curador ad litem en RUS. '
        'Se remite proyecto de resolución pidiendo cuenta del ingreso efectivo y correo al programa consultando fecha estimada de ingreso. '
        'Se sugiere asociarlo informáticamente.'
    )
    assert f.actions == ['programa_espera', 'PC_IE']  # el orden del texto no cambia las acciones


def test_informes_ordena_constancias_y_deja_la_accion_del_centro_al_final(tmp_path):
    columnas = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'FECHA VENCIMIENTO']
    filas = [fila(1, rit='C-1-2026') + ['2026-08-20'], fila(2, rit='C-1-2026') + ['2026-09-20']]
    filas[1][2] = filas[0][2]
    f = correr(tmp_path, 'INFORMES', columnas, filas, {}).rows[0]
    cuerpo = _cuerpo(f)
    assert cuerpo.index('Informe de avance vencido') < cuerpo.index('Registra 2 informes vencidos')
    assert cuerpo.index('Registra 2 informes vencidos') < cuerpo.index('Se deja constancia de que se remite proyecto')
    assert f.actions == ['programa_vencido', 'PC_INFO']


def test_texto_editado_con_otra_cantidad_de_oraciones_se_conserva_entero(tmp_path):
    config = deepcopy(defaults())
    config['textos']['ESPERA']['E05_SOLO_CORREO']['texto'] = 'Texto propio. Con tres. Oraciones.'
    f = correr(tmp_path, 'ESPERA', ESPERA, [fila(1, 'DCE UNO') + [45]], None, None, config).rows[0]
    assert 'Texto propio. Con tres. Oraciones.' in f.observation


def test_ficha_sin_registro_deja_la_sugerencia_al_tribunal_para_el_final(tmp_path):
    f = cum(tmp_path, cumplimiento(1, 'RTA UNO', egreso='2027-03-01', dpe=100), hist(('Prorroga', '2026-10-07', 3)))
    cuerpo = _cuerpo(f)
    assert cuerpo.index('se prorroga la medida') < cuerpo.index('No registra ficha individual en RUS.')
    assert cuerpo.endswith('Se sugiere confeccionarla.')


def test_reiterar_ausencia_de_ficha_es_gestion_del_tribunal_y_va_al_final(tmp_path):
    previa = bita(('2026-09-20', 'Persona X: Persona no registra ficha FAE en RUS. Se sugiere confeccionarla.'))
    f = cum(
        tmp_path,
        cumplimiento(1, 'FAE UNO', ingreso='2026-01-01', egreso='2027-03-01', dpe=100, fae=''),
        hist(('Prorroga', '2026-10-07', 3)),
        previa,
    )
    cuerpo = _cuerpo(f)
    assert cuerpo.index('se prorroga la medida') < cuerpo.index(
        'Se reitera la observación del 20 de septiembre de 2026'
    )
    assert cuerpo.endswith('sobre la ausencia de ficha FAE en RUS.')


def test_aclaracion_de_causa_en_cumplimiento_con_hermanos(tmp_path):
    h = hist(('Prorroga', '2026-10-07', 3))
    filas = [
        fila(1, rit='C-1-2026', rut='11111111-1') + cumplimiento(1, egreso='2027-03-01', dpe=100)[5:],
        fila(2, rit='C-1-2026', rut='22222222-2') + cumplimiento(1, egreso='2027-03-01', dpe=100)[5:],
    ]
    trabajo = correr(tmp_path, 'CUMPLIMIENTO', CUMPLIMIENTO, filas, {0: h, 1: h})
    for f in trabajo.rows:
        assert _cuerpo(f).endswith(ACLARACION) and _cuerpo(f).count(ACLARACION) == 1
        assert f.avisos == []  # hermanos no son duplicidad
    solo = correr(tmp_path, 'CUMPLIMIENTO', CUMPLIMIENTO, filas[:1], {0: h}).rows[0]
    assert ACLARACION not in solo.observation


def test_aclaracion_de_causa_en_informes_solo_con_varios_ninos_y_una_cita(tmp_path):
    columnas = ['RIT', 'TRIBUNAL', 'NOMBRE', 'RUT', 'NOMBRE CENTRO', 'FECHA VENCIMIENTO']
    filas = [
        fila(1, rit='C-1-2026', rut='11111111-1') + ['2026-09-20'],
        fila(2, rit='C-1-2026', rut='22222222-2') + ['2026-09-20'],
    ]
    h = hist(('Ingreso informe de cumplimiento X', '2026-09-25', 7))
    trabajo = correr(tmp_path, 'INFORMES', columnas, filas, {0: h, 1: h})
    assert all(_cuerpo(f).endswith(ACLARACION) for f in trabajo.rows)
    # Varios informes del mismo niño no son varios niños.
    mismo = [fila(1, rit='C-1-2026') + ['2026-08-20'], fila(1, rit='C-1-2026') + ['2026-09-20']]
    assert ACLARACION not in correr(tmp_path, 'INFORMES', columnas, mismo, {0: h, 1: h}).rows[0].observation


def test_egreso_posterior_a_la_mayoria_de_edad(tmp_path):
    # Mayoría el 1-dic-2026 (a 48 días): se informa la mayoría y luego el egreso, sin repetir la fecha.
    cerca = cum(tmp_path, cumplimiento(1, egreso='2027-03-01', dpe=100, nac='2008-12-01'))
    assert _cuerpo(cerca).startswith(
        'Persona alcanzará la mayoría de edad el 1 de diciembre de 2026. El egreso proyectado es posterior a la mayoría de edad.'
    )
    # Mayoría lejana: la frase trae la fecha, porque no se había informado.
    lejos = cum(tmp_path, cumplimiento(1, egreso='2027-09-01', dpe=100, nac='2009-05-01'))
    assert (
        'El egreso proyectado es posterior a la mayoría de edad, que se alcanza el 1 de mayo de 2027.'
        in lejos.observation
    )
    # Egreso anterior a la mayoría o sin fecha de nacimiento: nada.
    assert (
        'posterior a la mayoría'
        not in cum(tmp_path, cumplimiento(1, egreso='2027-03-01', dpe=100, nac='2009-05-01')).observation
    )
    assert 'posterior a la mayoría' not in cum(tmp_path, cumplimiento(1, egreso='2027-09-01', dpe=100)).observation


def test_informe_por_vencer_con_la_fecha_de_egreso_no_se_repite(tmp_path):
    from test_reglas_cumplimiento import CRUCE, _cumplimiento

    def ejecutar(fecha):
        cruce = [CRUCE, ['X-1-2026', '11111111-0', 'Persona Ficticia 1', 'LAJA', 'PRM UNO', fecha]]
        return _cumplimiento(tmp_path, [(5, 10, '2026-10-09', '2026-10-24')], cruce=cruce).rows[0]

    igual = ejecutar('2026-10-24')
    assert 'CUMPLIMIENTO.C10_HOJA2' not in igual.rules and 'Próximo informe de avance' not in igual.observation
    otra = ejecutar('2026-12-01')
    assert 'CUMPLIMIENTO.C10_HOJA2' in otra.rules


def test_oido_ya_dejado_en_la_bitacora_no_se_repite(tmp_path):
    enc = ['FEC. OIDO']
    previa = bita(('2026-10-02', 'Persona X: Oído el 1 de octubre de 2026.'))
    con_rus = esp(tmp_path, fila(1) + [3], extra=['2026-10-01'], encabezados=enc, bitacora=previa)
    assert 'Oído el' not in con_rus.observation
    otra_fecha = esp(tmp_path, fila(1) + [3], extra=['2026-10-03'], encabezados=enc, bitacora=previa)
    assert 'Oído el 3 de octubre de 2026.' in otra_fecha.observation
    desde_historia = esp(
        tmp_path,
        fila(1) + [3],
        hist(('Derecho a ser oído', '2026-10-01', 12)),
        extra=[''],
        encabezados=enc,
        bitacora=previa,
    )
    assert 'Oído el' not in desde_historia.observation


def test_servicio_solo_se_dice_en_espera(tmp_path):
    h = hist(('Ing.Serv. protección especializada y otras redes', '2026-10-02', 7))
    assert 'Tribunal dispone ingreso' in esp(tmp_path, fila(1) + [3], h).observation
    assert 'Tribunal dispone ingreso' not in cum(tmp_path, cumplimiento(1), h).observation


def _escribir(directorio, cambio):
    conf = Configuration(directorio)
    datos = json.loads(json.dumps(conf.data))
    cambio(datos)
    conf.path.write_text(json.dumps(datos, ensure_ascii=False), encoding='utf-8')
    return Configuration(directorio).data


def test_migracion_de_textos_rev4_cambia_solo_los_que_siguen_iguales(tmp_path):
    def antiguos(d):
        d['revision_textos'] = 4
        for clave, previo in LEGACY_TEXTS_REV4.items():
            ambito, nombre = clave.split('.', 1)
            d['textos'][ambito][nombre]['texto'] = previo
        d['textos']['INFORMES']['I01_VENCIDO_GENERAL']['texto'] = 'Texto propio {FECHA_VENCIMIENTO} {PROGRAMA}.'

    datos = _escribir(tmp_path / 'a', antiguos)
    base = defaults()['textos']
    assert datos['revision_textos'] == TEXT_REVISION == 7
    for clave in LEGACY_TEXTS_REV4:
        ambito, nombre = clave.split('.', 1)
        esperado = (
            'Texto propio {FECHA_VENCIMIENTO} {PROGRAMA}.'
            if clave == 'INFORMES.I01_VENCIDO_GENERAL'
            else base[ambito][nombre]['texto']
        )
        assert datos['textos'][ambito][nombre]['texto'] == esperado, clave
    # Las claves nuevas se completan en una configuración que no las tiene.
    assert 'C2_PRORROGA_HASTA' in datos['textos']['HISTORIA'] and 'R11_SIN_EFECTO' in datos['textos']['HISTORIA']


@pytest.mark.parametrize('clave', sorted(LEGACY_TEXTS_REV4))
def test_los_textos_antiguos_tienen_las_mismas_variables_que_los_nuevos(clave):
    import string

    ambito, nombre = clave.split('.', 1)
    campos = lambda texto: {c for _, c, _, _ in string.Formatter().parse(texto) if c}  # noqa: E731
    assert campos(LEGACY_TEXTS_REV4[clave]) == campos(defaults()['textos'][ambito][nombre]['texto'])


def test_migracion_de_textos_rev5_adopta_el_formato_de_resolucion_y_respeta_lo_editado(tmp_path):
    def antiguos(d):
        d['revision_textos'] = 5
        for clave, previo in LEGACY_TEXTS_REV5.items():
            ambito, nombre = clave.split('.', 1)
            d['textos'][ambito][nombre]['texto'] = previo
        d['textos']['HISTORIA']['CURADOR_DESIGNADO']['texto'] = 'Texto propio {FECHA} {FOLIO}.'
        for clave in ('R11_ORDEN', 'R11_SIN_EFECTO', 'R11_TERMINO', 'R11_PIDE_CUENTA'):
            d['textos']['HISTORIA'].pop(clave)
        d['umbrales'].pop('ventana_busqueda')

    datos = _escribir(tmp_path / 'r5', antiguos)
    base = defaults()['textos']['HISTORIA']
    for clave in LEGACY_TEXTS_REV5:
        nombre = clave.split('.', 1)[1]
        esperado = 'Texto propio {FECHA} {FOLIO}.' if nombre == 'CURADOR_DESIGNADO' else base[nombre]['texto']
        assert datos['textos']['HISTORIA'][nombre]['texto'] == esperado, nombre
    assert datos['umbrales']['ventana_busqueda'] == 60 and 'R11_ORDEN' in datos['textos']['HISTORIA']


@pytest.mark.parametrize('clave', sorted(LEGACY_TEXTS_REV5))
def test_los_textos_rev5_tienen_las_mismas_variables_que_los_nuevos(clave):
    import string

    ambito, nombre = clave.split('.', 1)
    campos = lambda texto: {c for _, c, _, _ in string.Formatter().parse(texto) if c}  # noqa: E731
    assert campos(LEGACY_TEXTS_REV5[clave]) == campos(defaults()['textos'][ambito][nombre]['texto'])


def _con_umbral(tmp_path, clave, valor, filas, historias, modo='CUMPLIMIENTO', encabezados=CUMPLIMIENTO):
    config = deepcopy(defaults())
    config['umbrales'][clave] = valor
    return correr(tmp_path, modo, encabezados, filas, historias, None, config).rows[0]


def test_ventana_de_ingreso_efectivo_es_un_umbral(tmp_path):
    fila_ = cumplimiento(1, dc=100, ingreso='2026-06-01')
    h = hist(('Ingreso efectivo', '2026-06-20', 5))  # 19 días después
    assert 'informada al Tribunal' in _cuerpo(_con_umbral(tmp_path, 'ventana_ingreso', 30, [fila_], {0: h}))
    assert 'informada al Tribunal' not in _cuerpo(_con_umbral(tmp_path, 'ventana_ingreso', 10, [fila_], {0: h}))


def test_prorroga_reciente_es_un_umbral(tmp_path):
    fila_ = cumplimiento(1, egreso='2026-10-29', dpe=15)
    h = hist(('Prorroga', '2026-08-14', 3))  # 61 días antes del corte
    assert _con_umbral(tmp_path, 'prorroga_reciente', 90, [fila_], {0: h}).avisos
    assert not _con_umbral(tmp_path, 'prorroga_reciente', 60, [fila_], {0: h}).avisos


def test_una_configuracion_sin_los_umbrales_nuevos_se_completa(tmp_path):
    def sin_umbrales(d):
        for clave in ('prorroga_reciente', 'ventana_ingreso', 'lista_espera_larga'):
            d['umbrales'].pop(clave)

    datos = _escribir(tmp_path / 'u', sin_umbrales)
    assert (datos['umbrales']['prorroga_reciente'], datos['umbrales']['ventana_ingreso']) == (90, 30)
    assert datos['umbrales']['lista_espera_larga'] == 100


def test_deja_sin_efecto_dr_y_mn_marcan_tambien_en_cumplimiento(tmp_path):
    base = cum(tmp_path, cumplimiento(1))
    f = cum(tmp_path, cumplimiento(1), hist(('Deja sin efecto derivación D.R.', '2026-09-15', 20)))
    assert (f.observation, f.rules) == (base.observation, base.rules) and [a[:2] for a in f.avisos] == [
        (2, 'DEJA_SIN_EFECTO_DR')
    ]


def test_ingreso_efectivo_de_la_historia_en_espera_no_cambia_el_texto(tmp_path):
    base = esp(tmp_path, fila(1) + [45])
    f = esp(tmp_path, fila(1) + [45], hist(('Ingreso efectivo', '2026-10-02', 41)))
    assert f.observation == base.observation and f.actions == base.actions == ['programa_espera', 'PC_IE']


@pytest.mark.parametrize(
    'crudo,limpio',
    [
        ('Persona Ficticia (Reservado)', 'Persona Ficticia'),
        ('Persona (A) Ficticia', 'Persona Ficticia'),
        ('Persona Ficticia )(', 'Persona Ficticia'),
        ('Persona Ficticia', 'Persona Ficticia'),
    ],
)
def test_el_nombre_no_lleva_parentesis_en_correos_ni_proyectos(crudo, limpio):
    from nurus.personal.outputs import value

    work = SimpleNamespace(mapping={'nombre': 'NOMBRE'})
    row = SimpleNamespace(values={'NOMBRE': crudo}, overrides={})
    assert value(work, row, 'nombre') == limpio


def test_un_dato_que_dos_reglas_dicen_se_escribe_una_sola_vez():
    from nurus.personal.motor.composicion import componer
    from nurus.personal.motor.textos import Frag

    def frag(texto, dato=''):
        f = Frag(texto)
        f.dato = dato
        return f

    texto = componer(
        'P: ',
        [
            frag('Ingreso efectivo registrado el 1 de octubre de 2026.', 'ingreso_efectivo'),
            frag('La fecha de ingreso efectivo informada corresponde al 1 de octubre de 2026.', 'ingreso_efectivo'),
            frag('Medida revisada.'),
        ],
    )
    assert texto == 'P: Ingreso efectivo registrado el 1 de octubre de 2026. Medida revisada.'


def test_pending_fields_no_relee_un_word_sin_cambios(tmp_path, monkeypatch):
    from docx import Document
    from nurus.personal import word_quality

    ruta = tmp_path / 'p.docx'
    documento = Document()
    documento.add_paragraph('Sin campos pendientes')
    documento.save(ruta)
    llamadas = []
    original = word_quality._pending_fields
    monkeypatch.setattr(word_quality, '_pending_fields', lambda d: llamadas.append(d) or original(d))
    word_quality._CACHE.clear()
    assert word_quality.pending_fields(ruta) == [] and word_quality.pending_fields(ruta) == []
    assert len(llamadas) == 1
    documento.add_paragraph('Falta {{NOMBRE}}')
    documento.save(ruta)
    assert word_quality.pending_fields(ruta) == ['{{NOMBRE}}'] and len(llamadas) == 2


def test_abrir_un_archivo_no_deja_pegada_la_ventana():
    import threading
    from nurus.personal.app_base import launch_without_blocking

    liberar = threading.Event()
    assert launch_without_blocking('x.xlsx', startfile=lambda _p: liberar.wait(5), wait=0.05) is True
    liberar.set()
    assert launch_without_blocking('x.xlsx', startfile=lambda _p: None, wait=1) is False
    with pytest.raises(ValueError, match='No se pudo abrir'):
        launch_without_blocking('x.xlsx', startfile=lambda _p: (_ for _ in ()).throw(OSError('bloqueado')), wait=1)


def test_migracion_de_textos_rev6_ajusta_la_redaccion_aprobada_y_respeta_lo_editado(tmp_path):
    def antiguos(d):
        d['revision_textos'] = 6
        for clave, previo in LEGACY_TEXTS_REV6.items():
            ambito, nombre = clave.split('.', 1)
            d['textos'][ambito][nombre]['texto'] = previo
        d['textos']['HISTORIA']['R11_TERMINO']['texto'] = 'Texto propio {FECHA} {FOLIO}.'

    datos = _escribir(tmp_path / 'r6', antiguos)
    base = defaults()['textos']['HISTORIA']
    assert datos['revision_textos'] == TEXT_REVISION == 7
    assert datos['textos']['HISTORIA']['R11_TERMINO']['texto'] == 'Texto propio {FECHA} {FOLIO}.'
    assert datos['textos']['HISTORIA']['R11_SIN_EFECTO']['texto'] == base['R11_SIN_EFECTO']['texto']
    assert datos['textos']['HISTORIA']['T5_INSTITUCIONES']['texto'].startswith('Atendido que en sistema')


def test_configuracion_de_revision_4_con_variables_antiguas_abre_y_repone_el_texto(tmp_path):
    antiguo = (
        'En la historia de la causa consta un trámite descrito como «{DESC}», de fecha {FECHA}, folio {FOLIO}, que pone '
        'término a la búsqueda decretada mediante «Orden de búsqueda» de fecha {FECHA_ORDEN}, folio {FOLIO_ORDEN}.'
    )

    def rev4(d):
        d['revision_textos'] = 4
        d['textos']['HISTORIA']['R11_TERMINO']['texto'] = antiguo
        d['textos']['HISTORIA']['R11_SIN_EFECTO']['texto'] = antiguo.replace('pone término a', 'deja sin efecto')

    datos = _escribir(tmp_path / 'r4', rev4)
    base = defaults()['textos']['HISTORIA']
    assert datos['textos']['HISTORIA']['R11_TERMINO']['texto'] == base['R11_TERMINO']['texto']
    assert datos['textos']['HISTORIA']['R11_SIN_EFECTO']['texto'] == base['R11_SIN_EFECTO']['texto']


def test_un_error_tardio_al_abrir_queda_disponible_para_la_ventana():
    import threading
    from nurus.personal.app_base import launch_without_blocking

    liberar = threading.Event()

    def lento(_ruta):
        liberar.wait(5)
        raise OSError('sin programa asociado')

    resultado = {}
    assert launch_without_blocking('x.docx', startfile=lento, wait=0.05, result=resultado) is True
    liberar.set()
    resultado['thread'].join(5)
    assert isinstance(resultado['error'], OSError)
