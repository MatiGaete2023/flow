"""Each GUI scenario owns its process and Tcl interpreter, like a CSMP launch."""

import os
from pathlib import Path
import subprocess
import sys

ROOT = Path(__file__).parents[1]


def run_gui(tmp_path, scenario):
    environment = dict(os.environ, PYTHONPATH=os.pathsep.join([str(ROOT / 'src'), str(ROOT / 'tests'), *sys.path]))
    script = (
        '''
from datetime import date
from pathlib import Path
import sys
from nurus.personal.config import Configuration
from nurus.personal.app import App
from nurus.personal.work import Work
from current_fixture import source
root=Path(sys.argv[1])
app=App(Configuration(root/'config'))
app.geometry('1024x650')
try:
'''
        + ''.join('    ' + line + '\n' for line in scenario.splitlines())
        + '''
finally:
    app.destroy()
'''
    )
    result = subprocess.run(
        [sys.executable, '-c', script, str(tmp_path)], env=environment, capture_output=True, text=True, timeout=45
    )
    assert result.returncode == 0, result.stdout + '\n' + result.stderr


def test_current_results_and_manual_editors_are_reachable(tmp_path):
    run_gui(
        tmp_path,
        """
app.tabs.select(app.pages['Resultados']);app.update()
assert set(app.pages)=={'Trabajo','Correos','Resoluciones','Resultados','Configuración'}
assert app.results_download_button.winfo_exists()
from nurus.personal.manual import open_editor
open_editor(app,'mail');open_editor(app,'word');app.update()
""",
    )


def test_save_mail_batch_with_excel_and_roster_files_excludes_only_invalid_mail(tmp_path):
    run_gui(
        tmp_path,
        """
from types import SimpleNamespace
from unittest.mock import patch
from shutil import copyfile
from nurus.personal.outputs import Draft
from nurus.personal.product_checks import check_products
from nurus.personal.current_tools import summary
app.work=Work(app.cfg.data).analyze(source(root),'ESPERA',as_of=date(2026,10,2))
app.work.export(root/'copia.xlsx',backend='portable',reduced_fidelity=True)
copyfile(app.work.path,root/'nomina.xlsx')
valid=Draft('Correo válido','Texto revisado',to='equipo@example.test',attachments=[str(root/'nomina.xlsx')],
            roster=[{'program':'Programa ficticio','generated':str(root/'nomina.xlsx'),'rows':[]}])
invalid=Draft('Correo pendiente','Texto pendiente',to='')
app.drafts=[valid,invalid];app.draft_index=None
checks=check_products(app.work,app.drafts,[])
assert any(c.kind=='Excel' and not c.issues and c.target is None for c in checks)
assert any(c.kind=='Nómina Excel' and not c.issues and c.target is None for c in checks)
confirmed=[];saved=[]
def confirm(parent,drafts,excluded):
    assert len(drafts)==1 and drafts[0] is valid
    assert len(excluded)==1 and 'Correo pendiente' in excluded[0]
    assert 'equipo@example.test' in summary(drafts)
    confirmed.append(True);return True
def save(product,**kwargs):
    assert confirmed and kwargs['confirmed'] is True
    saved.append(product);return SimpleNamespace(entry_id='ENTRY-FICTICIA',store_id='STORE-FICTICIA')
def run(label,operation,done):
    result=operation();assert not result['errors'],result['errors'];done(result)
with patch('nurus.personal.current_tools.confirm_drafts',side_effect=confirm), patch('nurus.personal.outputs.save_draft',side_effect=save), patch.object(app,'_capture_mail'), patch.object(app,'_save_session'), patch.object(app,'_run',side_effect=run):
    app._send_all()
    assert len(saved)==1 and saved[0].recipient=='equipo@example.test'
    assert '1 borradores guardados' in app.status.get()
    assert len(app.work.receipts)==1
    app._send_all()
    assert len(saved)==1 and '1 ya procesados' in app.status.get()
app.work.receipts.clear();confirmed.clear();saved.clear()
with patch('nurus.personal.current_tools.confirm_drafts',return_value=False), patch('nurus.personal.outputs.save_draft',side_effect=save), patch.object(app,'_capture_mail'), patch.object(app,'_run',side_effect=run):
    app._send_all()
    assert not saved and not app.work.receipts
app.drafts=[invalid]
with patch('nurus.personal.current_tools.confirm_drafts') as confirm, patch.object(app,'_capture_mail'):
    try:app._send_all();raise AssertionError('No debe ofrecer guardar archivos como correos')
    except ValueError as error:assert 'No hay borradores completos' in str(error)
    assert not confirm.called
""",
    )


def test_hidden_selection_and_view_preferences_survive_refresh(tmp_path):
    run_gui(
        tmp_path,
        """
app.work=Work(app.cfg.data).analyze(source(root),'ESPERA',as_of=date(2026,10,2));app._show_work()
id=app.work.rows[0].id
app.records.selection_set(id);app.work_search.set('no coincidencia');app._apply_work_filter();app.update()
from nurus.personal.selection import selected_ids
assert id in selected_ids(app,'records')
app._show_work();app.update();from nurus.personal.selection import selected_ids
assert id in selected_ids(app,'records')
app.records.configure(displaycolumns=('Nombre','RIT'));app.work_search.set('');app._apply_work_filter()
assert tuple(app.records['displaycolumns'])==('Nombre','RIT')
""",
    )


def test_themes_result_identity_and_command_palette(tmp_path):
    run_gui(
        tmp_path,
        """
from nurus.personal.outputs import Draft
from nurus.personal.commands import show_palette
for theme in ('Claro','Oscuro','Sistema'):
    app.theme_choice.set(theme);app.apply_appearance();app.update()
assert app.cfg.data['vista']['tema']=='Sistema'
app.mail_target.set('todos')
app._display_prepared_drafts([Draft('Igual','Primero',to='a@example.test'),Draft('Igual','Segundo',to='b@example.test')],'{count} preparados')
app.tabs.select(app.pages['Resultados']);app.update()
second='draft:'+app.drafts[1].product_id
app.results_tree.selection_set(second);app.open_result();app.update()
assert app.body.get('1.0','end-1c')=='Segundo'
show_palette(app);app.update();assert app.command_palette.winfo_exists()
""",
    )


def test_large_table_yields_and_retains_hidden_selection(tmp_path):
    run_gui(
        tmp_path,
        """
from copy import deepcopy
import time
app.work=Work(app.cfg.data).analyze(source(root),'ESPERA',as_of=date(2026,10,2))
base=app.work.rows[0]
app.work.rows=[]
for n in range(1100):
    row=deepcopy(base);row.id='row-'+str(n);row.values[app.work.mapping['rit']]='X-'+str(n);row.actions=[];app.work.rows.append(row)
app._show_work();assert app.rendering
heartbeat=[];app.after(0,lambda:heartbeat.append(True))
limit=time.monotonic()+10
while app.rendering and time.monotonic()<limit:app.update()
assert not app.rendering and heartbeat
assert len(app.records.get_children())==1100
app.records.selection_set('row-1099');app.work_search.set('no-match');app._apply_work_filter()
from nurus.personal.selection import selected_ids
assert 'row-1099' in selected_ids(app,'records')
app.work_search.set('');app._apply_work_filter();assert 'row-1099' in app.records.selection()
""",
    )


def test_25_mail_attachments_are_automatic_and_editing_one_is_optional(tmp_path):
    run_gui(
        tmp_path,
        """
from copy import deepcopy
from hashlib import sha256
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from nurus.personal.outputs import Draft
from nurus.personal.roster import snapshot,review
from nurus.personal.grid_edit import GridEditor
app.work=Work(app.cfg.data).analyze(source(root),'ESPERA',as_of=date(2026,10,2))
app.work.export(root/'copia.xlsx',backend='portable',reduced_fidelity=True)
def pending():
    return [Draft('Correo '+str(n),'Texto '+str(n),to='equipo@example.test',required=True,
            program='Programa '+str(n),kind='programa_espera',recipient_type='programas',
            roster=[snapshot(app.work,app.work.rows,'Programa '+str(n),'programa_espera')],roster_reviewed=False)
            for n in range(25)]
def run(label,operation,done):done(operation())
app.mail_target.set('todos')
with patch('nurus.personal.app.prepare_required_drafts',side_effect=lambda *a,**k:pending()), patch.object(app,'_selected_mail_modalities',return_value=['PRM']), patch.object(app,'_capture_mail'), patch.object(app,'_run',side_effect=run):
    app._prepare_all_mail()
assert len(app.drafts)==25 and all(d.roster_reviewed and len(d.attachments)==1 for d in app.drafts)
original_paths=[d.attachments[0] for d in app.drafts]
original_hashes=[sha256(Path(p).read_bytes()).hexdigest() for p in original_paths]
# Open only the selected editor and accept a real inline cell edit.
app.draft_index=9
with patch.object(app,'_capture_mail'), patch.object(app,'_save_session'), patch.object(GridEditor,'commit',autospec=True,side_effect=lambda editor:editor.group['rows'][0]['cells'].__setitem__(2,'Dato corregido')):
    review(app);app.update()
    window=next(w for w in app.winfo_children() if w.winfo_class()=='Toplevel')
    button=next(w for w in window.winfo_children() if w.winfo_class()=='TButton' and 'Aceptar' in w.cget('text'))
    button.invoke();app.update()
assert app.drafts[9].attachments[0]!=original_paths[9]
book=load_workbook(app.drafts[9].attachments[0]);assert book.active['C2'].value=='Dato corregido';book.close()
for n,draft in enumerate(app.drafts):
    if n!=9:
        assert draft.attachments[0]==original_paths[n]
        assert sha256(Path(draft.attachments[0]).read_bytes()).hexdigest()==original_hashes[n]
# Final batch confirmation works for all 25 without opening 24 editors.
saved=[]
def confirm(parent,drafts,excluded):
    assert len(drafts)==25 and not excluded
    assert all(d.roster_reviewed and Path(d.attachments[0]).is_file() for d in drafts)
    return True
with patch('nurus.personal.current_tools.confirm_drafts',side_effect=confirm), patch('nurus.personal.outputs.save_draft',side_effect=lambda product,**kw:saved.append(product) or SimpleNamespace(entry_id='FAKE',store_id='FAKE')), patch.object(app,'_capture_mail'), patch.object(app,'_save_session'), patch.object(app,'_run',side_effect=run):
    app._send_all()
    assert len(saved)==25 and '25 borradores guardados' in app.status.get()
    app._send_all();assert len(saved)==25 and '25 ya procesados' in app.status.get()
# Repreparation preserves the one edited cell, replaces automatic paths and
# preserves explicit body edits and a separately attached file.
app.drafts[9].body='Texto corregido manualmente'
manual=root/'manual.txt';manual.write_text('Adjunto opcional');app.drafts[9].attachments.append(str(manual))
with patch('nurus.personal.app.prepare_required_drafts',side_effect=lambda *a,**k:pending()), patch.object(app,'_selected_mail_modalities',return_value=['PRM']), patch.object(app,'_capture_mail'), patch.object(app,'_run',side_effect=run):
    app._prepare_all_mail()
assert app.drafts[9].body=='Texto corregido manualmente', (app.drafts[9].body,app.drafts[9].original)
assert app.drafts[9].roster_reviewed and len(app.drafts[9].attachments)==2
assert str(manual) in app.drafts[9].attachments
book=load_workbook(app.drafts[9].roster[0]['generated']);assert book.active['C2'].value=='Dato corregido';book.close()
# Old saved sessions with deferred rosters are materialized on Save too.
app.work.receipts.clear();app._draft_scope_source=[];app.drafts=pending();app.draft_index=None
with patch('nurus.personal.current_tools.confirm_drafts',side_effect=confirm), patch('nurus.personal.outputs.save_draft',side_effect=lambda product,**kw:SimpleNamespace(entry_id='FAKE',store_id='FAKE')), patch.object(app,'_capture_mail'), patch.object(app,'_save_session'), patch.object(app,'_run',side_effect=run):
    app._send_all()
assert len(app.work.receipts)==25 and all(d.roster_reviewed for d in app.drafts)
# Saving one draft from an older session also avoids a mandatory editor.
app.work.receipts.clear();app.drafts=pending();app.draft_index=9
with patch('nurus.personal.current_tools.confirm_drafts',return_value=True), patch('nurus.personal.outputs.save_draft',return_value=SimpleNamespace(entry_id='FAKE',store_id='FAKE')), patch.object(app,'_capture_mail'), patch.object(app,'_save_session'), patch.object(app,'_run',side_effect=run):
    app._send_draft()
assert len(app.work.receipts)==1 and app.drafts[9].roster_reviewed
""",
    )
