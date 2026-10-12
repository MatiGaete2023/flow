"""Sonda de la auditoría (AS3): la copia anónima no filtra datos y conserva las reglas."""

import subprocess
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment

from nurus.personal.anonymize import OMITTED, anonymize, check_digit, main
from nurus.personal.config import defaults
from nurus.personal.work import Work

ROOT = Path(__file__).resolve().parents[1]
ENCABEZADOS = [
    'RIT',
    'TRIBUNAL',
    'NOMBRE',
    'RUT',
    'FEC. NACIMIENTO',
    'NOMBRE CENTRO',
    'T ESPERA',
    'CURADOR',
    'OBSERVACION',
]


def rut(number, dotted=False):
    # Se calcula en ejecución: el archivo de prueba no contiene RUT con dígito válido.
    digits = f'{number:,}'.replace(',', '.') if dotted else str(number)
    return f'{digits}-{check_digit(number)}'


PERSONAS = [
    'Valentina Ignacia Soto Reyes',
    'Tomás Andrés Soto Reyes',
    'Martina Fuentes Lagos',
    'Benjamín Alarcón Pizarro',
]
CURADORA = 'Carolina Muñoz Herrera'
RUTS = [rut(15_284_731), rut(15_284_732, dotted=True), rut(21_406_377), rut(19_873_215)]
RUT_CURADORA = rut(13_552_918)
RIT_HERMANOS, RIT_OTRO, RIT_TERCERO = 'P-1534-2025', 'P-77-2026', 'X-902-2024'


def sonda(tmp_path):
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'ESPERA'
    hoja.append([f'Informe de espera · {PERSONAS[0]} {RUTS[0]} · causa {RIT_HERMANOS}'])
    hoja.append(ENCABEZADOS)
    filas = [
        [
            RIT_HERMANOS,
            'LAJA',
            PERSONAS[0],
            RUTS[0],
            datetime(2010, 3, 4),
            'PRM UNO',
            45,
            f'{CURADORA} ({RUT_CURADORA})',
            f'Hermana de {PERSONAS[1]}; ver {RIT_OTRO} y RUT {RUTS[1]}',
        ],
        [
            RIT_HERMANOS,
            'LAJA',
            PERSONAS[1],
            RUTS[1],
            datetime(2012, 7, 9),
            'PRM UNO',
            45,
            'NO POSEE',
            f'Curadora {CURADORA.split()[0]} {CURADORA.split()[1]} informa; madre: Soto',
        ],
        [
            RIT_OTRO,
            'LAJA',
            PERSONAS[2],
            RUTS[2],
            datetime(2008, 11, 2),
            'DCE PRUEBA',
            40,
            '---',
            f'Contactar a {PERSONAS[3].split()[0]} en {RIT_TERCERO}',
        ],
        [
            RIT_TERCERO,
            'LAJA',
            PERSONAS[3],
            RUTS[3],
            datetime(2009, 10, 20),
            'PRM UNO',
            10,
            'Institución: CAJ Biobío',
            '',
        ],
    ]
    for fila in filas:
        hoja.append(fila)
    hoja['C3'].comment = Comment(f'Revisar {PERSONAS[0]}', 'autor')
    hoja['D4'].hyperlink = f'https://ejemplo.invalid/{RUTS[1]}'
    libro.properties.creator = PERSONAS[3]
    libro.properties.title = f'Causa {RIT_HERMANOS}'
    origen = tmp_path / 'original.xlsx'
    libro.save(origen)
    return origen


def reglas(ruta):
    trabajo = Work(defaults()).analyze(ruta, 'ESPERA', as_of=date(2026, 10, 9))
    return [(sorted(map(str, fila.rules)), sorted(map(str, fila.actions)), fila.excluded) for fila in trabajo.rows]


def texto(ruta):
    with zipfile.ZipFile(ruta) as archivo:
        return '\n'.join(archivo.read(n).decode('utf-8', 'ignore') for n in archivo.namelist())


def test_sonda_sin_filtraciones_mismas_reglas_y_check_real_ids(tmp_path):
    origen = sonda(tmp_path)
    huella = origen.read_bytes()
    destino = tmp_path / 'copia.xlsx'
    anonymize(origen, destino)
    assert origen.read_bytes() == huella
    contenido = texto(destino)
    sensibles = [
        *PERSONAS,
        CURADORA,
        *RUTS,
        RUT_CURADORA,
        RIT_HERMANOS,
        RIT_OTRO,
        RIT_TERCERO,
        *(parte for nombre in [*PERSONAS, CURADORA] for parte in nombre.split()),
        *(r.replace('.', '') for r in RUTS),
    ]
    filtrados = [s for s in sensibles if s.lower() in contenido.lower()]
    assert filtrados == []
    copia = load_workbook(destino)
    assert all(c.comment is None and c.hyperlink is None for h in copia.worksheets for f in h.iter_rows() for c in f)
    assert reglas(origen) == reglas(destino)
    assert any(acciones for _, acciones, _ in reglas(origen))
    hoja = load_workbook(destino)['ESPERA']
    rits = [hoja.cell(fila, 1).value for fila in range(3, 7)]
    assert rits[0] == rits[1] != rits[2] and len(set(rits)) == 3
    # H1 (revisión externa): el texto libre fuera de las columnas del motor se omite, no se pseudonimiza.
    assert hoja.cell(5, 9).value == OMITTED
    assert 'Institución: CAJ Biobío' == hoja.cell(6, 8).value
    resultado = subprocess.run(
        [sys.executable, str(ROOT / 'tools' / 'check_real_ids.py'), str(destino)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert resultado.returncode == 0, resultado.stdout
    original = subprocess.run(
        [sys.executable, str(ROOT / 'tools' / 'check_real_ids.py'), str(origen)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert original.returncode == 1


def test_nunca_sobrescribe_ni_usa_el_original_como_destino(tmp_path):
    origen = sonda(tmp_path)
    huella = origen.read_bytes()
    existente = tmp_path / 'existe.xlsx'
    existente.write_bytes(b'no tocar')
    with pytest.raises(FileExistsError):
        anonymize(origen, existente)
    with pytest.raises(FileExistsError):
        anonymize(origen, origen)
    assert existente.read_bytes() == b'no tocar' and origen.read_bytes() == huella
    assert main([str(origen), str(existente)]) == 1


@pytest.mark.parametrize('texto_literal', ['=1+1', '# informe A', '#DIV/0!'])
def test_texto_literal_no_se_convierte_en_formula(tmp_path, texto_literal):
    origen = sonda(tmp_path)
    libro = load_workbook(origen)
    celda = libro['ESPERA']['I3']
    celda.value = texto_literal
    celda.data_type = 's'
    libro.save(origen)
    libro.close()
    destino = tmp_path / 'copia.xlsx'
    anonymize(origen, destino)
    copia = load_workbook(destino, data_only=False)
    try:
        celda = copia['ESPERA']['I3']
        # La columna es texto libre (H1): se omite; lo que importa aquí es que nunca sea fórmula.
        assert celda.value == OMITTED and celda.data_type == 's'
        assert not any(c.data_type == 'f' for h in copia for f in h for c in f)
    finally:
        copia.close()


def test_fallo_al_copiar_no_deja_destino_parcial(tmp_path, monkeypatch):
    import nurus.personal.anonymize as modulo

    origen = sonda(tmp_path)
    destino = tmp_path / 'copia.xlsx'
    abrir = Path.open

    class EscrituraFallida:
        def __enter__(self):
            self.archivo = abrir(destino, 'xb')
            return self

        def write(self, contenido):
            self.archivo.write(contenido[:10])
            raise OSError('disco sin espacio')

        def __exit__(self, *exc):
            self.archivo.close()

    def abrir_con_fallo(self, mode='r', *args, **kwargs):
        return EscrituraFallida() if self == destino and mode == 'xb' else abrir(self, mode, *args, **kwargs)

    def sin_enlaces(*args):
        raise OSError('el sistema de archivos no admite enlaces')

    monkeypatch.setattr(modulo.os, 'link', sin_enlaces)
    monkeypatch.setattr(Path, 'open', abrir_con_fallo)
    with pytest.raises(OSError, match='disco sin espacio'):
        anonymize(origen, destino)
    assert not destino.exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == [origen.name]


def test_copias_concurrentes_usan_temporales_distintos(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    origen = sonda(tmp_path)
    destinos = [tmp_path / 'copia_uno.xlsx', tmp_path / 'copia_dos.xlsx']
    guardar = Workbook.save
    barrera = Barrier(2)

    def guardar_a_la_vez(self, path):
        guardar(self, path)
        barrera.wait(timeout=10)

    monkeypatch.setattr(Workbook, 'save', guardar_a_la_vez)
    with ThreadPoolExecutor(max_workers=2) as ejecutor:
        informes = list(ejecutor.map(lambda ruta: anonymize(origen, ruta), destinos))
    assert [i.target for i in informes] == destinos
    for destino in destinos:
        copia = load_workbook(destino)
        copia.close()
    assert sorted(p.name for p in tmp_path.iterdir()) == sorted([origen.name, *(p.name for p in destinos)])


def test_rit_ficticios_no_colisionan_con_ninguna_causa_original(tmp_path):
    libro = Workbook()
    hoja = libro.active
    hoja.append(['RIT', 'NOMBRE', 'OBSERVACION'])
    hoja.append(['P-1-2026', 'Persona Uno', 'Ver P-3-2026'])
    hoja.append(['P-1-2026', 'Persona Dos', ''])
    hoja.append(['P-2-2026', 'Persona Tres', ''])
    origen, destino = tmp_path / 'original.xlsx', tmp_path / 'copia.xlsx'
    libro.save(origen)
    anonymize(origen, destino)
    contenido = texto(destino)
    assert all(rit not in contenido for rit in ('P-1-2026', 'P-2-2026', 'P-3-2026'))
    copia = load_workbook(destino)
    try:
        assert copia.active['A2'].value == copia.active['A3'].value != copia.active['A4'].value
    finally:
        copia.close()


def test_boton_configuracion_crea_copia_y_no_sobrescribe(tmp_path, monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import Mock

    from nurus.personal import app_base

    assert "text='Exportar caso de prueba anónimo'" in Path(app_base.__file__).read_text(encoding='utf-8')
    origen = sonda(tmp_path)
    huella = origen.read_bytes()
    destino = tmp_path / 'copia.xlsx'
    monkeypatch.setattr(app_base.filedialog, 'askopenfilename', lambda **kw: str(origen))
    monkeypatch.setattr(app_base.filedialog, 'asksaveasfilename', lambda **kw: str(destino))
    monkeypatch.setattr(app_base.messagebox, 'showinfo', Mock())
    app = SimpleNamespace(status=Mock())
    app_base.App._export_anonymous_case(app)
    assert destino.is_file() and origen.read_bytes() == huella
    copia = destino.read_bytes()
    with pytest.raises(FileExistsError):
        app_base.App._export_anonymous_case(app)
    monkeypatch.setattr(app_base.filedialog, 'asksaveasfilename', lambda **kw: str(origen))
    with pytest.raises(FileExistsError):
        app_base.App._export_anonymous_case(app)
    assert destino.read_bytes() == copia and origen.read_bytes() == huella


def test_persona_solo_en_texto_libre_no_sobrevive(tmp_path):
    # Revisión externa H1: un nombre que no figura en ninguna columna de identidad no tiene reemplazo conocido.
    from test_personal_flow import source

    origen = source(tmp_path)
    libro = load_workbook(origen)
    hoja = libro['informe_1']
    hoja.cell(3, 7, 'OBSERVACION')
    hoja.cell(4, 7, 'Contacto: ' + PERSONAS[3])
    libro.create_sheet('Notas').cell(1, 1, 'Visita de ' + PERSONAS[3])
    libro.save(origen)
    destino = tmp_path / 'anon.xlsx'
    anonymize(origen, destino)
    assert PERSONAS[3].lower() not in texto(destino).lower()
    copia = load_workbook(destino)
    assert copia['informe_1'].cell(4, 7).value == OMITTED
