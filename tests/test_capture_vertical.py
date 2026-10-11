"""Un lote original atraviesa productor, contrato y consumidor sin editar el ZIP."""

from pathlib import Path
import subprocess
import sys
import pytest


@pytest.mark.parametrize('mode', ['Espera', 'Informes'])
def test_downloader_original_to_assistant_and_excel(tmp_path, mode):
    root = Path(__file__).parents[1]
    script = '''
from datetime import date
from pathlib import Path
import sys
root=Path(sys.argv[1]);output=Path(sys.argv[2]);mode=sys.argv[3]
sys.path.insert(0,str(root/'src'));sys.path.insert(0,str(root/'descargador'))
from test_resultados import fixture
from resultados import load_lot
from export_csmp_capture import export_capture
from nurus.personal.capture_input import import_capture
from nurus.personal.capture_products import export_requests
from nurus.personal.config import defaults
from openpyxl import load_workbook
folder=fixture(output/'originales',mode=mode)
before={p.name:p.read_bytes() for p in folder.iterdir()}
package=export_capture([load_lot(folder)],output/'capture.zip')
config=defaults();config['contactos']={'programas':{},'tribunales':{}}
work=import_capture(package,config,as_of=date(2026,10,8))
assert len(work.rows)==2
assert all(r['identity_status']=='INCIERTA' for r in work.capture_sources.values())
assert all(c['class']=='NO_EVALUABLE' for c in work.capture_context.values())
work.rows[0].review['OBSERVACION']='Revisión humana ficticia'
work.export(output/'Revisable.xlsx',backend='portable',reduced_fidelity=True)
book=load_workbook(work.output);sheet=book['Trabajo']
columns={c.value:c.column for c in sheet[1]}
assert sheet.cell(2,columns['OBSERVACION']).value=='Revisión humana ficticia'
book.close()
export_requests(work,output/'solicitud.json')
assert {p.name:p.read_bytes() for p in folder.iterdir()}==before
'''
    result = subprocess.run(
        [sys.executable, '-c', script, str(root), str(tmp_path), mode],
        capture_output=True,
        text=True,
        encoding='utf-8',
        timeout=60,
    )
    assert result.returncode == 0, result.stderr
