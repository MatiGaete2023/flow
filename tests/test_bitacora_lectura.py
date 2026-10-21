"""Lector vigente de bitácoras (nurus.bitacora_html y nurus.capture.diaries): rechazos e integridad del lote.

Repone, contra el código de producción, las comprobaciones que antes pasaban por la auditoría retirada en WP-69."""

import json

import pytest
from bitacora_fixtures import PARAMS, manifest, popup
from openpyxl import load_workbook

from nurus.bitacora_html import read_snapshot
from nurus.capture.diaries import export_diaries, import_lote


def test_texto_completo_respuesta_e_identidad():
    texto = 'Línea uno & comillas "exactas"\nLínea dos ' * 12
    lectura = read_snapshot(popup(texto), PARAMS)
    entrada = lectura['entries'][0]
    assert entrada['texto'] == texto and entrada['entry_id'] == '44' and entrada['respuesta'] == ''
    assert entrada['fecha'] == '2026-08-12T10:13:00'
    assert (lectura['tribunal_codigo'], lectura['causa_id'], lectura['ingreso_id']) == ('11', '22', '33')


@pytest.mark.parametrize(
    'cambio',
    [
        lambda t: t.replace('name="ID_Ingreso" value="33"', 'name="ID_Ingreso" value="99"'),
        lambda t: t.replace('HOBS_Centro_44', 'HOBS_Centro_55'),
        lambda t: t.replace('Fecha Centro', 'Fecha desconocida'),
        lambda t: t.replace('</html>', ''),
        lambda t: t.replace('<img id="44"></td>', '<img id="45"></td>'),
        lambda t: t.replace('<img id="44" onclick="MuestraDiv(this.id,3)"><img id="44">', '<img id="45"><img id="45">'),
    ],
)
def test_contenido_no_verificable_se_rechaza(cambio):
    with pytest.raises(ValueError):
        read_snapshot(cambio(popup()), PARAMS)


def test_html_cambiado_queda_como_lectura_fallida(tmp_path):
    fuente = manifest(tmp_path)
    (tmp_path / 'captura.html').write_text('alterada')
    consulta = import_lote(fuente)['queries'][0]
    assert consulta['coverage'] == 'FALLIDA' and 'cambió' in consulta['error']


def test_bitacora_de_otro_ingreso_no_se_acepta(tmp_path):
    fuente = manifest(tmp_path)
    datos = json.loads(fuente.read_text())
    datos['registros'][0]['ingreso_id'] = '99'
    fuente.write_text(json.dumps(datos))
    consultas = import_lote(fuente)['queries']
    ingreso_99 = next(q for q in consultas if q['ingreso_id'] == '99')
    assert ingreso_99['coverage'] == 'FALLIDA' and 'otro ingreso' in ingreso_99['error']


def test_error_del_recorrido_queda_visible_en_el_excel(tmp_path):
    fuente = manifest(tmp_path)
    datos = json.loads(fuente.read_text())
    datos.update(registros=[], consultas=[], error='Actualizar extensión incompatible')
    fuente.write_text(json.dumps(datos))
    libro = load_workbook(export_diaries(fuente, tmp_path / 'control.xlsx'))
    try:
        filas = {fila[0]: fila[1] for fila in libro['Lote'].iter_rows(min_row=2, values_only=True)}
    finally:
        libro.close()
    assert filas['Error del recorrido'] == 'Actualizar extensión incompatible'
