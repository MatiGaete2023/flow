"""Pruebas exclusivamente sintéticas: Excel, ZIP, consolidado y rechazo seguro."""

from io import BytesIO
import json
from hashlib import sha256
from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from nurus.capture import read_capture, write_capture
from nurus.capture.fixtures import fixture
from nurus.personal.anonymize import OMITTED, anonymize
from nurus.personal.anonymize_package import UnsafePackage, anonymize_capture


def book_fixture(path):
    wb = Workbook()
    sh = wb.active
    sh.title = 'ESPERA'
    sh.append(['RIT', 'NOMBRE', 'TRIBUNAL', 'T ESPERA', 'OBSERVACION'])
    sh.append(
        [
            'P-432-2026',
            'Martina Lagos Durán',
            'LAJA',
            45,
            'Se reitera observación respecto de Martina Lagos Durán, RIT P-432-2026.',
        ]
    )
    wb.properties.creator = 'Martina Lagos Durán'
    wb.save(path)
    wb.close()


@pytest.mark.parametrize('mode', ['conservador', 'estricto'])
def test_excel_gui_motor_preserva_estructura_y_original(tmp_path, mode):
    source = tmp_path / 'con_datos.xlsx'
    target = tmp_path / 'copia.xlsx'
    book_fixture(source)
    data = source.read_bytes()
    anonymize(source, target, text_mode=mode)
    assert source.read_bytes() == data
    wb = load_workbook(target)
    sheet = wb['ESPERA']
    assert [c.value for c in sheet[1]] == ['RIT', 'NOMBRE', 'TRIBUNAL', 'T ESPERA', 'OBSERVACION']
    assert sheet['D2'].value == 45
    assert sheet['B2'].value != 'Martina Lagos Durán'
    if mode == 'estricto':
        assert sheet['E2'].value == OMITTED
    else:
        assert 'Se reitera observación' in sheet['E2'].value
        assert 'Martina Lagos Durán' not in sheet['E2'].value
    assert wb.properties.creator == 'anonimizado'
    wb.close()
    with zipfile.ZipFile(target) as archive:
        body = '\n'.join(archive.read(n).decode('utf8', 'ignore') for n in archive.namelist())
    assert 'Martina Lagos Durán' not in body


def zip_fixture(path, *, consolidated=False, originals=False):
    metadata, tables = fixture(count=4)
    metadata['contains_sensitive_data'] = True
    metadata['operator_email'] = 'usuario.privado@example.net'
    for row in tables['registros']:
        obj = json.loads(row['normalized_json'])
        obj['RIT'] = 'P-432-2026'
        obj['NOMBRE'] = 'Martina Lagos Durán'
        obj['RUT'] = '15284731-5'
        obj['OBSERVACION'] = 'Se reitera observación respecto de Martina Lagos Durán.'
        row['rit_original'] = 'P-432-2026'
        row['normalized_json'] = json.dumps(obj, ensure_ascii=False)
        row['raw_json'] = json.dumps(obj, ensure_ascii=False)
    for entry in tables['bitacoras']:
        entry['author'] = 'Martina Lagos Durán'
        entry['text'] = 'Se reitera observación a Martina Lagos Durán.'
    output = None
    if consolidated:
        wb = Workbook()
        sh = wb.active
        sh.title = 'ESPERA'
        sh.append(
            [
                'RIT',
                'NOMBRE',
                'RUT',
                'OBSERVACION',
                'SITFA_SHA256',
                'SITFA_ARCHIVO',
                'SITFA_CAUSA_RUS',
                'SITFA_INGRESO_RUS',
                'SITFA_FILA',
            ]
        )
        row = tables['registros'][0]
        sh.append(
            [
                'P-432-2026',
                'Martina Lagos Durán',
                '15284731-5',
                'Se reitera observación a Martina Lagos Durán.',
                row['source_sha256'],
                row['original_file'],
                row['cause_id'],
                row['ingreso_id'],
                int(row['original_row']),
            ]
        )
        wb.properties.creator = 'Martina Lagos Durán'
        output = BytesIO()
        wb.save(output)
        wb.close()
    if originals:
        metadata['originals_authorized'] = True
    write_capture(
        path,
        metadata,
        tables,
        consolidated=output.getvalue() if output else None,
        originals={'originales/documento.pdf': b'%PDF-1.4 FICTICIO'} if originals else None,
    )


@pytest.mark.parametrize('mode', ['conservador', 'estricto'])
@pytest.mark.parametrize('consolidated', [False, True])
def test_paquete_relaciones_y_consolidado(tmp_path, mode, consolidated):
    origin = tmp_path / 'captura.zip'
    destination = tmp_path / 'anon.zip'
    zip_fixture(origin, consolidated=consolidated)
    original_sha = sha256(origin.read_bytes()).hexdigest()
    original = read_capture(origin)
    anonymize_capture(origin, destination, mode=mode)
    assert sha256(origin.read_bytes()).hexdigest() == original_sha
    result = read_capture(destination)
    assert result.manifest['contains_sensitive_data'] is True
    assert result.manifest['originals_authorized'] is False
    assert 'operator_email' not in result.manifest
    assert result.manifest['capture_id'] != original.manifest['capture_id']
    assert len(result.rows('registros')) == 4
    record = result.rows('registros')[0]
    assert record['cause_id'] == result.rows('registros')[1]['cause_id']
    assert record['ingreso_id'] == result.rows('bitacoras')[0]['ingreso_id']
    assert record['source_sha256'] != original.rows('registros')[0]['source_sha256']
    assert record['original_file'] != original.rows('registros')[0]['original_file']
    doc = json.loads(record['normalized_json'])
    assert doc['NOMBRE'] != 'Martina Lagos Durán'
    assert doc['RIT'] == record['rit_original']
    assert doc['OBSERVACION'] == (OMITTED if mode == 'estricto' else doc['OBSERVACION'])
    if mode == 'conservador':
        assert 'Se reitera observación' in doc['OBSERVACION']
        assert 'Martina Lagos Durán' not in doc['OBSERVACION']
    if consolidated:
        assert result.consolidated
        wb = load_workbook(BytesIO(result.consolidated))
        row = [c.value for c in wb.active[2]]
        assert row[4] == record['source_sha256']
        assert row[5] == record['original_file']
        assert row[6] == record['cause_id']
        assert row[7] == record['ingreso_id']
        assert wb.properties.creator == 'anonimizado'
        wb.close()
    else:
        assert result.consolidated is None
    with zipfile.ZipFile(destination) as archive:
        contents = '\n'.join(archive.read(n).decode('utf8', 'ignore') for n in archive.namelist())
    assert 'Martina Lagos Durán' not in contents
    assert 'usuario.privado@example.net' not in contents


def test_rechaza_binarios_y_nunca_publica(tmp_path):
    origin = tmp_path / 'captura.zip'
    target = tmp_path / 'copia.zip'
    zip_fixture(origin, originals=True)
    checksum = sha256(origin.read_bytes()).hexdigest()
    with pytest.raises(UnsafePackage, match='binarios'):
        anonymize_capture(origin, target)
    assert not target.exists()
    assert sha256(origin.read_bytes()).hexdigest() == checksum


def test_nunca_sobrescribe(tmp_path):
    origin = tmp_path / 'captura.zip'
    target = tmp_path / 'existente.zip'
    zip_fixture(origin)
    target.write_bytes(b'no modificar')
    with pytest.raises(FileExistsError):
        anonymize_capture(origin, target)
    with pytest.raises(FileExistsError):
        anonymize_capture(origin, origin)
    assert target.read_bytes() == b'no modificar'


def test_zip_no_colisiona_rit_original_con_sustituto(tmp_path):
    metadata, tables = fixture(count=2)
    tables['registros'][0]['rit_original'] = 'P-1-2026'
    tables['registros'][1]['rit_original'] = 'P-2-2026'
    origin, target = tmp_path / 'colision.zip', tmp_path / 'colision_anon.zip'
    write_capture(origin, metadata, tables)
    anonymize_capture(origin, target)
    result = read_capture(target)
    rits = {row['rit_original'] for row in result.rows('registros')}
    assert not rits.intersection({'P-1-2026', 'P-2-2026'})
    assert len(rits) == 2


def test_excel_observacion_nombre_ajeno_a_campos_identidad(tmp_path):
    source, target = tmp_path / 'origen.xlsx', tmp_path / 'salida.xlsx'
    book_fixture(source)
    wb = load_workbook(source)
    wb.active['E2'] = 'Se reitera observación de Camila Aguilera Contreras sobre curador ad litem.'
    wb.save(source)
    wb.close()
    anonymize(source, target, text_mode='conservador')
    copy = load_workbook(target)
    observation = copy.active['E2'].value
    copy.close()
    assert 'Camila Aguilera Contreras' not in observation
    assert 'Se reitera observación' in observation


def test_zip_anonimizado_se_importa_en_asistente(tmp_path):
    from datetime import date

    from nurus.personal.capture_input import import_capture
    from nurus.personal.config import defaults

    origin, target = tmp_path / 'captura.zip', tmp_path / 'importable_anon.zip'
    zip_fixture(origin, consolidated=True)
    anonymize_capture(origin, target, mode='estricto')
    work = import_capture(target, defaults(), as_of=date(2026, 10, 8))
    assert len(work.rows) == 4
    assert all(row.id for row in work.rows)
