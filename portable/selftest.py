"""Autoprueba del binario: solo datos sintéticos, sin leer archivos del usuario."""
import hashlib
import json
from pathlib import Path
import tempfile
import time

from openpyxl import Workbook, load_workbook
from nurus.capture import read_capture, write_capture
from nurus.capture.fixtures import fixture
from nurus.personal.anonymize import OMITTED, anonymize
from nurus.personal.anonymize_package import UnsafePackage, anonymize_capture


def run(report_path, create_app):
    checks = []
    root = None
    try:
        # Apertura exclusiva: tampoco la autoprueba sobrescribe archivos.
        report = report_path.open('x', encoding='utf-8')
    except OSError:
        return 2
    try:
        with tempfile.TemporaryDirectory(prefix='csmp-autoprueba-') as folder:
            folder = Path(folder)
            source = folder / 'sintetico.xlsx'
            wb = Workbook()
            sh = wb.active
            sh.title = 'ESPERA'
            sh.append(['RIT', 'NOMBRE', 'TRIBUNAL', 'T ESPERA', 'OBSERVACION'])
            sh.append(['P-432-2026', 'Persona Ficticia Uno', 'LAJA', 45,
                       'Se reitera observación respecto de Persona Ficticia Uno.'])
            wb.properties.creator = 'Persona Ficticia Uno'
            wb.save(source)
            wb.close()
            metadata, tables = fixture(count=4)
            capture = folder / 'sintetico.zip'
            write_capture(capture, metadata, tables, consolidated=source.read_bytes())
            originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, capture)}
            root, app = create_app()
            root.withdraw()
            root.update_idletasks()
            assert root.winfo_height() >= root.winfo_reqheight()
            checks.append('ventana_tk_cargada')
            for mode in ('conservador', 'estricto'):
                for input_file in (source, capture):
                    before = set(folder.iterdir())
                    app.source = input_file
                    app.mode.set(mode)
                    app.start()
                    deadline = time.monotonic() + 30
                    while app.go.instate(['disabled']) and time.monotonic() < deadline:
                        root.update()
                        time.sleep(0.01)
                    assert not app.go.instate(['disabled'])
                    assert app.status.get().startswith('Copia creada')
                    generated = set(folder.iterdir()) - before
                    assert len(generated) == 1
                    output = generated.pop()
                    if output.suffix == '.xlsx':
                        result = load_workbook(output)
                        assert result.active['B2'].value != 'Persona Ficticia Uno'
                        assert result.active['D2'].value == 45
                        observation = result.active['E2'].value
                        assert observation == OMITTED if mode == 'estricto' else 'Se reitera' in observation
                        assert result.properties.creator == 'anonimizado'
                        result.close()
                    else:
                        result = read_capture(output)
                        assert len(result.rows('registros')) == 4
                        assert result.consolidated
                        assert result.rows('registros')[0]['ingreso_id'] == result.rows('bitacoras')[0]['ingreso_id']
                    checks.append(input_file.suffix[1:] + '_' + mode)
            for p, digest in originals.items():
                assert hashlib.sha256(p.read_bytes()).hexdigest() == digest
            checks.append('originales_intactos')
            for operation, input_file in ((anonymize, source), (anonymize_capture, capture)):
                try:
                    operation(input_file, input_file)
                except (FileExistsError, ValueError):
                    pass
                else:
                    raise AssertionError('sobrescritura')
                assert hashlib.sha256(input_file.read_bytes()).hexdigest() == originals[input_file]
            checks.append('sin_sobrescritura')
            metadata['originals_authorized'] = True
            unsafe = folder / 'binario.zip'
            write_capture(unsafe, metadata, tables, originals={'originales/sintetico.pdf': b'%PDF FICTICIO'})
            rejected = folder / 'no_debe_existir.zip'
            try:
                anonymize_capture(unsafe, rejected)
            except UnsafePackage:
                pass
            else:
                raise AssertionError('binario admitido')
            assert not rejected.exists()
            checks.append('rechazo_binarios')
        status, code = 'PASA', 0
    except Exception:
        # Nunca persistir excepciones o rutas; el informe solo usa etiquetas fijas.
        status, code = 'FALLA', 1
    finally:
        if root is not None:
            root.destroy()
    with report:
        json.dump({'resultado': status, 'pruebas': checks, 'datos': 'exclusivamente sinteticos'}, report, indent=2)
    return code
