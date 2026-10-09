"""Interfaz experimental: resultados offline, filtros y selección conservada."""
from pathlib import Path
import tempfile
import time
import tkinter as tk
import unittest
from gui_test_isolation import isolated_gui

from descargador import Application,TribunalPicker
from test_lotes import CATALOG,dispose_interface
from test_resultados import fixture
from ventana_resultados import ResultsWindow


class ExperienceTests(unittest.TestCase):
    @isolated_gui
    def test_search_preserves_hidden_choices_and_cancel(self):
        root=tk.Tk();root.withdraw();chosen=[]
        try:
            picker=TribunalPicker(root,CATALOG['tribunales'],('49',),chosen.append)
            picker.search.set('Laja');picker.set_visible(True)
            self.assertTrue(picker.choices['49'].get());self.assertTrue(picker.choices['113'].get())
            picker.destroy();self.assertEqual(chosen,[])
        finally:root.destroy()
    @isolated_gui
    def test_offline_results_load_and_generate(self):
        with tempfile.TemporaryDirectory() as temp:
            root=tk.Tk();root.withdraw();app=Application(root,start_worker=False)
            try:
                folder=fixture(Path(temp)/'lote');app.dest.set(temp);app.last_folder=folder
                popup=ResultsWindow(app);popup.withdraw()
                deadline=time.monotonic()+5
                while popup.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
                self.assertFalse(popup.busy);self.assertIsNotNone(popup.lot);self.assertEqual(len(popup.detail.get_children()),1)
                popup.busy=True;app.close_app();self.assertFalse(app.quitting);popup.busy=False
                from resultados import export_book
                popup.run(lambda:export_book([popup.lot],csmp=True))
                deadline=time.monotonic()+5
                while popup.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
                self.assertFalse(popup.busy);self.assertEqual(len(list(folder.glob('Para CSMP*.xlsx'))),1)
                popup.destroy()
            finally:dispose_interface(app,root)


if __name__=='__main__':unittest.main()


class PreviousDataExperienceTests(unittest.TestCase):
    @isolated_gui
    def test_previous_diaries_and_sheets_build_package_without_chrome(self):
        from unittest.mock import patch
        from preferencias import Preferences
        from archivo_local import remember
        from test_reutilizacion_local import diaries,bound_lot
        from nurus.capture import read_capture
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            base=Path(temp);manifest=diaries(base/'previous',3)
            first=bound_lot(base/'primary');second=bound_lot(base/'reports','Informes')
            root=tk.Tk();root.withdraw();app=Application(root,start_worker=False)
            def settle(window):
                deadline=time.monotonic()+15
                while window.busy and time.monotonic()<deadline:root.update();time.sleep(.02)
                self.assertFalse(window.busy,window.status.get())
            try:
                app.prefs=Preferences(base/'prefs.json');remember(app.prefs,manifest.parent)
                app.dest.set(str(base/'different_destination'));app.last_folder=second.folder
                window=ResultsWindow(app);window.withdraw();settle(window)
                self.assertTrue(any(window.history.item(i,'values')[1]=='Bitácoras' for i in window.history.get_children()))
                with patch('ventana_resultados.filedialog.askdirectory',return_value=str(first.folder)):window.choose()
                settle(window)
                with patch('ventana_resultados.filedialog.askopenfilenames',return_value=[str(manifest)]):window.add_diaries()
                settle(window)
                self.assertEqual(len(window.diary_capture['queries']),3)
                excel=base/'previous.xlsx'
                with patch('ventana_resultados.filedialog.asksaveasfilename',return_value=str(excel)):window.diary_excel()
                settle(window);book=load_workbook(excel);self.assertEqual(book['Copia íntegra'].max_row,4);book.close()
                target=base/'assistant.zip'
                with patch('ventana_resultados.filedialog.asksaveasfilename',return_value=str(target)):window.capture_package()
                settle(window);capture=read_capture(target)
                self.assertEqual(len(capture.rows('registros')),4);self.assertEqual(len(capture.rows('bitacoras')),2)
                self.assertIn('ajenas excluidas: 1',window.status.get())
                window.remove_package_source();self.assertEqual(window.package_diaries(),[])
                self.assertTrue(manifest.is_file());self.assertIn(str(manifest.resolve()),app.prefs.data['historial_bitacoras'])
                window.destroy()
            finally:dispose_interface(app,root)
