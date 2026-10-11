"""Recuperación y reutilización offline: todos los datos son ficticios."""

import hashlib
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from nurus.capture import read_capture
from nurus.capture.diaries import import_lotes, export_diaries
from export_csmp_capture import export_capture
from resultados import load_lot
from preferencias import Preferences
from archivo_local import remember, discover, load_sources
from test_resultados import fixture
from test_lectura_bitacoras import popup


def diaries(folder, count=1, text='Texto completo\nSegunda línea', when='2026-10-08T12:00:00Z'):
    folder.mkdir()
    records = []
    for index in range(count):
        params = {
            'tipo_popUp': '12',
            'COD_Tribunal': '113',
            'CRR_IdCausa': str(100 + index),
            'ID_Ingreso': str(201 + index),
            'COD_Etapa': '3',
            'TIP_Consulta': '2',
            'FLG_MejorNinez': '0',
        }
        raw = popup(params).replace('Texto completo\nSegunda línea', text).encode('utf-8')
        name = f'bitacora_{index}.html'
        (folder / name).write_bytes(raw)
        records.append(
            {
                'tribunal_codigo': '113',
                'causa_id': params['CRR_IdCausa'],
                'ingreso_id': params['ID_Ingreso'],
                'rit': 'X-1-2026',
                'nombre': 'Persona ficticia',
                'programa': 'PRM Centro ficticio',
                'estado': 'LEIDA',
                'params': params,
                'archivo': name,
                'sha256': hashlib.sha256(raw).hexdigest(),
                'capturada': when,
            }
        )
    path = folder / 'bitacoras.json'
    path.write_text(
        json.dumps(
            {'tipo': 'BITACORAS_LECTURA', 'version': 1, 'estado': 'COMPLETA', 'registros': records, 'consultas': []},
            ensure_ascii=False,
        ),
        encoding='utf-8',
    )
    return path


def bound_lot(folder, mode='Cumplimiento'):
    folder = fixture(folder, mode=mode, names=('Persona uno', 'Persona dos'), program='PRM Centro')
    path = folder / 'verificacion 001.json'
    data = json.loads(path.read_text())
    data['archivos'][0]['vinculos_ingreso'] = [
        {
            'tribunal_codigo': '113',
            'rit': 'X-1-2026',
            'rut': '00111111-1',
            'nombre': name,
            'programa': 'PRM Centro',
            'causa_id': str(100 + index),
            'ingreso_id': str(201 + index),
            'etapa': '3',
            'antiguo': '0',
        }
        for index, name in enumerate(('Persona uno', 'Persona dos'))
    ]
    path.write_text(json.dumps(data), encoding='utf-8')
    return load_lot(folder)


def test_800_previous_diaries_export_excel_and_reuse_only_verified_matches_offline(tmp_path):
    manifest = diaries(tmp_path / 'anteriores', 800)
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in manifest.parent.iterdir()}
    excel = export_diaries(manifest, tmp_path / '800_bitacoras.xlsx')
    book = load_workbook(excel, read_only=True)
    try:
        assert book['Copia íntegra'].max_row == 801 and book['Ingresos'].max_row == 801
    finally:
        book.close()
    lots = [bound_lot(tmp_path / 'principal'), bound_lot(tmp_path / 'informes', 'Informes')]
    path = export_capture(lots, tmp_path / 'asistente.zip', diaries=[manifest, manifest], match_diaries=True)
    capture = read_capture(path)
    assert len(capture.rows('registros')) == 4 and len(capture.rows('bitacoras')) == 2
    assert {r['ingreso_id'] for r in capture.rows('bitacoras')} == {'201', '202'}
    assert {r['captured_at'] for r in capture.rows('bitacoras')} == {'2026-10-08T12:00:00Z'}
    assert capture.manifest['table_states']['bitacoras'] == 'PARCIAL'
    note = next(i['message'] for i in capture.rows('incidencias') if i['code'] == 'BITACORAS_REUTILIZADAS')
    assert '798' in note and 'no se acredita actualización remota' in note
    assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in manifest.parent.iterdir()} == before


def test_same_rit_and_name_without_remote_binding_never_attach_previous_diaries(tmp_path):
    manifest = diaries(tmp_path / 'diaries')
    capture = read_capture(
        export_capture(
            [load_lot(fixture(tmp_path / 'unbound'))], tmp_path / 'unbound.zip', diaries=manifest, match_diaries=True
        )
    )
    assert not capture.rows('bitacoras')
    assert capture.manifest['table_states']['bitacoras'] == 'IDENTIDAD_INCIERTA'


def test_duplicate_and_newer_versions_merge_without_overwriting_originals(tmp_path):
    old = diaries(tmp_path / 'old', text='Versión anterior')
    same = diaries(tmp_path / 'same', text='Versión anterior')
    new = diaries(tmp_path / 'new', text='Versión más reciente', when='2026-10-09T12:00:00Z')
    result = import_lotes([old, same, new])
    assert len(result['queries']) == 1
    assert result['queries'][0]['entries'][0]['texto'] == 'Versión más reciente'
    assert result['reuse']['duplicates'] == 2 and len(result['sources']) == 3
    assert 'Versión anterior' in (old.parent / 'bitacora_0.html').read_text(encoding='utf-8')
    conflicting = diaries(tmp_path / 'conflict', text='Versión diferente sin fecha posterior')
    with pytest.raises(ValueError, match='fecha inequívoca'):
        import_lotes([old, conflicting])


def test_history_persists_across_destinations_and_recovers_previous_sources(tmp_path):
    manifest = diaries(tmp_path / 'previous')
    lot = fixture(tmp_path / 'old_download')
    prefs = Preferences(tmp_path / 'prefs.json')
    remember(prefs, manifest.parent)
    remember(prefs, lot)
    restored = Preferences(tmp_path / 'prefs.json')
    rows = discover(tmp_path / 'different_destination', restored)
    assert {kind for _, kind, _, _ in rows} == {'Bitácoras', 'Descarga'}
    assert len(import_lotes(restored.data['historial_bitacoras'])['queries']) == 1
    assert len(load_sources(restored.data['historial_lotes'])) == 1
    (lot / 'pagina.xls').write_bytes(b'alterado')
    with pytest.raises(Exception, match='cambió|hash|verificado'):
        load_sources([lot])


def test_joint_sources_load_all_phases_and_deduplicate_selected_phase(tmp_path):
    root = tmp_path / 'joint'
    root.mkdir()
    first = fixture(root / 'primary')
    second = fixture(root / 'reports', mode='Informes')
    (root / 'flujo.json').write_text(
        json.dumps(
            {'estado': 'COMPLETA', 'fases': {'principal': {'carpeta': 'primary'}, 'informes': {'carpeta': 'reports'}}}
        ),
        encoding='utf-8',
    )
    assert {lot.folder for lot in load_sources([root, first])} == {first.resolve(), second.resolve()}


def test_zip_uses_its_shared_reader_before_a_different_installed_release():
    import os
    import subprocess
    import sys

    root = Path(__file__).resolve().parent
    result = subprocess.run(
        [
            sys.executable,
            '-c',
            "from csmp_shared import parser; from pathlib import Path; import inspect; print(Path(inspect.getfile(parser())).resolve())",
        ],
        cwd=root,
        env={**os.environ, 'PYTHONPATH': ''},
        capture_output=True,
        text=True,
        timeout=15,
    )
    assert result.returncode == 0, result.stderr
    assert Path(result.stdout.strip()) == root.parent / 'src' / 'nurus' / 'bitacora_html.py'
