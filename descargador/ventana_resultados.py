"""Resumen e historial: acciones secundarias fuera del formulario de descarga."""

import os
from pathlib import Path
import queue
import subprocess
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import motor as m
from resultados import (
    load_lot,
    export_book,
    export_csv,
    export_ics,
    export_zip,
    compare_lots,
    inside,
    tables,
    compatibility,
)


class MissingDiariesWindow(tk.Toplevel):
    """Resolver referencias invisibles del historial sin borrar sus originales."""

    def __init__(self, owner, on_ready=None):
        super().__init__(owner)
        self.owner = owner
        self.on_ready = on_ready
        self.title('Bitácoras no disponibles para el paquete')
        self.geometry('1050x520')
        self.minsize(1000, 480)
        self.transient(owner)
        ttk.Label(
            self,
            text='Estos controles se recuerdan en el historial, pero no se encuentran en su ruta. Localízalos con sus HTML o exclúyelos solo de este paquete. El resto de las fuentes se conserva.',
            wraplength=990,
        ).pack(fill='x', padx=12, pady=12)
        table = ttk.Frame(self)
        table.pack(fill='both', expand=True, padx=12)
        table.rowconfigure(0, weight=1)
        table.columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(table, columns=('lote', 'ruta'), show='headings', selectmode='extended')
        self.tree.heading('lote', text='Lote NO_DISPONIBLE')
        self.tree.heading('ruta', text='Control que falta')
        self.tree.column('lote', width=270)
        self.tree.column('ruta', width=740)
        self.tree.grid(row=0, column=0, sticky='nsew')
        vertical = ttk.Scrollbar(table, orient='vertical', command=self.tree.yview)
        vertical.grid(row=0, column=1, sticky='ns')
        horizontal = ttk.Scrollbar(table, orient='horizontal', command=self.tree.xview)
        horizontal.grid(row=1, column=0, sticky='ew')
        self.tree.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        buttons = ttk.Frame(self)
        buttons.pack(fill='x', padx=12, pady=12)
        for label, action in [
            ('Localizar seleccionado…', self.locate),
            ('Excluir seleccionados', self.exclude_selected),
            ('Excluir todos los no disponibles', self.exclude_all),
            ('Continuar', self.finish),
            ('Cancelar', self.close),
        ]:
            ttk.Button(buttons, text=label, command=action).pack(side='left', padx=(0, 6))
        self.status = tk.StringVar()
        ttk.Label(self, textvariable=self.status, wraplength=990).pack(fill='x', padx=12, pady=(0, 12))
        self.protocol('WM_DELETE_WINDOW', self.close)
        self.refresh()

    def refresh(self):
        self.paths = self.owner.missing_package_diaries()
        self.tree.delete(*self.tree.get_children())
        for i, path in enumerate(self.paths):
            self.tree.insert('', 'end', iid=str(i), values=(path.parent.name, str(path)))
        self.owner.update_collection()
        self.status.set(
            f'{len(self.paths)} referencia(s) pendientes. Excluir no borra archivos ni preferencias; puedes añadir el lote nuevamente cuando lo encuentres.'
            if self.paths
            else 'Referencias resueltas. Pulsa Continuar para volver a la exportación.'
        )

    def locate(self):
        selected = self.tree.selection()
        if len(selected) != 1:
            self.status.set('Selecciona un solo lote para localizar su control.')
            return
        old = self.paths[int(selected[0])]
        path = filedialog.askopenfilename(
            parent=self,
            title='Localizar bitacoras.json con sus HTML originales',
            filetypes=[('Control de bitácoras', 'bitacoras.json')],
        )
        if not path:
            return
        try:
            from archivo_local import remember

            path = Path(path).resolve()
            if path.name != 'bitacoras.json' or not path.is_file():
                raise m.PocError('Selecciona bitacoras.json de la carpeta conservada con sus HTML.')
            remember(self.owner.app.prefs, path.parent)
            self.owner.excluded_diaries.discard(path)
            if old != path:
                self.owner.excluded_diaries.add(old)
            self.owner.reload(load_current=False)
            self.refresh()
        except (m.PocError, OSError, ValueError) as exc:
            self.status.set(str(exc))

    def exclude_selected(self):
        selected = self.tree.selection()
        if not selected:
            self.status.set('Selecciona las referencias que deseas excluir.')
            return
        self.owner.excluded_diaries.update(self.paths[int(i)] for i in selected)
        self.refresh()

    def exclude_all(self):
        self.owner.excluded_diaries.update(self.paths)
        self.refresh()

    def finish(self):
        if self.owner.missing_package_diaries():
            self.status.set('Todavía quedan referencias no disponibles. Localízalas o exclúyelas antes de continuar.')
            return
        ready = self.on_ready
        self.close()
        if ready:
            ready()

    def close(self):
        if self.owner:
            self.owner.missing_window = None
        self.on_ready = None
        self.destroy()
        # Liberar variables Tcl y referencias de callbacks en el hilo de interfaz.
        self.status = None
        self.owner = None


def open_local(path):
    path = Path(path).resolve()
    if not path.exists():
        raise m.PocError('El archivo fue movido o ya no existe.')
    if os.name == 'nt':
        os.startfile(str(path))
    else:
        subprocess.Popen(['xdg-open', str(path)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


class ResultsWindow(tk.Toplevel):
    def __init__(self, app):
        from csmp_shared import parser

        parser()
        super().__init__(app.root)
        self.app = app
        self.title('Resultados y herramientas · rama experimental')
        self.geometry('1100x780')
        self.minsize(900, 650)
        self.busy = False
        self.events = queue.Queue()
        self.folders = []
        self.lot = None
        self.package_folders = []
        self.diary_capture = None
        self.loaded_kind = None
        self.excluded_diaries = set()
        self.missing_window = None
        app.result_windows.append(self)
        self.partial = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value='Selecciona un lote del historial o abre su carpeta.')
        self.history = ttk.Treeview(
            self, columns=('fecha', 'tipo', 'estado', 'consultas'), show='headings', height=5, selectmode='extended'
        )
        for col, label in [
            ('fecha', 'Lote'),
            ('tipo', 'Tipo'),
            ('estado', 'Estado'),
            ('consultas', 'Consultas / bitácoras'),
        ]:
            self.history.heading(col, text=label)
        self.history.pack(fill='x', padx=12, pady=12)
        self.history.bind('<<TreeviewSelect>>', self.select_history)
        controls = ttk.Frame(self)
        controls.pack(fill='x', padx=12)
        ttk.Button(controls, text='Añadir descarga o lote anterior…', command=self.choose).pack(side='left')
        ttk.Button(controls, text='Añadir bitácoras anteriores…', command=self.add_diaries).pack(side='left', padx=8)
        ttk.Button(controls, text='Quitar lote del paquete', command=self.remove_package_source).pack(
            side='left', padx=8
        )
        ttk.Button(controls, text='Actualizar historial', command=self.reload).pack(side='left', padx=8)
        ttk.Checkbutton(
            controls, text='Permitir resultados INCOMPLETOS', variable=self.partial, command=self.load
        ).pack(side='left')
        self.package_consolidated = tk.BooleanVar(value=True)
        ttk.Checkbutton(controls, text='Paquete con XLSX consolidado', variable=self.package_consolidated).pack(
            side='left', padx=(8, 0)
        )
        self.detail = ttk.Treeview(self, columns=('consulta', 'estado', 'filas', 'aptitud'), show='headings', height=8)
        for col, label in [
            ('consulta', 'Consulta'),
            ('estado', 'Resultado'),
            ('filas', 'Registros'),
            ('aptitud', 'Compatibilidad CSMP'),
        ]:
            self.detail.heading(col, text=label)
        self.detail.column('consulta', width=350)
        self.detail.pack(fill='both', expand=True, padx=12, pady=12)
        self.detail.bind('<Double-1>', lambda _: self.open_source())
        row = ttk.Frame(self)
        row.pack(fill='x', padx=12)
        for text, action in [
            ('Consolidar XLSX', lambda: self.run(lambda: export_book([self.lot]))),
            ('Libro CSMP (Excel)', self.csmp),
            ('Añadir informes al cruce…', self.cross),
            ('Comparar con otro lote…', self.compare),
        ]:
            ttk.Button(row, text=text, command=action).pack(side='left', padx=(0, 6))
        row = ttk.Frame(self)
        row.pack(fill='x', padx=12, pady=8)
        for text, action in [
            ('CSV', lambda: self.run(lambda: export_csv(self.lot))),
            ('Calendario ICS', lambda: self.run(lambda: export_ics(self.lot))),
            ('ZIP de originales del lote', lambda: self.run(lambda: export_zip(self.lot))),
            ('Abrir carpeta', lambda: self.safe(lambda: open_local(self.lot.folder))),
            ('Preparar pendientes', self.pending),
        ]:
            ttk.Button(row, text=text, command=action).pack(side='left', padx=(0, 6))
        row = ttk.Frame(self)
        row.pack(fill='x', padx=12, pady=4)
        ttk.Button(row, text='Excel bitácoras autónomo…', command=self.diary_excel).pack(side='left', padx=(0, 6))
        ttk.Button(row, text='Crear paquete para Asistente (ZIP)…', command=self.capture_package).pack(
            side='left', padx=(0, 6)
        )
        ttk.Button(row, text='Ver descarga conjunta…', command=self.show_joint).pack(side='left', padx=(0, 6))
        self.collection = tk.StringVar()
        ttk.Label(self, textvariable=self.collection, wraplength=1050).pack(fill='x', padx=12, pady=4)
        ttk.Button(self, text='Archivar o borrar lotes antiguos', command=self.manage_old_batches).pack(
            anchor='w', padx=12, pady=4
        )
        ttk.Button(self, text='Revisar referencias no disponibles…', command=self.review_missing_diaries).pack(
            anchor='w', padx=12, pady=4
        )
        ttk.Label(
            self,
            text='Añade las carpetas de hojas y los controles bitacoras.json con sus HTML. El ZIP reúne las descargas añadidas; reutiliza bitácoras solo por tribunal, causa e ingreso. Importa ese ZIP en Trabajo → Importar paquete del Asistente.',
            wraplength=1050,
        ).pack(fill='x', padx=12, pady=4)
        ttk.Label(self, textvariable=self.status, wraplength=1050).pack(fill='x', padx=12, pady=8)
        self.protocol('WM_DELETE_WINDOW', self.close)
        self.reload()
        self.after(100, self.poll)

    def retention_protected(self):
        paths = [self.app.last_folder]
        for window in self.app.result_windows:
            if window.winfo_exists():
                paths.append(getattr(window, 'folder', None))
                paths.append(getattr(window, 'root_folder', None))
                paths.extend(getattr(window, 'package_folders', []))
                diaries = getattr(window, 'package_diaries', None)
                if callable(diaries):
                    paths.extend(Path(control).parent for control in diaries())
        if self.app.recovery:
            paths.append(self.app.recovery[0].folder if hasattr(self.app.recovery[0], 'folder') else None)
        return [path for path in paths if path]

    def manage_old_batches(self):
        if self.busy or self.app.busy:
            self.status.set('Espera a que termine la operación para gestionar lotes antiguos.')
            return
        from preferencias import old_batches, RETENCION_DIAS

        window = tk.Toplevel(self)
        window.title('Archivar o borrar lotes antiguos')
        window.geometry('820x420')
        window.transient(self)
        ttk.Label(
            window,
            text=f'Lotes con más de {RETENCION_DIAS} días en la carpeta de descarga. Se conservan los lotes en uso.',
            wraplength=780,
        ).pack(padx=12, pady=10)
        table = ttk.Treeview(
            window, columns=('lote', 'fecha', 'tamaño'), show='headings', selectmode='extended', height=9
        )
        for key, label in [('lote', 'Lote'), ('fecha', 'Último control'), ('tamaño', 'Tamaño')]:
            table.heading(key, text=label)
        table.pack(fill='both', expand=True, padx=12, pady=8)
        rows = {}

        def refresh():
            from datetime import datetime

            table.delete(*table.get_children())
            rows.clear()
            for index, item in enumerate(old_batches(self.app.dest.get(), protected=self.retention_protected())):
                iid = str(index)
                rows[iid] = item['folder']
                table.insert(
                    '',
                    'end',
                    iid=iid,
                    values=(
                        item['folder'].name,
                        datetime.fromtimestamp(item['modified']).strftime('%d/%m/%Y'),
                        f"{item['bytes'] / 1048576:.2f} MB",
                    ),
                )

        preserve = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            window, text='Al borrar, conservar Bitacoras_RUS.xlsx y diagnostico_bitacoras.json', variable=preserve
        ).pack(anchor='w', padx=12, pady=5)
        actions = ttk.Frame(window)
        actions.pack(fill='x', padx=12, pady=10)

        def act(archive=False):
            selected = [rows[iid] for iid in table.selection()]
            if not selected:
                self.status.set('Selecciona los lotes antiguos que quieres gestionar.')
                return
            destination = None
            if archive:
                destination = filedialog.askdirectory(parent=window, title='Carpeta donde archivar los lotes completos')
                if not destination:
                    return
            try:
                self.apply_batch_retention(selected, preserve=preserve.get(), archive=destination, parent=window)
                refresh()
                self.reload(load_current=False)
            except (m.PocError, OSError, ValueError) as exc:
                self.status.set(str(exc))

        ttk.Button(actions, text='Archivar seleccionados…', command=lambda: act(True)).pack(side='left', padx=5)
        ttk.Button(actions, text='Borrar seleccionados…', command=act).pack(side='left', padx=5)
        ttk.Button(actions, text='Actualizar', command=refresh).pack(side='left', padx=5)
        refresh()

    def apply_batch_retention(self, selected, *, preserve=True, archive=None, parent=None):
        if self.busy or self.app.busy:
            raise m.PocError('Hay una operación en curso; los lotes se conservan.')
        from preferencias import manage_old_batches, RETENCION_DIAS

        action = 'archivarán completos' if archive else 'borrarán'
        note = ' Se conservarán Bitacoras_RUS.xlsx y diagnostico_bitacoras.json.' if preserve and not archive else ''
        confirmed = messagebox.askyesno(
            'Confirmar gestión local',
            f'Se {action} {len(selected)} lotes de más de {RETENCION_DIAS} días.{note} ¿Continuar?',
            parent=parent or self,
        )
        if not confirmed:
            self.status.set('No se cambió ningún lote.')
            return []
        changed = manage_old_batches(
            self.app.dest.get(),
            selected,
            confirmed=True,
            preserve=preserve,
            archive=archive,
            protected=self.retention_protected(),
        )
        self.status.set(f'{len(changed)} lotes antiguos gestionados. Los lotes vigentes se conservan.')
        return changed

    def close(self):
        if self.busy:
            self.status.set('Espera a que termine la escritura local.')
            return
        self.destroy()

    def safe(self, action):
        try:
            if self.busy:
                raise m.PocError('Hay una operación de resultados en curso.')
            if not self.lot:
                raise m.PocError('Selecciona un lote primero.')
            action()
        except (m.PocError, OSError, ValueError) as exc:
            self.status.set(str(exc))

    def reload(self, load_current=True):
        if self.busy:
            return
        self.history.delete(*self.history.get_children())
        self.folders = []
        from archivo_local import discover, remember

        for folder, label, state, count in discover(self.app.dest.get(), self.app.prefs, self.app.last_folder):
            iid = str(len(self.folders))
            self.folders.append(folder)
            self.history.insert('', 'end', iid=iid, values=(folder.name, label, state, count))
            if label == 'Bitácoras' and state not in ('NO_DISPONIBLE', 'ERROR_CONTROL'):
                remember(self.app.prefs, folder)
        self.update_collection()
        if load_current and (getattr(self, 'folder', None) or self.app.last_folder):
            self.folder = getattr(self, 'folder', None) or self.app.last_folder
            self.load()

    def update_collection(self):
        names = '; '.join(p.name for p in self.package_folders)
        references = self.package_diaries()
        missing = self.missing_package_diaries()
        unavailable = (
            f' · {len(missing)} referencia(s) NO_DISPONIBLE: recuperar o excluir antes de exportar' if missing else ''
        )
        self.collection.set(
            f'Paquete: {len(self.package_folders)} descarga(s) añadida(s): {names} · {len(references) - len(missing)} lote(s) de bitácoras disponibles para este paquete{unavailable}. Las fechas originales se conservan; reutilizar no vuelve a consultar RUS.'
        )

    def package_diaries(self):
        from archivo_local import diary_references

        return diary_references(self.app.prefs, self.excluded_diaries)

    def missing_package_diaries(self):
        return [p for p in self.package_diaries() if not p.is_file()]

    def review_missing_diaries(self, on_ready=None):
        if self.busy:
            return
        self.update_collection()
        if not self.missing_package_diaries():
            self.status.set('Los controles de bitácoras están disponibles. Sus originales se verifican al exportar.')
            if on_ready:
                on_ready()
            return
        self.status.set(
            'Hay referencias antiguas no disponibles. Localiza sus controles o exclúyelas de este paquete; no se borran los lotes.'
        )
        if self.missing_window and self.missing_window.winfo_exists():
            self.missing_window.lift()
            return
        self.missing_window = MissingDiariesWindow(self, on_ready)

    def check_diary_references(self, manifests, on_ready):
        if any(not Path(p).is_file() for p in manifests):
            if self.missing_package_diaries():
                self.review_missing_diaries(on_ready)
            else:
                self.status.set(
                    'El lote de bitácoras seleccionado no está disponible. Localiza bitacoras.json con sus HTML o selecciona una descarga para consolidar los otros lotes disponibles.'
                )
            return False
        return True

    def remove_package_source(self):
        if self.busy:
            return
        selected = self.history.selection()
        targets = [(self.folders[int(i)], self.history.item(i, 'values')[1]) for i in selected]
        if not targets and getattr(self, 'folder', None):
            targets = [(self.folder, self.loaded_kind)]
        for folder, label in targets:
            if label == 'Bitácoras':
                self.excluded_diaries.add((folder / 'bitacoras.json').resolve())
            if folder in self.package_folders:
                self.package_folders.remove(folder)
        self.update_collection()
        self.status.set(
            'Lote retirado de este paquete; sus archivos e historial se conservan. Añadir de nuevo lo incorpora otra vez.'
        )

    def select_history(self, event=None):
        selected = self.history.selection()
        if selected and not self.busy:
            self.folder = self.folders[int(selected[0])]
            self.load()

    def choose(self):
        if self.busy:
            return
        folder = filedialog.askdirectory(parent=self, title='Carpeta del lote')
        if folder:
            from archivo_local import remember, kind

            try:
                self.folder = remember(self.app.prefs, folder)
                self.excluded_diaries.discard((self.folder / 'bitacoras.json').resolve())
                if kind(self.folder)[1] != 'Bitácoras' and self.folder not in self.package_folders:
                    self.package_folders.append(self.folder)
                self.reload()
                self.load()
            except (m.PocError, OSError, ValueError) as exc:
                self.status.set(str(exc))

    def add_diaries(self):
        if self.busy:
            return
        paths = filedialog.askopenfilenames(
            parent=self,
            title='Controles bitacoras.json con sus HTML originales',
            filetypes=[('Control de bitácoras', 'bitacoras.json')],
        )
        if not paths:
            return
        from archivo_local import remember

        try:
            for path in paths:
                remember(self.app.prefs, Path(path).parent)
                self.excluded_diaries.discard(Path(path).resolve())
            self.folder = Path(paths[0]).parent
            self.reload()
            self.update_collection()
        except (m.PocError, OSError, ValueError) as exc:
            self.status.set(str(exc))

    def load(self):
        if self.busy or not getattr(self, 'folder', None):
            return
        self.lot = None
        self.diary_capture = None
        from archivo_local import kind

        self.detail.delete(*self.detail.get_children())
        self.loaded_kind = None
        try:
            _, self.loaded_kind = kind(self.folder)
        except m.PocError:
            if (self.folder / 'bitacoras.json').resolve() in self.package_diaries():
                self.loaded_kind = 'Bitácoras'
            self.status.set(
                'NO_DISPONIBLE · '
                + str(self.folder)
                + ' · Añade la carpeta en su ubicación actual o usa Quitar lote del paquete.'
            )
            return
        if self.loaded_kind == 'Bitácoras':
            from nurus.capture.diaries import import_lotes

            self.start(lambda: import_lotes(self.folder / 'bitacoras.json'), 'diary_load')
            return
        if self.loaded_kind == 'Conjunto':
            if self.folder not in self.package_folders:
                self.package_folders.append(self.folder)
            self.update_collection()
            self.status.set(
                'Conjunto disponible. Crear paquete reúne sus fases; Ver descarga conjunta permite abrir el libro y recuperar pendientes.'
            )
            return
        if self.folder not in self.package_folders:
            self.package_folders.append(self.folder)
        self.update_collection()
        # Capturar variables Tk en el hilo de interfaz.
        partial = self.partial.get()
        folder = self.folder

        def worker_action():
            lot = load_lot(folder, partial)
            aptitudes = {}
            for query, selection, item, headers, rows in tables(lot):
                check = compatibility(headers, selection, rows)
                aptitudes.setdefault(query['consulta'], []).append(
                    check['estado'] + (' · faltan ' + ', '.join(check['faltan']) if check['faltan'] else '')
                )
            return lot, aptitudes

        self.start(worker_action, 'load')

    def start(self, action, kind='product'):
        if self.busy:
            return
        self.busy = True
        self.status.set('Comprobando archivos y procesando resultados localmente…')

        def work():
            try:
                self.events.put((kind, action(), None))
            except Exception as exc:
                self.events.put(
                    (
                        kind,
                        None,
                        str(exc)
                        if isinstance(exc, (m.PocError, OSError, ValueError))
                        else 'E_RESULTADOS: no se pudo generar el producto; los originales se conservan.',
                    )
                )

        threading.Thread(target=work, daemon=True).start()

    def run(self, action):
        self.safe(lambda: self.start(action))

    def poll(self):
        try:
            while True:
                kind, result, error = self.events.get_nowait()
                self.busy = False
                if error:
                    self.status.set(error)
                    continue
                try:
                    self.show_result(kind, result)
                except Exception as exc:
                    # Un resultado que no se puede mostrar no detiene la cola ni deja la ventana ocupada.
                    self.status.set(f'No se pudo completar la operación: {exc}')
        except queue.Empty:
            pass
        finally:
            self.after(100, self.poll)

    def show_result(self, kind, result):
        if kind == 'load':
            self.lot, aptitudes = result
            self.detail.delete(*self.detail.get_children())
            for i, (q, manifest) in enumerate(self.lot.queries):
                self.detail.insert(
                    '',
                    'end',
                    iid=str(i),
                    values=(
                        q['consulta'],
                        manifest['estado'],
                        sum(v['registros'] for v in manifest.get('archivos', [])),
                        '; '.join(dict.fromkeys(aptitudes.get(q['consulta'], ['Sin registros']))),
                    ),
                )
            self.status.set(
                ('INCOMPLETO · ' if self.lot.partial else '')
                + self.lot.folder.name
                + ' · Doble clic en una consulta para abrir su primer archivo verificado.'
            )
        elif kind == 'diary_load':
            self.diary_capture = result
            self.detail.delete(*self.detail.get_children())
            for i, query in enumerate(result['queries']):
                self.detail.insert(
                    '',
                    'end',
                    iid=str(i),
                    values=(
                        query['rit'] + ' · ' + query['persona'],
                        query['coverage'],
                        len(query['entries']),
                        'Identidad de tribunal/causa/ingreso conservada',
                    ),
                )
            self.status.set(
                f'{len(result["queries"])} ingresos recuperados de bitácoras anteriores. Excel bitácoras autónomo genera su consolidado sin Chrome.'
            )
        elif kind == 'capture':
            self.status.set(
                'Paquete creado: '
                + str(result['path'])
                + ' · '
                + result['reuse']
                + ' Importa este ZIP en Trabajo → Importar paquete del Asistente.'
            )
        else:
            self.status.set(
                'Creado: ' + ('; '.join(str(p) for p in result) if isinstance(result, list) else str(result))
            )
            if kind == 'csmp':
                self.launch_csmp(result)

    def open_source(self):
        if self.diary_capture:
            selected = self.detail.selection()
            if selected:
                query = self.diary_capture['queries'][int(selected[0])]
                try:
                    open_local(inside(Path(query['manifest_path']).parent, query['source']))
                except (m.PocError, OSError, KeyError) as exc:
                    self.status.set('No hay una copia HTML verificable para esta lectura: ' + str(exc))
            return

        def action():
            selected = self.detail.selection()
            if not selected:
                return
            _, manifest = self.lot.queries[int(selected[0])]
            items = manifest.get('archivos', []) or [manifest.get('evidencia', {})]
            if items and items[0].get('archivo'):
                open_local(inside(self.lot.folder, items[0]['archivo']))

        self.safe(action)

    def csmp(self):
        self.safe(lambda: self.start(lambda: export_book([self.lot], csmp=True)))

    def diary_excel(self):
        if self.busy:
            return
        manifests = [self.folder / 'bitacoras.json'] if self.loaded_kind == 'Bitácoras' else self.package_diaries()
        if not manifests:
            self.add_diaries()
            return
        if not self.check_diary_references(manifests, self.diary_excel):
            return
        target = filedialog.asksaveasfilename(parent=self, initialfile='Bitacoras_RUS.xlsx', defaultextension='.xlsx')
        if target:
            from nurus.capture.diaries import export_diaries

            self.start(lambda: export_diaries(manifests, target))

    def capture_package(self):
        if self.busy:
            return
        if not self.package_folders:
            self.status.set(
                'Añade primero una carpeta de descarga de hojas RUS. Las bitácoras anteriores seguirán disponibles.'
            )
            return
        if not self.check_diary_references(self.package_diaries(), self.capture_package):
            return
        target = filedialog.asksaveasfilename(parent=self, initialfile='Captura_CSMP.zip', defaultextension='.zip')
        if not target:
            return
        folders = list(self.package_folders)
        diaries = self.package_diaries()
        partial = self.partial.get()
        wanted = self.package_consolidated.get()
        from export_csmp_capture import export_capture, consolidated_note
        from archivo_local import load_sources

        def export():
            from nurus.capture import read_capture

            lots = load_sources(folders, partial)
            include, note = consolidated_note(lots, wanted)
            path = export_capture(lots, target, diaries=diaries or None, match_diaries=True, consolidated=include)
            capture = read_capture(path)
            summary = next(
                (i['message'] for i in capture.rows('incidencias') if i['code'] == 'BITACORAS_REUTILIZADAS'),
                'Sin bitácoras añadidas.',
            )
            return {'path': path, 'reuse': summary + note}

        self.start(export, 'capture')

    def show_joint(self):
        if self.busy:
            return
        if self.loaded_kind != 'Conjunto':
            self.status.set('Selecciona un conjunto del historial para abrir sus resultados.')
            return
        from ventana_flujo import JointWindow

        previous = self.app.last_folder
        self.app.last_folder = self.folder
        try:
            JointWindow(self.app)
        except (m.PocError, OSError, ValueError, KeyError) as exc:
            self.status.set('No se pudo abrir el conjunto: ' + str(exc))
        finally:
            self.app.last_folder = previous

    def cross(self):
        def action():
            folder = filedialog.askdirectory(parent=self, title='Lote de informes compatible para el cruce')
            if folder:
                current = self.lot

                def export():
                    other = load_lot(folder)
                    modes = {
                        compatibility(h, s, rows)['modo']
                        for lot in (current, other)
                        for _, s, _, h, rows in tables(lot)
                    }
                    if modes != {'CUMPLIMIENTO', 'INFORMES'}:
                        raise m.PocError(
                            'El cruce requiere un lote de Cumplimiento y otro de Informes, sin otras pantallas.'
                        )
                    return export_book([current, other], csmp=True)

                self.start(export)

        self.safe(action)

    def compare(self):
        def action():
            folder = filedialog.askdirectory(parent=self, title='Lote anterior del mismo alcance')
            if folder:
                current = self.lot
                self.start(lambda: compare_lots(load_lot(folder), current))

        self.safe(action)

    def pending(self):
        self.safe(lambda: self.app.prepare_pending(self.lot.folder))

    def launch_csmp(self, path):
        executable = self.app.prefs.data.get('csmp', '')
        if not executable:
            if not messagebox.askyesno(
                'Libro preparado',
                '¿Quieres abrirlo en la rama experimental de CSMP? Selecciona csmp-assistant.exe o pythonw.exe de su entorno.',
                parent=self,
            ):
                return
            executable = filedialog.askopenfilename(
                parent=self, title='Ejecutable de CSMP o Python de su entorno', filetypes=[('Ejecutable', '*.exe')]
            )
            if not executable:
                return
            self.app.prefs.data['csmp'] = executable
            self.app.prefs.save()
        try:
            args = [executable]
            if Path(executable).stem.lower() in ('python', 'pythonw'):
                args += ['-m', 'nurus.personal.app']
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=True)
            try:
                modes = [name for name in ('ESPERA', 'CUMPLIMIENTO', 'INFORMES') if name in book.sheetnames]
            finally:
                book.close()
            extra = ['--modo', modes[0], '--hoja', modes[0]] if len(modes) == 1 else []
            subprocess.Popen(
                [*args, '--archivo', str(path), *extra], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL
            )
        except OSError:
            self.status.set('Libro creado. No se pudo abrir CSMP: revisa su ejecutable configurado.')
