"""Flujo completo ficticio de cuatro modalidades, páginas y fallas recuperables."""

import base64
import hashlib
import html
import json
import threading
from pathlib import Path
import pytest
from perfiles import Selection, selection_profile
from lotes import Batch
from lectura_bitacoras import DiaryRunner
from test_lotes import CATALOG, FakeBridge, document
from test_vinculos import listing
from csmp_shared import parser
from contrato_integral import VERSION


def popup(params, wrong=False):
    from nurus.bitacora_html import HEADERS

    fields = {
        **params,
        'tipo_popUp': '',
        'RIT_Causa': 'X-1-2026',
        'GLS_Nombre': 'Persona ficticia',
        'GLS_Centro': 'PRM Centro ficticio',
    }
    if wrong:
        fields['ID_Ingreso'] = '999'
    hidden = {
        'HOBS_Centro_': 'Texto completo\nSegunda línea',
        'HOBS_Tribunal_': '---',
        'HCOD_Estado_': '6',
        'HFUN_Tribunal_': '',
        'HFLG_Envio_': '1',
    }
    headers = ''.join(f'<td>{h}</td>' for h in HEADERS)
    row = (
        ''.join(
            f'<td>{v}</td>'
            for v in [
                'Al Tribunal',
                'Cumplimiento',
                '12/08/2026 10:13',
                'Usuario ficticio',
                'Texto recortado',
                '---',
                '---',
                '---',
            ]
        )
        + '<td><img id="44"></td>'
    )
    # Los campos completos pertenecen a la tabla, tal como en RUS.
    return (
        f'<html><form name="InformesPpalForm" action="/SITFAWEB/InformesDAction.do">'
        + ''.join(f'<input type="hidden" name="{k}" value="{html.escape(v, quote=True)}">' for k, v in fields.items())
        + f'<table id="TablaInforme"><tr>{headers}</tr><tr>{row}</tr>'
        + ''.join(f'<input type="hidden" name="{k}44" value="{html.escape(v, quote=True)}">' for k, v in hidden.items())
        + '</table></form></html>'
    )


class Bridge(FakeBridge):
    connected = True

    def __init__(self, wrong=False, missing=False):
        super().__init__(total=2)
        self.wrong = wrong
        self.missing = missing
        self.openings = []

    def call(self, command, payload=None, timeout=75):
        if command == 'search':
            result = super().call(command, payload, timeout)
            profile = selection_profile(self.selection, CATALOG)
            source = base64.b64decode(result['body']).decode()
            start = source.index('<table')
            end = source.index('</table>', start) + 8
            from lxml import etree

            table = listing(court=self.selection.tribunal)
            table.set('id', profile.table_id)
            text = etree.tostring(table, encoding='unicode').replace("'200'", f"'{200 + self.number}'")
            kind = {'Espera': '11', 'Cumplimiento': '12', 'Informes': '13'}[self.selection.tab]
            text = text.replace("'12'", repr(kind))
            if self.missing:
                text = text.replace('ShowObservaciones', 'MissingObservaciones')
            # TD encabezado y RUT de la pestaña Espera observados en RUS.
            text = text.replace('<th>', '<td>').replace('</th>', '</td>').replace('>RUT<', '>RUT (-&gt;RCeI)<')
            data = (source[:start] + text + source[end:]).encode()
            result['body'] = base64.b64encode(data).decode()
            return result
        if command == 'diary_open':
            self.events.append((command, None))
            params = payload['params']
            self.openings.append(params)
            data = popup(params, self.wrong and params['ID_Ingreso'] == '202').encode()
            return {
                'status': 200,
                'type': 'text/html',
                'body': base64.b64encode(data).decode(),
                'digest': hashlib.sha256(data).hexdigest(),
                'params': params,
                'startedAt': '2026-10-08T12:00:00Z',
            }
        return super().call(command, payload, timeout)


def run(tmp_path, bridge, batch=None, resume=None, cancel=None, emit=None):
    parser()
    return DiaryRunner(bridge, CATALOG, emit or (lambda *a: None), cancel or threading.Event(), pause=0).run(
        batch or Batch(('49',), ('2',), ('Cumplimiento',)), tmp_path, resume
    )


def test_four_modalities_all_courts_and_two_tabs(tmp_path):
    bridge = Bridge()
    result = run(
        tmp_path, bridge, Batch(tuple(CATALOG['tribunales']), ('1', '2', '3', '4'), ('Espera', 'Cumplimiento'))
    )
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert result['estado'] == 'COMPLETA' and len(data['registros']) == 32
    assert {p['TIP_Consulta'] for p in bridge.openings} == {'1', '2', '3', '4'}
    assert all(q['paginas'] == 2 for q in data['consultas'])
    assert set(cmd for cmd, _ in bridge.events) <= {'capabilities', 'lock', 'unlock', 'prepare', 'search', 'diary_open'}
    assert bridge.events[-1][0] == 'unlock'


def test_same_rit_different_ingresos_and_failed_popup(tmp_path):
    bridge = Bridge(wrong=True)
    result = run(tmp_path, bridge)
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert [r['ingreso_id'] for r in data['registros']] == ['201', '202']
    assert result['leidas'] == 1 and result['fallidas'] == 1 and result['estado'] == 'INCOMPLETA'
    bridge = Bridge()
    recovered = run(tmp_path, bridge, resume=result['folder'])
    assert recovered['leidas'] == 2 and len(bridge.openings) == 1
    assert bridge.openings[0]['ID_Ingreso'] == '202'


def test_visible_identity_difference_is_warning_in_manifest_and_excel(tmp_path):
    from nurus.personal.bitacora_lote import import_lote

    class DisplayVariant(Bridge):
        def call(self, command, payload=None, timeout=75):
            result = super().call(command, payload, timeout)
            if command == 'diary_open':
                body = base64.b64decode(result['body']).decode('utf-8')
                body = body.replace('name="RIT_Causa" value="X-1-2026"', 'name="RIT_Causa" value="X-0001-2026"')
                body = body.replace(
                    'name="GLS_Nombre" value="Persona ficticia"', 'name="GLS_Nombre" value="Otra grafía del nombre"'
                )
                raw = body.encode('utf-8')
                result['body'] = base64.b64encode(raw).decode()
                result['digest'] = hashlib.sha256(raw).hexdigest()
            return result

    initial = run(tmp_path, DisplayVariant())
    assert initial['estado'] == 'COMPLETA' and initial['leidas'] == 2
    rows = json.loads(Path(initial['manifest']).read_text(encoding='utf-8'))['registros']
    assert all(row['estado'] == 'LEIDA' for row in rows)
    assert all('RIT visible' in row['error'] and 'nombre visible' in row['error'] for row in rows)
    imported = import_lote(initial['manifest'])
    assert len(imported['queries']) == 2
    assert all(len(q['warnings']) >= 2 for q in imported['queries'])
    sheet = next(rows for title, headers, rows in imported['extra_sheets'] if title == 'Lecturas')
    assert all(row[10] == 'LEIDA' and 'RIT visible' in row[11] for row in sheet)
    bridge = Bridge()
    resumed = run(tmp_path, bridge, resume=initial['folder'])
    assert resumed['leidas'] == 2 and not bridge.openings
    resumed_rows = json.loads(Path(resumed['manifest']).read_text(encoding='utf-8'))['registros']
    assert all('nombre visible' in r['error'] for r in resumed_rows)


def test_missing_link_never_invents_get_and_all_rows_counted(tmp_path):
    bridge = Bridge(missing=True)
    result = run(tmp_path, bridge)
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert len(data['registros']) == 2 and not bridge.openings
    assert all(r['estado'] == 'SIN_VINCULO' for r in data['registros'])


def test_espera_without_historia_and_rut_script_reads_verified_ingresos(tmp_path):
    import re

    class EsperaSinHistoria(Bridge):
        def call(self, command, payload=None, timeout=75):
            result = super().call(command, payload, timeout)
            if command == 'search':
                source = base64.b64decode(result['body']).decode('utf-8')
                source = re.sub(r'<a onclick="ShowHistoria[^"]*">Causa</a>', '', source)
                source = source.replace('11111111-1</td>', '11111111-1<script>document.write("Imagen");</script></td>')
                raw = source.encode('utf-8')
                result['body'] = base64.b64encode(raw).decode('ascii')
            return result

    bridge = EsperaSinHistoria()
    result = run(tmp_path, bridge, Batch(('49',), ('2',), ('Espera',)))
    records = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))['registros']
    assert result['estado'] == 'COMPLETA' and result['leidas'] == 2
    assert len(bridge.openings) == 2
    assert {r['ingreso_id'] for r in records} == {'201', '202'}
    assert all(r['rut'] == '11111111-1' and r['estado'] == 'LEIDA' for r in records)
    assert all('vinculo_diagnostico' not in r for r in records)


@pytest.mark.parametrize('version', ['2.7.0', '2.7.2'])
def test_version_incompatible_fails_before_query(tmp_path, version):
    class OlderExtension(Bridge):
        def call(self, command, payload=None, timeout=75):
            result = super().call(command, payload, timeout)
            if command == 'capabilities':
                result['version'] = version
            return result

    bridge = OlderExtension()
    result = run(tmp_path, bridge)
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert result['estado'] == 'INCOMPLETA' and not bridge.openings
    assert not data['consultas'] and not data['registros']
    assert 'versión distinta' in result['detalle']


def test_anonymous_diagnostic_has_totals_but_no_personal_data(tmp_path):
    from lectura_bitacoras import anonymous_diagnostic

    bridge = Bridge(missing=True)
    result = run(tmp_path, bridge)
    path = Path(result['diagnostico_anonimo'])
    assert path.is_file()
    source = path.read_text(encoding='utf-8')
    data = json.loads(source)
    assert data['tipo'] == 'BITACORAS_DIAGNOSTICO_ANONIMO'
    assert data['estados'] == {'SIN_VINCULO': 2}
    assert data['registros'] == 2
    assert data['descargador_version'] in ('NO_VERIFICADA', VERSION)
    assert data['extension_version'] == VERSION
    assert data['descargador_commit'] == 'NO_VERIFICADO' or len(data['descargador_commit']) == 40
    assert data['diagnosticos_sin_vinculo']['SIN_ACCION_OBSERVACIONES'] == 2
    assert data['consultas'][0]['registros'] == 2
    assert data['contiene_identificadores_personales'] is False
    assert all(
        value not in source
        for value in (
            'Persona ficticia',
            'X-1-2026',
            '11111111-1',
            'PRM Centro ficticio',
            '\\u0022rut\\u0022',
            '\\u0022nombre\\u0022',
        )
    )
    # Ni siquiera un manifiesto malicioso puede inyectar texto libre al reporte.
    malicious = {
        'registros': [
            {
                'estado': 'FALLIDA',
                'rit': 'X-1-2026',
                'nombre': 'Persona ficticia',
                'error': 'Texto privado NNA',
                'diagnostico': {'codigo_extension': 'Texto privado NNA'},
            }
        ],
        'consultas': [{'tribunal': 'Texto privado NNA', 'modality': '1', 'tab': 'Espera', 'estado': 'FALLIDA'}],
        'origen': {'version': 'Texto privado NNA'},
        'extension': {'version': '2.7.5'},
    }
    sanitized = json.dumps(anonymous_diagnostic(malicious), ensure_ascii=False)
    assert 'Texto privado NNA' not in sanitized
    assert 'Persona ficticia' not in sanitized


def test_anonymous_diagnostic_keeps_egresados_tab():
    from lectura_bitacoras import anonymous_diagnostic

    data = {
        'registros': [],
        'consultas': [{'tribunal': '1', 'modality': '1', 'tab': 'Egresados', 'estado': 'ENUMERADA'}],
    }
    assert anonymous_diagnostic(data)['consultas'][0]['pestana'] == 'Egresados'


def test_extension_declared_version_matches_manifest():
    from contrato_integral import VERSION

    folder = Path(__file__).parent / 'extension'
    manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
    script = (folder / 'pagina.js').read_text(encoding='utf-8')
    assert f"if(command==='capabilities') return {{tipo:'CSMP_RUS_CAPACIDADES',version:'{VERSION}'" in script
    assert manifest['version'] == VERSION


@pytest.mark.parametrize(
    'tab,report,kind',
    [
        ('Espera', '0', '11'),
        ('Cumplimiento', '0', '12'),
        ('Informes', '1', '13'),
        ('Informes', '2', '13'),
        ('Informes', '3', '13'),
    ],
)
def test_all_followup_observation_popup_kinds_are_preserved_and_excel_created(tmp_path, tab, report, kind):
    from openpyxl import load_workbook

    bridge = Bridge()
    result = run(tmp_path, bridge, Batch(('49',), ('1', '2', '3', '4'), (tab,), report=report))
    assert result['estado'] == 'COMPLETA' and result['leidas'] == 8
    assert {p['tipo_popUp'] for p in bridge.openings} == {kind}
    book = load_workbook(result['excel'], read_only=True)
    try:
        # Cuatro modalidades repiten los mismos dos ingresos: dos copias
        # semánticas, pero se conservan las ocho lecturas en el control.
        assert book['Copia íntegra'].max_row == 3 and book['Lecturas'].max_row == 9
    finally:
        book.close()


@pytest.mark.parametrize(
    'encoding,meta,bom',
    [
        ('cp1252', '', b''),
        ('cp1252', '<meta charset="windows-1252">', b''),
        ('cp1252', '<meta http-equiv="Content-Type" content="text/html; charset=iso-8859-1">', b''),
        ('utf-8', '<meta charset="utf-8">', b''),
        ('utf-8', '', b'\xef\xbb\xbf'),
    ],
)
def test_legacy_diary_download_resume_and_both_importers_keep_original_bytes(tmp_path, encoding, meta, bom):
    from nurus.capture.diaries import import_lote as import_shared
    from nurus.personal.bitacora_lote import import_lote as import_personal

    complete = 'Observación íntegra: Ángela pidió revisión — línea siguiente.\nSegunda línea'

    class EncodedBridge(Bridge):
        def call(self, command, payload=None, timeout=75):
            result = super().call(command, payload, timeout)
            if command == 'diary_open':
                text = popup(payload['params']).replace('<html>', '<html>' + meta)
                text = text.replace('Texto completo\nSegunda línea', complete)
                raw = bom + text.encode(encoding)
                result.update(body=base64.b64encode(raw).decode(), digest=hashlib.sha256(raw).hexdigest())
            return result

    result = run(tmp_path, EncodedBridge())
    assert result['estado'] == 'COMPLETA' and result['leidas'] == 2 and result['fallidas'] == 0
    records = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))['registros']
    for record in records:
        raw = (Path(result['folder']) / record['archivo']).read_bytes()
        assert raw == bom + popup(record['params']).replace('<html>', '<html>' + meta).replace(
            'Texto completo\nSegunda línea', complete
        ).encode(encoding)
        assert hashlib.sha256(raw).hexdigest() == record['sha256']
    for importer in (import_shared, import_personal):
        queries = importer(result['manifest'])['queries']
        assert len(queries) == 2
        for query, record in zip(queries, records):
            entry = query['entries'][0]
            assert (entry['texto'] if isinstance(entry, dict) else entry.texto) == complete
            assert query['html_sha256'] == record['sha256']
    bridge = EncodedBridge()
    resumed = run(tmp_path, bridge, resume=result['folder'])
    assert resumed['leidas'] == 2 and not bridge.openings


def test_legacy_encoding_does_not_bypass_remote_identity_or_accept_broken_bytes():
    params = {
        'tipo_popUp': '12',
        'COD_Tribunal': '49',
        'CRR_IdCausa': '100',
        'ID_Ingreso': '201',
        'COD_Etapa': '3',
        'TIP_Consulta': '2',
    }
    read = parser()
    with pytest.raises(ValueError, match='ingreso solicitado'):
        read(popup(params, wrong=True).encode('cp1252'), params)
    with pytest.raises(ValueError, match='decodificar íntegramente'):
        read(popup(params).encode('cp1252').replace(b'Texto completo', b'Texto\x81 completo'), params)
    with pytest.raises(ValueError, match='incompleta'):
        read(popup(params).encode('cp1252')[:-7], params)


def test_cancel_and_resume_keeps_verified_copy(tmp_path):
    cancel = threading.Event()

    def emit(name, value):
        if name == 'diary_progress':
            cancel.set()

    result = run(tmp_path, Bridge(), cancel=cancel, emit=emit)
    assert result['estado'] == 'CANCELADA' and result['leidas'] == 1
    bridge = Bridge()
    result = run(tmp_path, bridge, resume=result['folder'])
    assert result['estado'] == 'COMPLETA' and len(bridge.openings) == 1


def test_modified_file_reloaded_on_retry(tmp_path):
    result = run(tmp_path, Bridge())
    folder = Path(result['folder'])
    next(folder.glob('*.html')).write_text('cambio', encoding='utf-8')
    bridge = Bridge()
    result = run(tmp_path, bridge, resume=folder)
    # La copia alterada no se reutiliza ni se sobrescribe silenciosamente.
    assert len(bridge.openings) == 1 and result['fallidas'] == 1


def test_previous_reads_preserved_when_reenumeration_fails(tmp_path):
    result = run(tmp_path, Bridge())

    class FailedListing(Bridge):
        def call(self, command, payload=None, timeout=75):
            if command == 'search':
                import motor as m

                raise m.PocError('Listado no disponible')
            return super().call(command, payload, timeout)

    result = run(tmp_path, FailedListing(), resume=result['folder'])
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert result['estado'] == 'INCOMPLETA' and len(data['anteriores_no_reenumeradas']) == 2
    assert not data['registros']


def test_listing_uses_the_encoding_accepted_by_normal_download():
    from lectura_bitacoras import listing_rows
    from lxml import etree
    from perfiles import response_profile

    profile = selection_profile(Selection('49', 'Cumplimiento', '1'), CATALOG)
    table = listing(court='49')
    table.set('id', profile.table_id)
    source = document(Selection('49', 'Cumplimiento', '1'), 1, 1).decode()
    start = source.index('<table')
    end = source.index('</table>', start) + 8
    source = source[:start] + etree.tostring(table, encoding='unicode') + source[end:]
    source = source.replace('<html>', '<html><meta charset="iso-8859-1">').replace(
        'Persona ficticia', 'Persona de prueba áéñ'
    )
    body = source.encode('iso-8859-1')
    pairs = [*profile.target.items(), ('irAccion', 'Buscar medida')]
    accepted, current = response_profile(body, pairs, Selection('49', 'Cumplimiento', '1'), CATALOG)
    rows = listing_rows(body, accepted)
    assert current.records and len(rows) == 1 and rows[0]['ingreso_id'] == '200'
    assert rows[0]['nombre'] == 'Persona de prueba áéñ'


@pytest.mark.parametrize('failure_phase', ('prepare', 'search'))
def test_query_failure_counts_and_safe_diagnostic(tmp_path, failure_phase):
    class Failed(Bridge):
        def call(self, command, payload=None, timeout=75):
            if command == failure_phase:
                raise ValueError('Dato privado que no debe ir al diagnóstico')
            return super().call(command, payload, timeout)

    result = run(tmp_path, Failed())
    text = Path(result['manifest']).read_text(encoding='utf-8')
    data = json.loads(text)
    assert result['leidas'] == 0 and result['fallidas'] == 0 and result['consultas_fallidas'] == 1
    diagnostic = data['consultas'][0]['diagnostico']
    assert diagnostic['fase'] == ('PREPARAR_FORMULARIO' if failure_phase == 'prepare' else 'CONSULTAR_LISTADO')
    assert diagnostic['tipo'] == 'ValueError' and diagnostic['ubicacion']
    assert 'Dato privado' not in text and result['detalle']


def test_zero_reads_ui_reports_failed_queries(tmp_path):
    # Tk/Tcl reutilizado por muchos tests en Python 3.12 Windows puede fallar al
    # cargar init.tcl. Comprobar una instancia real en su propio proceso, igual
    # que el usuario, conservando todas las aserciones y sin omitir la prueba.
    import os, subprocess, sys, textwrap

    source = textwrap.dedent('''
        from pathlib import Path
        import sys,tkinter as tk
        import preferencias
        from descargador import Application
        from test_lotes import dispose_interface
        folder=Path(sys.argv[1]);preferencias.data_directory=lambda:folder
        root=tk.Tk();root.withdraw();app=Application(root,start_worker=False)
        try:
            app.process_event('diary_done',{'folder':str(folder),'leidas':0,'fallidas':0,
                'consultas_fallidas':3,'estado':'INCOMPLETA','detalle':'Falla en la consulta'})
            assert '3 consultas fallidas' in app.count.get()
            assert app.status.get().startswith('No se guardaron copias')
            assert 'Falla en la consulta' in app.status.get()
        finally:dispose_interface(app,root)
    ''')
    result = subprocess.run(
        [sys.executable, '-c', source, str(tmp_path)],
        cwd=Path(__file__).parent,
        env={**os.environ, 'PYTHONIOENCODING': 'utf-8'},
        capture_output=True,
        text=True,
        encoding='utf-8',
        timeout=45,
    )
    assert result.returncode == 0, result.stderr


def test_package_roundtrip_without_network(tmp_path):
    from prueba_bitacoras_paquete import check

    assert check(tmp_path)


def test_extension_260_without_diary_command_is_rejected_before_any_query(tmp_path):
    class Extension260(Bridge):
        def call(self, command, payload=None, timeout=75):
            if command == 'capabilities':
                return {'__sitfaError': 'SITFA_TASK'}
            return super().call(command, payload, timeout)

    bridge = Extension260()
    result = run(tmp_path, bridge)
    data = json.loads(Path(result['manifest']).read_text(encoding='utf-8'))
    assert not bridge.openings and not data['registros'] and not data['consultas']
    assert result['estado'] == 'INCOMPLETA' and 'Actualiza la extensión' in result['detalle']


def test_lectura_importa_nurus_solo_despues_de_ubicar_src(tmp_path):
    """Como en la CI del Descargador: sin paquete instalado, nurus existe solo tras parser()."""
    import subprocess
    import sys

    script = '''
import sys,tempfile,threading
from pathlib import Path
src=str(Path.cwd().parent/"src")
class Bloqueo:
    def find_spec(self,name,path=None,target=None):
        if name.split(".")[0]=="nurus" and src not in sys.path:
            raise ModuleNotFoundError("No module named 'nurus'",name="nurus")
sys.meta_path.insert(0,Bloqueo())
import lectura_bitacoras
from lotes import Batch
class Runner:
    cancel=threading.Event()
batch=Batch(("49",),("2",),("Cumplimiento",))
lectura_bitacoras._Lectura(Runner(),batch,tempfile.mkdtemp(),None)
'''
    result = subprocess.run(
        [sys.executable, '-c', script], cwd=Path(__file__).resolve().parent, capture_output=True, text=True, timeout=60
    )
    assert result.returncode == 0, result.stderr


def test_disk_failure_on_close_still_unlocks_sitfa(tmp_path, monkeypatch):
    import lectura_bitacoras

    def broken(path, data):
        if path.name == 'diagnostico_bitacoras.json':
            raise OSError('disco lleno ficticio')
        return real(path, data)

    real = lectura_bitacoras.write
    monkeypatch.setattr(lectura_bitacoras, 'write', broken)
    bridge = Bridge()
    with pytest.raises(OSError):
        run(tmp_path, bridge)
    assert ('unlock', None) in bridge.events
