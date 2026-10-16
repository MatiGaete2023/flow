"""D6: el paquete puede llevar solo el XLSX consolidado; sin HTML ni PDF."""

from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from openpyxl import load_workbook
from datetime import date
from nurus.capture import read_capture
from export_csmp_capture import export_capture
from nurus.personal.capture_input import import_capture
from nurus.personal.config import defaults
from nurus.personal.sync import refresh
from resultados import load_lot
from test_resultados import fixture


class ConsolidadoEnPaqueteTests(unittest.TestCase):
    def test_con_consolidado_1_1_0_con_los_mismos_registros_y_sin_html_ni_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            lot = load_lot(fixture(root / 'lot'))
            plain = read_capture(export_capture([lot], root / 'plain.zip'))
            full = read_capture(export_capture([lot], root / 'full.zip', consolidated=True))
            self.assertEqual(plain.manifest['schema_version'], '1.0.0')
            self.assertIsNone(plain.consolidated)
            self.assertEqual(full.manifest['schema_version'], '1.1.0')
            self.assertEqual(full.rows('registros'), plain.rows('registros'))
            book = load_workbook(BytesIO(full.consolidated), read_only=True)
            self.assertIn('Resumen', book.sheetnames)
            self.assertIn('SITFA_ARCHIVOS', book.sheetnames)
            names = zipfile.ZipFile(root / 'full.zip').namelist()
            self.assertFalse([n for n in names if n.lower().endswith(('.html', '.htm', '.pdf'))])

    def test_el_paquete_sin_consolidado_no_cambia_y_el_flujo_no_deja_archivos_en_el_lote(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder = fixture(root / 'lot')
            lot = load_lot(folder)
            before = sorted(p.name for p in folder.iterdir())
            a = export_capture([lot], root / 'a.zip')
            b = export_capture([lot], root / 'b.zip')
            self.assertEqual(a.read_bytes(), b.read_bytes())
            export_capture([lot], root / 'c.zip', consolidated=True)
            self.assertEqual(sorted(p.name for p in folder.iterdir()), before)

    def test_copia_parte_del_consolidado_real_y_la_sincronizacion_sigue_funcionando(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            lot = load_lot(fixture(root / 'lot'))
            work = import_capture(
                export_capture([lot], root / 'c.zip', consolidated=True), defaults(), as_of=date(2026, 10, 8)
            )
            ids = [row.id for row in work.rows]
            copy = work.export(root / 'copia.xlsx', backend='portable', reduced_fidelity=True)
            book = load_workbook(copy)
            self.assertEqual(book.sheetnames[:2], ['Resumen', 'SITFA_ARCHIVOS'])
            self.assertEqual(work.sheet, 'ESPERA')
            self.assertEqual(work.header, 1)
            sheet = book['ESPERA']
            titles = [c.value for c in sheet[1]]
            col = titles.index('NURUS_ID_REGISTRO') + 1
            self.assertEqual([sheet.cell(row.source_row, col).value for row in work.rows], ids)
            self.assertFalse(refresh(work))

    def test_si_la_sincronizacion_falla_el_trabajo_no_queda_a_medias(self):
        from unittest import mock

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            lot = load_lot(fixture(root / 'lot'))
            work = import_capture(
                export_capture([lot], root / 'c.zip', consolidated=True), defaults(), as_of=date(2026, 10, 8)
            )
            before = (work.sheet, work.header, work.source_hash, [row.source_row for row in work.rows])
            with mock.patch('nurus.personal.sync.remember_export', side_effect=ValueError('falla')):
                with self.assertRaises(ValueError):
                    work.export(root / 'copia.xlsx', backend='portable', reduced_fidelity=True)
            self.assertEqual((work.sheet, work.header, work.source_hash, [row.source_row for row in work.rows]), before)


if __name__ == '__main__':
    unittest.main()
