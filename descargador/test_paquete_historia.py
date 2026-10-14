"""WP-53.6: historia de la causa en Captura_CSMP.zip (alcance CAUSA, contrato aditivo, Asistente intacto)."""

import json
import subprocess
import sys
import threading
from hashlib import sha256
from datetime import date
from pathlib import Path

from descargador import package_current_lot
from lectura_historia import LectorHistoria
from nurus.capture import read_capture
from nurus.personal.capture_input import import_capture
from nurus.personal.config import defaults
from test_lectura_historia import LISTADO, PERFIL, QUERY, SELECCION, FakeBridge
from test_reutilizacion_local import bound_lot, diaries

ROOT = Path(__file__).resolve().parents[1]


def historia(tmp_path, bridge=None):
    carpeta = tmp_path / 'lote-historia'
    carpeta.mkdir()
    lector = LectorHistoria(bridge or FakeBridge(), lambda *a: None, threading.Event(), pause=0)
    lector.abrir(carpeta)
    lector.pagina(SELECCION, QUERY, 1, LISTADO, PERFIL)
    lector.cerrar()
    return carpeta / 'historia' / 'historias.json'


def paquete(tmp_path, nombre, **kw):
    lote = bound_lot(tmp_path / ('lote-' + nombre))
    manifiesto = diaries(tmp_path / ('Bitacoras ' + nombre), 2)
    return package_current_lot(lote.folder, manifiesto, tmp_path / (nombre + '.zip'), **kw)['path']


def validar(path):
    r = subprocess.run(
        [sys.executable, str(ROOT / 'tools' / 'flow_capture_cli.py'), 'validate', str(path)],
        capture_output=True,
        text=True,
        cwd=ROOT,
        check=False,
    )
    assert r.returncode == 0, r.stdout + r.stderr


def test_historia_en_paquete_alcance_causa(tmp_path):
    ruta = paquete(tmp_path, 'con', historia=historia(tmp_path))
    captura = read_capture(ruta)
    filas = captura.rows('historia_tramites')
    assert [f['folio'] for f in filas] == ['9120', '9119', '9117']
    assert all(
        f['scope'] == 'CAUSA'
        and f['ingreso_id'] == ''
        and f['identity_status'] == 'CAUSA'
        and f['source_system'] == 'RUS'
        and f['cause_id'] == '100'
        and f['annex_visible'] in ('true', 'false')
        and f['original_file'].startswith('respaldo/lote-historia/historia/historia-')
        for f in filas
    )
    assert json.loads(filas[0]['raw_json'])['Fecha Tram'] == '13/10/2099'
    assert captura.manifest['table_states']['historia_tramites'] == 'COMPLETA'
    consulta = next(q for q in captura.manifest['queries'] if q['screen'] == 'historia_tramites')
    filtros = json.loads(consulta['filters_json'])
    assert 'remote_identity' not in filtros and filtros['historia_causa']['cause_id'] == '100'
    assert filtros['ventana']['criterio'] == 'ventana_4_meses'
    assert any(i['code'] == 'HISTORIA_FILA_INCOMPLETA' for i in captura.rows('incidencias'))
    validar(ruta)


def test_sin_historia_el_paquete_no_cambia(tmp_path):
    ruta = paquete(tmp_path, 'sin')
    captura = read_capture(ruta)
    assert captura.manifest['table_states']['historia_tramites'] == 'NO_SOLICITADA'
    assert len(captura.rows('historia_tramites')) == 0
    assert all(q['screen'] != 'historia_tramites' for q in captura.manifest['queries'])
    import zipfile

    with zipfile.ZipFile(ruta) as z:
        assert 'data/historia_tramites.csv' not in z.namelist()


def test_historia_fallida_no_es_vacia_y_el_paquete_se_crea(tmp_path):
    ruta = paquete(tmp_path, 'fallida', historia=historia(tmp_path, FakeBridge(b'<html><table></table></html>')))
    captura = read_capture(ruta)
    assert len(captura.rows('historia_tramites')) == 0
    assert captura.manifest['table_states']['historia_tramites'] == 'FALLIDA'
    assert any(i['code'] == 'HISTORIA_FALLIDA' for i in captura.rows('incidencias'))
    validar(ruta)


def test_asistente_produce_el_mismo_trabajo_con_y_sin_historia(tmp_path):
    """Aislamiento: la historia de tribunal 49 no se asocia a registros de tribunal 113."""
    con = import_capture(paquete(tmp_path, 'a', historia=historia(tmp_path)), defaults(), as_of=date(2026, 10, 8))
    sin = import_capture(paquete(tmp_path, 'b'), defaults(), as_of=date(2026, 10, 8))
    assert len(con.rows) == len(sin.rows)
    assert [(r.observation, r.review, r.decisions) for r in con.rows] == [
        (r.observation, r.review, r.decisions) for r in sin.rows
    ]


def test_productor_historia_asistente_y_productos_con_identidad_alineada(tmp_path):
    from docx import Document
    from openpyxl import load_workbook
    from perfiles import Selection, selection_profile
    from test_lotes import CATALOG
    from test_lectura_historia import VENTANA
    from nurus.personal.config import BASE
    from nurus.personal.outputs import prepare_required_drafts
    from nurus.personal.resolutions import prepare_projects, generate_projects

    lote = bound_lot(tmp_path / 'originales', 'Espera')
    originales = {p.name: sha256(p.read_bytes()).hexdigest() for p in lote.folder.iterdir() if p.is_file()}
    listado = LISTADO.replace(b'49', b'113').replace(b'Cumplimiento', b'Espera').replace(b'C-99999-2099', b'X-1-2026')
    cuerpo = VENTANA.replace(b'C-99999-2099', b'X-1-2026').replace(b'2099', b'2026')
    cuerpo = cuerpo.replace(b'Tengase presente canario', b'Ingreso efectivo')

    class HistoriaFicticia(FakeBridge):
        def call(self, *args, **kwargs):
            result = super().call(*args, **kwargs)
            result['startedAt'] = '2026-10-14T12:00:00Z'
            return result

    bridge = HistoriaFicticia(cuerpo)
    seleccion = Selection('113', 'Espera', '2')
    lector = LectorHistoria(bridge, lambda *a: None, threading.Event(), pause=0)
    lector.abrir(lote.folder)
    lector.pagina(seleccion, QUERY, 1, listado, selection_profile(seleccion, CATALOG))
    lector.cerrar()
    manifiesto = diaries(tmp_path / 'bitacoras', 2)
    ruta = package_current_lot(
        lote.folder, manifiesto, tmp_path / 'Captura.zip', historia=lector.carpeta / 'historias.json'
    )['path']
    validar(ruta)
    trabajo = import_capture(ruta, defaults(), as_of=date(2026, 10, 14))
    primera, segunda = trabajo.rows
    assert len(bridge.calls) == 1 and 'HISTORIA.R3_CIERRE' in primera.rules and primera.actions == []
    assert 'HISTORIA.R3_CIERRE' not in segunda.rules and segunda.actions == ['programa_espera', 'PC_IE']
    trabajo.export(tmp_path / 'Observaciones.xlsx', backend='portable', reduced_fidelity=True)
    book = load_workbook(trabajo.output)
    try:
        sheet = book['Trabajo']
        headers = {c.value: c.column for c in sheet[1]}
        assert sheet.cell(2, headers['OBSERVACION']).value == primera.observation
        assert sheet.cell(3, headers['OBSERVACION']).value == segunda.observation
    finally:
        book.close()
    drafts = prepare_required_drafts(trabajo, modalities='Ficticias', period='octubre 2026', defer_attachments=True)
    programa = [d for d in drafts if d.kind == 'programa_espera']
    assert programa and all(d.record_ids == [segunda.id] for d in programa)
    projects, errors = prepare_projects(trabajo, [(segunda.id, 'PC_IE')], BASE / 'plantillas_word')
    assert projects and not errors
    word = tmp_path / 'Proyectos.docx'
    generate_projects(trabajo, projects, word)
    assert 'X-1-2026' in '\n'.join(p.text for p in Document(word).paragraphs)
    assert all(sha256((lote.folder / n).read_bytes()).hexdigest() == h for n, h in originales.items())


def test_e3_identidad_incierta_viaja_como_incidencia_sin_filas(tmp_path):
    fila = LISTADO.decode().split('<tr><td><a href="#" onclick="ShowHistoria')[1].split('</tr>')[0]
    cuerpo = LISTADO.replace(
        b'</table></form>',
        (
            '<tr><td><a href="#" onclick="ShowHistoria' + fila.replace(',100)', ',101)') + '</tr></table></form>'
        ).encode(),
    )
    carpeta = tmp_path / 'lote-historia'
    carpeta.mkdir()
    bridge = FakeBridge()
    lector = LectorHistoria(bridge, lambda *a: None, threading.Event(), pause=0)
    lector.abrir(carpeta)
    lector.pagina(SELECCION, QUERY, 1, cuerpo, PERFIL)
    lector.cerrar()
    assert bridge.calls == []
    ruta = paquete(tmp_path, 'incierta', historia=carpeta / 'historia' / 'historias.json')
    validar(ruta)
    captura = read_capture(ruta)
    assert captura.rows('historia_tramites') == []
    assert {i['code'] for i in captura.rows('incidencias')} >= {'HISTORIA_IDENTIDAD_INCIERTA'}


def test_e3_error_enumeracion_degrada_cobertura_con_y_sin_causas(tmp_path):
    for leida in (True, False):
        base = tmp_path / str(leida)
        base.mkdir()
        ruta = historia(base)
        datos = json.loads(ruta.read_text())
        datos['estado'] = 'INCOMPLETA'
        datos['errores_listado'] = [{'pagina': 2, 'error': 'HISTORIA_LISTADO_ERROR'}]
        if not leida:
            datos['causas'] = {}
            datos['vinculos'] = []
        ruta.write_text(json.dumps(datos))
        captura = read_capture(paquete(base, 'incompleta', historia=ruta))
        assert captura.manifest['table_states']['historia_tramites'] == ('PARCIAL' if leida else 'FALLIDA')
        assert any(
            q['query_id'].startswith('hist-listado-') and q['coverage'] == 'FALLIDA'
            for q in captura.manifest['queries']
        )
        assert any(i['code'] == 'HISTORIA_ENUMERACION_FALLIDA' for i in captura.rows('incidencias'))
